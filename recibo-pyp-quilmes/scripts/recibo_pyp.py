#!/usr/bin/env python3
"""recibo_pyp.py — helper de recibo-pyp-quilmes.js.

  python recibo_pyp.py leer <recibo.pdf>          -> JSON con los datos del recibo
  python recibo_pyp.py excel <datos.json> <salida.xlsx>

leer: el recibo de "Derechos de Publicidad y Propaganda" de ARQui es un PDF de
una pagina con el talon repetido tres veces. Del cuerpo salen CUIT,
destinatario, direccion y la clase/area de cada cartel; del talon, cuenta,
recibo, vencimiento, importe y mes.

excel: arma el resumen (hoja Recibos y, si viene, hoja Control planilla) a
partir del JSON que escribe el script.
"""

import json
import re
import sys


def numero(s):
    """'1.234,56' / '34694,00' -> 1234.56"""
    return float(s.replace('.', '').replace(',', '.'))


def leer(pdf):
    from pypdf import PdfReader

    texto = PdfReader(pdf).pages[0].extract_text() or ''
    lineas = [l.strip() for l in texto.splitlines() if l.strip()]
    out = {'archivo': pdf}

    if not any('PUBLICIDAD Y PROPAGANDA' in l for l in lineas):
        out['error'] = 'el PDF no es un recibo de Derechos de Publicidad y Propaganda'
        return out

    talon = re.search(
        r'CUENTA:\nRECIBO:\n(\d+-\d)\n(\d+)\nVENCIMIENTO:\nIMPORTE:\n([\d.,]+)\n(\d{2}/\d{2}/\d{4})\n(\d{1,2})\nMES\n(\d{4})',
        '\n'.join(lineas))
    if not talon:
        out['error'] = 'no encontre el talon (cuenta, recibo, vencimiento, importe) en el PDF'
        return out
    cuenta, recibo, importe, venc, mes, anio = talon.groups()
    out.update({
        'cuenta': cuenta,
        'recibo': recibo,
        'importe': numero(importe),
        'vencimiento': venc,
        'periodo': f'{anio}{int(mes):02d}',
    })

    # Cuerpo: CUIT (11 digitos) y, debajo, destinatario, calle, numero,
    # localidad, codigo postal y nomenclatura.
    i = next((k for k, l in enumerate(lineas) if re.fullmatch(r'\d{11}', l)), None)
    if i is not None:
        out['cuit'] = lineas[i]
        sig = lineas[i + 1:i + 7]
        if sig:
            out['destinatario'] = sig[0]
        if len(sig) >= 3:
            out['direccion'] = f'{sig[1]} {sig[2]}'
        if len(sig) >= 4:
            out['localidad'] = sig[3]

    # Carteles: entre "ID" y el primer "IMPORTE" van, por cartel, la clase (una
    # o dos lineas), el area (con coma) y el ID; al final, el importe.
    try:
        a = lineas.index('ID') + 1
        b = lineas.index('IMPORTE', a)
        bloque = lineas[a:b - 1]  # sin el importe
        carteles, clase = [], []
        k = 0
        while k < len(bloque):
            l = bloque[k]
            if re.fullmatch(r'\d+,\d+', l) and k + 1 < len(bloque) and re.fullmatch(r'\d+', bloque[k + 1]):
                carteles.append({'clase': ' '.join(clase), 'area': numero(l), 'id': bloque[k + 1]})
                clase = []
                k += 2
                continue
            clase.append(l)
            k += 1
        out['carteles'] = carteles
    except ValueError:
        out['carteles'] = []
    return out


COLUMNAS = [
    ('Sociedad', 'sociedad', 22), ('CUIT', 'cuit', 14), ('Cuenta', 'cuenta', 13),
    ('Destinatario', 'destinatario', 28), ('Clase', 'clase', 34), ('Área', 'area', 9),
    ('Período', 'periodo', 9), ('Vencimiento', 'vencimiento', 12), ('Vencido', 'vencido', 8),
    ('Importe', 'importe', 14), ('Nº recibo', 'recibo', 15), ('Estado', 'estado', 14),
    ('Observación', 'observacion', 50), ('Archivo', 'archivo', 60),
]


def excel(datos_json, salida):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    with open(datos_json, encoding='utf-8') as f:
        datos = json.load(f)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Recibos'
    negrita = Font(bold=True)
    gris = PatternFill('solid', fgColor='DDDDDD')
    ws.append([c[0] for c in COLUMNAS])
    for r in datos.get('recibos', []):
        fila = []
        for _, k, _ in COLUMNAS:
            v = r.get(k)
            if k == 'vencido':
                v = 'Sí' if v else ('No' if v is False else '')
            fila.append(v)
        ws.append(fila)
    for n, (_, k, ancho) in enumerate(COLUMNAS, 1):
        c = ws.cell(row=1, column=n)
        c.font, c.fill = negrita, gris
        ws.column_dimensions[c.column_letter].width = ancho
        if k in ('importe',):
            for fila in ws.iter_rows(min_row=2, min_col=n, max_col=n):
                fila[0].number_format = '"$" #,##0.00'
    ws.freeze_panes = 'A2'

    control = datos.get('control')
    if control:
        wc = wb.create_sheet('Control planilla')
        wc.append(['Cuenta', 'CUIT en ARQui', 'Destinatario', 'Planilla: nombre', 'Planilla: CUIT', 'Resultado'])
        for r in control:
            wc.append([r.get('cuenta'), r.get('cuit_arqui'), r.get('destinatario'),
                       r.get('nombre_planilla'), r.get('cuit_planilla'), r.get('resultado')])
        for n, ancho in enumerate([13, 14, 30, 22, 14, 60], 1):
            c = wc.cell(row=1, column=n)
            c.font, c.fill = negrita, gris
            wc.column_dimensions[c.column_letter].width = ancho
        wc.freeze_panes = 'A2'

    wb.save(salida)
    return {'xlsx': salida}


if __name__ == '__main__':
    if len(sys.argv) >= 3 and sys.argv[1] == 'leer':
        print(json.dumps(leer(sys.argv[2]), ensure_ascii=False))
    elif len(sys.argv) >= 4 and sys.argv[1] == 'excel':
        print(json.dumps(excel(sys.argv[2], sys.argv[3]), ensure_ascii=False))
    else:
        print(json.dumps({'error': 'uso: recibo_pyp.py leer <pdf> | excel <json> <xlsx>'}))
        sys.exit(1)
