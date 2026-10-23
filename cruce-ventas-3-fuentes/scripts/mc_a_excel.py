#!/usr/bin/env python3
"""
mc_a_excel.py — convierte el CSV de Mis Comprobantes de ARCA (emitidos o
recibidos) a Excel.

    python3 mc_a_excel.py <entrada.csv> <salida.xlsx> [--tipo=emitidos|recibidos]

Es el CSV tal cual: las mismas columnas en el mismo orden, una fila por
comprobante, solo que con tipos de dato reales (fechas, enteros, importes).
Sin fila de TOTAL: en Mis Comprobantes las notas de credito vienen con importe
POSITIVO, asi que una suma de la columna mezclaria facturas y NC (el cruce de
ventas las resta).

Hoja: "Mis Comprobantes Emitidos" (o "Recibidos").
Imprime un JSON con el resumen por stdout.
"""

import csv
import json
import re
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FUENTE = 'Arial'
AZUL = '1F4E79'
FMT_NUM = '#,##0.00;-#,##0.00;-'
FMT_ENT = '0'
FMT_FECHA = 'DD/MM/YYYY'

RE_ENTERO = re.compile(r'^(Tipo de Comprobante|Punto de Venta|N.mero Desde|N.mero Hasta|Tipo Doc\.)', re.I)
# Cod. de autorizacion (CAE de 14 digitos) y nro. de documento van como texto:
# como numero Excel los muestra en notacion cientifica o les come digitos.
RE_TEXTO = re.compile(r'^(C.d\. Autorizaci.n|Nro\. Doc\.|Denominaci.n|Moneda)', re.I)
RE_FECHA = re.compile(r'^Fecha', re.I)


def a_numero(s):
    """'13215,00' / '1.234,56' -> float · '' -> None."""
    s = (s or '').strip()
    if not s:
        return None
    try:
        return float(s.replace('.', '').replace(',', '.')) if ',' in s else float(s)
    except ValueError:
        return None


def a_entero(s):
    v = a_numero(s)
    return int(v) if v is not None else None


def a_fecha(s):
    s = (s or '').strip()
    for fmt in ('%Y-%m-%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def tipo_de(h):
    if RE_FECHA.match(h):
        return 'fecha'
    if RE_ENTERO.match(h):
        return 'entero'
    if RE_TEXTO.match(h):
        return 'texto'
    return 'numero'


def leer(csv_path):
    # Mis Comprobantes baja en UTF-8 (con o sin BOM); el Portal IVA en latin-1.
    for enc in ('utf-8-sig', 'latin-1'):
        try:
            with open(csv_path, encoding=enc, newline='') as fh:
                return list(csv.reader(fh, delimiter=';'))
        except UnicodeDecodeError:
            continue
    raise SystemExit('No pude leer el CSV')


def convertir(csv_path, xlsx_path, tipo):
    filas = leer(csv_path)
    if not filas:
        raise SystemExit('El CSV esta vacio.')
    encabezados = [h.strip() for h in filas[0]]
    datos = [f for f in filas[1:] if any((c or '').strip() for c in f)]
    tipos = [tipo_de(h) for h in encabezados]

    wb = Workbook()
    ws = wb.active
    ws.title = f'Mis Comprobantes {tipo.capitalize()}'

    for j, h in enumerate(encabezados, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(name=FUENTE, size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=AZUL)
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[1].height = 30

    for i, fila in enumerate(datos, start=2):
        for j, t in enumerate(tipos, start=1):
            crudo = fila[j - 1] if j - 1 < len(fila) else ''
            if t == 'fecha':
                v, fmt = a_fecha(crudo), FMT_FECHA
            elif t == 'entero':
                v, fmt = a_entero(crudo), FMT_ENT
            elif t == 'texto':
                v, fmt = ((crudo or '').strip() or None), None
            else:
                v, fmt = a_numero(crudo), FMT_NUM
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=FUENTE, size=10)
            if fmt:
                c.number_format = fmt

    for j, (h, t) in enumerate(zip(encabezados, tipos), start=1):
        ancho = {'fecha': 12, 'entero': 10, 'texto': 18}.get(t, 15)
        if re.match(r'^Denominaci', h, re.I):
            ancho = 40
        ws.column_dimensions[get_column_letter(j)].width = ancho
    ws.freeze_panes = 'A2'
    if datos:
        ws.auto_filter.ref = f'A1:{get_column_letter(len(encabezados))}{1 + len(datos)}'
    wb.save(xlsx_path)

    i_total = next((k for k, h in enumerate(encabezados) if re.match(r'^Imp\. Total$', h, re.I)), None)
    return {
        'ok': True, 'filas': len(datos), 'columnas': len(encabezados),
        'hoja': ws.title, 'archivo': xlsx_path,
        'suma_imp_total_sin_signo': round(sum(a_numero(f[i_total]) or 0 for f in datos), 2) if i_total is not None else None,
    }


def main():
    pos = [a for a in sys.argv[1:] if not a.startswith('--')]
    opts = dict(a[2:].split('=', 1) for a in sys.argv[1:] if a.startswith('--') and '=' in a)
    if len(pos) != 2:
        raise SystemExit(__doc__)
    tipo = opts.get('tipo', 'emitidos').lower()
    print(json.dumps(convertir(pos[0], pos[1], tipo), ensure_ascii=False))


if __name__ == '__main__':
    main()
