#!/usr/bin/env python3
"""
cruzar_ventas_3.py — cruza las ventas de un mes de TRES fuentes y explica las
diferencias:

    Mis Comprobantes (ARCA)  lo que ARCA tiene autorizado / registrado
    Portal IVA (libro)       lo que va (o fue) en la DDJJ de IVA
    Zetti (Subdiario 5.6.5)  lo que registro el sistema de la farmacia

    # Carpeta IVA del mes (lo normal): lee los tres archivos que dejan
    # mis-comprobantes-emitidos.js, portal-iva-descarga.js y zetti-iva-ventas.js
    python3 cruzar_ventas_3.py --cuit=<CUIT> --periodo=09/2026
        lee   [IVA]/AAAAMM - MIS COMPROBANTES - EMITIDOS.xlsx
              [IVA]/AAAAMM - PORTAL IVA - VENTAS.xlsx
              [IVA]/AAAAMM - ZETTI - VENTAS.xlsx
        deja  [IVA]/AAAAMM - CRUCE VENTAS MC PORTAL ZETTI.xlsx

    # Archivos sueltos (alcanza con dos de los tres)
    python3 cruzar_ventas_3.py --mc=<x.xlsx> --portal=<y.xlsx> --zetti=<z.xlsx> --salida=<cruce.xlsx>

Opciones: --tolerancia=0.5 (por componente) | --si-existe=abortar (default) o
reemplazar | --salida=<ruta> tambien en el modo carpeta | --incluir-nc-controlador.

Como se cruza:
  * Facturas, NC y ND con numero propio: por (tipo ARCA, punto de venta, numero).
    El TC/letra de Zetti (FV A, NC B, ...) se traduce al codigo de ARCA; si no
    aparece asi, se busca por (punto de venta, numero) si es unico.
  * Tiques del controlador fiscal: Mis Comprobantes y el Portal los traen uno por
    uno (tipo 83, 81, 82, 110...), Zetti trae un Z por dia. Se comparan sumados
    por (fecha, punto de venta) y, como el Z a veces queda con la fecha del dia
    siguiente, se marca el "corrimiento de fecha" cuando dos dias se compensan.
  * Signos: Mis Comprobantes trae las NC en positivo; se dan vuelta para que todo
    quede como en el Portal y en Zetti (NC en negativo, total neto).
  * IVA del Portal: suma de "Importe IVA x%" (el "Total IVA" del CSV viene
    redondeado por ARCA a 1 decimal en los tiques).
  * NC B a consumidor final del controlador fiscal viejo que Zetti lista pero que
    ya estan dentro del Z (ver separar_nc_controlador en cruzar_ventas.py): van a
    la hoja "NC controlador (fuera)".

Hojas: Resumen · Diferencias (solo lo que no cierra, con el motivo) · Detalle
(todo, para filtrar) · NC controlador (fuera).
Imprime un JSON con el resumen por stdout.
"""
import argparse
import json
import os
import re
import sys
from collections import defaultdict, OrderedDict
from datetime import date, datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from cruzar_ventas import (norm, r2, to_int, find_sheet, parse_zetti, separar_nc_controlador,  # noqa: E402
                           carpeta_iva, nombre_sociedad, parse_periodo, leer_libro, pesos, fallar, MESES)

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

FUENTES = ("mc", "portal", "zetti")
NOMBRE_FUENTE = {"mc": "Mis Comprobantes", "portal": "Portal IVA", "zetti": "Zetti"}
COMP = ("exen", "grav", "iva", "total")

TIPOS = {
    1: "Factura A", 2: "Nota de Debito A", 3: "Nota de Credito A", 4: "Recibo A",
    6: "Factura B", 7: "Nota de Debito B", 8: "Nota de Credito B", 9: "Recibo B",
    11: "Factura C", 12: "Nota de Debito C", 13: "Nota de Credito C", 15: "Recibo C",
    19: "Factura E", 20: "Nota de Debito E", 21: "Nota de Credito E",
    51: "Factura M", 52: "Nota de Debito M", 53: "Nota de Credito M",
    81: "Tique Factura A", 82: "Tique Factura B", 83: "Tique", 90: "Nota de Credito Otros",
    110: "Tique Nota de Credito", 111: "Tique Factura C", 112: "Tique NC A", 113: "Tique NC B",
    114: "Tique NC C", 115: "Tique ND A", 116: "Tique ND B", 117: "Tique ND C",
    118: "Tique Factura M", 119: "Tique NC M", 120: "Tique ND M",
    201: "FCE MiPyME A", 202: "ND MiPyME A", 203: "NC MiPyME A",
    206: "FCE MiPyME B", 207: "ND MiPyME B", 208: "NC MiPyME B",
    211: "FCE MiPyME C", 212: "ND MiPyME C", 213: "NC MiPyME C",
}
NOTAS_CREDITO = {3, 8, 13, 21, 53, 90, 110, 112, 113, 114, 119, 203, 208, 213}
TIQUES = {81, 82, 83, 110, 111, 112, 113, 114, 115, 116, 117, 118, 119, 120}
# TC + letra de Zetti -> codigo de ARCA
ZETTI_A_ARCA = {
    ("FV", "A"): 1, ("ND", "A"): 2, ("NC", "A"): 3,
    ("FV", "B"): 6, ("ND", "B"): 7, ("NC", "B"): 8,
    ("FV", "C"): 11, ("ND", "C"): 12, ("NC", "C"): 13,
    ("FV", "E"): 19, ("ND", "E"): 20, ("NC", "E"): 21,
    ("FV", "M"): 51, ("ND", "M"): 52, ("NC", "M"): 53,
}


def nombre_tipo(cod):
    return TIPOS.get(cod, f"Tipo {cod}") if cod is not None else ""


def grupo(cod):
    if cod in TIQUES:
        return "Tiques"
    if cod in NOTAS_CREDITO:
        return "Notas de credito"
    return "Facturas y ND"


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0


def como_fecha(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


# ---------------------------------------------------------------------------
# Lectores de los dos formatos de ARCA
# ---------------------------------------------------------------------------

def _encabezado(ws, requeridos):
    for r in range(1, 6):
        vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        normed = [norm(v) for v in vals]
        if all(x in normed for x in requeridos):
            return r, normed
    raise ValueError(f'No encontre el encabezado en la hoja "{ws.title}" (busco {", ".join(requeridos)}).')


def parse_arca(ws, fuente):
    """Lee la hoja de Mis Comprobantes Emitidos o de Ventas Portal IVA (CSV de
    ARCA pasado a Excel). Devuelve (filas, info). Cada fila:
    {fecha, tipo, punto, numero, hasta, doc, cliente, exen, grav, iva, total}
    con signo de libro (NC en negativo)."""
    if fuente == "mc":
        req = ["fecha de emision", "tipo de comprobante", "punto de venta", "numero desde"]
        mapa = {"fecha de emision": "fecha", "tipo de comprobante": "tipo", "punto de venta": "punto",
                "numero desde": "numero", "numero hasta": "hasta", "nro. doc. receptor": "doc",
                "denominacion receptor": "cliente", "imp. neto gravado total": "grav",
                "imp. neto no gravado": "no_grav", "imp. op. exentas": "exento",
                "total iva": "iva_total", "imp. total": "total"}
        prefijo_iva = re.compile(r"^iva \d")
    else:
        req = ["fecha de emision", "tipo de comprobante", "punto de venta", "numero de comprobante"]
        mapa = {"fecha de emision": "fecha", "tipo de comprobante": "tipo", "punto de venta": "punto",
                "numero de comprobante": "numero", "numero de comprobante hasta": "hasta",
                "nro. doc. comprador": "doc", "denominacion comprador": "cliente",
                "total neto gravado": "grav", "importe no gravado": "no_grav", "importe exento": "exento",
                "total iva": "iva_total", "importe total": "total"}
        prefijo_iva = re.compile(r"^importe iva \d")
    fila_enc, normed = _encabezado(ws, req)
    col = {}
    cols_iva = []
    for i, h in enumerate(normed):
        if h in mapa:
            col[mapa[h]] = i
        elif prefijo_iva.match(h):
            cols_iva.append(i)

    filas = []
    info = {"descartadas": [], "nc_con_signo_positivo": 0}
    for r, vals in enumerate(ws.iter_rows(min_row=fila_enc + 1, values_only=True), start=fila_enc + 1):
        g = lambda k: vals[col[k]] if k in col and col[k] < len(vals) else None  # noqa: E731
        fecha = como_fecha(g("fecha"))
        if fecha is None:
            if isinstance(g("total"), (int, float)) and not any(isinstance(v, str) and norm(v).startswith("total") for v in vals[:3]):
                info["descartadas"].append(r)
            continue
        tipo = to_int(g("tipo"))
        iva = sum(num(vals[i]) for i in cols_iva) if cols_iva else num(g("iva_total"))
        fila = {
            "fecha": fecha, "tipo": tipo, "punto": to_int(g("punto")), "numero": to_int(g("numero")),
            "hasta": to_int(g("hasta")), "doc": g("doc"),
            "cliente": (g("cliente") or "").strip() if isinstance(g("cliente"), str) else "",
            "exen": num(g("no_grav")) + num(g("exento")), "grav": num(g("grav")), "iva": iva,
            "total": num(g("total")), "fila": r,
        }
        if tipo in NOTAS_CREDITO:
            if fuente == "mc" and fila["total"] > 0:
                for k in COMP:
                    fila[k] = -fila[k]
            elif fuente == "portal" and fila["total"] > 0:
                # el Portal las trae negativas; si alguna viene positiva se avisa
                info["nc_con_signo_positivo"] += 1
        filas.append(fila)
    return filas, info


def zetti_a_filas(zetti_rows):
    out = []
    for z in zetti_rows:
        tc = (z["tc"] or "").strip().upper()
        m = (z["m"] or "").strip().upper() if isinstance(z["m"], str) else ""
        out.append({**z, "tipo": ZETTI_A_ARCA.get((tc, m)), "tc_zetti": f"{tc} {m}".strip()})
    return out


# ---------------------------------------------------------------------------
# Cruce
# ---------------------------------------------------------------------------

def vacio():
    return {k: 0.0 for k in COMP}


def sumar(acc, fila):
    for k in COMP:
        acc[k] += num(fila[k])


def cruzar(datos, tolerancia):
    """datos: {fuente: [filas]} (las fuentes que esten). Devuelve la lista de
    renglones del detalle, uno por comprobante o por (dia, PV) de tiques."""
    presentes = [f for f in FUENTES if f in datos]
    zetti = datos.get("zetti", [])

    # Puntos de venta que se cruzan sumados por dia: los que en Zetti tienen Z, y
    # los que en ARCA solo tienen tiques.
    pv_z = {r["punto"] for r in zetti if (r.get("tc") or "").upper() == "Z"}
    pv_arca = defaultdict(set)
    for f in ("mc", "portal"):
        for r in datos.get(f, []):
            pv_arca[r["punto"]].add(r["tipo"])
    pv_tiques = {pv for pv, tipos in pv_arca.items() if tipos and tipos <= TIQUES}
    pv_agregados = pv_z | pv_tiques

    renglones = OrderedDict()

    def renglon(clave, base):
        if clave not in renglones:
            renglones[clave] = {**base, "fuentes": {}, "clientes": {}}
        return renglones[clave]

    # --- comprobantes con numero propio ---
    indice_pv_nro = defaultdict(list)  # (pv, nro) -> claves ARCA, para el fallback de Zetti
    for f in ("mc", "portal"):
        for r in datos.get(f, []):
            if r["punto"] in pv_agregados:
                continue
            clave = ("c", r["tipo"], r["punto"], r["numero"])
            rg = renglon(clave, {"modo": "comprobante", "fecha": r["fecha"], "tipo": r["tipo"],
                                  "punto": r["punto"], "numero": r["numero"]})
            acc = rg["fuentes"].setdefault(f, {**vacio(), "n": 0, "fecha": r["fecha"]})
            sumar(acc, r); acc["n"] += 1
            if r["cliente"]:
                rg["clientes"].setdefault(f, r["cliente"])
            if clave not in indice_pv_nro[(r["punto"], r["numero"])]:
                indice_pv_nro[(r["punto"], r["numero"])].append(clave)

    for r in zetti:
        if r["punto"] in pv_agregados:
            continue
        clave = ("c", r["tipo"], r["punto"], r["numero"])
        if clave not in renglones:
            alt = indice_pv_nro.get((r["punto"], r["numero"]), [])
            if len(alt) == 1:
                clave = alt[0]
        rg = renglon(clave, {"modo": "comprobante", "fecha": r["fecha"], "tipo": r["tipo"],
                              "punto": r["punto"], "numero": r["numero"]})
        acc = rg["fuentes"].setdefault("zetti", {**vacio(), "n": 0, "fecha": r["fecha"]})
        sumar(acc, r); acc["n"] += 1
        acc["tc_zetti"] = r.get("tc_zetti")
        if r.get("cliente"):
            rg["clientes"].setdefault("zetti", r["cliente"])

    # --- tiques: por (fecha, PV) ---
    for f in presentes:
        for r in datos[f]:
            if r["punto"] not in pv_agregados:
                continue
            clave = ("z", r["fecha"], r["punto"])
            rg = renglon(clave, {"modo": "tiques", "fecha": r["fecha"], "tipo": None,
                                  "punto": r["punto"], "numero": None})
            acc = rg["fuentes"].setdefault(f, {**vacio(), "n": 0, "fecha": r["fecha"], "rangos": {}})
            sumar(acc, r); acc["n"] += 1
            lo, hi = r["numero"], r.get("hasta") or r["numero"]
            if f != "zetti" and isinstance(lo, int) and isinstance(hi, int):
                # cada tipo tiene su numeracion (83 tique, 81 tique factura A...)
                rgo = acc["rangos"].setdefault(r["tipo"], [lo, hi, 0])
                rgo[0], rgo[1], rgo[2] = min(rgo[0], lo), max(rgo[1], hi), rgo[2] + 1

    filas = list(renglones.values())
    for rg in filas:
        clasificar(rg, presentes, tolerancia)
    marcar_corrimientos(filas, presentes, tolerancia)
    marcar_corte_de_mes(filas, presentes)
    filas.sort(key=lambda x: (x["fecha"], 0 if x["modo"] == "comprobante" else 1,
                              x["punto"] or 0, x["tipo"] or 0, str(x["numero"] or "")))
    return filas, sorted(pv_agregados)


def _difiere(a, b, tol):
    return any(abs(r2(a[k] - b[k])) > tol for k in COMP)


def clasificar(rg, presentes, tol):
    fu = rg["fuentes"]
    en = [f for f in presentes if f in fu]
    faltan = [f for f in presentes if f not in fu]
    for f in presentes:
        rg.setdefault("valores", {})[f] = fu.get(f, vacio())
    v = rg["valores"]
    rg["dif_mc_portal"] = r2(v["mc"]["total"] - v["portal"]["total"]) if {"mc", "portal"} <= set(presentes) else None
    rg["dif_zetti_portal"] = r2(v["zetti"]["total"] - v["portal"]["total"]) if {"zetti", "portal"} <= set(presentes) else None
    rg["dif_zetti_mc"] = r2(v["zetti"]["total"] - v["mc"]["total"]) if {"zetti", "mc"} <= set(presentes) else None
    rg["en"] = en

    distintas = [f for f in en if any(_difiere(fu[f], fu[g], tol) for g in en if g != f)]
    tiques = rg["modo"] == "tiques"
    nombre = lambda fs: " y ".join(NOMBRE_FUENTE[f] for f in fs)  # noqa: E731

    if not faltan and not distintas:
        rg["estado"], rg["obs"] = "OK", ""
        return
    if len(en) == 1:
        solo = en[0]
        rg["estado"] = f"Solo en {NOMBRE_FUENTE[solo]}"
        rg["obs"] = {
            "zetti": ("Z de Zetti sin tiques en ARCA ese dia" if tiques else
                      "Registrado en Zetti sin comprobante en ARCA: sin CAE, anulado o mal numerado en Zetti"),
            "mc": ("Tiques en ARCA que no estan en el libro ni en Zetti" if tiques else
                   "Autorizado en ARCA pero no esta en el libro IVA ni en Zetti: el Portal se importo antes o se borro, y Zetti no lo tiene"),
            "portal": ("Tiques en el libro que no estan en ARCA ni en Zetti" if tiques else
                       "Solo en el libro IVA: cargado a mano en el Portal (no esta en Mis Comprobantes ni en Zetti)"),
        }[solo]
        return
    if faltan:
        falta = faltan[0] if len(faltan) == 1 else None
        rg["estado"] = f"Falta en {nombre(faltan)}"
        if falta == "portal":
            rg["obs"] = "Autorizado en ARCA y registrado en Zetti, pero no esta en el libro IVA: reimportar en el Portal (se importo antes de que estuviera) o se borro del borrador"
        elif falta == "mc":
            rg["obs"] = "Esta en el libro y en Zetti pero no en Mis Comprobantes: cargado a mano en el Portal, o comprobante de controlador fiscal viejo que ARCA no informa uno por uno"
        elif falta == "zetti":
            rg["obs"] = ("Tiques del dia en ARCA y en el libro sin Z en Zetti" if tiques else
                         "Emitido en ARCA y en el libro, pero no registrado en Zetti: facturado por fuera del sistema (RCEL / otro sistema)")
        else:
            rg["obs"] = ""
        if distintas:
            rg["obs"] += " · ademas difieren los importes entre " + nombre(en)
        return
    # estan en todas pero con importes distintos
    if len(en) == 3 and len(distintas) == 3:
        # cual es la que no coincide con las otras dos
        for f in en:
            otras = [g for g in en if g != f]
            if not _difiere(fu[otras[0]], fu[otras[1]], tol):
                distintas = [f]
                break
    if len(distintas) == 1:
        rg["estado"] = f"Importe distinto en {NOMBRE_FUENTE[distintas[0]]}"
    else:
        rg["estado"] = "Importes distintos"
    comps = [k for k in COMP if len({r2(fu[f][k]) for f in en}) > 1]
    dif_total = r2(max(fu[f]["total"] for f in en) - min(fu[f]["total"] for f in en))
    rg["obs"] = "Difiere: " + ", ".join({"exen": "exento/no gravado", "grav": "gravado", "iva": "IVA", "total": "total"}[k] for k in comps)
    if comps and set(comps) <= {"exen", "grav"} and abs(sum(r2(fu[en[0]][k] - fu[en[1]][k]) for k in ("exen", "grav"))) <= tol:
        rg["obs"] += " (misma base, distinta reparticion exento/gravado)"
    if dif_total >= 1000 and dif_total % 1000 == 0:
        rg["obs"] += f" · diferencia redonda de $ {pesos(dif_total)}: posible error de carga"


def marcar_corte_de_mes(filas, presentes):
    """Un comprobante que le falta a una fuente y cuyo numero queda ANTES del
    primero (o DESPUES del ultimo) que esa fuente tiene en el mes para el mismo
    tipo y PV casi seguro esta en esa fuente con fecha del mes de al lado: Zetti
    fecha la venta y ARCA el CAE (caso real 09/2026: 36 facturas de un PV que
    Zetti tiene el 31/08 y ARCA el 01/09, mismos importes)."""
    rango = defaultdict(lambda: [None, None])
    for rg in filas:
        if rg["modo"] != "comprobante" or not isinstance(rg["numero"], int):
            continue
        for f in rg["fuentes"]:
            x = rango[(f, rg["tipo"], rg["punto"])]
            x[0] = rg["numero"] if x[0] is None else min(x[0], rg["numero"])
            x[1] = rg["numero"] if x[1] is None else max(x[1], rg["numero"])
    for rg in filas:
        if rg["modo"] != "comprobante" or not isinstance(rg["numero"], int) or not rg["estado"].startswith(("Falta", "Solo")):
            continue
        faltan = [f for f in presentes if f not in rg["fuentes"]]
        # Solo Zetti contra ARCA: Mis Comprobantes y el Portal tienen la misma
        # fecha (la del CAE); lo que le falta al Portal es que se importo antes.
        if faltan != ["zetti"] and list(rg["fuentes"]) != ["zetti"]:
            continue
        lados = set()
        for f in faltan:
            lo, hi = rango.get((f, rg["tipo"], rg["punto"]), (None, None))
            if lo is None:
                lados = set()
                break
            lados.add("anterior" if rg["numero"] < lo else "siguiente" if rg["numero"] > hi else None)
        if len(lados) == 1 and None not in lados:
            lado = lados.pop()
            quien = " y ".join(NOMBRE_FUENTE[f] for f in faltan)
            rg["estado"] = f"Corte de mes ({quien} en el mes {lado})"
            rg["obs"] = (f"El numero queda {'antes del primero' if lado == 'anterior' else 'despues del ultimo'} "
                         f"que {quien} tiene en el mes para este tipo y PV: lo tiene con fecha del mes {lado} "
                         f"(revisar el subdiario/libro de ese mes)")


def marcar_corrimientos(filas, presentes, tol):
    """Tiques: un dia que le sobra a una fuente y el dia siguiente/anterior del
    mismo PV que le falta lo mismo = el Z quedo con otra fecha."""
    if "zetti" not in presentes:
        return
    ref = "portal" if "portal" in presentes else "mc"
    tiq = [f for f in filas if f["modo"] == "tiques" and f["estado"] != "OK"]
    for f in tiq:
        d = r2(f["valores"]["zetti"]["total"] - f["valores"][ref]["total"])
        f["_d"] = d
        if abs(d) <= tol * 4:
            f["obs"] = (f["obs"] + " · " if f["obs"] else "") + "redondeo"
    for f in tiq:
        if abs(f["_d"]) <= tol * 4:
            continue
        compensa = [g for g in tiq if g is not f and g["punto"] == f["punto"] and g["fecha"] != f["fecha"]
                    and abs(g["_d"] + f["_d"]) <= max(2.0, tol * 4) and abs((g["fecha"] - f["fecha"]).days) <= 3]
        if compensa:
            f["estado"] = "Corrimiento de fecha (tiques)"
            f["obs"] = ("El Z de Zetti y los tiques de ARCA quedaron en dias distintos; se compensa con "
                        + ", ".join(g["fecha"].strftime("%d/%m") for g in compensa))


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

AZUL = "1F4E79"
F_TIT = Font(name="Arial", size=12, bold=True)
F_B = Font(name="Arial", size=10, bold=True)
F_N = Font(name="Arial", size=10)
F_H = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FILL_H = PatternFill("solid", fgColor=AZUL)
FILL_FU = {"mc": PatternFill("solid", fgColor="DDEBF7"), "portal": PatternFill("solid", fgColor="E2EFDA"),
           "zetti": PatternFill("solid", fgColor="FFF2CC"), "dif": PatternFill("solid", fgColor="FCE4D6")}
FILL_ROJO = PatternFill("solid", fgColor="F8CBAD")
FMT = "#,##0.00;-#,##0.00;-"
FMT_F = "DD/MM/YYYY"
BORDE = Border(bottom=Side(style="thin", color="BFBFBF"))


def concepto(rg):
    if rg["modo"] == "tiques":
        rangos = next((v["rangos"] for f, v in rg["fuentes"].items() if f != "zetti" and v.get("rangos")), {})
        n = {f: v["n"] for f, v in rg["fuentes"].items()}
        txt = f"Tiques PV {rg['punto']}"
        if rangos:
            tipo, (lo, hi, _) = max(rangos.items(), key=lambda x: x[1][2])
            txt += f" ({nombre_tipo(tipo)} {lo}-{hi}" + (f" y {len(rangos) - 1} tipo(s) mas" if len(rangos) > 1 else "") + ")"
        partes = []
        if "mc" in n or "portal" in n:
            partes.append(f"{max(n.get('mc', 0), n.get('portal', 0))} tiques")
        if "zetti" in n:
            partes.append(f"{n['zetti']} Z")
        return txt + (" · " + ", ".join(partes) if partes else "")
    t = nombre_tipo(rg["tipo"]) or (rg["fuentes"].get("zetti", {}).get("tc_zetti") or "")
    return f"{t} {rg['punto']:05d}-{rg['numero']:08d}" if isinstance(rg["punto"], int) and isinstance(rg["numero"], int) \
        else f"{t} {rg['punto']}-{rg['numero']}"


def cliente(rg):
    c = rg["clientes"]
    return c.get("mc") or c.get("portal") or c.get("zetti") or ""


def hoja_detalle(wb, titulo_hoja, filas, presentes, titulo):
    ws = wb.create_sheet(titulo_hoja)
    ws.cell(row=1, column=1, value=titulo).font = F_TIT
    fijas = ["Fecha", "Comprobante", "Cliente", "Estado", "Observacion"]
    nf = len(fijas)
    col = nf + 1
    # fila 2: grupo de fuente; fila 3: componente
    for j, h in enumerate(fijas, start=1):
        c = ws.cell(row=3, column=j, value=h); c.font = F_H; c.fill = FILL_H
        ws.cell(row=2, column=j).fill = FILL_H
    cols_fuente = {}
    for f in presentes:
        ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + 3)
        c = ws.cell(row=2, column=col, value=NOMBRE_FUENTE[f].upper()); c.font = F_B; c.fill = FILL_FU[f]
        c.alignment = Alignment(horizontal="center")
        for k, h in enumerate(("Exento/No grav.", "Gravado", "IVA", "Total")):
            c = ws.cell(row=3, column=col + k, value=h); c.font = F_B; c.fill = FILL_FU[f]
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        cols_fuente[f] = col
        col += 4
    difs = []
    if {"mc", "portal"} <= set(presentes):
        difs.append(("dif_mc_portal", "MC - Portal"))
    if {"zetti", "portal"} <= set(presentes):
        difs.append(("dif_zetti_portal", "Zetti - Portal"))
    if {"zetti", "mc"} <= set(presentes):
        difs.append(("dif_zetti_mc", "Zetti - MC"))
    if difs:
        ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + len(difs) - 1)
        c = ws.cell(row=2, column=col, value="DIFERENCIA DE TOTAL"); c.font = F_B; c.fill = FILL_FU["dif"]
        c.alignment = Alignment(horizontal="center")
    for k, (_, h) in enumerate(difs):
        c = ws.cell(row=3, column=col + k, value=h); c.font = F_B; c.fill = FILL_FU["dif"]
        c.alignment = Alignment(horizontal="center")
    ultima_col = col + len(difs) - 1

    r = 4
    for rg in filas:
        ws.cell(row=r, column=1, value=rg["fecha"]).number_format = FMT_F
        ws.cell(row=r, column=2, value=concepto(rg))
        ws.cell(row=r, column=3, value=cliente(rg))
        ce = ws.cell(row=r, column=4, value=rg["estado"])
        if rg["estado"] != "OK":
            ce.fill = FILL_ROJO
        ws.cell(row=r, column=5, value=rg["obs"])
        for f in presentes:
            if f not in rg["fuentes"]:
                continue
            for k, comp in enumerate(COMP):
                c = ws.cell(row=r, column=cols_fuente[f] + k, value=r2(rg["valores"][f][comp]))
                c.number_format = FMT
        for k, (clave, _) in enumerate(difs):
            c = ws.cell(row=r, column=col + k, value=rg[clave]); c.number_format = FMT
        for j in range(1, ultima_col + 1):
            ws.cell(row=r, column=j).font = F_N
        r += 1
    # totales
    ws.cell(row=r, column=1, value="TOTAL").font = F_B
    for j in range(cols_fuente[presentes[0]], ultima_col + 1):
        L = get_column_letter(j)
        c = ws.cell(row=r, column=j, value=f"=SUBTOTAL(9,{L}4:{L}{r - 1})" if r > 4 else 0)
        c.number_format = FMT; c.font = F_B
    for j, w in enumerate([11, 42, 30, 26, 60], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for j in range(nf + 1, ultima_col + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14
    ws.freeze_panes = "C4"
    if r > 4:
        ws.auto_filter.ref = f"A3:{get_column_letter(ultima_col)}{r - 1}"
    return ws


def totales_por_fuente(datos):
    out = {}
    for f, filas in datos.items():
        g = OrderedDict((k, {**vacio(), "n": 0}) for k in ("Facturas y ND", "Notas de credito", "Tiques"))
        tot = {**vacio(), "n": 0}
        for r in filas:
            if f == "zetti":
                tc = (r.get("tc") or "").upper()
                k = "Tiques" if tc == "Z" else ("Notas de credito" if tc == "NC" else "Facturas y ND")
            else:
                k = grupo(r["tipo"])
            sumar(g[k], r); g[k]["n"] += 1
            sumar(tot, r); tot["n"] += 1
        out[f] = {"grupos": g, "total": tot}
    return out


def hoja_resumen(wb, titulo, presentes, tot, filas, archivos, separadas, avisos):
    ws = wb.active
    ws.title = "Resumen"
    ws.cell(row=1, column=1, value=titulo).font = F_TIT
    r = 2
    for f in presentes:
        ws.cell(row=r, column=1, value=f"{NOMBRE_FUENTE[f]}: {archivos.get(f, '')}").font = F_N
        r += 1
    r += 1

    def encabezado(fila, textos):
        for j, h in enumerate(textos, start=1):
            c = ws.cell(row=fila, column=j, value=h); c.font = F_H; c.fill = FILL_H
            c.alignment = Alignment(horizontal="center", wrap_text=True)

    ws.cell(row=r, column=1, value="1. Totales de cada fuente (NC en negativo)").font = F_B
    r += 1
    encabezado(r, ["Fuente", "Grupo", "Cantidad", "Exento/No grav.", "Gravado", "IVA", "Total"])
    r += 1
    for f in presentes:
        for gname, v in list(tot[f]["grupos"].items()) + [("TOTAL", tot[f]["total"])]:
            ws.cell(row=r, column=1, value=NOMBRE_FUENTE[f]).fill = FILL_FU[f]
            ws.cell(row=r, column=2, value=gname)
            ws.cell(row=r, column=3, value=v["n"])
            for k, comp in enumerate(COMP):
                ws.cell(row=r, column=4 + k, value=r2(v[comp])).number_format = FMT
            for j in range(1, 8):
                ws.cell(row=r, column=j).font = F_B if gname == "TOTAL" else F_N
            r += 1
    if separadas:
        s = {**vacio()}
        for x in separadas:
            sumar(s, x)
        ws.cell(row=r, column=1, value="Zetti").fill = FILL_FU["zetti"]
        ws.cell(row=r, column=2, value="NC controlador fuera del cruce")
        ws.cell(row=r, column=3, value=len(separadas))
        for k, comp in enumerate(COMP):
            ws.cell(row=r, column=4 + k, value=r2(s[comp])).number_format = FMT
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="2. Diferencias de total entre fuentes").font = F_B
    r += 1
    encabezado(r, ["Comparacion", "", "", "Exento/No grav.", "Gravado", "IVA", "Total"])
    r += 1
    # tot["zetti"] ya viene sin las NC del controlador (se separan antes del cruce)
    pares = [("mc", "portal"), ("zetti", "portal"), ("zetti", "mc")]
    for a, b in pares:
        if a not in presentes or b not in presentes:
            continue
        ws.cell(row=r, column=1, value=f"{NOMBRE_FUENTE[a]} - {NOMBRE_FUENTE[b]}")
        if a == "zetti" and separadas:
            ws.cell(row=r, column=2, value="(sin NC controlador)")
        for k, comp in enumerate(COMP):
            ws.cell(row=r, column=4 + k, value=r2(tot[a]["total"][comp] - tot[b]["total"][comp])).number_format = FMT
        for j in range(1, 8):
            ws.cell(row=r, column=j).font = F_N
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="3. Que explica las diferencias (renglones de la hoja Diferencias)").font = F_B
    r += 1
    cab = ["Estado", "Renglones", ""]
    if {"mc", "portal"} <= set(presentes):
        cab.append("MC - Portal")
    if {"zetti", "portal"} <= set(presentes):
        cab.append("Zetti - Portal")
    if {"zetti", "mc"} <= set(presentes):
        cab.append("Zetti - MC")
    encabezado(r, cab)
    r += 1
    por_estado = OrderedDict()
    for rg in filas:
        if rg["estado"] == "OK":
            continue
        e = por_estado.setdefault(rg["estado"], {"n": 0, "dif_mc_portal": 0.0, "dif_zetti_portal": 0.0, "dif_zetti_mc": 0.0})
        e["n"] += 1
        for k in ("dif_mc_portal", "dif_zetti_portal", "dif_zetti_mc"):
            e[k] += rg[k] or 0
    claves = [k for k, h in (("dif_mc_portal", "MC - Portal"), ("dif_zetti_portal", "Zetti - Portal"), ("dif_zetti_mc", "Zetti - MC")) if h in cab]
    for estado, e in sorted(por_estado.items(), key=lambda x: -x[1]["n"]):
        ws.cell(row=r, column=1, value=estado)
        ws.cell(row=r, column=2, value=e["n"])
        for k, clave in enumerate(claves):
            ws.cell(row=r, column=4 + k, value=r2(e[clave])).number_format = FMT
        for j in range(1, 8):
            ws.cell(row=r, column=j).font = F_N
        r += 1
    if not por_estado:
        ws.cell(row=r, column=1, value="Sin diferencias: las fuentes coinciden comprobante por comprobante.").font = F_N
        r += 1
    ok = sum(1 for rg in filas if rg["estado"] == "OK")
    ws.cell(row=r, column=1, value=f"OK (coinciden): {ok} renglones").font = F_N
    r += 2
    if avisos:
        ws.cell(row=r, column=1, value="Avisos").font = F_B
        r += 1
        for a in avisos:
            ws.cell(row=r, column=1, value=a).font = F_N
            r += 1
    for j, w in enumerate([34, 30, 11, 16, 16, 16, 16], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w


def hoja_nc(wb, separadas):
    ws = wb.create_sheet("NC controlador (fuera)")
    ws.cell(row=1, column=1, value="NC B a Consumidor Final de Zetti en puntos de venta sin NC B en ARCA: "
            "quedan fuera del cruce (ya estan dentro del cierre Z)").font = F_B
    for j, h in enumerate(["Fecha", "PV", "Numero", "Cliente", "Exento", "Gravado", "IVA", "Total"], start=1):
        c = ws.cell(row=3, column=j, value=h); c.font = F_H; c.fill = FILL_H
    r = 4
    for x in separadas:
        for j, v in enumerate([x["fecha"], x["punto"], x["numero"], x.get("cliente"), r2(x["exen"]), r2(x["grav"]), r2(x["iva"]), r2(x["total"])], start=1):
            c = ws.cell(row=r, column=j, value=v); c.font = F_N
            c.number_format = FMT_F if j == 1 else (FMT if j >= 5 else "General")
        r += 1
    ws.cell(row=r, column=1, value="TOTAL").font = F_B
    for j in range(5, 9):
        L = get_column_letter(j)
        c = ws.cell(row=r, column=j, value=f"=SUM({L}4:{L}{r - 1})" if r > 4 else 0); c.number_format = FMT; c.font = F_B
    for j, w in enumerate([11, 6, 10, 30, 14, 14, 14, 14], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def leer_fuente(f, ruta):
    wb = leer_libro(ruta)
    if f == "zetti":
        nombre = find_sheet(wb, ["zetti"]) or wb.sheetnames[0]
        rows, _, info = parse_zetti(wb[nombre])
        return zetti_a_filas(rows), {"razon_social": info.get("razon_social"), "descartadas": info["filas_descartadas"]}
    nombre = (find_sheet(wb, ["comprobantes"]) if f == "mc" else find_sheet(wb, ["portal"])) or wb.sheetnames[0]
    return parse_arca(wb[nombre], f)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--cuit"); ap.add_argument("--periodo")
    ap.add_argument("--mc"); ap.add_argument("--portal"); ap.add_argument("--zetti")
    ap.add_argument("--salida"); ap.add_argument("--sociedad")
    ap.add_argument("--tolerancia", type=float, default=0.5)
    ap.add_argument("--si-existe", default="abortar", choices=["abortar", "reemplazar"])
    ap.add_argument("--incluir-nc-controlador", action="store_true")
    a = ap.parse_args()

    rutas = {}
    sociedad = a.sociedad
    aaaamm = parse_periodo(a.periodo) if a.periodo else None
    if a.cuit:
        if not aaaamm:
            fallar("Con --cuit hace falta --periodo.")
        soc, iva = carpeta_iva(a.cuit, aaaamm)
        sociedad = sociedad or nombre_sociedad(soc)
        for f, nombre in (("mc", "MIS COMPROBANTES - EMITIDOS"), ("portal", "PORTAL IVA - VENTAS"), ("zetti", "ZETTI - VENTAS")):
            p = os.path.join(iva, f"{aaaamm} - {nombre}.xlsx")
            if getattr(a, f):
                rutas[f] = getattr(a, f)
            elif os.path.exists(p):
                rutas[f] = p
        salida = a.salida or os.path.join(iva, f"{aaaamm} - CRUCE VENTAS MC PORTAL ZETTI.xlsx")
    else:
        for f in FUENTES:
            if getattr(a, f):
                rutas[f] = getattr(a, f)
        if not a.salida:
            fallar("Sin --cuit hace falta --salida.")
        salida = a.salida
    if len(rutas) < 2:
        fallar(f"Hacen falta al menos dos fuentes; encontre: {', '.join(NOMBRE_FUENTE[f] for f in rutas) or 'ninguna'}.",
               buscadas=list(rutas.values()))
    if os.path.exists(salida) and a.si_existe == "abortar":
        fallar(f"Ya existe {salida}. Usa --si-existe=reemplazar.", code="ya_existe")

    presentes = [f for f in FUENTES if f in rutas]
    datos, infos, avisos = {}, {}, []
    for f in presentes:
        try:
            datos[f], infos[f] = leer_fuente(f, rutas[f])
        except Exception as e:  # noqa: BLE001
            fallar(f"No pude leer {NOMBRE_FUENTE[f]} ({rutas[f]}): {e}")
        if infos[f].get("descartadas"):
            avisos.append(f"{NOMBRE_FUENTE[f]}: filas con importe y sin fecha/numero que no entraron al cruce: {infos[f]['descartadas'][:20]}")
        if infos[f].get("nc_con_signo_positivo"):
            avisos.append(f"{NOMBRE_FUENTE[f]}: {infos[f]['nc_con_signo_positivo']} NC con importe positivo (se cruzan como vienen).")

    # periodo: el de la mayoria de las fechas, si no vino por argumento
    if not aaaamm:
        fechas = [r["fecha"] for f in presentes for r in datos[f]]
        if fechas:
            c = defaultdict(int)
            for d in fechas:
                c[f"{d.year}{d.month:02d}"] += 1
            aaaamm = max(c, key=c.get)
    for f in presentes:
        fuera = [r for r in datos[f] if aaaamm and f"{r['fecha'].year}{r['fecha'].month:02d}" != aaaamm]
        if fuera:
            avisos.append(f"{NOMBRE_FUENTE[f]}: {len(fuera)} comprobantes con fecha fuera de {aaaamm[4:]}/{aaaamm[:4]} (se cruzan igual).")

    separadas = []
    if "zetti" in datos and not a.incluir_nc_controlador:
        arca = datos.get("portal", []) + datos.get("mc", [])
        datos["zetti"], separadas = separar_nc_controlador(datos["zetti"], arca)

    filas, pv_agregados = cruzar(datos, a.tolerancia)
    tot = totales_por_fuente(datos)

    mes_txt = f"{MESES[int(aaaamm[4:]) - 1]} {aaaamm[:4]}" if aaaamm else ""
    razon = (infos.get("zetti") or {}).get("razon_social")
    titulo = f"CRUCE DE VENTAS - Mis Comprobantes / Portal IVA / Zetti - {sociedad or ''}" + \
             (f" ({razon})" if razon else "") + (f" - {mes_txt}" if mes_txt else "")
    if pv_agregados:
        avisos.append("Puntos de venta de tiques (cruzados por dia): " + ", ".join(str(p) for p in pv_agregados))

    wb = Workbook()
    hoja_resumen(wb, titulo, presentes, tot, filas, {f: os.path.basename(p) for f, p in rutas.items()}, separadas, avisos)
    difs = [rg for rg in filas if rg["estado"] != "OK"]
    hoja_detalle(wb, "Diferencias", difs, presentes, titulo + " - solo diferencias")
    hoja_detalle(wb, "Detalle", filas, presentes, titulo + " - todo")
    if separadas:
        hoja_nc(wb, separadas)
    os.makedirs(os.path.dirname(os.path.abspath(salida)), exist_ok=True)
    wb.save(salida)

    def t(f):
        return {k: r2(v) for k, v in tot[f]["total"].items()} if f in tot else None

    estados = defaultdict(int)
    for rg in difs:
        estados[rg["estado"]] += 1
    print(json.dumps({
        "ok": True, "salida": salida, "sociedad": sociedad, "periodo": aaaamm,
        "fuentes": {f: os.path.basename(p) for f, p in rutas.items()},
        "totales": {f: t(f) for f in presentes},
        "nc_controlador_fuera": {"cantidad": len(separadas), "total": r2(sum(num(x["total"]) for x in separadas))},
        "renglones": len(filas), "renglones_ok": len(filas) - len(difs), "renglones_con_diferencia": len(difs),
        "por_estado": dict(estados), "avisos": avisos,
        "diferencias": [{"fecha": rg["fecha"].isoformat(), "comprobante": concepto(rg), "cliente": cliente(rg),
                         "estado": rg["estado"], "obs": rg["obs"],
                         "total": {f: r2(rg["valores"][f]["total"]) for f in presentes if f in rg["fuentes"]}}
                        for rg in difs[:200]],
    }, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
