#!/usr/bin/env python3
"""
Cliente de la planilla de seguimiento de pedidos Estudio <-> Cliente (Hoja de cálculo de Google).

La planilla es una Google Sheet con el Apps Script de apps-script/Codigo.gs publicado como
app web. Este script le habla por HTTPS con PEDIDOS_URL y PEDIDOS_TOKEN del .env. Solo usa la
biblioteca estándar de Python (openpyxl hace falta únicamente para `importar`).

Subcomandos (devuelven JSON por stdout, salvo `pendientes --formato=texto|slack`):

  probar                                           Verifica la conexión con la planilla.
  estado     --email=<mail> | --persona=<nombre> | --slack=<ID>
                                                   Quién corre, desde cuándo leer, configuración y
                                                   pedidos ya cargados (para no duplicar).
  aplicar    <ops.json> [--solo-revisar]           Agrega pedidos, actualiza estados y suma reclamos.
  pendientes --persona=<nombre> [--formato=json|texto|slack] [--hoy=AAAA-MM-DD]
                                                   Lo abierto de esa persona, vencidos primero.
  importar   --xlsx=<planilla anterior.xlsx>       Migra la planilla Excel anterior (solo con la
                                                   Google Sheet vacía).

Formato de <ops.json>:
  {"persona": "<Nombre>", "desde": "AAAA-MM-DD", "hasta": "AAAA-MM-DD",
   "ops": [
     {"accion": "agregar", "hoja": "EC", "clave_origen": "gmail:<threadId>#1",
      "datos": {"fecha_pedido": "AAAA-MM-DD", "a_quien": "<persona del cliente>", "que": "...", ...}},
     {"accion": "actualizar", "id": "E-003", "datos": {"estado": "Completado", "fecha_recibido": "AAAA-MM-DD"}},
     {"accion": "reclamo", "id": "E-003", "fecha": "AAAA-MM-DD", "detalle": "Reclamado por mail"}
   ]}
  Hojas: EC = Estudio -> Cliente, CE = Cliente -> Estudio, TI = Tareas internas.
"""
import json
import os
import re
import sys
import urllib.error
import urllib.request
from datetime import date, datetime

PLUGIN = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HOJAS = {'EC': 'Estudio → Cliente', 'CE': 'Cliente → Estudio', 'TI': 'Tareas internas'}


def fail(msg, **extra):
    print(json.dumps({'ok': False, 'error': msg, **extra}, ensure_ascii=False, indent=1))
    sys.exit(1)


# ── .env (mismo orden que el resto de los plugins) ──────────────────────────
def cargar_env():
    for ruta in (os.environ.get('FISCO_ENV_FILE'), os.path.join(PLUGIN, '.env'),
                 os.path.join(os.path.expanduser('~'), '.fisco-ar', '.env')):
        if not ruta or not os.path.isfile(ruta):
            continue
        with open(ruta, encoding='utf-8') as f:
            for linea in f:
                m = re.match(r'\s*([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.*)\s*$', linea)
                if not m:
                    continue
                v = m.group(2).strip().strip('"').strip("'")
                if v and not os.environ.get(m.group(1)):
                    os.environ[m.group(1)] = v


# ── conexión con la app web ─────────────────────────────────────────────────
def llamar(accion, **datos):
    url = (os.environ.get('PEDIDOS_URL') or '').strip()
    token = (os.environ.get('PEDIDOS_TOKEN') or '').strip()
    if not url or not token:
        fail('Faltan PEDIDOS_URL y/o PEDIDOS_TOKEN en el .env (~/.fisco-ar/.env). '
             'Son la URL de la app web del Apps Script de la planilla y su token: ver el README del plugin.')
    cuerpo = json.dumps({'token': token, 'accion': accion, **datos}, ensure_ascii=False).encode('utf-8')
    req = urllib.request.Request(url, data=cuerpo, method='POST',
                                 headers={'Content-Type': 'text/plain;charset=utf-8'})
    try:
        # Apps Script contesta el POST con un 302 a googleusercontent.com: urllib lo sigue como GET.
        with urllib.request.urlopen(req, timeout=180) as r:
            texto = r.read().decode('utf-8')
    except urllib.error.HTTPError as e:
        fail(f'La planilla respondió HTTP {e.code}. Revisá PEDIDOS_URL (tiene que terminar en /exec).')
    except urllib.error.URLError as e:
        fail(f'No pude conectarme con la planilla: {e.reason}')
    try:
        res = json.loads(texto)
    except ValueError:
        fail('La planilla no devolvió JSON. ¿La app web está publicada con acceso "Cualquier persona"?',
             respuesta=texto[:300])
    if not res.get('ok'):
        fail(res.get('error') or 'Error en la planilla.', **{k: v for k, v in res.items() if k not in ('ok', 'error')})
    return res


def imprimir(obj):
    print(json.dumps(obj, ensure_ascii=False, indent=1))


# ── subcomandos ─────────────────────────────────────────────────────────────
def cmd_estado(a):
    if not any(a.get(k) for k in ('email', 'persona', 'slack')):
        fail('Pasá --email, --slack o --persona.')
    imprimir(llamar('estado', email=a.get('email') or '', slack=a.get('slack') or '',
                    persona=a.get('persona') or ''))


def cmd_aplicar(a, archivo):
    try:
        with open(archivo, encoding='utf-8') as f:
            pedido = json.load(f)
    except (OSError, ValueError) as e:
        fail(f'No pude leer {archivo}: {e}')
    imprimir(llamar('aplicar', pedido=pedido, solo_revisar=bool(a.get('solo-revisar'))))


def cmd_pendientes(a):
    if not a.get('persona'):
        fail('Pasá --persona.')
    res = llamar('pendientes', persona=a['persona'], hoy=a.get('hoy') or '')
    fmt = a.get('formato', 'json')
    if fmt == 'json':
        imprimir(res)
        return
    print(formatear(res, slack=(fmt == 'slack')))


def formatear(res, slack):
    items = res.get('items', [])
    hoy = datetime.fromisoformat(res['hoy'])
    b = (lambda s: f'**{s}**') if slack else (lambda s: s.upper())  # el conector de Slack usa markdown estándar

    def ddmm(s):
        return datetime.fromisoformat(s).strftime('%d/%m')

    def renglon(x):
        partes = [x['id'], x.get('texto') or '']
        if x.get('con'):
            partes.append(f'({x["con"]}{" · " + x["sociedad"] if x.get("sociedad") else ""})')
        extra = []
        if x['vencido']:
            extra.append(f'venció {ddmm(x["fecha_maxima"])}')
        elif x.get('fecha_maxima'):
            extra.append(f'vence {ddmm(x["fecha_maxima"])}')
        if x.get('dias'):
            extra.append(f'{x["dias"]} día{"s" if x["dias"] > 1 else ""} abierto')
        if x.get('reclamos'):
            extra.append(f'{x["reclamos"]} reclamo{"s" if x["reclamos"] > 1 else ""}')
        if x.get('estado') not in ('Pendiente', None):
            extra.append(x['estado'].lower())
        if slack and x['vencido']:
            marca = ':red_circle: '
        elif slack and x.get('vence_pronto'):
            marca = ':large_yellow_circle: '
        else:
            marca = '• '
        return f'{marca}{" ".join(partes)}' + (f' — {", ".join(extra)}' if extra else '')

    lineas = [b(f'Tus pendientes al {hoy:%d/%m/%Y}') + f' ({len(items)})']
    if not items:
        lineas.append('No tenés nada abierto en la planilla de seguimiento. :tada:' if slack else 'No hay pendientes.')
    for k, titulo in (('EC', 'Le pediste al cliente (hacer seguimiento)'),
                      ('CE', 'Te pidió el cliente (entregar)'), ('TI', 'Tareas internas')):
        grupo = [x for x in items if x['hoja'] == k]
        if grupo:
            lineas += ['', b(titulo)] + [renglon(x) for x in grupo]
    if res.get('link'):
        lineas += ['', f'Planilla: {res["link"]}']
    return '\n'.join(lineas)


def cmd_probar(_a):
    imprimir(llamar('probar'))


def cmd_importar(a):
    """Lee la planilla Excel anterior (misma estructura) y la carga en la Google Sheet vacía."""
    ruta = a.get('xlsx')
    if not ruta or not os.path.isfile(ruta):
        fail('Pasá --xlsx=<ruta de la planilla anterior>.')
    try:
        from openpyxl import load_workbook
    except ImportError:
        fail('Para importar hace falta openpyxl: pip install openpyxl')
    wb = load_workbook(ruta, data_only=True)

    def val(v):
        if isinstance(v, datetime):
            return v.date().isoformat()
        if isinstance(v, date):
            return v.isoformat()
        return v

    datos = {'equipo': [], 'cliente': [], 'config': {}, 'filas': {}, 'corridas': []}
    if 'Equipo' in wb.sheetnames:
        for f in wb['Equipo'].iter_rows(min_row=2, values_only=True):
            if f and f[0]:
                datos['equipo'].append({'nombre': f[0], 'email': f[1] or '', 'slack': f[2] or '',
                                        'activo': str(f[3] or 'Sí').strip().lower() not in ('no', 'n', '0')})
    if 'Listas' in wb.sheetnames:
        ws = wb['Listas']
        enc = [c.value for c in ws[1]]
        if 'Cliente' in enc:
            j = enc.index('Cliente') + 1
            datos['cliente'] = [ws.cell(r, j).value for r in range(2, ws.max_row + 1) if ws.cell(r, j).value]
    if 'Config' in wb.sheetnames:
        for f in wb['Config'].iter_rows(min_row=2, values_only=True):
            if f and f[0] and f[1] not in (None, '') and f[0] != 'Link de la planilla':
                datos['config'][f[0]] = f[1]
    for k, titulo in HOJAS.items():
        if titulo not in wb.sheetnames:
            continue
        ws = wb[titulo]
        enc = [c.value for c in ws[2]]
        filas = []
        for f in ws.iter_rows(min_row=3, values_only=True):
            if not f or not f[0]:
                continue
            filas.append({'_enc': {enc[j]: val(v) for j, v in enumerate(f) if enc[j] and v not in (None, '')}})
        datos['filas'][k] = filas
    if 'Corridas' in wb.sheetnames:
        for f in wb['Corridas'].iter_rows(min_row=2, values_only=True):
            if f and f[1]:
                datos['corridas'].append([str(f[0] or '')] + [val(v) for v in f[1:8]])
    # La app web recibe las filas por clave: se traducen los encabezados con el mapa de estado.
    mapa = llamar('probar').get('columnas')
    if mapa:
        for k, filas in datos['filas'].items():
            inv = {enc: clave for clave, enc in mapa[k].items()}
            datos['filas'][k] = [{inv[e]: v for e, v in f['_enc'].items() if e in inv and inv[e] != 'dias_abiertos'}
                                 for f in filas]
    imprimir(llamar('importar', datos=datos))


def parse(argv):
    a, pos = {}, []
    for x in argv:
        if x.startswith('--'):
            k, _, v = x[2:].partition('=')
            a[k] = v if v else True
        else:
            pos.append(x)
    return a, pos


def main():
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    cargar_env()
    a, pos = parse(sys.argv[1:])
    if not pos:
        print(__doc__)
        sys.exit(1)
    cmd = pos[0]
    if cmd == 'estado':
        cmd_estado(a)
    elif cmd == 'aplicar':
        if len(pos) < 2:
            fail('Uso: pedidos.py aplicar <ops.json> [--solo-revisar]')
        cmd_aplicar(a, pos[1])
    elif cmd == 'pendientes':
        cmd_pendientes(a)
    elif cmd == 'probar':
        cmd_probar(a)
    elif cmd == 'importar':
        cmd_importar(a)
    else:
        fail(f'Subcomando desconocido "{cmd}" (probar, estado, aplicar, pendientes, importar)')


if __name__ == '__main__':
    main()
