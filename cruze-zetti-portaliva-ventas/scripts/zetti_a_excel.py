"""Subdiario de IVA Ventas de Zetti (reporte 5.6.5, CSV) -> Excel limpio.

    python3 zetti_a_excel.py <entrada.csv> <salida.xlsx> [--periodo=AAAAMM]

El CSV de T&S Web viene en UTF-8 (a veces latin-1), separado por ";" y con columnas vacias
intercaladas:
    filas 1-5  encabezado (T&S Web, fecha de emision, "Subdiario de IVA Ventas",
               Fecha des/has, Nodo, Razon social, C.U.I.T.)
    fila 6     titulos (Fecha;TC;M;Nro. Comp.;;Cliente;;;;CUIT;;;RESP;;;Exen;...)
    datos      una fila por comprobante (los tiques van agrupados en una fila Z
               por cierre de caja)
    ultima     ";;;...Total:;;;<Exen>;;<Grav>;<IVA>;;;<P.IB>;;<P.IVA>;<Total>"

Salida: hoja "Ventas Zetti" con el formato de la solapa del papel de trabajo
mensual, para copiarla tal cual:
    fila 1  "<razon social>  -  CUIT NN-NNNNNNNN-N  -  Subdiario de IVA Ventas MM/AAAA"
    fila 2  Fecha | TC | M | Nro. Comp. (D = punto de venta, E = numero) | Cliente
            | CUIT | RESP | Exen | Grav | IVA | P.IB | Total
    datos   desde la fila 3, ordenados por TC, letra y fecha; sin fila de totales.
P.IVA y el codigo 901/902 del reporte no van (el papel no los tiene); si P.IVA
viniera con importe se avisa en el JSON.

Imprime por stdout un JSON con razon social, CUIT, fechas, cantidad por TC,
totales y el control contra la fila "Total:" del reporte. Sale con 1 (y el
JSON con "error") si el Total no coincide al centavo (o alguna otra columna
difiere en mas de $0,05), si hay filas que no
reconoce o si hay fechas fuera del periodo pedido: en ese caso no hay que
archivar.
"""
import csv
import json
import re
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# Titulo del reporte -> clave interna. El codigo final (901/902) no tiene titulo.
COLUMNAS = [
    ('Fecha', 'fecha'), ('TC', 'tc'), ('M', 'm'), ('Nro. Comp.', 'nro'),
    ('Cliente', 'cliente'), ('CUIT', 'cuit'), ('RESP', 'resp'),
    ('Exen', 'exen'), ('Grav', 'grav'), ('IVA', 'iva'),
    ('P.IB', 'pib'), ('P.IVA', 'piva'), ('Total', 'total'),
]
IMPORTES = ['exen', 'grav', 'iva', 'pib', 'piva', 'total']
TOLERANCIA = 0.05
TEXTO = ['tc', 'm', 'nro', 'cliente', 'cuit', 'resp', 'cod']


def fallar(msg, **extra):
    print(json.dumps({'error': msg, **extra}, ensure_ascii=False))
    sys.exit(1)


def importe(s):
    s = (s or '').strip()
    if not s:
        return 0.0
    try:
        return float(s.replace('.', '').replace(',', '.'))
    except ValueError:
        fallar(f'Importe que no entiendo: "{s}"')


def fecha(s):
    s = (s or '').strip()
    for fmt in ('%d/%m/%y', '%d/%m/%Y'):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None


def tras(fila, etiqueta):
    """Primer valor no vacio a la derecha de la celda que dice `etiqueta`."""
    for i, c in enumerate(fila):
        if c.strip().lower().startswith(etiqueta.lower()):
            for v in fila[i + 1:]:
                if v.strip():
                    return v.strip()
    return ''


def main():
    pos = [a for a in sys.argv[1:] if not a.startswith('--')]
    opts = dict(a[2:].split('=', 1) for a in sys.argv[1:] if a.startswith('--') and '=' in a)
    if len(pos) != 2:
        sys.exit(__doc__)
    entrada, salida = pos
    periodo = opts.get('periodo')

    # T&S Web lo baja en UTF-8 (la skill manual decia latin-1: con latin-1 la Ñ
    # sale "Ã\x91"). Si no es UTF-8 valido, latin-1.
    crudo = open(entrada, 'rb').read()
    try:
        texto = crudo.decode('utf-8-sig')
    except UnicodeDecodeError:
        texto = crudo.decode('latin-1')
    filas = list(csv.reader(texto.splitlines(), delimiter=';'))

    i_tit = next((i for i, f in enumerate(filas) if f and f[0].strip() == 'Fecha' and 'TC' in [c.strip() for c in f]), None)
    if i_tit is None:
        fallar('No encontre la fila de titulos (Fecha;TC;...) en el CSV.')
    titulos = [c.strip() for c in filas[i_tit]]
    idx = {}
    for nombre, clave in COLUMNAS:
        if nombre not in titulos:
            fallar(f'Falta la columna "{nombre}" en el CSV.', titulos=[t for t in titulos if t])
        idx[clave] = titulos.index(nombre)
    idx['cod'] = idx['total'] + 1

    encabezado = filas[:i_tit]
    razon = cuit_soc = desde = hasta = nodo = ''
    for f in encabezado:
        razon = razon or tras(f, 'Raz')
        cuit_soc = cuit_soc or tras(f, 'C.U.I.T')
        desde = desde or tras(f, 'Fecha des')
        hasta = hasta or tras(f, 'Fecha has')
        nodo = nodo or tras(f, 'Nodo')
    cuit_soc = re.sub(r'\D', '', cuit_soc)

    datos, total_reporte, raras = [], None, []
    for n, f in enumerate(filas[i_tit + 1:], start=i_tit + 2):
        if not any(c.strip() for c in f):
            continue
        f = f + [''] * (idx['cod'] + 1 - len(f))
        d = fecha(f[idx['fecha']])
        if d is None:
            if any(c.strip().lower().startswith('total') for c in f):
                total_reporte = {k: importe(f[idx[k]]) for k in IMPORTES}
            else:
                raras.append({'fila': n, 'texto': ';'.join(c for c in f if c.strip())[:120]})
            continue
        r = {'fecha': d}
        for k in TEXTO:
            r[k] = f[idx[k]].strip()
        for k in IMPORTES:
            r[k] = importe(f[idx[k]])
        datos.append(r)

    if raras:
        fallar('Hay filas que no son comprobantes ni el total: revisar el CSV antes de archivar.', filas=raras)
    if not datos:
        fallar('El reporte no trae comprobantes.', cuit=cuit_soc, desde=desde, hasta=hasta)
    if total_reporte is None:
        fallar('El CSV no trae la fila "Total:": puede haber bajado cortado.')

    fuera = []
    if periodo:
        fuera = sorted({r['fecha'].strftime('%d/%m/%Y') for r in datos if r['fecha'].strftime('%Y%m') != periodo})

    suma = {k: round(sum(r[k] for r in datos), 2) for k in IMPORTES}
    # El reporte suma con decimales internos: aparecen centavos cruzados entre
    # Exen y Grav que se compensan. El Total tiene que dar exacto; en las demas
    # columnas se tolera TOLERANCIA y se informa como aviso.
    difs = {k: round(suma[k] - total_reporte[k], 2) for k in IMPORTES if abs(suma[k] - total_reporte[k]) >= 0.005}
    graves = {k: v for k, v in difs.items() if k == 'total' or abs(v) > TOLERANCIA}

    # Orden de la solapa del papel de trabajo: TC, letra y fecha. sorted() es
    # estable: dentro de una misma fecha queda el orden del reporte.
    datos.sort(key=lambda r: (r['tc'], r['m'], r['fecha']))

    # "0006-00000070" -> punto de venta 6 y numero 70, como en el papel.
    sin_separar = 0
    for r in datos:
        m = re.fullmatch(r'(\d+)-(\d+)', r['nro'])
        if m:
            r['pv'], r['numero'] = int(m.group(1)), int(m.group(2))
        else:
            r['pv'], r['numero'] = r['nro'], None
            sin_separar += 1

    wb = Workbook()
    ws = wb.active
    ws.title = 'Ventas Zetti'
    cuit_fmt = f'{cuit_soc[:2]}-{cuit_soc[2:10]}-{cuit_soc[10:]}' if len(cuit_soc) == 11 else cuit_soc
    d = fecha(desde)
    mes = d.strftime('%m/%Y') if d else ''
    ws['A1'] = f'{razon}  -  CUIT {cuit_fmt}  -  Subdiario de IVA Ventas {mes}'

    # Columnas de la solapa "Ventas Zetti" del papel de trabajo. "Nro. Comp."
    # abarca D (punto de venta) y E (numero), E sin titulo propio. El papel no
    # lleva P.IVA ni el codigo 901/902, ni fila de totales.
    columnas = [('Fecha', 'fecha'), ('TC', 'tc'), ('M', 'm'), ('Nro. Comp.', 'pv'), (None, 'numero'),
                ('Cliente', 'cliente'), ('CUIT', 'cuit'), ('RESP', 'resp'),
                ('Exen', 'exen'), ('Grav', 'grav'), ('IVA', 'iva'), ('P.IB', 'pib'), ('Total', 'total')]
    contable = '_-"$"\\ * #,##0.00_-;\\-"$"\\ * #,##0.00_-;_-"$"\\ * "-"??_-;_-@_-'
    chica = Font(size=10)
    for j, (t, k) in enumerate(columnas, start=1):
        c = ws.cell(2, j, t)
        if k in IMPORTES:
            c.number_format = contable
            c.font = chica

    for i, r in enumerate(datos, start=3):
        for j, (_, k) in enumerate(columnas, start=1):
            v = r[k]
            c = ws.cell(i, j, v if v != '' else None)
            if k == 'fecha':
                c.number_format = 'dd/mm/yyyy'
            elif k in IMPORTES:
                c.number_format = contable
                c.font = chica

    anchos = {'fecha': 11, 'tc': 5, 'm': 4, 'pv': 6, 'numero': 9, 'cliente': 40, 'cuit': 15, 'resp': 12,
              'exen': 17, 'grav': 16, 'iva': 15, 'pib': 13, 'total': 17}
    for j, (_, k) in enumerate(columnas, start=1):
        ws.column_dimensions[get_column_letter(j)].width = anchos[k]
    wb.save(salida)

    por_tc = {}
    for r in datos:
        por_tc[r['tc']] = por_tc.get(r['tc'], 0) + 1
    res = {
        'razon_social': razon, 'cuit': cuit_soc, 'nodo': nodo, 'desde': desde, 'hasta': hasta,
        'comprobantes': len(datos), 'por_tc': por_tc, 'totales': suma,
        'control_total_reporte': 'ok' if not difs else ('redondeo' if not graves else 'no coincide'),
        **({'diferencias_total_reporte': difs} if difs else {}),
        'fechas_fuera_de_periodo': fuera, 'xlsx': salida,
    }
    # P.IVA no tiene columna en el papel: si viniera con importe, el Total no
    # cerraria con Exen+Grav+IVA+P.IB. Se avisa (hasta ahora siempre vino en 0).
    if suma['piva']:
        res['aviso_piva'] = f'Hay P.IVA por {suma["piva"]:.2f} que no tiene columna en la hoja (esta incluida en Total).'
    if sin_separar:
        res['aviso_nro'] = f'{sin_separar} comprobantes con numero que no es PV-NUMERO: quedaron enteros en la columna D.'
    if graves:
        res['error'] = 'La suma de los comprobantes no coincide con el Total del reporte.'
    elif fuera:
        res['error'] = f'Hay comprobantes fuera del periodo {periodo}.'
    print(json.dumps(res, ensure_ascii=False))
    sys.exit(1 if 'error' in res else 0)


if __name__ == '__main__':
    main()
