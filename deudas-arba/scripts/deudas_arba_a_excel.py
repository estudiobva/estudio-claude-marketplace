#!/usr/bin/env python3
"""
deudas_arba_a_excel.py — arma el Excel consolidado de deudas y DDJJ sin
presentar en ARBA (lo que releva deudas-arba.js).

    python deudas_arba_a_excel.py <lote.json> <salida.xlsx>

Hojas:
  - Control: una fila por sociedad del archivo de claves, con los totales.
  - Situación por objeto: IIBB, agente de retención/percepción, sellos… tal
    como los devuelve ARBA (deuda prejudicial/judicial, saldo a favor, DDJJ
    sin presentar).
  - Agentes por período: DDJJ de agente con saldo (declarado − pagado).
  - DDJJ sin presentar: períodos vencidos sin DDJJ cerrada y DDJJ iniciadas
    sin enviar.

Si una hoja no tiene nada, se escribe igual una fila "Sin ...": saltearla en
silencio hace que después nadie sepa si se revisó o no.
"""

import json
import re
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FUENTE = 'Arial'
AZUL = '1F4E79'
ROJO = '9C0006'
GRIS = '808080'
AMARILLO_BG = 'FFF2CC'
FMT_PESOS = '#,##0.00;-#,##0.00;-'


def _encabezado(ws, columnas, anchos):
    for j, h in enumerate(columnas, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(name=FUENTE, size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=AZUL)
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        ws.column_dimensions[get_column_letter(j)].width = anchos[j - 1]
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'


def _filas(ws, columnas, anchos, filas, pesos=(), fechas=(), resaltar=None, vacio='Sin datos'):
    """filas: lista de listas. pesos/fechas: índices (1-based) con formato."""
    _encabezado(ws, columnas, anchos)
    if not filas:
        c = ws.cell(row=2, column=1, value=vacio)
        c.font = Font(name=FUENTE, size=10, italic=True, color=GRIS)
        return
    for i, vals in enumerate(filas, start=2):
        rojo = bool(resaltar and resaltar(vals))
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=FUENTE, size=10, bold=rojo, color=ROJO if rojo else None)
            if j in pesos:
                c.number_format = FMT_PESOS
            if j in fechas:
                c.number_format = 'DD/MM/YYYY'
    ultima = len(filas) + 1
    if pesos:
        ws.cell(row=ultima + 2, column=min(pesos) - 1, value='TOTAL').font = Font(name=FUENTE, size=10, bold=True)
        for j in pesos:
            L = get_column_letter(j)
            t = ws.cell(row=ultima + 2, column=j, value=f'=SUM({L}2:{L}{ultima})')
            t.font = Font(name=FUENTE, size=10, bold=True)
            t.number_format = FMT_PESOS
    ws.auto_filter.ref = f'A1:{get_column_letter(len(columnas))}{ultima}'


def fecha(s):
    s = str(s or '')
    for fmt in ('%Y-%m-%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            pass
    return None


def suma(objs, cod=None, campo='deudaTotal'):
    """None si ningún objeto de ese impuesto trae el dato (ARBA no lo pudo
    calcular): un 0 ahí se leería como "sin deuda"."""
    vals = [o.get(campo) for o in objs if (cod is None or o.get('codImpuesto') == cod) and o.get(campo) is not None]
    return round(sum(vals), 2) if vals else None


def estado_agente(s):
    a = s.get('agenteEstado') or {}
    if a.get('estado'):
        return a['estado'] + (' (SUSPENDIDO)' if a.get('suspendido') else '')
    return None


def main(entrada, salida):
    with open(entrada, encoding='utf-8') as f:
        lote = json.load(f)
    socs = lote.get('sociedades') or []
    wb = Workbook()

    # ── Control ─────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = 'Control'
    filas = []
    for s in socs:
        objs = s.get('objetos') or []
        dfe = sum(a.get('cantidad') or 0 for a in (s.get('alertas') or []) if 'DFE' in (a.get('descripcion') or ''))
        if not s.get('ok'):
            rel = 'Error'
        elif s.get('sinCit'):
            rel = 'Sin CIT'
        else:
            rel = 'OK'
        obs = list(s.get('observaciones') or [])
        if s.get('error'):
            obs.insert(0, s['error'])
        if s.get('sinCit'):
            obs.insert(0, 'Sin CIT de ARBA en Claves_Organismos: solo se consultó el estado del agente')
        errores_obj = [f"{o.get('impuesto')}{' ' + o['agente'] if o.get('agente') else ''}: {o['error']}"
                       for o in objs if o.get('error') and not re.search(r'IMPUESTO INVALIDO|proceso de actualizaci', o['error'])]
        obs += errores_obj
        inmob = round((suma(objs, 0) or 0) + (suma(objs, 10) or 0), 2)
        if inmob:
            obs.append(f'La deuda total incluye $ {inmob:,.2f} de partidas inmobiliarias'
                       .replace(',', 'X').replace('.', ',').replace('X', '.'))
        a_vencer = round(sum(p.get('saldo') or 0 for p in (s.get('periodos') or []) if p.get('estado') == 'A vencer'), 2)
        if a_vencer:
            obs.append(f'La deuda de agente incluye $ {a_vencer:,.2f} de períodos que todavía no vencieron'
                       .replace(',', 'X').replace('.', ',').replace('X', '.'))
        filas.append([
            s.get('pos'), s.get('sociedad'), s.get('cuit'), rel,
            suma(objs, 2), suma(objs, 7), suma(objs), suma(objs, 2, 'saldoAFavor'),
            suma(objs, None, 'ddjjSinPresentar'),
            len(s.get('faltantes') or []) or None,
            'Sí' if any(o.get('planDePago') for o in objs) else ('No' if suma(objs) is not None else None),
            dfe if objs else None,
            estado_agente(s),
            ' · '.join(dict.fromkeys(obs)) or None,
        ])
    _filas(ws, ['Pos', 'Sociedad', 'CUIT', 'Relevamiento', 'Deuda IIBB', 'Deuda agente', 'Deuda total',
                'Saldo a favor IIBB', 'DDJJ sin presentar (ARBA)', 'Períodos sin DDJJ (detalle)', 'Plan de pago',
                'Notif. DFE (aviso de ARBA)', 'Estado agente (ARSuspension)', 'Observaciones'],
           [6, 30, 14, 13, 15, 15, 15, 16, 11, 11, 9, 10, 18, 70], filas, pesos=(5, 6, 7, 8),
           resaltar=lambda v: (v[6] or 0) > 0 or (v[8] or 0) > 0 or (v[9] or 0) > 0 or v[3] == 'Error'
           or 'SUSPENDIDO' in str(v[12] or ''),
           vacio='Sin sociedades relevadas')
    for i, v in enumerate(filas, start=2):
        if v[3] == 'Sin CIT':
            for j in range(1, 15):
                ws.cell(row=i, column=j).fill = PatternFill('solid', fgColor=AMARILLO_BG)

    # ── Situación por objeto ────────────────────────────────────────────────
    ws = wb.create_sheet('Situación por objeto')
    filas = []
    for s in socs:
        for o in s.get('objetos') or []:
            obs = o.get('error')
            if obs and 'IMPUESTO INVALIDO' in obs:
                obs = ('Convenio Multilateral: la deuda se ve en SIFERE' if 'CONVENIO' in (o.get('vinculo') or '')
                       else 'ARBA no pudo calcular la situación del objeto (en la web: "No se pudo calcular la situación")')
            elif o.get('codImpuesto') in (0, 10) and (o.get('deudaTotal') or 0) > 0:
                obs = 'Partida inmobiliaria: puede incluir cuotas del año en curso; si hay cotitulares, la misma partida aparece en cada uno'
            elif not o.get('consultada') and o.get('codImpuesto') in (2, 7):
                obs = 'ARBA no devolvió la deuda de este objeto'
            elif not o.get('consultada'):
                obs = 'No consultado (Gestionar solo calcula IIBB y agentes al entrar)'
            impuesto = 'Inmobiliario' if o.get('codImpuesto') == 0 else o.get('impuesto')
            filas.append([
                s.get('sociedad'), s.get('cuit'), impuesto, o.get('agente') or None, o.get('objeto'),
                o.get('vinculo') or None, o.get('deudaPrejudicial'), o.get('deudaJudicial'), o.get('deudaTotal'),
                o.get('saldoAFavor'), o.get('ddjjSinPresentar'),
                ('Sí' if o.get('planDePago') else 'No') if o.get('consultada') and not o.get('error') else None,
                obs,
            ])
    _filas(ws, ['Sociedad', 'CUIT', 'Impuesto', 'Agente', 'Objeto', 'Vínculo', 'Deuda prejudicial',
                'Deuda judicial', 'Deuda total', 'Saldo a favor', 'DDJJ sin presentar', 'Plan de pago', 'Observación'],
           [30, 14, 22, 11, 20, 18, 15, 14, 15, 15, 11, 9, 60], filas, pesos=(7, 8, 9, 10),
           resaltar=lambda v: (v[8] or 0) > 0 or (v[10] or 0) > 0, vacio='Sin objetos relevados')

    # ── Agentes por período ─────────────────────────────────────────────────
    ws = wb.create_sheet('Agentes por período')
    filas = []
    for s in socs:
        for p in s.get('periodos') or []:
            filas.append([
                s.get('sociedad'), s.get('cuit'), p.get('actividad'), p.get('regimen'), p.get('periodo'),
                p.get('comprobante'), p.get('rectificativa'), p.get('declarado'), p.get('pagado'), p.get('saldo'),
                fecha(p.get('ultimoPago')), fecha(p.get('vencimientoAprox')), p.get('estado'),
            ])
    _filas(ws, ['Sociedad', 'CUIT', 'Act.', 'Régimen', 'Período', 'Nº comprobante', 'Rectif.', 'Declarado',
                'Pagado', 'Saldo', 'Último pago', 'Venc. aprox.', 'Estado'],
           [30, 14, 6, 44, 12, 14, 7, 15, 15, 15, 12, 12, 10], filas, pesos=(8, 9, 10), fechas=(11, 12),
           resaltar=lambda v: v[12] == 'Vencido',
           vacio=f"Sin DDJJ de agente con saldo impago desde {lote.get('desdeAnio', '')}")

    # ── DDJJ sin presentar ──────────────────────────────────────────────────
    ws = wb.create_sheet('DDJJ sin presentar')
    filas = []
    for s in socs:
        for f in s.get('faltantes') or []:
            reg = f"{f.get('actividad')} - {f.get('regimen')}" if f.get('actividad') else None
            obs = f.get('nota') or 'Vencida sin DDJJ cerrada en ARBA'
            if f.get('pago'):
                obs += ' (hay pago registrado del período)'
            filas.append([s.get('sociedad'), s.get('cuit'), f.get('impuesto'), reg, f.get('periodo'), obs])
        for txt in s.get('iniciadas') or []:
            filas.append([s.get('sociedad'), s.get('cuit'), 'Agente de recaudación', None, None,
                          f'DDJJ iniciada sin enviar: {txt}'])
    _filas(ws, ['Sociedad', 'CUIT', 'Impuesto', 'Actividad / régimen', 'Período', 'Observación'],
           [30, 14, 22, 46, 12, 60], filas, vacio='Sin DDJJ sin presentar')

    wb.save(salida)
    print(json.dumps({'ok': True, 'archivo': salida, 'sociedades': len(socs)}))


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit('uso: deudas_arba_a_excel.py <lote.json> <salida.xlsx>')
    main(sys.argv[1], sys.argv[2])
