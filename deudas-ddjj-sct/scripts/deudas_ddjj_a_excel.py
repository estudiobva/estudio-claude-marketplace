#!/usr/bin/env python3
"""
deudas_ddjj_a_excel.py — arma el Excel de deudas (que no son anticipos) y DDJJ
pendientes de presentacion del SCT.

    python deudas_ddjj_a_excel.py <datos.json> <salida.xlsx>
    python deudas_ddjj_a_excel.py --consolidado <lote.json> <salida.xlsx>

Un archivo por sociedad con dos hojas:
  - Deudas: lo de la solapa Deudas que no es anticipo (saldos de DDJJ, multas,
    intereses, retenciones). Capital | Int. Resarcitorio | Int. Punitorio | Total.
  - DDJJ pendientes: la solapa "DDJJ pendientes de presentación". No tiene
    importes: son presentaciones que el SCT da por omitidas.

Los anticipos quedan afuera a proposito: los releva anticipos-sct.

Si una hoja no tiene nada, se escribe igual una fila "Sin ...": saltearla en
silencio hace que despues nadie sepa si se reviso o no.
"""

import json
import re
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FUENTE = 'Arial'
AZUL = '1F4E79'
ROJO = '9C0006'
AMARILLO_BG = 'FFF2CC'
FMT_PESOS = '#,##0.00;-#,##0.00;-'

COLS_DEUDAS = ['Impuesto', 'Concepto', 'Subconcepto', 'Período', 'Capital',
               'Int. Resarcitorio', 'Int. Punitorio', 'Total', 'Vencimiento',
               'Estado', 'Observación']
ANCHOS_DEUDAS = [34, 30, 30, 10, 15, 15, 14, 15, 13, 13, 36]

COLS_PEND = ['Impuesto', 'Concepto', 'Subconcepto', 'Período', 'Vencimiento',
             'Días de atraso']
ANCHOS_PEND = [36, 40, 40, 10, 13, 13]


def a_numero(s):
    """'$ 2.933,14' -> 2933.14 · '' -> None."""
    s = re.sub(r'[^\d,.-]', '', str(s or ''))
    if not s:
        return None
    try:
        return float(s.replace('.', '').replace(',', '.'))
    except ValueError:
        return None


def a_fecha(s):
    m = re.search(r'(\d{2})/(\d{2})/(\d{4})', str(s or ''))
    if not m:
        return None
    return datetime(int(m.group(3)), int(m.group(2)), int(m.group(1))).date()


def periodo(s):
    """'202508' -> '08/2025' · '2024' -> '2024'. Lo demas, tal cual."""
    s = str(s or '').strip()
    if re.fullmatch(r'\d{6}', s) and 1 <= int(s[4:]) <= 12:
        return f'{s[4:]}/{s[:4]}'
    if re.fullmatch(r'\d{8}', s) and s.endswith('00') and 1 <= int(s[4:6]) <= 12:
        return f'{s[4:6]}/{s[:4]}'
    if re.fullmatch(r'\d{4}0000', s):
        return s[:4]
    return s


def montos(d):
    return (a_numero(d.get('capital')), a_numero(d.get('intResarcitorio')),
            a_numero(d.get('intPunitorio')))


def observacion(d):
    """Aclaraciones que cambian como se lee la fila."""
    cap, res, pun = montos(d)
    total = (cap or 0) + (res or 0) + (pun or 0)
    obs = []
    if re.search(r'multa', f"{d.get('concepto')} {d.get('subconcepto')}", re.I):
        obs.append('Multa')
    if not cap and (res or 0) + (pun or 0) > 0:
        obs.append('Capital cancelado: quedan intereses')
    if 0 < total < 1:
        obs.append('Diferencia menor a $ 1')
    if not a_fecha(d.get('vencimiento')):
        obs.append('Sin fecha de vencimiento')
    return ' · '.join(obs) or None


def _encabezado(ws, columnas, anchos):
    for j, h in enumerate(columnas, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(name=FUENTE, size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=AZUL)
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        ws.column_dimensions[get_column_letter(j)].width = anchos[j - 1]
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = 'A2'


def _hoja_deudas(ws, filas, prefijo=(), anchos_prefijo=()):
    """Escribe las deudas; `prefijo` son columnas fijas por fila (Sociedad, CUIT
    en el consolidado). Devuelve la ultima fila con datos."""
    cols = [p[0] for p in prefijo] + COLS_DEUDAS
    _encabezado(ws, cols, list(anchos_prefijo) + ANCHOS_DEUDAS)
    k = len(prefijo)
    c_cap, c_tot, c_ven = k + 5, k + 8, k + 9
    i = 2
    for pre, d in filas:
        cap, res, pun = montos(d)
        L0, L2 = get_column_letter(c_cap), get_column_letter(c_cap + 2)
        vals = [*pre, d.get('impuesto'), d.get('concepto'), d.get('subconcepto'),
                periodo(d.get('periodo')), cap, res, pun, f'=SUM({L0}{i}:{L2}{i})',
                a_fecha(d.get('vencimiento')), d.get('estado'), observacion(d)]
        vencido = d.get('estado') == 'Vencido'
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=FUENTE, size=10, bold=vencido, color=ROJO if vencido else None)
            if c_cap <= j <= c_tot:
                c.number_format = FMT_PESOS
            if j == c_ven:
                c.number_format = 'DD/MM/YYYY'
        i += 1
    ultima = i - 1
    if filas:
        ws.cell(row=ultima + 2, column=c_cap - 1, value='TOTAL').font = Font(name=FUENTE, size=10, bold=True)
        for j in range(c_cap, c_tot + 1):
            L = get_column_letter(j)
            t = ws.cell(row=ultima + 2, column=j, value=f'=SUM({L}2:{L}{ultima})')
            t.font = Font(name=FUENTE, size=10, bold=True)
            t.number_format = FMT_PESOS
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{max(ultima, 1)}'
    return ultima


def _hoja_pendientes(ws, filas, prefijo=(), anchos_prefijo=()):
    cols = [p[0] for p in prefijo] + COLS_PEND
    _encabezado(ws, cols, list(anchos_prefijo) + ANCHOS_PEND)
    k = len(prefijo)
    i = 2
    for pre, p in filas:
        vals = [*pre, p.get('impuesto'), p.get('concepto'), p.get('subconcepto'),
                periodo(p.get('periodo')), a_fecha(p.get('vencimiento')), p.get('diasAtraso')]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=FUENTE, size=10)
            c.fill = PatternFill('solid', fgColor=AMARILLO_BG)
            if j == k + 5:
                c.number_format = 'DD/MM/YYYY'
            if j == k + 6:
                c.number_format = '#,##0'
        i += 1
    ultima = i - 1
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{max(ultima, 1)}'
    return ultima


def _vacia(ws, fila, texto):
    c = ws.cell(row=fila, column=1, value=texto)
    c.font = Font(name=FUENTE, size=10, italic=True)
    return fila


def _nota(ws, fila, texto):
    ws.cell(row=fila, column=1, value=texto).font = Font(name=FUENTE, size=9, italic=True)


def escribir(datos, salida):
    wb = Workbook()
    fuente = (f'Relevado el {datos.get("fecha")} · {datos.get("sociedad")} · '
              f'CUIT {datos.get("cuit")} · Fuente: Sistema de Cuentas Tributarias de ARCA')

    ws = wb.active
    ws.title = 'Deudas'
    deudas = datos.get('deudas') or []
    ultima = _hoja_deudas(ws, [((), d) for d in deudas])
    if not deudas:
        ultima = _vacia(ws, 2, 'Sin deudas que no sean anticipos')
    nota = f'{fuente}, solapa Deudas. No incluye anticipos (ver anticipos-sct).'
    if datos.get('anticiposExcluidos'):
        nota += f' Anticipos excluidos: {datos["anticiposExcluidos"]}.'
    _nota(ws, ultima + (3 if deudas else 2), nota)

    wp = wb.create_sheet('DDJJ pendientes')
    pend = datos.get('pendientes') or []
    ultima = _hoja_pendientes(wp, [((), p) for p in pend])
    if not pend:
        ultima = _vacia(wp, 2, 'Sin DDJJ pendientes de presentación')
    _nota(wp, ultima + 2, f'{fuente}, solapa DDJJ pendientes de presentación.')

    wb.save(salida)
    return {'archivo': salida, 'deudas': len(deudas), 'pendientes': len(pend)}


def total_deuda(deudas):
    return round(sum(sum(x or 0 for x in montos(d)) for d in deudas), 2)


def resultado(s):
    nd = len(s.get('deudas') or [])
    npd = len(s.get('pendientes') or [])
    if nd and npd:
        return 'Deudas y DDJJ pendientes'
    if nd:
        return 'Con deudas'
    if npd:
        return 'DDJJ pendientes'
    return 'Sin observaciones'


def consolidar(lote, salida):
    """Todas las sociedades del lote: hojas Deudas y DDJJ pendientes (una fila
    por registro) y hoja Control (una fila por sociedad, incluidas las que
    fallaron — asi se ve de un vistazo que falta relevar)."""
    wb = Workbook()
    ok = [s for s in lote.get('sociedades', []) if s.get('ok')]
    pre = (('Sociedad',), ('CUIT',))
    anchos = (30, 14)

    ws = wb.active
    ws.title = 'Deudas'
    filas = [((s.get('sociedad'), s.get('cuit')), d) for s in ok for d in s.get('deudas') or []]
    _hoja_deudas(ws, filas, pre, anchos)

    wp = wb.create_sheet('DDJJ pendientes')
    filas = [((s.get('sociedad'), s.get('cuit')), p) for s in ok for p in s.get('pendientes') or []]
    _hoja_pendientes(wp, filas, pre, anchos)

    wc = wb.create_sheet('Control')
    cc = ['Pos.', 'Sociedad', 'CUIT', 'Resultado', 'Deudas', 'Deudas vencidas',
          'Total deudas', 'DDJJ pendientes', 'Detalle']
    _encabezado(wc, cc, [6, 30, 14, 24, 9, 10, 16, 11, 60])
    r = 2
    for s in lote.get('sociedades', []):
        if s.get('ok'):
            deudas = s.get('deudas') or []
            vals = [s.get('pos'), s.get('sociedad'), s.get('cuit'), resultado(s), len(deudas),
                    sum(1 for d in deudas if d.get('estado') == 'Vencido'), total_deuda(deudas),
                    len(s.get('pendientes') or []),
                    f"{s.get('anticiposExcluidos', 0)} anticipo(s) excluidos"]
        else:
            vals = [s.get('pos'), s.get('sociedad'), s.get('cuit'), 'ERROR',
                    None, None, None, None, s.get('error')]
        err = vals[3] == 'ERROR'
        for j, v in enumerate(vals, start=1):
            c = wc.cell(row=r, column=j, value=v)
            c.font = Font(name=FUENTE, size=10, bold=err or vals[3] != 'Sin observaciones',
                          color=ROJO if err else None)
            if j == 7:
                c.number_format = FMT_PESOS
        r += 1
    for o in lote.get('omitidos', []):
        vals = [o.get('pos'), o.get('titular'), None, 'Omitida', None, None, None, None, o.get('motivo')]
        for j, v in enumerate(vals, start=1):
            wc.cell(row=r, column=j, value=v).font = Font(name=FUENTE, size=10, italic=True)
        r += 1
    wc.auto_filter.ref = f'A1:{get_column_letter(len(cc))}{max(r - 1, 1)}'
    # El Control primero: es lo que se mira antes que nada.
    wb.move_sheet('Control', offset=-2)
    wb.active = 0
    wb.save(salida)
    return {'archivo': salida, 'sociedades': r - 2,
            'deudas': sum(len(s.get('deudas') or []) for s in ok),
            'pendientes': sum(len(s.get('pendientes') or []) for s in ok)}


def main():
    if len(sys.argv) == 4 and sys.argv[1] == '--consolidado':
        with open(sys.argv[2], encoding='utf-8') as fh:
            lote = json.load(fh)
        print(json.dumps(consolidar(lote, sys.argv[3]), ensure_ascii=False))
        return
    if len(sys.argv) != 3:
        raise SystemExit('Uso: deudas_ddjj_a_excel.py <datos.json> <salida.xlsx>\n'
                         '     deudas_ddjj_a_excel.py --consolidado <lote.json> <salida.xlsx>')
    with open(sys.argv[1], encoding='utf-8') as fh:
        datos = json.load(fh)
    print(json.dumps(escribir(datos, sys.argv[2]), ensure_ascii=False))


if __name__ == '__main__':
    main()
