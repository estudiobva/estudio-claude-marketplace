#!/usr/bin/env python3
# clasificar-compras.py — clasifica las compras del Portal IVA de una sociedad con
# el Maestro de Proveedores del estudio: a cada comprobante le pone la cuenta
# contable de su proveedor (por CUIT) y el rubro del EERR, y lista aparte los
# proveedores que no estan clasificados en firme.
#
#   python scripts/clasificar-compras.py --nombre="<sociedad>" --periodo=08/2026
#   python scripts/clasificar-compras.py --cuit=<CUIT> --desde=01/2026 --hasta=08/2026
#   python scripts/clasificar-compras.py --nombres="<A>;<B>" --periodo=08/2026
#   python scripts/clasificar-compras.py --todas --periodo=08/2026 --solo-revisar
#
# Argumentos:
#   --cuit / --cuits     CUIT de la sociedad (o varios separados por ";").
#   --nombre / --nombres Nombre (o parte) de la carpeta de la sociedad (o varios con ";").
#   --todas              Todas las sociedades de la unidad compartida.
#   --periodo            MM/AAAA o AAAAMM. O un rango con --desde y --hasta.
#   --fuente             auto (default) | sin-ncr | portal. Ver "Fuente" abajo.
#   --compras            Archivo de compras a usar (solo con una sociedad y un mes).
#   --maestro            Maestro de Proveedores. Default: <unidad>/0 - ESTUDIO - 00000000000/
#                        Maestro de Proveedores.xlsx (o BVA_MAESTRO_PROVEEDORES en el .env).
#   --si-existe          error (default) | reemplazar (el anterior va a _anteriores/) | saltear
#   --solo-revisar       Clasifica y reporta, pero no escribe nada.
#
# Fuente: con "auto" usa la salida de conciliacion-ncr ([AAAAMM]-COMPRAS-SIN-NCR-*.xlsx,
# hoja "Compras S.NCR") si existe, porque es el libro ya depurado de las NCR; si no,
# el libro de compras del Portal IVA ([AAAAMM] - PORTAL IVA - COMPRAS.xlsx / .csv / .zip)
# y avisa que las NCR no estan depuradas. "sin-ncr" exige la conciliacion.
#
# Salida: [AAAAMM] - COMPRAS CLASIFICADAS.xlsx en la carpeta IVA del mes, y un JSON por
# stdout (logs por stderr). Si algo no cierra imprime {"error": ...} y termina con 1.
#
# Reglas (las mismas con que se armaron los EERR 2026):
#   - El Maestro manda. "Propuesta" usa la cuenta propuesta; "Revisar" va a GASTOS VARIOS;
#     "Bloqueado" (CUIT limitada RG 3832/16) y COMPRA DE EQUIPOS quedan EXCLUIDOS del
#     resultado; un CUIT que no esta en el Maestro va a GASTOS VARIOS como SIN MAESTRO.
#   - Importe imputado: con IVA discriminado, no gravado + exento + netos gravados; sin IVA
#     discriminado (B, C, tiques), el total menos percepciones e impuestos. Las notas de
#     credito vienen en negativo del Portal IVA, asi que restan solas.
#   - El Maestro no se toca: los pendientes se informan para darlos de alta o confirmarlos.
#
# Solo usa openpyxl (sin pandas) para que corra en cualquier maquina del estudio.

import csv, datetime, io, json, os, re, shutil, sys, time, unicodedata, zipfile
from collections import OrderedDict, defaultdict

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print(json.dumps({'error': 'Falta openpyxl: pip install openpyxl'}))
    sys.exit(1)

RAIZ_DEFECTO = r'I:\Unidades compartidas\BVA - Sociedades Farmaceuticas'
PLUGIN = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
MAESTRO_REL = os.path.join('0 - ESTUDIO - 00000000000', 'Maestro de Proveedores.xlsx')

# Columnas del Portal IVA que NO son numericas (el resto si).
TEXTO = {'nro. doc. vendedor', 'denominacion vendedor', 'moneda original'}
ENTEROS = {'tipo de comprobante', 'punto de venta', 'numero de comprobante', 'tipo doc. vendedor'}
FECHA = 'fecha de emision'
NETOS = ['importe no gravado', 'importe exento', 'neto gravado iva 0%', 'neto gravado iva 2,5%',
         'neto gravado iva 5%', 'neto gravado iva 10,5%', 'neto gravado iva 21%', 'neto gravado iva 27%']
PERCEPCIONES = ['importe de per. o pagos a cta. de otros imp. nac.', 'importe de percepciones de ingresos brutos',
                'importe de impuestos municipales', 'importe de percepciones o pagos a cuenta de iva',
                'importe de impuestos internos', 'importe otros tributos']

# Cuenta del Maestro -> fila del EERR (etiqueta de la columna A de la hoja EERR).
RUBRO = {
    'CMV': 'CMV', 'AGUA': 'Agua', 'GAS': 'Metrogas', 'LUZ': 'Edesur',
    'GASTOS BANCARIOS': 'Gastos Bancarios', 'INTERESES PAGADOS': 'Intereses de Prestamos',
    'TARJETA DE CRÉDITO': 'GASTOS DE TARJETAS', 'ALQUILERES': 'ALQUILERES',
    'GASTOS DE COMERCIALIZACION': 'GASTOS DE COMERCIALIZACION',
    'COMISION MERCADO LIBRE': 'COMISIONES DE MERCADO LIBRE', 'VENTA ON LINE': 'COMISIONES DE MERCADO LIBRE',
    'MANTEN. ESTRUCTURA': 'GASTOS DE MANTENIMIENTO', 'MANTEN. RODADOS': 'GASTOS DE MANTENIMIENTO',
    'GASTOS DE MANTENIMIENTO': 'GASTOS DE MANTENIMIENTO',
    'MANTEN. EQ. DE COMPUTACION': 'GASTOS DE MANTENIMIENTO BS . USO',
    'GASTOS DE SISTEMAS': 'GASTOS DE SISTEMAS', 'SERVICIOS DE INFORMATICA': 'GASTOS DE SISTEMAS',
    'GASTOS DE TRANSPORTE': 'GASTOS DE TRANSPORTE', 'GASTOS DE MENSAJERIA, MOVILIDAD, VIATICOS': 'GASTOS DE TRANSPORTE',
    'GASTOS VARIOS': 'GASTOS VARIOS', 'ARCHIVO': 'GASTOS VARIOS', 'GASTOS DE ADMINISTRACION': 'GASTOS VARIOS',
    'A DEFINIR': 'GASTOS VARIOS',
    'HONORARIOS POR SERVICIOS': 'HONORARIOS POR SERVICIOS',
    'IMPRENTA Y LIBRERÍA': 'IMPRENTA Y PAPELERIA', 'PAPELERA': 'IMPRENTA Y PAPELERIA',
    'LIMPIEZA': 'GASTOS DE LIMPIEZA', 'INTERNET / TELEFONIA': 'INTERNET/TELEFONIA', 'REFRIGERIO': 'REFRIGERIO',
    'GASTOS DE REPRESENTACION - VIAJE': 'GASTOS DE REPRESENTACION',
    'GASTOS DE REPRESENTACION - EVENTOS': 'GASTOS DE REPRESENTACION',
    'GASTOS DE REPRESENTACION - ALOJAMIENTO': 'GASTOS DE REPRESENTACION',
    'REGALOS EMPRESARIALES': 'GASTOS DE REPRESENTACION',
    'RESIDUOS PATOGENICOS': 'RESIDUOS PATOGENICOS', 'SEGURIDAD': 'SEGURIDAD', 'GASTOS FINANCIEROS': 'SERVICIOS FINANCIEROS',
    'SEGUROS': 'SEGUROS', 'GASTOS MEDICOS': 'SERVICIOS MEDICOS', 'SERVICIOS DE EMERGENCIAS Y TRASLADOS': 'SERVICIOS MEDICOS',
    'UNIFORMES': 'UNIFORMES', 'MATRICULA COLEGIO FARMACEUTICOS': 'Matricula Colegio Farmaceuticos',
    'PUBLICIDAD - CARTELERIA': 'Publicidad', 'PUBLICIDAD - FOLLETOS': 'Publicidad',
}
EXCL_BLOQ = 'EXCLUIDO - PROVEEDOR BLOQUEADO'
EXCL_BU = 'EXCLUIDO - BIEN DE USO'
SIN_MAESTRO = 'SIN MAESTRO'

FUENTE_XL, NEGRITA = Font(name='Arial', size=10), Font(name='Arial', size=10, bold=True)
AMARILLO = PatternFill('solid', fgColor='FFF2CC')


def log(m):
    sys.stderr.write(f'[clasificar-compras] {m}\n')


class Falla(Exception):
    pass


def fail(msg, **extra):
    print(json.dumps({'error': msg, **extra}, ensure_ascii=False, default=str))
    sys.exit(1)


def args():
    out = {}
    for a in sys.argv[1:]:
        k, _, v = a.lstrip('-').partition('=')
        out[k] = v
    return out


def norm(s):
    s = unicodedata.normalize('NFD', str(s or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', s).strip().upper()


def digitos(s):
    return re.sub(r'\D', '', str(s or ''))


def periodo_aaaamm(p):
    s = str(p or '').strip()
    m = re.match(r'^(\d{1,2})[/\-.](\d{4})$', s)
    if m:
        return m.group(2) + m.group(1).zfill(2)
    if re.match(r'^\d{6}$', s):
        return s
    return None


def meses(desde, hasta):
    a, m = int(desde[:4]), int(desde[4:])
    out = []
    while f'{a}{m:02d}' <= hasta:
        out.append(f'{a}{m:02d}')
        a, m = (a + 1, 1) if m == 12 else (a, m + 1)
    return out


# ── .env (mismo orden que lib/env.js de los plugins de Node) ────────────────
def cargar_env():
    for ruta in (os.environ.get('FISCO_ENV_FILE'), os.path.join(PLUGIN, '.env'),
                 os.path.join(os.path.expanduser('~'), '.fisco-ar', '.env')):
        if not ruta or not os.path.isfile(ruta):
            continue
        for linea in open(ruta, encoding='utf-8', errors='replace'):
            m = re.match(r'^\s*(?:export\s+)?([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.*?)\s*$', linea)
            if not m or linea.lstrip().startswith('#'):
                continue
            v = m.group(2)
            if len(v) >= 2 and v[0] == v[-1] and v[0] in '"\'':
                v = v[1:-1]
            if v and not os.environ.get(m.group(1)):
                os.environ[m.group(1)] = v


# ── carpetas (misma logica que lib/rutas-bva.js) ────────────────────────────
def raiz():
    r = (os.environ.get('BVA_UNIDAD_PATH') or '').strip() or RAIZ_DEFECTO
    if not os.path.isdir(r):
        fail(f'No encuentro la unidad compartida en "{r}". Defini BVA_UNIDAD_PATH en el .env.')
    return r


def sociedades():
    """{cuit: carpeta} de las carpetas '<n> - <Nombre> - <CUIT>' de la unidad."""
    base = raiz()
    out = {}
    for d in sorted(os.listdir(base)):
        m = re.match(r'^\s*\d+\s*-\s*(.+?)\s*-\s*(\d{11})\s*$', unicodedata.normalize('NFC', d))
        if m and os.path.isdir(os.path.join(base, d)):
            out[m.group(2)] = os.path.join(base, d)
    return out


def por_cuit(socs, cuit):
    hits = [c for c in socs if c == cuit]
    if not hits:
        raise Falla(f'No hay ninguna carpeta de sociedad con el CUIT {cuit}.')
    return socs[hits[0]]


def por_nombre(socs, nombre):
    q = norm(nombre)
    nom = {c: nombre_sociedad(p) for c, p in socs.items()}
    exactas = [c for c, n in nom.items() if n == q]
    hits = exactas or [c for c, n in nom.items() if re.search(rf'(^|\W){re.escape(q)}(\W|$)', n)]
    if not hits:
        raise Falla(f'No hay ninguna sociedad que se llame "{nombre}".')
    if len(hits) > 1:
        raise Falla(f'"{nombre}" coincide con varias sociedades: {" | ".join(nom[c] for c in hits)}. '
                    'Escribi el nombre completo o pasa --cuit.')
    return socs[hits[0]]


def nombre_sociedad(soc):
    n = unicodedata.normalize('NFC', os.path.basename(soc))
    n = re.sub(r'\d{11}', '', n)
    n = re.sub(r'^\s*\d+\s*-\s*', '', n)
    return norm(n.strip(' -'))


def carpeta_mes(mens, per):
    suelto, anio = os.path.join(mens, per), os.path.join(mens, per[:4])
    return os.path.join(anio, per) if not os.path.isdir(suelto) and os.path.isdir(anio) else suelto


def carpeta_iva(soc, per):
    """Como carpetaIva() de lib/rutas-bva.js, pero sin crear nada: None si no existe."""
    mens = os.path.join(soc, '01-Impuestos Mensuales')
    mes = carpeta_mes(mens, per)
    for d in (os.path.join(mes, f'{per}-IVA'), os.path.join(mens, f'{per}-IVA'), os.path.join(mes, 'IVA')):
        if os.path.isdir(d):
            return d
    return mes if os.path.isdir(mes) else None


def carpetas_busqueda(soc, per):
    """Donde buscar compras y conciliacion: las carpetas ncr/ y las de IVA del mes."""
    mens = os.path.join(soc, '01-Impuestos Mensuales')
    mes = carpeta_mes(mens, per)
    out = []
    for d in (os.path.join(mes, f'{per}-IVA', 'ncr'), os.path.join(mens, f'{per}-IVA', 'ncr'),
              os.path.join(mes, 'IVA', 'ncr'), os.path.join(mes, 'ncr'),
              os.path.join(mes, f'{per}-IVA'), os.path.join(mens, f'{per}-IVA'), os.path.join(mes, 'IVA'), mes):
        if os.path.isdir(d) and d not in out:
            out.append(d)
    return out


# ── lectura del libro de compras ────────────────────────────────────────────
def numero_ar(s):
    s = str(s).strip()
    if s == '':
        return None
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    return float(s)


def fecha(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    s = str(v or '').strip()
    for f in ('%Y-%m-%d', '%d/%m/%Y', '%Y-%m-%d %H:%M:%S', '%m-%d-%y', '%d-%m-%Y'):
        try:
            return datetime.datetime.strptime(s[:19] if ' ' in s else s, f).date()
        except ValueError:
            pass
    return None


def tipar(cab, filas, origen):
    claves = [norm(c).lower() for c in cab]
    falt = [c for c in (FECHA, 'tipo de comprobante', 'nro. doc. vendedor', 'denominacion vendedor',
                        'importe total', 'total iva') if c not in claves]
    if falt:
        raise Falla(f'{origen} no tiene las columnas {falt}: no parece un libro de compras del Portal IVA.')
    out = []
    for n, f in enumerate(filas, start=2):
        if not any(v not in (None, '') for v in f):
            continue
        r = {}
        for c, k, v in zip(cab, claves, f):
            try:
                if k == FECHA:
                    r[c] = fecha(v)
                elif k in TEXTO:
                    r[c] = '' if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v).strip())
                elif isinstance(v, (int, float)):
                    r[c] = int(v) if k in ENTEROS else float(v)
                else:
                    x = numero_ar(v) if v is not None else None
                    r[c] = (int(x) if k in ENTEROS and x is not None else x)
            except ValueError:
                raise Falla(f'{origen}, fila {n}: "{c}" no es numerico ({v!r}).')
        out.append(r)
    return out


def decodificar(b):
    for enc in ('utf-8-sig', 'ISO-8859-1'):
        try:
            return b.decode(enc)
        except UnicodeDecodeError:
            pass


def abrir_xlsx(ruta):
    # La unidad de Drive en streaming a veces tarda en bajar el archivo: se reintenta.
    for intento in range(3):
        try:
            return load_workbook(ruta, read_only=True, data_only=True)
        except (OSError, zipfile.BadZipFile) as e:
            if intento == 2:
                raise Falla(f'No pude abrir {os.path.basename(ruta)}: {e}. Si esta en la unidad de Drive, '
                            'abrilo una vez para que se descargue y volve a correr.')
            time.sleep(3)


def leer_compras(ruta, hoja=None):
    nombre = os.path.basename(ruta)
    ext = os.path.splitext(ruta)[1].lower()
    if ext in ('.csv', '.zip'):
        if ext == '.csv':
            texto = decodificar(open(ruta, 'rb').read())
        else:
            with zipfile.ZipFile(ruta) as z:
                csvs = [n for n in z.namelist() if n.lower().endswith('.csv')]
                if len(csvs) != 1:
                    raise Falla(f'El zip {nombre} tiene {len(csvs)} CSV (se espera uno).')
                texto = decodificar(z.read(csvs[0]))
        filas = list(csv.reader(io.StringIO(texto), delimiter=';'))
        return filas[0], tipar(filas[0], filas[1:], nombre)
    if ext != '.xlsx':
        raise Falla(f'No se leer compras en formato {ext}.')
    wb = abrir_xlsx(ruta)
    hojas = [wb[hoja]] if hoja and hoja in wb.sheetnames else wb.worksheets
    for ws in hojas:
        filas = list(ws.iter_rows(values_only=True))
        if filas and FECHA in [norm(c).lower() for c in filas[0] if c is not None]:
            cab = [str(c).strip() if c is not None else '' for c in filas[0]]
            while cab and cab[-1] == '':
                cab.pop()
            filas = [f[:len(cab)] for f in filas[1:]]
            # El Excel de portal-iva-descarga cierra con una fila en blanco y abajo TOTAL + nota.
            vacia = next((i for i, f in enumerate(filas) if not any(v not in (None, '') for v in f)), None)
            if vacia is not None:
                filas = filas[:vacia]
            return cab, tipar(cab, filas, nombre)
    raise Falla(f'{nombre} no tiene ninguna hoja con el formato del Portal IVA.')


def es_portal(nombre, per):
    n = norm(nombre)
    if nombre.startswith('~$') or 'SIN-NCR' in n or 'SIN NCR' in n or 'CLASIFICADAS' in n or 'ELIMINACION' in n:
        return False
    ext = os.path.splitext(nombre)[1].lower()
    if ext in ('.csv', '.zip') and re.search(rf'COMPROBANTES_PERIODO_{per}_COMPRAS', n):
        return True
    return ext in ('.csv', '.zip', '.xlsx') and 'COMPRAS' in n and 'PORTAL' in n.replace(' ', '') and per in n


def ubicar_fuente(soc, per, modo):
    """(ruta, hoja, descripcion) de las compras a clasificar."""
    dirs = carpetas_busqueda(soc, per)
    if not dirs:
        raise Falla(f'No existe la carpeta del mes {per} en 01-Impuestos Mensuales.')
    re_sin = re.compile(rf'^{per}-COMPRAS-SIN-NCR-.*\.xlsx$', re.I)
    if modo in ('auto', 'sin-ncr'):
        sin = [os.path.join(d, f) for d in dirs for f in sorted(os.listdir(d))
               if re_sin.match(f) and not f.startswith('~$') and os.path.isfile(os.path.join(d, f))]
        if len(sin) > 1:
            raise Falla(f'Hay {len(sin)} conciliaciones NCR de {per}: {" | ".join(sin)}. Indica cual con --compras.')
        if sin:
            return sin[0], 'Compras S.NCR', 'Compras sin NCR (conciliacion-ncr)'
        if modo == 'sin-ncr':
            raise Falla(f'No hay conciliacion NCR de {per} ({per}-COMPRAS-SIN-NCR-*.xlsx). '
                        'Corre primero conciliacion-ncr o usa --fuente=auto.')
    for d in dirs:
        hits = sorted(f for f in os.listdir(d) if os.path.isfile(os.path.join(d, f)) and es_portal(f, per))
        xlsx = [f for f in hits if f.lower().endswith('.xlsx')]
        hits = xlsx or hits        # si estan el Excel y el CSV del mismo libro, el Excel
        if len(hits) == 1:
            return os.path.join(d, hits[0]), None, 'Libro compras Portal IVA (sin depurar NCR)'
        if len(hits) > 1:
            raise Falla(f'Hay {len(hits)} libros de compras en {d}: {" | ".join(hits)}. Indica cual con --compras.')
    raise Falla(f'No encontre las compras del Portal IVA de {per}. Bajalas con portal-iva-descarga '
                '(--libros=compras) o pasa --compras.')


# ── Maestro de Proveedores ──────────────────────────────────────────────────
def leer_maestro(ruta):
    if not os.path.isfile(ruta):
        fail(f'No encuentro el Maestro de Proveedores en "{ruta}". Pasa --maestro o defini BVA_MAESTRO_PROVEEDORES.')
    try:
        wb = abrir_xlsx(ruta)
    except Falla as e:
        fail(str(e))
    if 'MAESTRO' not in wb.sheetnames:
        fail('El Maestro de Proveedores no tiene la hoja MAESTRO.')
    filas = wb['MAESTRO'].iter_rows(values_only=True)
    cab = None
    for _ in range(5):
        f = next(filas, None)
        if f and 'CUIT' in [norm(c) for c in f]:
            cab = [norm(c) for c in f]
            break
    if not cab:
        fail('No encuentro el encabezado (CUIT, CUENTA CONTABLE, ESTADO) en la hoja MAESTRO.')
    falt = [c for c in ('CUIT', 'RAZON SOCIAL', 'CUENTA CONTABLE', 'ESTADO') if c not in cab]
    if falt:
        fail(f'A la hoja MAESTRO le faltan las columnas {falt}.')
    i = {c: cab.index(c) for c in cab if c}
    out = {}
    for f in filas:
        c = digitos(f[i['CUIT']]) if f and f[i['CUIT']] is not None else ''
        if len(c) != 11:
            continue
        out[c] = {'razon': f[i['RAZON SOCIAL']], 'cuenta': (str(f[i['CUENTA CONTABLE']] or '').strip() or None),
                  'estado': str(f[i['ESTADO']] or '').strip(),
                  'obs': f[i['OBSERVACIONES']] if 'OBSERVACIONES' in i else None}
    return out


def imputar(m):
    """(cuenta, rubro EERR, estado) de un proveedor del Maestro (o None si no esta)."""
    if not m:
        return SIN_MAESTRO, 'GASTOS VARIOS', SIN_MAESTRO
    cuenta, estado = m['cuenta'], m['estado']
    if estado == 'Bloqueado':
        return cuenta, EXCL_BLOQ, estado
    if estado == 'Revisar':
        return cuenta, 'GASTOS VARIOS', estado
    if cuenta and cuenta.startswith('COMPRA DE EQUIPOS'):
        return cuenta, EXCL_BU, estado
    if not cuenta:
        return None, 'GASTOS VARIOS', estado or 'Sin cuenta'
    return cuenta, RUBRO.get(cuenta, 'GASTOS VARIOS'), estado


def importe_imputado(r, k):
    v = lambda c: (r.get(k[c]) or 0) if c in k else 0
    netos = sum(v(c) for c in NETOS)
    if abs(netos) > 0.004 or abs(v('total iva')) > 0.004:
        return round(netos, 2)
    return round(v('importe total') - sum(v(c) for c in PERCEPCIONES), 2)


def accion(estado):
    if estado == SIN_MAESTRO:
        return 'Dar de alta en el Maestro con su cuenta (skill padron-proveedores)'
    if estado == 'Bloqueado':
        return 'CUIT limitada: no computar credito fiscal ni deducir el gasto'
    if estado == 'Revisar':
        return 'Definir la cuenta en el Maestro (hoy va a GASTOS VARIOS)'
    if estado.startswith('Propuesta'):
        return 'Confirmar la cuenta propuesta en el Maestro'
    return 'Completar la cuenta en el Maestro'


# ── Excel de salida ─────────────────────────────────────────────────────────
def formato(ws, anchos=None):
    for fila in ws.iter_rows():
        for cel in fila:
            cel.font = NEGRITA if cel.row == 1 else FUENTE_XL
            if isinstance(cel.value, datetime.date):
                cel.number_format = 'DD/MM/YYYY'
            elif isinstance(cel.value, float):
                cel.number_format = '#,##0.00'
    for i in range(1, ws.max_column + 1):
        h = ws.cell(1, i).value
        ws.column_dimensions[get_column_letter(i)].width = (anchos or {}).get(i) or min(max(len(str(h or '')) + 2, 12), 40)
    ws.freeze_panes = 'A2'
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def escribir(destino, cab, filas, por_prov, resumen, pendientes, control):
    wb = Workbook()
    wb.remove(wb.active)
    extra = ['Cuenta contable', 'Rubro EERR', 'Estado Maestro', 'Importe imputado']
    ws = wb.create_sheet('Compras clasificadas')
    ws.append(cab + extra)
    for r in filas:
        ws.append([r.get(c) for c in cab] + [r['_cuenta'], r['_rubro'], r['_estado'], r['_imp']])
    formato(ws)
    for i in range(len(cab) + 1, len(cab) + 5):
        ws.column_dimensions[get_column_letter(i)].width = 28

    ws = wb.create_sheet('Resumen por rubro')
    ws.append(['Rubro EERR', 'Cuenta contable', 'Comprobantes', 'Importe imputado', 'Total IVA', 'Importe Total'])
    for x in resumen:
        ws.append(x)
    formato(ws, {1: 34, 2: 40, 3: 14, 4: 18, 5: 16, 6: 18})
    for cel in ws[ws.max_row]:
        cel.font = NEGRITA

    ws = wb.create_sheet('Por proveedor')
    ws.append(['CUIT', 'Proveedor', 'Cuenta contable', 'Rubro EERR', 'Estado Maestro', 'Comprobantes',
               'Importe imputado', 'Total IVA', 'Importe Total'])
    for x in por_prov:
        ws.append(x)
    formato(ws, {1: 14, 2: 40, 3: 34, 4: 30, 5: 18})

    ws = wb.create_sheet('Pendientes Maestro')
    ws.append(['CUIT', 'Proveedor', 'Cuenta usada', 'Rubro EERR', 'Estado Maestro', 'Comprobantes',
               'Importe imputado', 'Que hacer', 'Observaciones del Maestro'])
    for x in pendientes:
        ws.append(x)
    formato(ws, {1: 14, 2: 40, 3: 34, 4: 30, 5: 18, 8: 60, 9: 60})
    for fila in ws.iter_rows(min_row=2):
        if fila[4].value == SIN_MAESTRO:
            for cel in fila:
                cel.fill = AMARILLO

    ws = wb.create_sheet('Control')
    ws.append(['Concepto', 'Comprobantes', 'Importe'])
    for x in control:
        ws.append(x)
    ws.append([])
    col_imp = get_column_letter(len(cab) + 4)
    col_tot = get_column_letter(cab.index(next(c for c in cab if norm(c).lower() == 'importe total')) + 1)
    ws.append(['Recalculado en Excel: Importe Total del detalle', "=COUNTA('Compras clasificadas'!A:A)-1",
               f"=SUM('Compras clasificadas'!{col_tot}:{col_tot})"])
    ws.append(['Recalculado en Excel: Importe imputado del detalle', None,
               f"=SUM('Compras clasificadas'!{col_imp}:{col_imp})"])
    for fila in ws.iter_rows():
        for cel in fila:
            cel.font = NEGRITA if cel.row == 1 else FUENTE_XL
            if cel.column == 3 and isinstance(cel.value, (int, float)):
                cel.number_format = '#,##0.00'
    ws.column_dimensions['A'].width, ws.column_dimensions['B'].width, ws.column_dimensions['C'].width = 52, 16, 60
    wb.save(destino)


# ── una sociedad, un mes ────────────────────────────────────────────────────
def clasificar(soc, per, maestro, maestro_ruta, a, solo_revisar, si_existe):
    nombre = nombre_sociedad(soc)
    cuit = digitos(os.path.basename(soc))[-11:]
    base = {'sociedad': nombre, 'cuit': cuit, 'periodo': per}
    destino_dir = carpeta_iva(soc, per)
    if not destino_dir:
        raise Falla(f'No existe la carpeta del mes {per} en 01-Impuestos Mensuales.')
    salida = os.path.join(destino_dir, f'{per} - COMPRAS CLASIFICADAS.xlsx')
    if os.path.exists(salida) and not solo_revisar:
        if si_existe == 'saltear':
            return {**base, 'ok': True, 'salteada': True, 'salida': salida,
                    'motivo': 'ya estaba clasificado (--si-existe=saltear)'}
        if si_existe == 'error':
            raise Falla(f'Ya existe {os.path.basename(salida)}. Usa --si-existe=reemplazar '
                        '(el anterior va a _anteriores/) o --si-existe=saltear.')

    if a.get('compras'):
        ruta, hoja, fuente = os.path.abspath(a['compras']), None, 'Archivo indicado con --compras'
        if re.search(r'COMPRAS-SIN-NCR', os.path.basename(ruta), re.I):
            hoja, fuente = 'Compras S.NCR', 'Compras sin NCR (conciliacion-ncr)'
    else:
        ruta, hoja, fuente = ubicar_fuente(soc, per, a.get('fuente') or 'auto')
    cab, filas = leer_compras(ruta, hoja)
    log(f'{nombre} {per}: {os.path.basename(ruta)} · {len(filas)} comprobantes')
    k = {norm(c).lower(): c for c in cab}
    C_CUIT, C_DEN, C_TOT, C_IVA, C_FECHA = (k['nro. doc. vendedor'], k['denominacion vendedor'],
                                            k['importe total'], k['total iva'], k[FECHA])

    avisos = []
    if 'sin depurar' in fuente:
        avisos.append('Las NCR de Suizo y Monroe no estan depuradas (no hay conciliacion NCR del mes): '
                      'el CMV incluye esas notas de credito.')
    fuera = [r for r in filas if r[C_FECHA] and r[C_FECHA].strftime('%Y%m') > per]
    if fuera:
        avisos.append(f'{len(fuera)} comprobantes con fecha posterior al periodo: revisar que sea el libro de {per}.')

    prov = OrderedDict()
    for r in filas:
        c = digitos(r[C_CUIT])
        m = maestro.get(c)
        r['_cuenta'], r['_rubro'], r['_estado'] = imputar(m)
        r['_imp'] = importe_imputado(r, k)
        p = prov.setdefault(c, {'den': (m or {}).get('razon') or r[C_DEN], 'cuenta': r['_cuenta'], 'rubro': r['_rubro'],
                                'estado': r['_estado'], 'obs': (m or {}).get('obs'), 'n': 0, 'imp': 0.0, 'iva': 0.0,
                                'tot': 0.0})
        p['n'] += 1; p['imp'] += r['_imp']; p['iva'] += r[C_IVA] or 0; p['tot'] += r[C_TOT] or 0

    orden_rubro = lambda rub: (rub.startswith('EXCLUIDO'), rub != 'CMV', rub)
    por_prov = [[c, p['den'], p['cuenta'], p['rubro'], p['estado'], p['n'], round(p['imp'], 2), round(p['iva'], 2),
                 round(p['tot'], 2)] for c, p in sorted(prov.items(), key=lambda x: (orden_rubro(x[1]['rubro']), -x[1]['imp']))]
    acc = defaultdict(lambda: [0, 0.0, 0.0, 0.0])
    for p in prov.values():
        x = acc[(p['rubro'], p['cuenta'] or '')]
        x[0] += p['n']; x[1] += p['imp']; x[2] += p['iva']; x[3] += p['tot']
    resumen = [[rub, cta, x[0], round(x[1], 2), round(x[2], 2), round(x[3], 2)]
               for (rub, cta), x in sorted(acc.items(), key=lambda kv: (orden_rubro(kv[0][0]), kv[0][1]))]
    resumen.append(['TOTAL', None, len(filas), round(sum(x[3] for x in resumen), 2),
                    round(sum(x[4] for x in resumen), 2), round(sum(x[5] for x in resumen), 2)])
    pend = [[c, p['den'], p['cuenta'], p['rubro'], p['estado'], p['n'], round(p['imp'], 2), accion(p['estado']), p['obs']]
            for c, p in sorted(prov.items(), key=lambda x: (x[1]['estado'] != SIN_MAESTRO, -abs(x[1]['imp'])))
            if p['estado'] != 'Clasificado']

    total = round(sum(r[C_TOT] or 0 for r in filas), 2)
    imputado = round(sum(r['_imp'] for r in filas), 2)
    excluido = round(sum(r['_imp'] for r in filas if r['_rubro'].startswith('EXCLUIDO')), 2)
    sin_m = [p for p in prov.values() if p['estado'] == SIN_MAESTRO]
    verif = round(sum(x[3] for x in resumen[:-1]) - imputado, 2)
    control = [
        ['Comprobantes clasificados', len(filas), None],
        ['Importe Total (con IVA y percepciones)', None, total],
        ['Importe imputado (sin IVA ni percepciones)', None, imputado],
        ['  de eso, CMV', None, round(sum(r['_imp'] for r in filas if r['_rubro'] == 'CMV'), 2)],
        ['  de eso, gastos', None, round(imputado - excluido - sum(r['_imp'] for r in filas if r['_rubro'] == 'CMV'), 2)],
        ['  de eso, excluido del resultado (bloqueados y bienes de uso)', None, excluido],
        ['Verificacion: resumen por rubro - detalle (tiene que dar 0)', None, verif],
        [],
        ['Proveedores', len(prov), None],
        ['  sin Maestro (dar de alta)', len(sin_m), round(sum(p['imp'] for p in sin_m), 2)],
        ['  pendientes de confirmar en el Maestro', len(pend) - len(sin_m), None],
        [],
        ['Fuente', None, f'{fuente}: {os.path.basename(ruta)}'],
        ['Maestro de Proveedores', None, f'{os.path.basename(maestro_ruta)} (modificado '
                                         f'{datetime.datetime.fromtimestamp(os.path.getmtime(maestro_ruta)):%d/%m/%Y %H:%M})'],
        ['Clasificado el', None, f'{datetime.datetime.now():%d/%m/%Y %H:%M}'],
    ] + [['Aviso', None, x] for x in avisos]
    if verif:
        raise Falla(f'El resumen por rubro no cierra con el detalle (diferencia {verif}).')

    res = {**base, 'ok': True, 'fuente': fuente, 'compras': ruta, 'comprobantes': len(filas),
           'importe_total': total, 'importe_imputado': imputado, 'excluido': excluido,
           'cmv': control[3][2], 'gastos': control[4][2], 'proveedores': len(prov),
           'sin_maestro': [{'cuit': c, 'proveedor': p['den'], 'comprobantes': p['n'], 'importe': round(p['imp'], 2)}
                           for c, p in prov.items() if p['estado'] == SIN_MAESTRO],
           'pendientes_confirmar': len(pend) - len(sin_m), 'avisos': avisos}
    if solo_revisar:
        return {**res, 'solo_revisar': True}
    if os.path.exists(salida):
        ant = os.path.join(destino_dir, '_anteriores')
        os.makedirs(ant, exist_ok=True)
        movido = os.path.join(ant, f'{per} - COMPRAS CLASIFICADAS - {datetime.datetime.now():%Y%m%d-%H%M%S}.xlsx')
        shutil.move(salida, movido)
        res['anterior'] = movido
    escribir(salida, cab, filas, por_prov, resumen, pend, control)
    return {**res, 'salida': salida}


# ── principal ───────────────────────────────────────────────────────────────
def main():
    cargar_env()
    a = args()
    if a.get('periodo'):
        p = periodo_aaaamm(a['periodo'])
        pers = [p] if p else None
    else:
        d, h = periodo_aaaamm(a.get('desde')), periodo_aaaamm(a.get('hasta'))
        pers = meses(d, h) if d and h and d <= h else None
    if not pers:
        fail('Falta --periodo (MM/AAAA) o --desde y --hasta.')
    si_existe = a.get('si-existe') or 'error'
    if si_existe not in ('error', 'reemplazar', 'saltear'):
        fail('--si-existe tiene que ser error, reemplazar o saltear.')
    if (a.get('fuente') or 'auto') not in ('auto', 'sin-ncr', 'portal'):
        fail('--fuente tiene que ser auto, sin-ncr o portal.')

    socs = sociedades()
    elegidas, errores = [], []
    if 'todas' in a:
        elegidas = list(socs.values())
    for c in filter(None, (a.get('cuits') or a.get('cuit') or '').split(';')):
        try:
            elegidas.append(por_cuit(socs, digitos(c)))
        except Falla as e:
            errores.append({'sociedad': c, 'error': str(e)})
    for n in filter(None, (a.get('nombres') or a.get('nombre') or '').split(';')):
        try:
            elegidas.append(por_nombre(socs, n.strip()))
        except Falla as e:
            errores.append({'sociedad': n, 'error': str(e)})
    elegidas = list(OrderedDict.fromkeys(elegidas))
    if not elegidas and not errores:
        fail('Indica la sociedad: --nombre, --cuit, --nombres, --cuits o --todas.')
    if a.get('compras') and (len(elegidas) != 1 or len(pers) != 1):
        fail('--compras sirve solo para una sociedad y un mes.')

    maestro_ruta = a.get('maestro') or os.environ.get('BVA_MAESTRO_PROVEEDORES') or os.path.join(raiz(), MAESTRO_REL)
    maestro = leer_maestro(maestro_ruta)
    log(f'Maestro: {len(maestro)} proveedores')

    resultados, sin_libro = [], []
    for soc in elegidas:
        for per in pers:
            try:
                resultados.append(clasificar(soc, per, maestro, maestro_ruta, a, 'solo-revisar' in a, si_existe))
            except Falla as e:
                # Con --todas, las carpetas sin libro del mes (personas, sociedades sin IVA) no son fallas.
                if 'todas' in a and re.match(r'No (existe la carpeta del mes|encontre las compras)', str(e)):
                    sin_libro.append(f'{nombre_sociedad(soc)} {per}')
                    continue
                errores.append({'sociedad': nombre_sociedad(soc), 'periodo': per, 'error': str(e)})
                log(f'{nombre_sociedad(soc)} {per}: {e}')
    if len(resultados) == 1 and not errores and not sin_libro:
        print(json.dumps(resultados[0], ensure_ascii=False, default=str, indent=1))
        return
    if len(errores) == 1 and not resultados:
        fail(errores[0]['error'], **{k: v for k, v in errores[0].items() if k != 'error'})
    print(json.dumps({'ok': not errores, 'clasificadas': len([r for r in resultados if not r.get('salteada')]),
                      'resultados': resultados, 'errores': errores, 'sin_libro_del_mes': sin_libro},
                     ensure_ascii=False, default=str, indent=1))
    if errores and not resultados:
        sys.exit(1)


if __name__ == '__main__':
    main()
