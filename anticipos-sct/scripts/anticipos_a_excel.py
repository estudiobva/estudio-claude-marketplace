#!/usr/bin/env python3
"""
anticipos_a_excel.py — arma el Excel de anticipos del SCT.

    python anticipos_a_excel.py <datos.json> <salida.xlsx>

Formato fijado por la skill anticipos-sct: un archivo por sociedad, columnas
Sociedad | Impuesto | Nro. Anticipo | Capital | Int. Resarcitorio |
Int. Punitorio | Total | Fecha de Vencimiento | Estado.

Los intereses van porque un anticipo pagado fuera de termino queda con capital
$ 0 y "Vencido": sin ellos no se entiende por que sigue en Deudas.

Si la sociedad no tiene anticipos, se escribe igual una fila "Sin anticipos":
saltearla en silencio hace que despues nadie sepa si se reviso o no.
"""

import json
import re
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FUENTE = 'Arial'
AZUL = '1F4E79'
ROJO = '9C0006'
ROJO_BG = 'FFC7CE'
FMT_PESOS = '#,##0.00;-#,##0.00;-'

COLUMNAS = ['Sociedad', 'Impuesto', 'Nro. Anticipo', 'Capital',
            'Int. Resarcitorio', 'Int. Punitorio', 'Total',
            'Fecha de Vencimiento', 'Estado']


def a_numero(s):
    """'$ 2.933,14' -> 2933.14 · '' -> None."""
    s = re.sub(r'[^\d,.-]', '', str(s or ''))
    if not s:
        return None
    try:
        return float(s.replace('.', '').replace(',', '.'))
    except ValueError:
        return None


def montos(a):
    """(capital, resarcitorio, punitorio) del anticipo. Los intereses quedan
    en None si el dato es de una corrida anterior a que se leyeran."""
    return (a_numero(a.get('importe')), a_numero(a.get('intResarcitorio')),
            a_numero(a.get('intPunitorio')))


def estado_de(a):
    """'Vencido' con capital cancelado y solo intereses pendientes se aclara,
    para que no parezca un anticipo impago."""
    est = a.get('estado')
    cap, res, pun = montos(a)
    if est == 'Vencido' and not cap and ((res or 0) + (pun or 0)) > 0:
        return 'Vencido - solo intereses'
    return est


def a_fecha(s):
    m = re.search(r'(\d{2})/(\d{2})/(\d{4})', str(s or ''))
    if not m:
        return None
    return datetime(int(m.group(3)), int(m.group(2)), int(m.group(1))).date()


def escribir(datos, salida):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Anticipos SCT'

    for j, h in enumerate(COLUMNAS, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(name=FUENTE, size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=AZUL)
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[1].height = 28

    soc = datos.get('sociedad') or datos.get('cuit')
    # Los inesperados van en la misma hoja, marcados: son una novedad real y
    # esconderlos en otra solapa es la forma mas facil de que nadie los vea.
    filas = [(a, False) for a in datos.get('anticipos', [])] + \
            [(a, True) for a in datos.get('inesperados', [])]

    i = 2
    if not filas:
        ws.cell(row=2, column=1, value=soc).font = Font(name=FUENTE, size=10)
        c = ws.cell(row=2, column=2, value='Sin anticipos')
        c.font = Font(name=FUENTE, size=10, italic=True)
        i = 3
    else:
        for a, inesperado in filas:
            vals = [soc, a.get('impuesto'), a.get('nroAnticipo'), *montos(a),
                    f'=SUM(D{i}:F{i})', a_fecha(a.get('vencimiento')),
                    estado_de(a)]
            for j, v in enumerate(vals, start=1):
                c = ws.cell(row=i, column=j, value=v)
                c.font = Font(name=FUENTE, size=10,
                              bold=(a.get('estado') == 'Vencido'))
                if 4 <= j <= 7:
                    c.number_format = FMT_PESOS
                if j == 8:
                    c.number_format = 'DD/MM/YYYY'
                if a.get('estado') == 'Vencido':
                    c.font = Font(name=FUENTE, size=10, bold=True, color=ROJO)
                if inesperado:
                    c.fill = PatternFill('solid', fgColor=ROJO_BG)
            i += 1

    ultima = i - 1
    if filas:
        tot = ws.cell(row=ultima + 2, column=3, value='TOTAL')
        tot.font = Font(name=FUENTE, size=10, bold=True)
        for j in range(4, 8):
            L = get_column_letter(j)
            s = ws.cell(row=ultima + 2, column=j, value=f'=SUM({L}2:{L}{ultima})')
            s.font = Font(name=FUENTE, size=10, bold=True)
            s.number_format = FMT_PESOS

    nota = (f'Relevado el {datos.get("fecha")} · CUIT {datos.get("cuit")} · '
            f'Fuente: Sistema de Cuentas Tributarias de ARCA, solapa Deudas.')
    if datos.get('inesperados'):
        nota += (f' ATENCION: {len(datos["inesperados"])} anticipo(s) de un impuesto '
                 f'distinto de Ganancias/Bienes Personales, marcados en rojo.')
    r = datos.get('resumen') or {}
    if r.get('saldoTotal'):
        nota += (f' Saldo total del SCT: $ {r["saldoTotal"]}'
                 f' (capital vencido $ {r.get("capitalVencido", "?")}).')
    n = ws.cell(row=ultima + 3, column=1, value=nota)
    n.font = Font(name=FUENTE, size=9, italic=True)

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNAS))}{ultima}'
    for j, ancho in enumerate([34, 34, 13, 16, 16, 14, 16, 19, 22], start=1):
        ws.column_dimensions[get_column_letter(j)].width = ancho

    wb.save(salida)
    return {'archivo': salida, 'filas': max(0, ultima - 1),
            'anticipos': len(datos.get('anticipos', [])),
            'inesperados': len(datos.get('inesperados', []))}


def _encabezado(ws, columnas, anchos):
    for j, h in enumerate(columnas, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(name=FUENTE, size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=AZUL)
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        ws.column_dimensions[get_column_letter(j)].width = anchos[j - 1]
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = 'A2'


def consolidar(lote, salida):
    """Todas las sociedades del lote: hoja Anticipos (una fila por anticipo) y
    hoja Control (una fila por sociedad, incluidas las que fallaron — asi se ve
    de un vistazo que falta relevar)."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Anticipos'
    cols = ['Sociedad', 'CUIT', 'Impuesto', 'Nro. Anticipo', 'Capital',
            'Int. Resarcitorio', 'Int. Punitorio', 'Total',
            'Fecha de Vencimiento', 'Estado', 'Novedad']
    _encabezado(ws, cols, [30, 14, 34, 13, 16, 16, 14, 16, 19, 22, 30])
    i = 2
    for s in lote.get('sociedades', []):
        if not s.get('ok'):
            continue
        filas = [(a, False) for a in s.get('detalle') or []] + \
                [(a, True) for a in s.get('inesperados') or []]
        for a, inesperado in filas:
            vals = [s.get('sociedad'), s.get('cuit'), a.get('impuesto'), a.get('nroAnticipo'),
                    *montos(a), f'=SUM(E{i}:G{i})', a_fecha(a.get('vencimiento')),
                    estado_de(a), 'Impuesto distinto de Ganancias/BP' if inesperado else None]
            for j, v in enumerate(vals, start=1):
                c = ws.cell(row=i, column=j, value=v)
                vencido = a.get('estado') == 'Vencido'
                c.font = Font(name=FUENTE, size=10, bold=vencido, color=ROJO if vencido else None)
                if 5 <= j <= 8:
                    c.number_format = FMT_PESOS
                if j == 9:
                    c.number_format = 'DD/MM/YYYY'
                if inesperado:
                    c.fill = PatternFill('solid', fgColor=ROJO_BG)
            i += 1
    ultima = max(i - 1, 1)
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{ultima}'
    if i > 2:
        ws.cell(row=ultima + 2, column=4, value='TOTAL').font = Font(name=FUENTE, size=10, bold=True)
        for j in range(5, 9):
            L = get_column_letter(j)
            t = ws.cell(row=ultima + 2, column=j, value=f'=SUM({L}2:{L}{ultima})')
            t.font = Font(name=FUENTE, size=10, bold=True)
            t.number_format = FMT_PESOS

    wc = wb.create_sheet('Control')
    cc = ['Pos.', 'Sociedad', 'CUIT', 'Resultado', 'Anticipos', 'Vencidos',
          'No Vencidos', 'Otro impuesto', 'Saldo total SCT', 'Detalle']
    _encabezado(wc, cc, [6, 30, 14, 16, 11, 11, 12, 13, 17, 70])
    r = 2
    for s in lote.get('sociedades', []):
        if s.get('ok'):
            res = 'Con anticipos' if s.get('anticipos') else 'Sin anticipos'
            det = f"{s.get('otrasDeudas', 0)} deuda(s) que no son anticipos"
            vals = [s.get('pos'), s.get('sociedad'), s.get('cuit'), res, s.get('anticipos'),
                    s.get('vencidos'), s.get('noVencidos'), len(s.get('inesperados') or []),
                    a_numero((s.get('resumenSct') or {}).get('saldoTotal')), det]
        else:
            vals = [s.get('pos'), s.get('sociedad'), s.get('cuit'), 'ERROR',
                    None, None, None, None, None, s.get('error')]
        for j, v in enumerate(vals, start=1):
            c = wc.cell(row=r, column=j, value=v)
            c.font = Font(name=FUENTE, size=10, bold=(vals[3] == 'ERROR'),
                          color=ROJO if vals[3] == 'ERROR' else None)
            if j == 9:
                c.number_format = FMT_PESOS
        r += 1
    for o in lote.get('omitidos', []):
        vals = [o.get('pos'), o.get('titular'), None, 'Omitida', None, None, None, None, None, o.get('motivo')]
        for j, v in enumerate(vals, start=1):
            wc.cell(row=r, column=j, value=v).font = Font(name=FUENTE, size=10, italic=True)
        r += 1
    wc.auto_filter.ref = f'A1:{get_column_letter(len(cc))}{max(r - 1, 1)}'
    wb.save(salida)
    return {'archivo': salida, 'anticipos': i - 2, 'sociedades': r - 2}


def main():
    if len(sys.argv) == 4 and sys.argv[1] == '--consolidado':
        with open(sys.argv[2], encoding='utf-8') as fh:
            lote = json.load(fh)
        print(json.dumps(consolidar(lote, sys.argv[3]), ensure_ascii=False))
        return
    if len(sys.argv) != 3:
        raise SystemExit('Uso: anticipos_a_excel.py <datos.json> <salida.xlsx>\n'
                         '     anticipos_a_excel.py --consolidado <lote.json> <salida.xlsx>')
    with open(sys.argv[1], encoding='utf-8') as fh:
        datos = json.load(fh)
    print(json.dumps(escribir(datos, sys.argv[2]), ensure_ascii=False))


if __name__ == '__main__':
    main()
