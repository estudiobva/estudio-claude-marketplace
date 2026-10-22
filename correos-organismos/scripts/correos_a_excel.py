#!/usr/bin/env python3
"""
correos_a_excel.py — arma el Excel del relevamiento de correos declarados en
los organismos (lo genera correos-organismos.js).

    python correos_a_excel.py <lote.json> <salida.xlsx>

Hojas:
  - Resumen: una fila por sociedad con los correos de ARCA y ARBA, su estado,
    la columna de AGIP para completar a mano y las observaciones.
  - Detalle: una fila por correo encontrado (organismo, pantalla de donde sale,
    tipo/caracter, verificacion, dominio, si el dominio recibe correo).
  - Correos únicos: cada direccion distinta con en cuantas sociedades aparece,
    para confirmar de una vez las que son de terceros.
  - Control: que se pudo relevar y que no, por organismo.
"""

import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FUENTE = 'Arial'
AZUL = '1F4E79'
ROJO_BG = 'FFC7CE'
AMARILLO_BG = 'FFF2CC'
VERDE_BG = 'C6EFCE'
GRIS_BG = 'E7E6E6'

ESTADO_BG = {'Revisar': ROJO_BG, 'Ver observaciones': AMARILLO_BG, 'OK': VERDE_BG, 'Sin relevar': GRIS_BG}


def hoja(ws, titulo, columnas, anchos, filas, claves, envolver=()):
    ws.title = titulo
    for i, (c, w) in enumerate(zip(columnas, anchos), 1):
        cel = ws.cell(row=1, column=i, value=c)
        cel.font = Font(name=FUENTE, bold=True, color='FFFFFF')
        cel.fill = PatternFill('solid', fgColor=AZUL)
        cel.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, f in enumerate(filas, 2):
        for i, k in enumerate(claves, 1):
            v = f.get(k, '')
            if k in ('pos', 'cantidad'):
                try:
                    v = int(v)
                except (TypeError, ValueError):
                    pass
            cel = ws.cell(row=r, column=i, value=v)
            cel.font = Font(name=FUENTE)
            cel.alignment = Alignment(vertical='top', wrap_text=k in envolver)
    ws.freeze_panes = 'A2'
    if filas:
        ws.auto_filter.ref = f'A1:{get_column_letter(len(columnas))}{len(filas) + 1}'
    if not filas:
        ws.cell(row=2, column=1, value='Sin datos').font = Font(name=FUENTE, italic=True)


def main():
    lote = json.load(open(sys.argv[1], encoding='utf-8'))
    destino = sys.argv[2]
    wb = Workbook()

    ws = wb.active
    hoja(ws, 'Resumen',
         ['#', 'Sociedad', 'CUIT', 'Tipo', 'Estado', 'ARCA - correo(s)', 'ARCA - verificación',
          'ARCA - DFE', 'ARCA - avisos DFE a', 'ARBA - correo(s) (carácter)', 'ARBA - verificación',
          'ARBA - correo de la CIT', 'AGIP - correo (completar)', 'Teléfonos', 'Observaciones'],
         [5, 26, 13, 11, 16, 32, 14, 22, 30, 38, 22, 28, 26, 24, 70],
         lote['resumen'],
         ['pos', 'sociedad', 'cuit', 'tipo', 'estado', 'arcaCorreos', 'arcaVerif', 'arcaDfe', 'arcaAvisos',
          'arbaCorreos', 'arbaVerif', 'arbaCit', 'agip', 'telefonos', 'observaciones'],
         envolver=('arcaCorreos', 'arcaAvisos', 'arbaCorreos', 'arbaVerif', 'telefonos', 'observaciones'))
    for r, f in enumerate(lote['resumen'], 2):
        bg = ESTADO_BG.get(f.get('estado'))
        if bg:
            ws.cell(row=r, column=5).fill = PatternFill('solid', fgColor=bg)
        if f.get('agip'):
            ws.cell(row=r, column=13).fill = PatternFill('solid', fgColor=AMARILLO_BG)

    hoja(wb.create_sheet(), 'Detalle',
         ['#', 'Sociedad', 'CUIT', 'Organismo', 'Dónde figura', 'Correo', 'Tipo / carácter',
          'Estado / verificación', 'Dominio', 'Clasificación', 'El dominio recibe correo'],
         [5, 26, 13, 10, 34, 34, 28, 20, 22, 24, 16],
         lote['detalle'],
         ['pos', 'sociedad', 'cuit', 'organismo', 'fuente', 'correo', 'tipo', 'estado', 'dominio',
          'clasificacion', 'recibeCorreo'])

    hoja(wb.create_sheet(), 'Correos únicos',
         ['Correo', 'Clasificación', 'El dominio recibe correo', 'Cant. sociedades', 'Organismos',
          '¿Sigue activo? (completar)', 'Sociedades'],
         [36, 24, 16, 10, 14, 18, 50],
         lote['unicos'],
         ['correo', 'clasificacion', 'recibeCorreo', 'cantidad', 'organismos', '_', 'sociedades'],
         envolver=('sociedades',))

    hoja(wb.create_sheet(), 'Control',
         ['#', 'Sociedad', 'CUIT', 'ARCA', 'ARBA', 'AGIP', 'Relevado'],
         [5, 26, 13, 40, 40, 26, 22],
         lote['control'],
         ['pos', 'sociedad', 'cuit', 'arca', 'arba', 'agip', 'cuando'],
         envolver=('arca', 'arba'))
    if lote.get('omitidos'):
        ws = wb['Control']
        r = ws.max_row + 2
        ws.cell(row=r, column=1, value='Omitidos del archivo de claves').font = Font(name=FUENTE, bold=True)
        for o in lote['omitidos']:
            r += 1
            ws.cell(row=r, column=1, value=o.get('pos'))
            ws.cell(row=r, column=2, value=o.get('titular'))
            ws.cell(row=r, column=4, value=o.get('motivo'))

    wb.save(destino)
    print(json.dumps({'ok': True, 'archivo': destino}))


if __name__ == '__main__':
    main()
