"""Carga la clave de Zetti (T&S Web) en Claves_Organismos.xlsx.

Pide la clave por teclado sin mostrarla y la escribe en la columna "Clave" del
grupo ZETTI, en la fila "Zetti T&S Web (todas las sociedades)".

    python3 scripts/cargar-clave-zetti.py "<ruta a Claves_Organismos.xlsx>"
"""
import getpass
import sys

import openpyxl


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    ruta = sys.argv[1]
    wb = openpyxl.load_workbook(ruta)
    ws = wb.worksheets[0]

    col_clave = None
    for c in range(1, ws.max_column + 1):
        if str(ws.cell(2, c).value or '').strip().upper() == 'ZETTI':
            if str(ws.cell(3, c + 1).value or '').strip().lower() == 'clave':
                col_clave = c + 1
    fila = next((r for r in range(4, ws.max_row + 1)
                 if str(ws.cell(r, 2).value or '').lower().startswith('zetti')), None)
    if not col_clave or not fila:
        sys.exit('No encontre el grupo ZETTI o la fila "Zetti T&S Web" en la planilla.')

    clave = getpass.getpass('Clave de Zetti (no se ve al escribir): ')
    if not clave:
        sys.exit('Clave vacia: no se cambio nada.')
    if getpass.getpass('Repetila: ') != clave:
        sys.exit('No coinciden: no se cambio nada.')

    ws.cell(fila, col_clave).value = clave
    wb.save(ruta)
    print(f'Listo: clave cargada en la fila {fila}.')


if __name__ == '__main__':
    main()
