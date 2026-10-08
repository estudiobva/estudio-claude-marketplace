"""Subdiario de IVA Ventas de Zetti (reporte 5.6.5, CSV) -> Excel limpio.

    python3 zetti_a_excel.py <entrada.csv> <salida.xlsx> [--periodo=AAAAMM]

El CSV de T&S Web viene en latin-1, separado por ";" y con columnas vacias
intercaladas:
    filas 1-5  encabezado (T&S Web, fecha de emision, "Subdiario de IVA Ventas",
               Fecha des/has, Nodo, Razon social, C.U.I.T.)
    fila 6     titulos (Fecha;TC;M;Nro. Comp.;;Cliente;;;;CUIT;;;RESP;;;Exen;...)
    datos      una fila por comprobante (los tiques van agrupados en una fila Z
               por cierre de caja)
    ultima     ";;;...Total:;;;<Exen>;;<Grav>;<IVA>;;;<P.IB>;;<P.IVA>;<Total>"

Salida (hoja "Ventas Zetti"):
    fila 1 T&S Web | fila 2 razon social | fila 3 Fecha des / Fecha has
    fila 4 Subdiario de IVA Ventas + C.U.I.T. (texto) | fila 6 titulos
    datos desde la fila 7, columnas contiguas A..N, ordenados por TC (FV, NC,
    ND, Z...) manteniendo el orden por fecha dentro de cada TC, y fila TOTAL
    con SUM.

Imprime por stdout un JSON con razon social, CUIT, fechas, cantidad por TC,
totales y el control contra la fila "Total:" del reporte. Sale con 1 (y el
JSON con "error") si el Total no coincide al centavo (o alguna otra columna
difiere en mas de $0,05), si hay filas que no
reconoce o si hay fechas fuera del periodo pedido: en ese caso no hay que
archivar.
"""
import csv
import json
import re
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Titulo del reporte -> clave interna. El codigo final (901/902) no tiene titulo.
COLUMNAS = [
    ('Fecha', 'fecha'), ('TC', 'tc'), ('M', 'm'), ('Nro. Comp.', 'nro'),
    ('Cliente', 'cliente'), ('CUIT', 'cuit'), ('RESP', 'resp'),
    ('Exen', 'exen'), ('Grav', 'grav'), ('IVA', 'iva'),
    ('P.IB', 'pib'), ('P.IVA', 'piva'), ('Total', 'total'),
]
IMPORTES = ['exen', 'grav', 'iva', 'pib', 'piva', 'total']
TOLERANCIA = 0.05
TEXTO = ['tc', 'm', 'nro', 'cliente', 'cuit', 'resp', 'cod']


def fallar(msg, **extra):
    print(json.dumps({'error': msg, **extra}, ensure_ascii=False))
    sys.exit(1)


def importe(s):
    s = (s or '').strip()
    if not s:
        return 0.0
    try:
        return float(s.replace('.', '').replace(',', '.'))
    except ValueError:
        fallar(f'Importe que no entiendo: "{s}"')


def fecha(s):
    s = (s or '').strip()
    for fmt in ('%d/%m/%y', '%d/%m/%Y'):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None


def tras(fila, etiqueta):
    """Primer valor no vacio a la derecha de la celda que dice `etiqueta`."""
    for i, c in enumerate(fila):
        if c.strip().lower().startswith(etiqueta.lower()):
            for v in fila[i + 1:]:
                if v.strip():
                    return v.strip()
    return ''


def main():
    pos = [a for a in sys.argv[1:] if not a.startswith('--')]
    opts = dict(a[2:].split('=', 1) for a in sys.argv[1:] if a.startswith('--') and '=' in a)
    if len(pos) != 2:
        sys.exit(__doc__)
    entrada, salida = pos
    periodo = opts.get('periodo')

    with open(entrada, encoding='latin-1', newline='') as fh:
        filas = list(csv.reader(fh, delimiter=';'))

    i_tit = next((i for i, f in enumerate(filas) if f and f[0].strip() == 'Fecha' and 'TC' in [c.strip() for c in f]), None)
    if i_tit is None:
        fallar('No encontre la fila de titulos (Fecha;TC;...) en el CSV.')
    titulos = [c.strip() for c in filas[i_tit]]
    idx = {}
    for nombre, clave in COLUMNAS:
        if nombre not in titulos:
            fallar(f'Falta la columna "{nombre}" en el CSV.', titulos=[t for t in titulos if t])
        idx[clave] = titulos.index(nombre)
    idx['cod'] = idx['total'] + 1

    encabezado = filas[:i_tit]
    razon = cuit_soc = desde = hasta = nodo = ''
    for f in encabezado:
        razon = razon or tras(f, 'Raz')
        cuit_soc = cuit_soc or tras(f, 'C.U.I.T')
        desde = desde or tras(f, 'Fecha des')
        hasta = hasta or tras(f, 'Fecha has')
        nodo = nodo or tras(f, 'Nodo')
    cuit_soc = re.sub(r'\D', '', cuit_soc)

    datos, total_reporte, raras = [], None, []
    for n, f in enumerate(filas[i_tit + 1:], start=i_tit + 2):
        if not any(c.strip() for c in f):
            continue
        f = f + [''] * (idx['cod'] + 1 - len(f))
        d = fecha(f[idx['fecha']])
        if d is None:
            if any(c.strip().lower().startswith('total') for c in f):
                total_reporte = {k: importe(f[idx[k]]) for k in IMPORTES}
            else:
                raras.append({'fila': n, 'texto': ';'.join(c for c in f if c.strip())[:120]})
            continue
        r = {'fecha': d}
        for k in TEXTO:
            r[k] = f[idx[k]].strip()
        for k in IMPORTES:
            r[k] = importe(f[idx[k]])
        datos.append(r)

    if raras:
        fallar('Hay filas que no son comprobantes ni el total: revisar el CSV antes de archivar.', filas=raras)
    if not datos:
        fallar('El reporte no trae comprobantes.', cuit=cuit_soc, desde=desde, hasta=hasta)
    if total_reporte is None:
        fallar('El CSV no trae la fila "Total:": puede haber bajado cortado.')

    fuera = []
    if periodo:
        fuera = sorted({r['fecha'].strftime('%d/%m/%Y') for r in datos if r['fecha'].strftime('%Y%m') != periodo})

    suma = {k: round(sum(r[k] for r in datos), 2) for k in IMPORTES}
    # El reporte suma con decimales internos: aparecen centavos cruzados entre
    # Exen y Grav que se compensan. El Total tiene que dar exacto; en las demas
    # columnas se tolera TOLERANCIA y se informa como aviso.
    difs = {k: round(suma[k] - total_reporte[k], 2) for k in IMPORTES if abs(suma[k] - total_reporte[k]) >= 0.005}
    graves = {k: v for k, v in difs.items() if k == 'total' or abs(v) > TOLERANCIA}

    # Orden: por TC de la A a la Z; sorted() es estable, asi que dentro de cada
    # TC queda el orden del reporte (por fecha).
    datos.sort(key=lambda r: r['tc'])

    wb = Workbook()
    ws = wb.active
    ws.title = 'Ventas Zetti'
    negrita = Font(bold=True)
    ws['A1'] = 'T&S Web'
    ws['A2'] = razon
    ws['A3'] = 'Fecha des:'
    ws['B3'] = fecha(desde) or desde
    ws['C3'] = 'Fecha has:'
    ws['D3'] = fecha(hasta) or hasta
    for c in ('B3', 'D3'):
        ws[c].number_format = 'DD/MM/YYYY'
        ws[c].alignment = Alignment(horizontal='left')
    ws['A4'] = 'Subdiario de IVA Ventas'
    ws['C4'] = 'C.U.I.T.:'
    ws['D4'] = cuit_soc
    ws['D4'].number_format = '@'
    for c in ('A1', 'A2', 'A4'):
        ws[c].font = negrita

    FILA_TIT = 6
    cabecera = [n for n, _ in COLUMNAS] + ['Cod.']
    claves = [k for _, k in COLUMNAS] + ['cod']
    relleno = PatternFill('solid', fgColor='FFD9E1F2')
    borde = Border(bottom=Side(style='thin'))
    for j, t in enumerate(cabecera, start=1):
        c = ws.cell(FILA_TIT, j, t)
        c.font = negrita
        c.fill = relleno
        c.border = borde

    for i, r in enumerate(datos, start=FILA_TIT + 1):
        for j, k in enumerate(claves, start=1):
            c = ws.cell(i, j, r[k])
            if k == 'fecha':
                c.number_format = 'DD/MM/YYYY'
            elif k in IMPORTES:
                c.number_format = '#,##0.00'
            else:
                c.number_format = '@'

    ultima = FILA_TIT + len(datos)
    fila_total = ultima + 1
    ws.cell(fila_total, 1, 'TOTAL').font = negrita
    for j, k in enumerate(claves, start=1):
        if k in IMPORTES:
            col = get_column_letter(j)
            c = ws.cell(fila_total, j, f'=SUM({col}{FILA_TIT + 1}:{col}{ultima})')
            c.number_format = '#,##0.00'
            c.font = negrita
            c.border = Border(top=Side(style='thin'))

    anchos = {'fecha': 12, 'tc': 6, 'm': 4, 'nro': 17, 'cliente': 42, 'cuit': 15, 'resp': 13, 'cod': 7}
    for j, k in enumerate(claves, start=1):
        ws.column_dimensions[get_column_letter(j)].width = anchos.get(k, 16)
    ws.column_dimensions['A'].width = max(ws.column_dimensions['A'].width, 24)
    ws.freeze_panes = ws.cell(FILA_TIT + 1, 1)
    wb.save(salida)

    por_tc = {}
    for r in datos:
        por_tc[r['tc']] = por_tc.get(r['tc'], 0) + 1
    res = {
        'razon_social': razon, 'cuit': cuit_soc, 'nodo': nodo, 'desde': desde, 'hasta': hasta,
        'comprobantes': len(datos), 'por_tc': por_tc, 'totales': suma,
        'control_total_reporte': 'ok' if not difs else ('redondeo' if not graves else 'no coincide'),
        **({'diferencias_total_reporte': difs} if difs else {}),
        'fechas_fuera_de_periodo': fuera, 'xlsx': salida,
    }
    if graves:
        res['error'] = 'La suma de los comprobantes no coincide con el Total del reporte.'
    elif fuera:
        res['error'] = f'Hay comprobantes fuera del periodo {periodo}.'
    print(json.dumps(res, ensure_ascii=False))
    sys.exit(1 if 'error' in res else 0)


if __name__ == '__main__':
    main()
