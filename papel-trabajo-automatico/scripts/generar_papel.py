# Confecciona el papel de trabajo mensual (formato unificado "202600-liquidaciones-IVA, IIBB y Tasa Municipal-<SOC>.xlsx")
# de los meses YA PRESENTADOS, desde las fuentes oficiales, y lo guarda APARTE ("... - AUTOMATICO.xlsx", al lado del papel
# del equipo, que no se toca). Pedido del usuario del 08/10/2026.
#
# Molde: el papel de la propia sociedad (regla de oro de la skill papeles-trabajo-unificados: no reconstruir el modelo,
# copiar y reemplazar valores). Por mes se ubican los bloques por etiqueta, se agrandan si hace falta, se cargan los datos
# (azul) y se regeneran todas las fórmulas.
#
# Fuentes por bloque:
#   Ventas por tipo de comprobante ......... libro "AAAAMM - PORTAL IVA - VENTAS.xlsx" (no gravado de tiques va en Exento)
#   Compras por tipo x clasificación ....... libro "AAAAMM - PORTAL IVA - COMPRAS.xlsx" + Maestro de Proveedores
#                                            (cuenta CMV -> MERCADERIA; el resto -> SERVICIOS AL x%)
#   IVA: retenciones y saldos anteriores ... PDF de la DDJJ F.2051 archivado en la carpeta del mes (*IVA-DDJJ*.pdf);
#                                            meses siguientes encadenados con el mes anterior
#   IIBB ARBA .............................. constancia de la DJ (la que baja constancias-anticipo-arba en
#                                            ~/Documents/BVA-salidas/arba-ddjj/<CUIT>/<AAAAMM>.pdf, o la archivada en la
#                                            carpeta del mes *IIBB*DJ*.pdf): alícuotas, retenciones, bancarias, percepciones,
#                                            saldo anterior; el ajuste de redondeo cierra contra el saldo de la DJ
#   Seg. e Higiene ......................... base = ventas gravadas; alícuota y concepto fijo del papel (Avellaneda:
#                                            importe fijo mensual común a todas)
#   Publicidad y propaganda ................ se conserva lo cargado por el equipo (es manual)
# Lo que no tiene fuente queda en AMARILLO con un comentario que dice qué falta.
#
#   python3 generar_papel.py --cuit=<CUIT> [--hasta=AAAAMM] [--salida=ruta.xlsx]
#   python3 generar_papel.py --todas [--hasta=AAAAMM] [--dir=carpeta]   (--dir: deja cada papel ahí en vez de al lado del original)
# --hasta: último mes a cargar (default: el mes anterior); el año del papel es el de --hasta.
# --todas: las sociedades que tienen papel unificado propio. Las que no, se piden de a una con --cuit (usan de molde el
#          papel de otra sociedad del mismo municipio, vaciado).
import openpyxl, glob, os, re, sys, json, pickle, collections, datetime, unicodedata
from copy import copy
from openpyxl.styles import Font, PatternFill
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as CL, range_boundaries

PLUGIN = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')

# .env (mismo orden que lib/env.js de los plugins de Node)
def cargar_env():
    for ruta in (os.environ.get('FISCO_ENV_FILE'), os.path.join(PLUGIN, '.env'),
                 os.path.join(os.path.expanduser('~'), '.fisco-ar', '.env')):
        if not ruta or not os.path.isfile(ruta):
            continue
        for linea in open(ruta, encoding='utf-8', errors='replace'):
            m = re.match(r'^\s*(?:export\s+)?([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.*?)\s*$', linea)
            if not m or linea.lstrip().startswith('#'):
                continue
            v = m.group(2)
            if len(v) >= 2 and v[0] == v[-1] and v[0] in '"\'':
                v = v[1:-1]
            if v and not os.environ.get(m.group(1)):
                os.environ[m.group(1)] = v

cargar_env()
U = (os.environ.get('BVA_UNIDAD_PATH') or r'I:\Unidades compartidas\BVA - Sociedades Farmaceuticas').strip()
if not os.path.isdir(U): sys.exit(f'No encuentro la unidad compartida en "{U}": revisá que Google Drive esté montado o definí BVA_UNIDAD_PATH.')
SALIDAS = os.environ.get('BVA_SALIDAS_PATH') or os.path.join(os.path.expanduser('~'), 'Documents', 'BVA-salidas')
ARBA = os.path.join(SALIDAS, 'arba-ddjj')
ANIO = '2026'                                    # lo fija --hasta
BLUE, GREEN = 'FF0000FF', 'FF006100'
AMARILLO = PatternFill('solid', fgColor='FFF2CC')
MESES = {1: 'Enero', 2: 'Febrero', 3: 'Marzo', 4: 'Abril', 5: 'Mayo', 6: 'Junio', 7: 'Julio', 8: 'Agosto', 9: 'Septiembre', 10: 'Octubre', 11: 'Noviembre', 12: 'Diciembre'}
TIPOS = {1: 'Factura A', 2: 'Nota de Débito A', 3: 'Nota de Crédito A', 6: 'Factura B', 7: 'Nota de Débito B', 8: 'Nota de Crédito B',
         11: 'Factura C', 12: 'Nota de Débito C', 13: 'Nota de Crédito C', 51: 'Factura M', 52: 'Nota de Débito M', 53: 'Nota de Crédito M',
         63: 'Liquidación A', 64: 'Liquidación B', 81: 'Ticket Z', 82: 'Ticket Z', 83: 'Ticket Z', 110: 'Nota de Crédito Z', 111: 'Tique Factura C',
         112: 'Tique Nota de Crédito A', 113: 'Tique Nota de Crédito B', 114: 'Tique Nota de Crédito C', 115: 'Tique Nota de Débito A',
         116: 'Tique Nota de Débito B', 201: 'Factura de Crédito Electrónica A', 202: 'Nota de Débito FCE A', 203: 'Nota de Crédito FCE A',
         206: 'Factura de Crédito Electrónica B', 207: 'Nota de Débito FCE B', 208: 'Nota de Crédito FCE B', 211: 'Factura de Crédito Electrónica C',
         212: 'Nota de Débito FCE C', 213: 'Nota de Crédito FCE C'}
NC = {3, 8, 13, 53, 110, 112, 113, 114, 203, 208, 213}
ALICS = [('2,5%', 'Neto Gravado IVA 2,5%', 'Importe IVA 2,5%'), ('5%', 'Neto Gravado IVA 5%', 'Importe IVA 5%'),
         ('10,5%', 'Neto Gravado IVA 10,5%', 'Importe IVA 10,5%'), ('21%', 'Neto Gravado IVA 21%', 'Importe IVA 21%'),
         ('27%', 'Neto Gravado IVA 27%', 'Importe IVA 27%')]
PERC = ['Importe de Per. o Pagos a Cta. de Otros Imp. Nac.', 'Importe de Percepciones de Ingresos Brutos', 'Importe de Impuestos Municipales',
        'Importe de Percepciones o Pagos a Cuenta de IVA', 'Importe de Impuestos Internos', 'Importe Otros Tributos']

def num(v): return isinstance(v, (int, float)) and not isinstance(v, bool)
def nfc(s): return unicodedata.normalize('NFC', str(s or '')).strip()

def abrir(f, **kw):
    import time
    for i in range(4):
        try: return openpyxl.load_workbook(f, **kw)
        except Exception:
            if i == 3: raise
            time.sleep(5 * (i + 1))

carpeta = {}
for d in os.listdir(U):
    m = re.match(r'^\d+ - (.+) - (\d{11})$', d)
    if m: carpeta[m.group(2)] = d

# ------------------------------------------------------------------ fuentes
_maestro = None
def maestro():
    global _maestro
    if _maestro is None:
        ws = abrir(f"{U}/0 - ESTUDIO - 00000000000/Maestro de Proveedores.xlsx", read_only=True, data_only=True)['MAESTRO']
        _maestro = {str(r[0]).strip(): (r[4], r[5]) for r in ws.iter_rows(min_row=3, values_only=True) if r[0]}
    return _maestro

def libro(cuit, per, cual):
    fs = glob.glob(f"{U}/{carpeta.get(cuit, '_')}/01-Impuestos Mensuales/**/{per} - PORTAL IVA - {cual}.xlsx", recursive=True)
    return sorted(fs)[0] if fs else None

def filas_libro(f):
    ws = abrir(f, read_only=True, data_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True); h = {k: i for i, k in enumerate(next(it))}
    for r in it:
        if r[0] is None or r[1] is None: break
        g = lambda k: float(r[h[k]]) if k in h and num(r[h[k]]) else 0.0
        yield dict(tipo=int(r[h['Tipo de Comprobante']]), cuit=str(r[h.get('Nro. Doc. Vendedor', h.get('Nro. Doc. Comprador', 5))] or '').strip(),
                   nograv=g('Importe No Gravado'), exento=g('Importe Exento'), total=g('Importe Total'), perc=sum(g(k) for k in PERC),
                   neto={a: g(n) for a, n, _ in ALICS}, iva={a: g(i) for a, _, i in ALICS}, neto0=g('Neto Gravado IVA 0%'),
                   tot_neto=g('Total Neto Gravado'), tot_iva=g('Total IVA'))

def ventas_mes(f):
    # por tipo: [neto gravado, no gravado(0), exento + no gravado, IVA]
    t = collections.OrderedDict()
    for x in filas_libro(f):
        neto = x['tot_neto'] or sum(x['neto'].values())
        iva = x['tot_iva'] or sum(x['iva'].values())
        r = t.setdefault(x['tipo'], [0.0, 0.0, 0.0, 0.0])
        r[0] += neto; r[2] += x['exento'] + x['nograv'] + x['neto0']; r[3] += iva
    nombre = lambda k: f"{k} - {TIPOS.get(k, 'Comprobante')}"
    pos = [[nombre(k)] + v for k, v in sorted(t.items()) if k not in NC]
    neg = [[nombre(k)] + v for k, v in sorted(t.items()) if k in NC]
    return pos, neg

def compras_mes(f):
    # filas (etiqueta, neto, no gravado, exento, IVA, clasificación). Clasificación con el Maestro: CMV -> MERCADERIA.
    M = maestro(); t = collections.OrderedDict(); sin = set()
    for x in filas_libro(f):
        cuenta = (M.get(x['cuit']) or (None, None))[0]
        if cuenta is None: sin.add(x['cuit'])
        merc = cuenta == 'CMV'
        ivas = {a: v for a, v in x['iva'].items() if abs(v) > 0.004 or abs(x['neto'][a]) > 0.004}
        extra_iva = x['tot_iva'] - sum(x['iva'].values()) if x['tot_iva'] else 0.0
        sin_discr = not ivas and abs(x['nograv']) + abs(x['exento']) + abs(x['neto0']) < 0.005
        nograv = (x['total'] - x['perc']) if sin_discr else x['nograv'] + x['neto0']
        principal = max(ivas, key=lambda a: abs(x['neto'][a])) if ivas else None
        def fila(alic):
            if merc: return (x['tipo'], 'MERCADERIA', alic, f"{x['tipo']} - {TIPOS.get(x['tipo'], 'Comprobante')}" + (f" al {alic}" if alic else ''), 'MERCADERIA')
            if alic: return (x['tipo'], 'SERVICIOS', alic, f"{x['tipo']} - {TIPOS.get(x['tipo'], 'Comprobante')}", f"SERVICIOS AL {alic}")
            return (x['tipo'], 'SERVICIOS', None, f"{x['tipo']} - {TIPOS.get(x['tipo'], 'Comprobante')}", 'SERVICIOS')
        for a in ivas:
            r = t.setdefault(fila(a), [0.0] * 4); r[0] += x['neto'][a]; r[3] += x['iva'][a] + (extra_iva if a == principal else 0)
        r = t.setdefault(fila(principal), [0.0] * 4); r[1] += nograv; r[2] += x['exento']
        if not ivas and abs(extra_iva) > 0.004: r[3] += extra_iva
    filas = [(k, v) for k, v in t.items() if any(abs(z) > 0.004 for z in v)]
    orden = lambda kv: (kv[0][0], kv[0][1] != 'MERCADERIA', ['2,5%', '5%', '10,5%', '21%', '27%', None].index(kv[0][2]))
    pos = [[k[3]] + v + [k[4]] for k, v in sorted(filas, key=orden) if k[0] not in NC]
    # en las NC el papel clasifica solo MERCADERIA / SERVICIOS (las fórmulas usan SUMIF exacto)
    neg = [[k[3]] + v + ['MERCADERIA' if k[1] == 'MERCADERIA' else 'SERVICIOS'] for k, v in sorted(filas, key=orden) if k[0] in NC]
    return pos, neg, sin

def montos(s): return [float(x.replace('.', '').replace(',', '.')) for x in re.findall(r'\$\s?(-?[\d.]+,\d{2})', s)]

DDJJ_IVA = os.path.join(SALIDAS, 'ddjj-iva')

def ddjj_iva(cuit, per):
    # 1) lo bajado con el plugin ddjj-iva (JSON por mes); 2) el PDF F.2051 archivado por el equipo en la carpeta del mes
    fj = f"{DDJJ_IVA}/{cuit}/{per}.json"
    if os.path.exists(fj):
        d = json.load(open(fj))
        if num(d.get('df')): return dict(archivo=f"ARCA, DDJJ F.2051 {d.get('secuencia') or ''} (ddjj-iva)", df=d['df'], cf=d.get('cf'),
                                          st_ant=d.get('st_ant') or 0, sld_ant=d.get('sld_ant') or 0, ret=d.get('ret') or 0)
    fs = [f for f in glob.glob(f"{U}/{carpeta.get(cuit, '_')}/01-Impuestos Mensuales/**/{per}*IVA*DDJJ*.pdf", recursive=True)]
    for f in sorted(fs, key=os.path.getmtime, reverse=True):
        try:
            import pypdf
            t = '\n'.join(p.extract_text() for p in pypdf.PdfReader(f).pages)
        except Exception: continue
        if f"{per}" not in t.replace(' ', ''): continue
        g = lambda pat: (montos(re.search(pat + r'[^\n]*', t).group(0)) or [None])[0] if re.search(pat, t) else None
        return dict(archivo=os.path.relpath(f, U), df=g(r'Total del débito fiscal del período'), cf=g(r'Total del crédito fiscal del período'),
                    st_ant=g(r'Saldo técnico a favor del contribuyente del período anterior'),
                    sld_ant=g(r'Saldo a favor de libre disponibilidad del período anterior'),
                    ret=g(r'Total de retenciones, percepciones y pagos a cuenta'))
    return None

def ddjj_arba(cuit, per):
    # 1) lo bajado con el plugin constancias-anticipo-arba; 2) la DJ archivada en la carpeta del mes (sin texto = escaneada: no sirve)
    import pypdf
    cands = [f"{ARBA}/{cuit}/{per}.pdf"] + sorted(glob.glob(f"{U}/{carpeta.get(cuit, '_')}/01-Impuestos Mensuales/**/{per}*IIBB*DJ*.pdf", recursive=True),
                                                  key=os.path.getmtime, reverse=True)
    f = t = None
    for c in cands:
        if not os.path.exists(c): continue
        try: t = '\n'.join(p.extract_text() or '' for p in pypdf.PdfReader(c).pages)
        except Exception: continue
        if 'Deducciones del contribuyente' in t: f = c; break
    if not f: return None
    o = {'archivo': f}
    i = t.find('Agregadas Eliminadas Total')
    if i >= 0:
        filas = [montos(l) for l in t[i:].split('\n')[1:6]]
        if all(len(x) == 4 for x in filas): o['ded'], o['ret'], o['banc'], o['perc'], o['aduana'] = [x[3] for x in filas]
    else:                                                   # rectificativas: una sola columna "Total"
        j0 = t.find('Deducciones del contribuyente')
        lineas = t[j0:].split('\n')[1:12] if j0 >= 0 else []
        if lineas and lineas[0].strip() == 'Total':
            vals = [montos(l) for l in lineas[1:6]]
            if all(len(x) == 1 for x in vals): o['ded'], o['ret'], o['banc'], o['perc'], o['aduana'] = [x[0] for x in vals]
    o['act'] = [(c, float(b.replace('.', '').replace(',', '.')), float(a.replace(',', '.')) / 100)
                for c, b, a in re.findall(r'^(\d{6}) \$ ([\d.]+,\d{2}) ([\d,]+) \$', t, re.M)]
    j = t.find('Deducciones del contribuyente')
    vals = montos(t[j:]) if j > 0 else []
    for a, b, c in zip(vals, vals[1:], vals[2:]):
        if abs((a - b) - c) < 0.15 and (abs(b) > 0 or abs(a) > 0) and a != c: o['saldo_dj'], o['ant'], o['saldo'] = a, b, c
    return o

# ------------------------------------------------------------------ hoja
def etiquetas(ws, col):
    b = collections.defaultdict(list)
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str): b[v.strip()].append(r)
    return b

def locate(ws):
    b, j = etiquetas(ws, 2), etiquetas(ws, 10)
    one = lambda lab, i=0: b[lab][i] if len(b.get(lab, [])) > i else None
    hdr = b['Tipo de Comprobante']
    k = dict(hv=hdr[0], hc=hdr[1], sdf=one('Subtotal Débito Fiscal'), sncv=one('Subtotal Notas de Crédito s/Ventas'),
             vn=one('Ventas Netas (s/ Notas de Crédito)'), pg=one('Venta Gravada'), pe=one('Venta Exenta'),
             scf=one('Subtotal Crédito Fiscal'), sncc=one('Subtotal Notas de Crédito s/Compras'), totc=one('TOTALES COMPRAS (Crédito Fiscal + NC)'),
             tvg=one('Total Venta Gravada'), tve=one('Total Venta Exenta'), tdf=one('Total Débito Fiscal'),
             ncm=one('Más NC Compras Mercadería'), ncs=one('Más NC Compras Servicios'), sdf2=one('SUBTOTAL DÉBITO FISCAL'),
             tcg=one('Total Compras Gravadas'), tce=one('Total Compras Exentas'), cfm=one('Total Crédito Fiscal Mercadería'),
             cfs=one('Total Crédito Fiscal Servicios (Prorrateo)'), subcf=one('SUBTOTAL CRÉDITO FISCAL'), ncdf=one('Más NC Débito Fiscal'),
             tcf=one('TOTAL CRÉDITO FISCAL'), stec=one('Saldo Técnico'), ret=one('Menos Retenciones'),
             sldp=one('Saldo de Libre Disponibilidad del Período'), st=one('Saldo Técnico a Favor del Período Anterior'),
             sld=one('Saldo de Libre Disponibilidad del Período Anterior'), mas=one('MÁS SALDO A FAVOR PERÍODO ANTERIOR'),
             siva=one('SALDO DE IVA + A PAGAR / - A FAVOR'), ret_ib=one('Retenciones de IB'), banc=one('Retenciones Bancarias'),
             farm=one('Farmacia (exento en IVA)'), perf=one('Perfumería / Otros (gravado en IVA)'), impdet=one('TOTAL / IMPUESTO DETERMINADO'),
             sant=one('Saldo a Favor DDJJ Anterior'), pag=one('Menos Percepciones Agentes'), psir=one('Menos Retenciones SIRTAC'),
             pra=one('Menos Retenciones Agentes'), prb=one('Menos Retenciones Bancarias'), aj=one('Ajuste de Redondeo (s/ DDJJ)'),
             siibb=one('SALDO A FAVOR / (A PAGAR)'), tbase=one('Base Imponible (Gravado)'), talic=one('Alícuota'),
             tsh=one('IMPORTE A PAGAR'), tpp=one('IMPORTE A PAGAR', 1), pyp=one('PYP'), inte=one('Intereses'),
             kcomp=(j.get('Total IVA Computable (Servicios)') or [None])[0], knoc=(j.get('Total IVA No Computable (Servicios)') or [None])[0],
             kmerc=(j.get('Crédito Fiscal Mercadería') or [None])[0], ksub=(j.get('SUBTOTAL CF (Merc.+Servicios)') or [None])[0],
             k105=(j.get('IVA Computable al 10,5%') or [None])[0], k21=(j.get('IVA Computable al 21%') or [None])[0], k27=(j.get('IVA Computable al 27%') or [None])[0])
    k['v0'], k['v1'] = k['hv'] + 1, k['sdf'] - 1
    k['ncv0'], k['ncv1'] = k['sdf'] + 1, k['sncv'] - 1
    k['c0'], k['c1'] = k['hc'] + 1, k['scf'] - 1
    k['ncc0'], k['ncc1'] = k['scf'] + 1, k['sncc'] - 1
    k['fijos'] = [r for r in range(k['talic'] + 1, k['tsh']) if isinstance(ws.cell(r, 2).value, str)] if k['talic'] and k['tsh'] and k['tsh'] > k['talic'] else []
    return k

def reparar(ws):
    # filas de retenciones de IIBB borradas en el papel del equipo (pasó en un papel de 06/2026): reponer etiquetas debajo del título
    b = etiquetas(ws, 2)
    if 'Retenciones de IB' not in b and 'Retenciones y Percepciones de IB' in b:
        r = b['Retenciones y Percepciones de IB'][0]
        if all(ws.cell(r + i, 2).value is None for i in (1, 2)):
            ws.cell(r + 1, 2, 'Retenciones de IB'); ws.cell(r + 1, 4, 'Percepciones de IB')
            ws.cell(r + 2, 2, 'Retenciones Bancarias'); ws.cell(r + 2, 4, 'Retenciones SIRTAC')
            return True
    return False

def insert_rows_styled(ws, idx, n, src_row):
    # openpyxl.insert_rows no mueve celdas combinadas ni altos de fila (de la skill papeles-trabajo-unificados)
    heights = {kk: v.height for kk, v in ws.row_dimensions.items() if v.height is not None}
    merged = [str(r) for r in ws.merged_cells.ranges]
    styles = [copy(ws.cell(src_row, c)._style) for c in range(1, 12)]
    ws.insert_rows(idx, n)
    for kk in list(ws.row_dimensions.keys()): del ws.row_dimensions[kk]
    for kk, h in heights.items(): ws.row_dimensions[kk + n if kk >= idx else kk].height = h
    for mr in list(ws.merged_cells.ranges): ws.merged_cells.ranges.remove(mr)
    for mr in merged:
        c1, r1, c2, r2 = range_boundaries(mr)
        if r1 >= idx: r1 += n; r2 += n
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    for jj in range(n):
        for c in range(1, 12): ws.cell(idx + jj, c)._style = copy(styles[c - 1])

def asegurar_filas(ws, necesita):
    # necesita: {bloque: filas}; se agranda de abajo hacia arriba para no correr los otros bloques
    for blq in ('ncc', 'c', 'ncv', 'v'):
        k = locate(ws)
        a, b = k[blq + '0'], k[blq + '1']
        falta = necesita[blq] - (b - a + 1)
        if falta > 0: insert_rows_styled(ws, b + 1, falta, b if b >= a else a - 1)
    return locate(ws)

def put(ws, r, c, v, color=BLUE, fmt=None):
    cell = ws.cell(r, c); cell.value = v
    if color:
        f = cell.font; cell.font = Font(name=f.name, size=f.sz, bold=f.bold, italic=f.i, color=color)
    if fmt: cell.number_format = fmt
    return cell

def pendiente(ws, r, c, texto):
    cell = ws.cell(r, c); cell.fill = AMARILLO; cell.comment = Comment(texto, 'papel automático')

def regen(ws, k, prev, kp):
    S = lambda r, c, f: ws.cell(r, c).__setattr__('value', f) if r else None
    for blq in ('v', 'ncv', 'c', 'ncc'):
        for r in range(k[blq + '0'], k[blq + '1'] + 1): S(r, 7, f'=SUM(C{r}:F{r})')
    for c in range(3, 8):
        X = CL(c)
        S(k['sdf'], c, f'=SUM({X}{k["v0"]}:{X}{k["v1"]})'); S(k['sncv'], c, f'=SUM({X}{k["ncv0"]}:{X}{k["ncv1"]})')
        S(k['vn'], c, f'=+{X}{k["sdf"]}+{X}{k["sncv"]}')
        S(k['scf'], c, f'=SUM({X}{k["c0"]}:{X}{k["c1"]})'); S(k['sncc'], c, f'=SUM({X}{k["ncc0"]}:{X}{k["ncc1"]})')
        S(k['totc'], c, f'=+{X}{k["scf"]}+{X}{k["sncc"]}')
    vn = k['vn']
    S(k['pg'], 3, f'=IF((C{vn}+E{vn})=0,0,C{vn}/(C{vn}+E{vn}))'); S(k['pe'], 3, f'=IF((C{vn}+E{vn})=0,0,E{vn}/(C{vn}+E{vn}))')
    rng, frg = f'H{k["c0"]}:H{k["c1"]}', f'F{k["c0"]}:F{k["c1"]}'
    S(k['k105'], 11, f'=C{k["pg"]}*SUMIF({rng},"SERVICIOS*10,5*%",{frg})')
    S(k['k21'], 11, f'=C{k["pg"]}*SUMIF({rng},"SERVICIOS*21*%",{frg})')
    S(k['k27'], 11, f'=C{k["pg"]}*SUMIF({rng},"SERVICIOS*27*%",{frg})')
    if k['k105'] and k['k27']: S(k['kcomp'], 11, f'=SUM(K{k["k105"]}:K{k["k27"]})')
    S(k['knoc'], 11, f'=C{k["pe"]}*SUMIF({rng},"SERVICIOS*",{frg})')
    S(k['kmerc'], 11, f'=SUMIF({rng},"MERCADERIA",{frg})'); S(k['ksub'], 11, f'=K{k["kmerc"]}+K{k["kcomp"]}')
    ncr, ncf = f'H{k["ncc0"]}:H{k["ncc1"]}', f'F{k["ncc0"]}:F{k["ncc1"]}'
    S(k['tvg'], 3, f'=C{k["sdf"]}'); S(k['tve'], 3, f'=E{k["sdf"]}'); S(k['tdf'], 3, f'=F{k["sdf"]}')
    S(k['ncm'], 3, f'=-SUMIF({ncr},"MERCADERIA",{ncf})'); S(k['ncs'], 3, f'=-SUMIF({ncr},"SERVICIOS",{ncf})')
    S(k['sdf2'], 3, f'=C{k["tdf"]}+C{k["ncm"]}+C{k["ncs"]}')
    S(k['tcg'], 3, f'=C{k["scf"]}'); S(k['tce'], 3, f'=D{k["scf"]}+E{k["scf"]}')
    S(k['cfm'], 3, f'=K{k["kmerc"]}'); S(k['cfs'], 3, f'=K{k["kcomp"]}')
    S(k['subcf'], 3, f'=C{k["cfm"]}+C{k["cfs"]}'); S(k['ncdf'], 3, f'=-F{k["sncv"]}'); S(k['tcf'], 3, f'=C{k["subcf"]}+C{k["ncdf"]}')
    S(k['stec'], 3, f'=+C{k["tcf"]}+C{k["st"]}-C{k["sdf2"]}'); S(k['sldp'], 3, f'=+C{k["sld"]}+C{k["ret"]}')
    S(k['mas'], 3, f'=-(C{k["st"]}+C{k["sld"]})'); S(k['siva'], 3, f'=C{k["sdf2"]}-C{k["tcf"]}-C{k["ret"]}+C{k["mas"]}')
    if prev:
        put(ws, k['st'], 3, f"=MAX('{prev}'!C{kp['stec']},0)", GREEN)
        put(ws, k['sld'], 3, f"='{prev}'!C{kp['sldp']}", GREEN)
    if k['farm']:
        for r, col in ((k['farm'], 'E'), (k['perf'], 'C')):
            put(ws, r, 3, f'={col}{vn}', GREEN); S(r, 4, f'=ROUND(C{r},0)'); S(r, 6, f'=D{r}*E{r}')
        for c in (3, 4, 6): S(k['impdet'], c, f'=SUM({CL(c)}{k["farm"]}:{CL(c)}{k["perf"]})')
        S(k['pag'], 3, f'=-E{k["ret_ib"]}'); S(k['psir'], 3, f'=-E{k["banc"]}'); S(k['pra'], 3, f'=-C{k["ret_ib"]}'); S(k['prb'], 3, f'=-C{k["banc"]}')
        S(k['siibb'], 3, f'=F{k["impdet"]}+C{k["sant"]}+C{k["pag"]}+C{k["psir"]}+C{k["pra"]}+C{k["prb"]}+C{k["aj"]}')
        if prev and kp.get('siibb'): put(ws, k['sant'], 3, f"=MIN('{prev}'!C{kp['siibb']},0)", GREEN)
    if k['tbase']:
        put(ws, k['tbase'], 3, f'=C{k["perf"]}' if k['perf'] else f'=C{vn}', GREEN)
        S(k['tsh'], 3, f'=C{k["tbase"]}*C{k["talic"]}' + ''.join(f'+C{r}' for r in k['fijos']))
    if k['tpp'] and k['pyp']: S(k['tpp'], 3, f'=+C{k["pyp"]}+C{k["inte"]}' if k['inte'] else f'=+C{k["pyp"]}')

# ------------------------------------------------------------------ armado
def papel_equipo(cuit):
    fs = sorted(set(glob.glob(f"{U}/{carpeta.get(cuit, '_')}/01-Impuestos Mensuales/*liquidacion*.xlsx") +
                    glob.glob(f"{U}/{carpeta.get(cuit, '_')}/01-Impuestos Mensuales/*/*liquidacion*.xlsx")))
    fs = [f for f in fs if 'AUTOMATICO' not in f.upper()]
    return fs[0] if fs else None

def municipio(cuit):
    ws = abrir(f"{U}/0 - ESTUDIO - 00000000000/Maestro de Sociedades.xlsx", read_only=True, data_only=True)['MAESTRO']
    for r in ws.iter_rows(min_row=3, values_only=True):
        if str(r[1]) == cuit: return (r[3], r[26])
    return (None, None)

def molde_ajeno(cuit):
    # sin papel unificado propio: el de otra sociedad con la misma jurisdicción y municipio (se vacía todo lo suyo)
    jur, mu = municipio(cuit)
    ws = abrir(f"{U}/0 - ESTUDIO - 00000000000/Maestro de Sociedades.xlsx", read_only=True, data_only=True)['MAESTRO']
    for r in ws.iter_rows(min_row=3, values_only=True):
        c = str(r[1])
        if c != cuit and r[3] == jur and r[26] == mu and c in carpeta:
            f = papel_equipo(c)
            if f and es_unificado(f): return c, f
    return None, None

def es_unificado(f):
    try:
        wb = abrir(f, read_only=True)
        return f'01-{ANIO}' in wb.sheetnames and 'Subtotal Débito Fiscal' in etiquetas(wb[f'01-{ANIO}'], 2)
    except Exception: return False

def setv(ws, ref, v):
    c = ws[ref] if isinstance(ref, str) else ws.cell(*ref)
    if type(c).__name__ != 'MergedCell': c.value = v

def limpiar_mes(ws, soc, cuit, h):
    k = locate(ws)
    for blq in ('v', 'ncv', 'c', 'ncc'):
        for r in range(k[blq + '0'], k[blq + '1'] + 1):
            for c in (2, 3, 4, 5, 6, 8): setv(ws, (r, c), None)
    for key, cols in (('ret', (3,)), ('st', (3,)), ('sld', (3,)), ('ret_ib', (3, 5)), ('banc', (3, 5)), ('sant', (3, 4)), ('aj', (3,)), ('pyp', (3,)), ('inte', (3,))):
        if k.get(key):
            for c in cols: setv(ws, (k[key], c), None)
    for r in k['fijos']: setv(ws, (r, 3), None)
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 2).value
        if isinstance(v, str) and v.startswith('Cuenta municipal'): ws.cell(r, 2).value = 'Cuenta municipal: (completar)'
    setv(ws, 'A1', f"{soc.upper()}   |   CUIT {cuit[:2]}-{cuit[2:10]}-{cuit[10]}   |   Período {MESES[int(h[:2])]} {ANIO}")

def fijo_hoja(ws):
    # Seg. e Higiene: entre "Alícuota" y el primer "IMPORTE A PAGAR" puede haber conceptos fijos con nombres
    # distintos ("Recargo Fijo / Otros Conceptos", "Importe Fijo Mensual", "OTROS"): se suman como fijo
    alic = False; fijo = None
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 200), max_col=3):
        lab = r[1].value
        if not isinstance(lab, str): continue
        lab = lab.strip()
        if lab == 'Alícuota': alic = True; continue
        if lab == 'IMPORTE A PAGAR': break
        if alic and num(r[2].value): fijo = (fijo or 0) + r[2].value
    return fijo

def fijo_avellaneda():
    # importe fijo mensual de la TSH de Avellaneda (común a todas): el valor más repetido por mes en los papeles del equipo
    cnt = collections.defaultdict(collections.Counter)
    for c in carpeta:
        f = papel_equipo(c)
        if not f: continue
        try: wb = abrir(f, read_only=True, data_only=True)
        except Exception: continue
        for h in wb.sheetnames:
            m = re.match(r'^\s*(\d{2})-(\d{4})\s*$', h)
            if not m: continue
            v = fijo_hoja(wb[h])
            if num(v) and v > 0: cnt[m.group(2) + m.group(1)][round(v, 2)] += 1
    return {per: c.most_common(1)[0][0] for per, c in cnt.items() if c.most_common(1)[0][1] >= 3}

def generar(cuit, hasta, salida=None, fijos=None, log=print):
    src = papel_equipo(cuit); ajeno = None
    if not src or not es_unificado(src):
        ajeno, src = molde_ajeno(cuit)
        if not src: raise SystemExit(f'{cuit}: no hay papel unificado propio ni de otra sociedad del mismo municipio para usar de molde')
    eq_val = abrir(src, data_only=True)
    wb = abrir(src)
    soc = nfc(carpeta[cuit]).split(' - ')[1]
    if ajeno:                                                  # vaciar todo lo de la otra sociedad
        for h in [x for x in wb.sheetnames if re.match(rf'^\d{{2}}-{ANIO}$', x)]:
            limpiar_mes(wb[h], soc, cuit, h); limpiar_mes(eq_val[h], soc, cuit, h)
        if 'Sueldos' in wb.sheetnames:
            ws = wb['Sueldos']; setv(ws, 'F2', soc.upper()); setv(ws, 'F3', f"{cuit[:2]}-{cuit[2:10]}-{cuit[10]}")
            for row in ws.iter_rows(min_row=9, max_row=24):
                for c in row:
                    if num(c.value) and type(c).__name__ != 'MergedCell': c.value = None
        if 'EERR' in wb.sheetnames: setv(wb['EERR'], 'E2', soc.upper()); setv(wb['EERR'], 'E3', f"{cuit[:2]}-{cuit[2:10]}-{cuit[10]}")
        if not salida: salida = os.path.join(os.path.dirname(papel_equipo(cuit) or os.path.join(U, carpeta[cuit], '01-Impuestos Mensuales', 'x')),
                                             f"{ANIO}00-liquidaciones-IVA, IIBB y Tasa Municipal- {soc.upper()} - AUTOMATICO.xlsx")
    meses = [f'{m:02d}-{ANIO}' for m in range(1, 13) if f'{ANIO}{m:02d}' <= hasta and f'{m:02d}-{ANIO}' in wb.sheetnames]
    informe = [('molde', f"papel de {nfc(carpeta[ajeno]).split(' - ')[1]} (vaciado): la sociedad no tiene papel unificado propio")] if ajeno else []
    prev = None; kp = None
    for h in meses:
        per = h[3:] + h[:2]; ws = wb[h]; wv = eq_val[h]
        fv, fc = libro(cuit, per, 'VENTAS'), libro(cuit, per, 'COMPRAS')
        if not fv or not fc:
            informe.append((h, 'sin libros del Portal IVA en la unidad: el mes queda como estaba')); prev, kp = h, locate(ws); continue
        vpos, vneg = ventas_mes(fv); cpos, cneg, sin = compras_mes(fc)
        reparado = reparar(ws)
        k = asegurar_filas(ws, dict(v=len(vpos), ncv=len(vneg), c=len(cpos), ncc=len(cneg)))
        kv = locate(wv) if wv.max_row > 20 else None
        # limpiar datos de los bloques y celdas de carga
        for blq in ('v', 'ncv', 'c', 'ncc'):
            for r in range(k[blq + '0'], k[blq + '1'] + 1):
                for c in (2, 3, 4, 5, 6, 8): ws.cell(r, c).value = None
        for r in range(k['totc'] + 1, k['tvg'] - 2):           # notas "NCR eliminadas / NC final" del equipo: el libro presentado ya viene sin NCR
            for c in (5, 6, 7, 8):
                x = ws.cell(r, c)
                if type(x).__name__ != 'MergedCell' and (num(x.value) or (isinstance(x.value, str) and (x.value.startswith('=') or 'NC' in x.value.upper()))):
                    x.value = None
        def bloque(filas, r0, cls):
            for i, f in enumerate(filas):
                r = r0 + i; ws.cell(r, 2).value = f[0]
                for j, c in enumerate((3, 4, 5, 6)):
                    if abs(f[1 + j]) > 0.004: put(ws, r, c, round(f[1 + j], 2))
                if cls: ws.cell(r, 8).value = f[5]
        bloque(vpos, k['v0'], False); bloque(vneg, k['ncv0'], False); bloque(cpos, k['c0'], True); bloque(cneg, k['ncc0'], True)
        ws['A1'] = f"{soc.upper()}   |   CUIT {cuit[:2]}-{cuit[2:10]}-{cuit[10]}   |   Período {MESES[int(h[:2])]} {ANIO}"
        notas = []
        if sin: notas.append(f'{len(sin)} proveedores sin Maestro (tomados como servicio)')
        if reparado: notas.append('se repusieron las filas de retenciones de IIBB que faltaban en el papel del equipo')
        # ---- IVA: retenciones y saldos
        dj = ddjj_iva(cuit, per)
        if dj and num(dj.get('ret')): put(ws, k['ret'], 3, dj['ret'])
        else:
            ws.cell(k['ret'], 3).value = None; pendiente(ws, k['ret'], 3, 'Falta la DDJJ de IVA (F.2051) del mes para tomar las retenciones y percepciones.'); notas.append('retenciones IVA pendientes')
        regen(ws, k, prev, kp)
        if dj and num(dj.get('st_ant')):                        # con DDJJ, los saldos anteriores salen de ahí (manda lo presentado)
            put(ws, k['st'], 3, dj['st_ant']); put(ws, k['sld'], 3, dj['sld_ant'] or 0)
            if prev: ws.cell(k['st'], 3).comment = Comment(f"Tomado de la DDJJ F.2051 ({dj['archivo']}).", 'papel automático')
        elif not prev:                                          # primer mes del archivo sin DDJJ
                for r in (k['st'], k['sld']): ws.cell(r, 3).value = None; pendiente(ws, r, 3, 'Saldo del período anterior: falta la DDJJ de IVA del mes.')
                notas.append('saldos IVA de apertura pendientes')
        # ---- IIBB ARBA
        if k['farm']:
            ar = ddjj_arba(cuit, per)
            if ar and 'ret' in ar:
                farm = next((a for c, b, a in ar['act'] if c.startswith('47731')), None)
                perf = next((a for c, b, a in ar['act'] if not c.startswith('47731') and b > 0), None) or \
                       next((a for c, b, a in ar['act'] if not c.startswith('47731')), None)
                if farm is not None: put(ws, k['farm'], 5, farm)
                if perf is not None: put(ws, k['perf'], 5, perf)
                put(ws, k['ret_ib'], 3, ar['ret']); put(ws, k['ret_ib'], 5, ar['perc'])
                put(ws, k['banc'], 3, ar['banc']); put(ws, k['banc'], 5, 0)
                ant_cadena = min(informe_saldo_iibb.get(prev), 0) if prev and informe_saldo_iibb.get(prev) is not None else None
                if 'ant' in ar:                                 # manda la DJ: saldo a favor anterior que tomó ARBA
                    put(ws, k['sant'], 3, -ar['ant'])
                    if ant_cadena is not None and abs(-ar['ant'] - ant_cadena) > 1:
                        ws.cell(k['sant'], 3).comment = Comment(f"Saldo anterior según la DJ ARBA. El saldo del mes anterior daba {ant_cadena:,.2f} (compensaciones o pagos de ARBA).", 'papel automático')
                        notas.append(f'IIBB: el saldo anterior de la DJ difiere del mes anterior en {(-ar["ant"]) - ant_cadena:,.2f}')
                elif not prev or not kp.get('siibb'): put(ws, k['sant'], 3, 0)
                fa = farm if farm is not None else (wv.cell(k['farm'], 5).value or 0)
                pe = perf if perf is not None else (wv.cell(k['perf'], 5).value or 0)
                # ajuste de redondeo: residuo contra el saldo de la DJ (con el saldo anterior que deja la fórmula)
                vgr = sum(f[1] for f in vpos + vneg); vex = sum(f[3] for f in vpos + vneg)
                imp = round(vex) * fa + round(vgr) * pe
                ant = -ar['ant'] if 'ant' in ar else (ant_cadena or 0)
                calc = imp + ant - ar['perc'] - ar['ret'] - ar['banc']
                aj = round(ar['saldo'] - calc, 4) if 'saldo' in ar else 0
                put(ws, k['aj'], 3, aj)
                informe_saldo_iibb[h] = ar.get('saldo', calc + aj)
                if abs(aj) > 1: notas.append(f'IIBB: ajuste contra la DJ ARBA de {aj:,.2f} (revisar alícuotas o saldo anterior)')
            else:
                for r, c in ((k['ret_ib'], 3), (k['ret_ib'], 5), (k['banc'], 3), (k['banc'], 5), (k['aj'], 3)):
                    ws.cell(r, c).value = None; pendiente(ws, r, c, 'Sin constancia de la DJ de ARBA del mes (CM/AGIP o sin CIT): cargar a mano.')
                informe_saldo_iibb[h] = None; notas.append('IIBB pendiente (sin DJ ARBA)')
        # ---- Seg. e Higiene: concepto fijo
        for r in k['fijos']:
            lab = str(ws.cell(r, 2).value).upper()
            if fijos and per in fijos and ('FIJO' in lab or 'RECARGO' in lab) and 'AVELLANEDA' in ' '.join(str(ws.cell(x, 2).value or '') for x in range(k['tbase'] - 4, k['tbase'])).upper():
                put(ws, r, 3, fijos[per])
            elif not num(wv.cell(r, 3).value):
                pendiente(ws, r, 3, 'Concepto fijo de la tasa: no hay dato del equipo para este mes.'); notas.append('TSH concepto fijo pendiente')
        informe.append((h, '; '.join(notas) or 'ok'))
        prev, kp = h, k
    # solapa EERR del papel: reapuntar a las filas nuevas
    if 'EERR' in wb.sheetnames:
        we = wb['EERR']; filas = {nfc(we.cell(r, 1).value): r for r in range(1, 60) if we.cell(r, 1).value}
        for h in meses:
            k = locate(wb[h]); col = int(h[:2]) + 1
            for lab, ref in (('Ventas exentas', f"E{k['vn']}"), ('Ventas gravadas', f"C{k['vn']}"), ('IIBB', f"F{k['impdet']}" if k['impdet'] else None),
                             ('IVA Prorrateo', f"K{k['knoc']}"), ('Seguridad e Higiene', f"C{k['tsh']}" if k['tsh'] else None),
                             ('Publicidad', f"C{k['tpp']}" if k['tpp'] else None)):
                if ref and lab in filas: we.cell(filas[lab], col).value = f"=+'{h}'!{ref}"
    if not salida:
        salida = re.sub(r'\.xlsx$', ' - AUTOMATICO.xlsx', src)
    # vínculos externos que quedaron del papel del equipo (rutas G:\ de Windows): ninguna fórmula los usa y Excel en Mac no abre el libro
    if wb._external_links:
        informe.append(('vínculos', f'se quitaron {len(wb._external_links)} vínculos externos sin uso del papel del equipo')); wb._external_links = []
    wb.save(salida)
    return salida, informe

informe_saldo_iibb = {}

if __name__ == '__main__':
    a = dict(x[2:].split('=', 1) if '=' in x else (x[2:], True) for x in sys.argv[1:] if x.startswith('--'))
    hoy = datetime.date.today(); ant = (hoy.replace(day=1) - datetime.timedelta(days=1))
    hasta = str(a.get('hasta') or f'{ant.year}{ant.month:02d}')
    if not re.match(r'^\d{6}$', hasta): sys.exit('--hasta tiene que ser AAAAMM')
    ANIO = hasta[:4]
    if 'cuit' in a: cuits = [re.sub(r'\D', '', str(a['cuit']))]
    elif 'todas' in a: cuits = [c for c in sorted(carpeta, key=lambda c: nfc(carpeta[c])) if (lambda f: f and es_unificado(f))(papel_equipo(c))]
    else: sys.exit('Falta --cuit=<CUIT> o --todas')
    fijos = fijo_avellaneda()
    for c in cuits:
        informe_saldo_iibb.clear()
        try:
            sal = a.get('salida') if 'cuit' in a else None
            if 'dir' in a: sal = os.path.join(a['dir'], f"{nfc(carpeta[c]).split(' - ')[1]}.xlsx")
            out, inf = generar(c, hasta, sal, fijos)
            print(out)
            for h, n in inf: print('  ', h, n)
        except SystemExit as e: print(e)
        except Exception as e: print(c, 'ERROR', type(e).__name__, e)
