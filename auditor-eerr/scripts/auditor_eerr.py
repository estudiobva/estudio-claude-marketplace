# Auditor de EERR: corre 21 controles sobre cada "EERR ... .xlsx" de la carpeta EERR 2026 y documenta lo que encuentra
# en la hoja "Auditorias" del mismo archivo (pedido del usuario 09/10/2026).
#
# Controles (de lo general a lo específico; KPIs de auditoría del equipo + NCR vs NC comerciales + compras = Portal IVA):
#   Archivo ........ 1 Período del ejercicio · 2 Meses completos / valores repetidos · 3 Fórmulas
#   Números grandes  4 Resultado / ventas · 5 Margen bruto · 6 Cada gasto / ventas
#   Ingresos ....... 7 Ventas vs Portal IVA (y cruce Zetti: hoja "cruze" del papel del mes)
#   Costo .......... 8 CMV y existencias · 9 NCR vs NC comerciales · 10 Compras clasificadas = Portal IVA compras
#   Impuestos ...... 11 Sueldos vs F.931 · 12 IIBB vs DJ ARBA · 13 IVA prorrateo · 14 Tasas (S&H, matrícula, publicidad)
#   Gastos ......... 15 Gastos fijos y financieros
#   Papel de trabajo 16 Débito / crédito fiscal (y vs DDJJ) · 17 Signo de las NC · 18 Cómputo de NRF (NCR)
#   (hoja MM-AAAA)   19 Prorrateo del papel · 20 Venta vs compra exenta y gravada · 21 NC de compras vs libro (NCR + NC)
#   Lista de Meli, Ezequiel y Jaz (#automatizaciones-ai 14/10/2026); una sola auditoría con excepciones por sociedad
#   (criterio de Javi): CFG['excepciones'] = {"<sociedad o CUIT>": [números de control que no corren]}.
#
# Hoja "Auditorias": una fila por hallazgo (fecha, nivel, control, mes, descripción, valores, link a la celda del EERR,
# link al documento fuente en Drive, estado Pendiente/Resuelto/Aprobado, observaciones) + historial de ejecuciones.
#   - Hallazgo nuevo -> Pendiente.  Pendiente que ya no aparece -> Resuelto (solo si el control pudo correr ese mes).
#   - Aprobado -> no se vuelve a informar mientras el valor sea el mismo; si el valor cambia, se informa como hallazgo nuevo.
#   - Resuelto que vuelve a aparecer -> se reabre como Pendiente.
# La clave de cada hallazgo (columna oculta "Clave") es control|mes|concepto.
#
# Meses auditados: los del ejercicio hasta --hasta (por defecto el último mes con la DDJJ de IVA vencida: el mes anterior
# desde el día 26, si no el anterior a ese).
#
#   python3 auditor_eerr.py --nombre="<sociedad>" [--hasta=202608] [--simular]
#   python3 auditor_eerr.py --todas [--hasta=202608] [--simular] [--procesos=4]
# --simular: no toca la unidad; deja la copia auditada en <BVA_SALIDAS_PATH>/auditor-eerr/simulacion/.
# Excel recalcula cada archivo antes de devolverlo a la unidad, para que queden los valores de las fórmulas
# (Windows: PowerShell + Excel; Mac: AppleScript + Excel).
# Variables (entorno o ~/.fisco-ar/.env, como los demás plugins): BVA_UNIDAD_PATH (default I:\Unidades compartidas\...),
# BVA_SALIDAS_PATH (default ~/Documents/BVA-salidas). Umbrales propios: ~/.fisco-ar/auditor_eerr_config.json.
import os, re, sys, glob, json, shutil, datetime, statistics, subprocess, unicodedata, collections, time
from concurrent.futures import ProcessPoolExecutor
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as CL

WIN, MAC = os.name == 'nt', sys.platform == 'darwin'
FISCO = os.path.expanduser('~/.fisco-ar')


def _env():
    # el .env del estudio (FISCO_ENV_FILE o ~/.fisco-ar/.env); la variable de entorno pisa la del archivo
    f = os.environ.get('FISCO_ENV_FILE') or os.path.join(FISCO, '.env'); out = {}
    if os.path.exists(f):
        for l in open(f, encoding='utf-8', errors='replace'):
            m = re.match(r'^\s*(?:export\s+)?([A-Za-z_]\w*)\s*=\s*(.*?)\s*$', l)
            if m and not l.lstrip().startswith('#'): out[m[1]] = m[2].strip('"').strip("'")
    return lambda k: os.environ.get(k) or out.get(k) or ''
var = _env()

NOMBRE_UNIDAD = 'BVA - Sociedades Farmaceuticas'
def raiz_unidad():
    r = var('BVA_UNIDAD_PATH').strip()
    if r:
        if not os.path.isdir(r): sys.exit(f'BVA_UNIDAD_PATH apunta a "{r}", que no existe. Verificá que Google Drive esté abierto.')
        return r
    if WIN and os.path.isdir(f'I:\\Unidades compartidas\\{NOMBRE_UNIDAD}'): return f'I:\\Unidades compartidas\\{NOMBRE_UNIDAD}'
    c = glob.glob(os.path.expanduser(f'~/Library/CloudStorage/GoogleDrive-*/Unidades compartidas/{NOMBRE_UNIDAD}'))
    if len(c) == 1: return c[0]
    sys.exit('No encuentro la unidad compartida: definí BVA_UNIDAD_PATH en ~/.fisco-ar/.env.')

U = raiz_unidad()
EERR_DIR = os.path.join(U, '0 - ESTUDIO - 00000000000', 'EERR 2026')
MAESTRO = os.path.join(U, '0 - ESTUDIO - 00000000000', 'Maestro de Proveedores.xlsx')
SP = os.path.dirname(os.path.abspath(__file__))
SALIDAS = var('BVA_SALIDAS_PATH') or os.path.expanduser('~/Documents/BVA-salidas')
SALIDA = os.path.join(SALIDAS, 'auditor-eerr')
ARBA_LOCAL = os.path.join(SALIDAS, 'arba-ddjj')          # copia local de las DJ de ARBA, si existe
HOJA = 'Auditorias'

# Umbrales (se pueden pisar con auditor_eerr_config.json al lado del script)
CFG = dict(
    ventas_tol_pct=0.01,          # 7: diferencia ventas vs Portal IVA / Zetti mayor al 1% = rojo 
    margen_banda=[0.20, 0.40],    # 5: margen bruto esperado (provisorio hasta que el equipo fije el estándar)
    margen_desvio_pts=0.05,       # 5: desvío de ±5 puntos vs mes anterior o mediana del ejercicio 
    gasto_salto_pct=0.20,         # 6: un gasto / ventas que se aparta más de 20% de su mediana 
    materialidad_pct_ventas=0.002,  # 6, 15: y además la diferencia supera el 0,2% de las ventas del mes...
    materialidad_min=50000,       # ... y $50.000
    resultado_max_pct=0.25,       # 4: resultado / ventas por encima de 25% = utilidad atípica 
    perdidas_seguidas=2,          # 4: pérdidas en 2 o más meses seguidos 
    iibb_alic=[0.015, 0.05],      # 12: IIBB / ventas razonable 
    sh_alic=[0.002, 0.03],        # 14: Seguridad e Higiene / ventas gravadas razonable
    gastos_varios_max_pct=0.10,   # 10: GASTOS VARIOS sobre el total de gastos que vienen de compras
    ncr_desvio_pct=0.30,          # 9: NCR / (NCR + NC comerciales) que se aparta más de 30% de su mediana 
    compras_tol=100,              # 10: tolerancia en pesos compras imputadas vs Portal IVA
    prorrateo_tol_pct=0.10, prorrateo_tol_min=5000,  # 13
    fijos=['alquileres', 'edesur', 'metrogas', 'agua', 'internet/telefonia', 'seguridad', 'gastos de sistemas', 'residuos patogenicos'],
    fijos_obligatorios=['edesur', 'internet/telefonia'],
    fijos_presencia=0.75,         # 15: si un gasto fijo está en al menos el 75% de los meses, los meses sin él son falta
    duplicado_factor=1.8,         # 15: un proveedor de gasto fijo que factura 1,8 veces su mediana = posible duplicado
    financieros=['gastos bancarios', 'intereses de prestamos', 'comisiones de tarjetas', 'intereses por facilidades de pago', 'rdos financieros (fci)'],
    cf_df_desvio_pct=0.25,        # 16: crédito / débito fiscal que se aparta más de 25% de su mediana
    prorrateo_pct_tol=0.01,       # 19: % de venta exenta del papel vs Portal IVA ventas (1 punto)
    mix_exento_pts=0.15,          # 20: % exento de las ventas vs % exento de las compras de mercadería (15 puntos)
    compra_venta_desvio_pct=0.30, # 20: compras / ventas (exentas o gravadas) que se aparta más de 30% de su mediana
    nc_tol=100,                   # 21: tolerancia en pesos NC de compras del papel vs libro
    excepciones={},               # {"<sociedad o CUIT>": [controles que no corren]}, p. ej. {"<sociedad>": [19]}
)
for _f in (os.path.join(SP, 'auditor_eerr_config.json'), os.path.join(FISCO, 'auditor_eerr_config.json')):
    if os.path.exists(_f): CFG.update(json.load(open(_f, encoding='utf-8')))

CONTROLES = {1: 'Período del ejercicio', 2: 'Meses completos', 3: 'Fórmulas', 4: 'Resultado / ventas', 5: 'Margen bruto',
             6: 'Gastos / ventas', 7: 'Ventas vs Portal IVA', 8: 'CMV y existencias', 9: 'NCR vs NC comerciales',
             10: 'Compras clasificadas = Portal IVA', 11: 'Sueldos vs F.931', 12: 'IIBB vs DJ ARBA', 13: 'IVA prorrateo',
             14: 'Tasas municipales', 15: 'Gastos fijos y financieros', 16: 'Débito / crédito fiscal', 17: 'Signo de las NC',
             18: 'Cómputo de NRF', 19: 'Prorrateo del papel', 20: 'Venta vs compra exenta/gravada', 21: 'NC de compras vs libro'}
ALICS = ['2,5%', '5%', '10,5%', '21%', '27%']
PERC = ['Importe de Per. o Pagos a Cta. de Otros Imp. Nac.', 'Importe de Percepciones de Ingresos Brutos', 'Importe de Impuestos Municipales',
        'Importe de Percepciones o Pagos a Cuenta de IVA', 'Importe de Impuestos Internos', 'Importe Otros Tributos']


def num(v): return isinstance(v, (int, float)) and not isinstance(v, bool)
def ndoc(v): return str(int(v)) if num(v) else re.sub(r'\D', '', str(v or ''))
def nfc(s): return unicodedata.normalize('NFC', str(s or '')).strip()
def norm(s): return re.sub(r'\s+', ' ', ''.join(c for c in unicodedata.normalize('NFD', nfc(s).lower()) if unicodedata.category(c) != 'Mn')).strip()
def mm(per): return f"{per[4:]}/{per[:4]}"
def pesos(v): return f"$ {v:,.0f}".replace(',', '.')
def pct(v): return f"{v * 100:.1f}%".replace('.', ',')


def abrir(f, **kw):
    for i in range(4):
        try: return openpyxl.load_workbook(f, **kw)
        except Exception:
            if i == 3: raise
            time.sleep(5 * (i + 1))


def hasta_default(hoy=None):
    hoy = hoy or datetime.date.today()
    y, m = hoy.year, hoy.month - (1 if hoy.day >= 26 else 2)
    while m < 1: y, m = y - 1, m + 12
    return f"{y}{m:02d}"


carpetas = {}
for d in os.listdir(U):
    x = re.match(r'^\d+ - (.+) - (\d{11})$', nfc(d))
    if x: carpetas[x.group(2)] = d

_ids = {}
def drive_url(path):
    # Mac: link web de Drive (id que guarda Google Drive para escritorio en los atributos extendidos).
    # Windows u otro: link al archivo en la unidad (file:///I:/...), que abre igual en las PC del estudio.
    if not path: return None
    if path not in _ids:
        i = None
        if MAC:
            try:
                r = subprocess.run(['xattr', '-p', 'com.google.drivefs.item-id#S', path], capture_output=True, text=True, timeout=10)
                i = r.stdout.strip() or None
            except Exception: pass
        if i: _ids[path] = (f"https://drive.google.com/drive/folders/{i}" if os.path.isdir(path) else f"https://drive.google.com/file/d/{i}/view")
        else:
            from pathlib import Path
            try: _ids[path] = Path(path).resolve().as_uri()
            except Exception: _ids[path] = None
    return _ids[path]


# ------------------------------------------------------------------ fuentes
def filas_libro(f):
    ws = abrir(f, read_only=True, data_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True); h = {nfc(k): i for i, k in enumerate(next(it)) if k}
    for r in it:
        if r[0] is None or r[1] is None: break
        g = lambda k: float(r[h[k]]) if k in h and num(r[h[k]]) else 0.0
        n = lambda k: int(r[h[k]]) if k in h and num(r[h[k]]) else None
        tipo = re.match(r'\s*(\d+)', str(r[h['Tipo de Comprobante']] if 'Tipo de Comprobante' in h else ''))
        yield dict(cuit=ndoc(r[h.get('Nro. Doc. Vendedor', h.get('Nro. Doc. Comprador', 5))]),
                   tipo=int(tipo[1]) if tipo else None, pv=n('Punto de Venta'), nro=n('Número de Comprobante'),
                   nograv=g('Importe No Gravado'), exento=g('Importe Exento'), total=g('Importe Total'), perc=sum(g(k) for k in PERC),
                   neto=sum(g(f'Neto Gravado IVA {a}') for a in ALICS), iva=sum(g(f'Importe IVA {a}') for a in ALICS),
                   neto0=g('Neto Gravado IVA 0%'), tot_neto=g('Total Neto Gravado'), tot_iva=g('Total IVA'))

def ventas_libro(f):
    gr = ex = 0.0
    for x in filas_libro(f):
        gr += x['tot_neto'] or x['neto']; ex += x['exento'] + x['nograv'] + x['neto0']
    return gr, ex

def compras_libro(f):
    # misma base que "Compras imputadas": no gravado + exento + netos; sin IVA discriminado, total - percepciones
    t = 0.0
    for x in filas_libro(f):
        neto = x['tot_neto'] or x['neto']; iva = x['tot_iva'] or x['iva']
        sin_discr = abs(iva) < 0.005 and abs(neto) < 0.005 and abs(x['nograv']) + abs(x['exento']) + abs(x['neto0']) < 0.005
        t += (x['total'] - x['perc']) if sin_discr else neto + x['nograv'] + x['exento'] + x['neto0']
    return t

NC_COD = {3, 8, 13, 53, 110, 112, 113, 114, 203, 208, 213}     # códigos de ARCA de notas de crédito
def nc_libro(f):
    # NC del libro de compras con la misma base que la columna Total del papel (netos + no gravado + exento + IVA, sin
    # percepciones) y las claves (CUIT, PV, número) para reconocer las NCR
    t, claves = 0.0, set()
    for x in filas_libro(f):
        if x['tipo'] not in NC_COD: continue
        neto = x['tot_neto'] or x['neto']; iva = x['tot_iva'] or x['iva']
        base = neto + x['nograv'] + x['exento'] + x['neto0'] + iva
        t += base if abs(base) >= 0.005 else x['total'] - x['perc']
        claves.add((x['cuit'], x['pv'], x['nro']))
    return t, claves

def pdf_texto(f):
    import pypdf
    return '\n'.join(p.extract_text() or '' for p in pypdf.PdfReader(f).pages)

ACT = re.compile(r'^(\d{6}) \$ ?(-?[\d\.]+,\d{2}) (?:\* )?(\d+(?:,\d+)?) \$ ?(-?[\d\.]+,\d{2})$')
def dj_arba(f):
    # impuesto determinado de la DJ de anticipo (total después de las actividades; si no se lee, base x alícuota)
    n = lambda s: float(s.replace('$', '').replace('.', '').replace(',', '.').strip())
    L = [l.strip() for l in pdf_texto(f).splitlines()]
    acts, last = [], None
    for i, l in enumerate(L):
        m = ACT.match(l)
        if m: acts.append((n(m[2]), n(m[3]))); last = i
    if last is None: return None
    if last + 1 < len(L) and L[last + 1].startswith('$'): return n(L[last + 1])
    return round(sum(b * a / 100 for b, a in acts), 2)

def f931(f):
    t = pdf_texto(f).split('\n'); t = '\n'.join(t)
    N = r'(-?[\d\.]+,\d{2})'
    g = lambda pat: (lambda m: float(m.group(1).replace('.', '').replace(',', '.')) if m else None)(re.search(pat, t))
    rect = int((re.search(r'Rect\. \(1/9\):\s*(\d)', t) or [0, 0])[1])
    d = dict(rect=rect, rem=g(r'Remun\. con ART\s+' + N), c351=g(r'351 - Contribuciones de Seguridad Social\s+' + N),
             c352=g(r'352 - Contribuciones de Obra Social\s+' + N), c312=g(r'312 - L\.R\.T\.\s+' + N),
             c028=g(r'028 - Seguro Colectivo de Vida Obligatorio\s+' + N))
    if rect > 0 and all((d[k] or 0) == 0 for k in ('c351', 'c352', 'c312', 'c028')):
        d.update(c351=g(r'Contribuciones S\.S\. a pagar\s+' + N), c352=g(r'Contribuciones O\.S\. a pagar\s+' + N),
                 c312=g(r'L\.R\.T\. total a pagar\s+' + N), c028=g(r'S\.C\.V\.O\. a Pagar:\s*' + N))
    if d['rem'] is None: return None
    return d['rem'] + sum(d[k] or 0 for k in ('c351', 'c352', 'c312', 'c028')), d

_maestro = None
def maestro():
    global _maestro
    if _maestro is None:
        ws = abrir(MAESTRO, read_only=True, data_only=True)['MAESTRO']
        _maestro = {str(r[0]).strip(): (r[4], r[5]) for r in ws.iter_rows(min_row=3, values_only=True) if r[0]}
    return _maestro


# ------------------------------------------------------------------ EERR
class EERR:
    def __init__(self, f):
        self.f = f
        self.wbf = abrir(f); self.wbv = abrir(f, data_only=True)
        self.wf, self.wv = self.wbf['EERR'], self.wbv['EERR']
        hdr = next(r for r in range(1, 15) if norm(self.wf.cell(r, 1).value) == 'concepto')
        self.hdr = hdr
        self.col = {}
        for c in range(2, 40):
            d = self.wv.cell(hdr, c).value
            if hasattr(d, 'year'): self.col[f"{d.year}{d.month:02d}"] = c
        self.meses = sorted(self.col)
        self.filas = [(r, norm(self.wf.cell(r, 1).value)) for r in range(hdr + 1, self.wf.max_row + 1) if self.wf.cell(r, 1).value]
        L = [l for _, l in self.filas]
        def fila(lab, despues=None, antes=None):
            for r, l in self.filas:
                if l == lab and (despues is None or r > despues) and (antes is None or r < antes): return r
        self.r_costo, self.r_gastos = fila('costo'), fila('gastos')
        self.r_ing = fila('ingresos')
        self.r_cmv = fila('cmv', self.r_costo, self.r_gastos)
        self.r_tcosto = fila('total', self.r_costo, self.r_gastos)
        self.r_cmvf = fila('cmv', self.r_tcosto or self.r_cmv, self.r_gastos)
        self.r_ei, self.r_ef = fila('existencia inicial'), fila('existencia final')
        self.r_tgastos = fila('total', self.r_gastos)
        self.r_amort = fila('amortizacion')
        self.gastos = [(r, l) for r, l in self.filas if self.r_gastos and self.r_tgastos and self.r_gastos < r < self.r_tgastos]
        self.ix = {}
        for r, l in self.filas: self.ix.setdefault(l, r)
        self.ix['cmv'] = self.r_cmv
        cuit = None
        for r in range(1, hdr):
            for c in range(1, 10):
                v = re.sub(r'\D', '', str(self.wv.cell(r, c).value or ''))
                if len(v) == 11: cuit = v
        self.cuit = cuit
        # denominación: primera celda con texto a la derecha de "DENOMINACIÓN DE LA SOCIEDAD:"
        self.denominacion = ''
        for r in range(1, hdr):
            for c in range(1, 10):
                if norm(self.wv.cell(r, c).value).startswith('denominacion'):
                    self.denominacion = next((nfc(self.wv.cell(r, k).value) for k in range(c + 1, 15) if self.wv.cell(r, k).value), '')

    def v(self, fila, per):
        if fila is None or per not in self.col: return 0.0
        x = self.wv.cell(fila, self.col[per]).value
        return float(x) if num(x) else 0.0
    def val(self, lab, per): return self.v(self.ix.get(lab), per)
    def ref(self, fila, per=None):
        if fila is None: return None
        c = self.col.get(per, 1) if per else 1
        return f"'EERR'!{CL(c)}{fila}"
    def cmv_ef(self, per):
        # CMV del mes: el de la fila final si tiene dato, si no Total costo, si no compras CMV
        for r in (self.r_cmvf, self.r_tcosto, self.r_cmv):
            x = self.v(r, per)
            if x: return x
        return 0.0


class Ctx:
    def __init__(self, f, hasta):
        self.f, self.hasta = f, hasta
        self.e = EERR(f)
        self.carpeta = f"{U}/{carpetas[self.e.cuit]}" if self.e.cuit in carpetas else None
        self.meses = [p for p in self.e.meses if p <= hasta]
        self.h = []; self.scope = set(); self.log = []
        self._papel, self._conc, self.papel_avisado = None, {}, False
        self.idx = collections.defaultdict(list)
        if self.carpeta:
            for sub in ('01-Impuestos Mensuales', '02-Sueldos y F.931'):
                for raiz, _, fs in os.walk(f"{self.carpeta}/{sub}"):
                    for x in fs: self.idx[nfc(x)].append(os.path.join(raiz, x))

    def buscar(self, patron, per=None):
        rx = re.compile(patron, re.I); out = []
        for n, ps in self.idx.items():
            if rx.match(n): out += ps
        if per: out.sort(key=lambda p: (per not in p.replace(os.path.basename(p), ''), len(p)))
        return out

    def en_mes(self, per, patron):
        # archivos de la carpeta del mes (01-Impuestos Mensuales/AAAAMM o .../2026/AAAAMM, con sus subcarpetas IVA, ncr)
        rx = re.compile(patron, re.I)
        return sorted(p for n, ps in self.idx.items() if rx.search(n) for p in ps
                      if per in os.path.relpath(os.path.dirname(p), self.carpeta).split(os.sep))

    def add(self, control, mes, concepto, nivel, desc, valor=None, control_val=None, dif=None, celda=None, fuente=None,
            comparar=None, es_pct=False):
        if dif is None and num(valor) and num(control_val): dif = valor - control_val
        self.h.append(dict(control=control, mes=mm(mes) if mes and mes != '-' else '-', concepto=concepto, nivel=nivel, desc=desc,
                           valor=valor, control_val=control_val, dif=dif, celda=celda, fuente=fuente, es_pct=es_pct,
                           comparar=comparar if comparar is not None else (dif if num(dif) else valor)))
    def ok(self, control, mes='*'): self.scope.add((control, mm(mes) if mes not in ('*', '-') else mes))


# ------------------------------------------------------------------ controles
def c01(x):
    e = x.e
    cols = e.meses
    seguidos = len(cols) == 12 and all(
        (int(b[:4]) * 12 + int(b[4:])) - (int(a[:4]) * 12 + int(a[4:])) == 1 for a, b in zip(cols, cols[1:]))
    if not seguidos:
        x.add(1, '-', 'Columnas de meses', 'Rojo', f"Las columnas no son 12 meses seguidos ({', '.join(mm(p) for p in cols)}).",
              celda=f"'EERR'!B{e.hdr}")
    m = re.search(r'EERR (\d{2})-(\d{4}) a (\d{2})-(\d{4})', nfc(os.path.basename(x.f)))
    if m and cols and (f"{m[2]}{m[1]}" != cols[0] or f"{m[4]}{m[3]}" != cols[-1]):
        x.add(1, '-', 'Nombre del archivo', 'Rojo', f"El nombre dice {m[1]}-{m[2]} a {m[3]}-{m[4]} pero las columnas van de {mm(cols[0])} a {mm(cols[-1])}.",
              celda=f"'EERR'!B{e.hdr}")
    if not e.cuit:
        x.add(1, '-', 'CUIT', 'Rojo', 'No figura el CUIT de la sociedad en el encabezado del EERR.', celda="'EERR'!A1")
    elif not x.carpeta:
        x.add(1, '-', 'Carpeta de la sociedad', 'Rojo', f"No se encontró en la unidad la carpeta del CUIT {e.cuit}: no se pueden cruzar las fuentes.")
    for p in cols:
        if p <= x.hasta: continue
        con = [l for r, l in e.filas if r > e.hdr and e.v(r, p) and l not in ('total',)]
        if con:
            x.add(1, p, 'Mes no cerrado con datos', 'Amarillo', f"{mm(p)} todavía no está cerrado (DDJJ de IVA sin vencer) y ya tiene datos en {len(con)} filas.",
                  celda=e.ref(e.r_ing, p), comparar=len(con))
        x.ok(1, p)
    x.ok(1)


DUP = ['ventas exentas y no gravadas', 'ventas gravadas', 'cmv', 'sueldos y cargas sociales', 'iibb', 'seguridad e higiene', 'iva prorrateo']
def c02(x):
    e = x.e
    hay_sueldos = any(e.val('sueldos y cargas sociales', p) for p in e.meses) or bool(x.buscar(r'^\d{6}-F931-'))
    for p in x.meses:
        if e.val('ingresos', p) <= 0:
            x.add(2, p, 'Ventas', 'Rojo', f"{mm(p)} no tiene ventas cargadas.", 0, celda=e.ref(e.r_ing, p))
        if e.v(e.r_cmv, p) <= 0:
            x.add(2, p, 'CMV', 'Rojo', f"{mm(p)} no tiene compras de mercadería (CMV) cargadas.", 0, celda=e.ref(e.r_cmv, p))
        if hay_sueldos and e.val('sueldos y cargas sociales', p) <= 0:
            x.add(2, p, 'Sueldos', 'Rojo', f"{mm(p)} no tiene sueldos cargados.", 0, celda=e.ref(e.ix.get('sueldos y cargas sociales'), p))
        x.ok(2, p)
    for lab in DUP:
        r = e.ix.get(lab); vistos = collections.defaultdict(list)
        for p in x.meses:
            v = round(e.v(r, p), 2)
            if abs(v) >= 1000: vistos[v].append(p)
        for v, ps in vistos.items():
            if len(ps) > 1:
                for p in ps[1:]:
                    x.add(2, p, f"{nfc(e.wf.cell(r, 1).value)} repetido", 'Rojo',
                          f"El mismo importe ({pesos(v)}) aparece en {', '.join(mm(q) for q in ps)}: posible copia de otro mes.",
                          v, celda=e.ref(r, p))
    x.ok(2)


def c03(x):
    e, wf, wv = x.e, x.e.wf, x.e.wv
    for (r, l) in e.filas:
        for p, c in e.col.items():
            v = wv.cell(r, c).value
            if isinstance(v, str) and v.startswith('#'):
                x.add(3, p, f"Error {CL(c)}{r}", 'Rojo', f"La celda {CL(c)}{r} ({nfc(wf.cell(r, 1).value)}) da {v}.", celda=e.ref(r, p))
    for r, nombre in ((e.r_ing, 'Ingresos'), (e.r_tcosto, 'Total costo'), (e.r_tgastos, 'TOTAL de gastos'),
                      (e.ix.get('rdo antes de impuesto'), 'Rdo antes de Impuesto'), (e.ix.get('rdo del ejercicio'), 'Rdo del Ejercicio')):
        if r is None:
            x.add(3, '-', f"Falta fila: {nombre}", 'Rojo', f"El EERR no tiene la fila {nombre}."); continue
        sin = [p for p, c in e.col.items() if not str(wf.cell(r, c).value or '').startswith('=')]
        if sin:
            x.add(3, '-', f"Sin fórmula: {nombre}", 'Rojo',
                  f"La fila {nombre} no tiene fórmula en {len(sin)} de {len(e.col)} meses" +
                  (" (no se calcula el resultado)." if 'Rdo' in nombre else f" ({', '.join(mm(p) for p in sin[:4])}{'…' if len(sin) > 4 else ''})."),
                  comparar=len(sin), celda=e.ref(r, sin[0]))
    if e.r_tgastos and e.gastos:
        c0 = e.col[e.meses[0]]
        m = re.match(r'^=\s*SUM\(\s*[A-Z]+(\d+)\s*:\s*[A-Z]+(\d+)\s*\)$', str(wf.cell(e.r_tgastos, c0).value or ''), re.I)
        a, b = e.gastos[0][0], e.gastos[-1][0]
        if m and (int(m[1]) > a or int(m[2]) < b):
            x.add(3, '-', 'Rango del TOTAL de gastos', 'Rojo',
                  f"El TOTAL suma las filas {m[1]} a {m[2]} pero los gastos van de la fila {a} ({nfc(wf.cell(a, 1).value)}) a la {b} ({nfc(wf.cell(b, 1).value)}).",
                  celda=e.ref(e.r_tgastos, e.meses[0]))
        for p in x.meses:
            calc = sum(e.v(r, p) for r, _ in e.gastos); tot = e.v(e.r_tgastos, p)
            if abs(calc - tot) > 1:
                x.add(3, p, 'TOTAL de gastos', 'Rojo', f"El TOTAL de gastos guardado ({pesos(tot)}) no coincide con la suma de las filas ({pesos(calc)}).",
                      tot, calc, celda=e.ref(e.r_tgastos, p))
            x.ok(3, p)
    x.ok(3)


def resultado(e, p):
    tg = e.v(e.r_tgastos, p) or sum(e.v(r, p) for r, _ in e.gastos)
    return e.val('ingresos', p) - e.cmv_ef(p) - tg + e.val('rdos financieros (fci)', p) - e.val('impuesto a las ganancias', p)


def c04(x):
    e = x.e; celda_r = e.ix.get('rdo del ejercicio') or e.r_tgastos
    racha = []
    def cerrar():
        if len(racha) >= CFG['perdidas_seguidas']:
            for p in racha:
                x.add(4, p, 'Pérdida', 'Amarillo', f"Pérdida de {pesos(-resultado(e, p))} en {mm(p)}; hay {len(racha)} meses seguidos en pérdida ({mm(racha[0])} a {mm(racha[-1])}).",
                      resultado(e, p), celda=e.ref(celda_r, p), comparar=round(resultado(e, p)))
    for p in x.meses:
        ing = e.val('ingresos', p)
        if ing <= 0: continue
        r = resultado(e, p)
        if r < 0: racha.append(p)
        else: cerrar(); racha = []
        if r / ing > CFG['resultado_max_pct']:
            x.add(4, p, 'Utilidad atípica', 'Amarillo', f"El resultado de {mm(p)} es {pct(r / ing)} de las ventas (más de {pct(CFG['resultado_max_pct'])}): revisar costos o gastos sin cargar.",
                  round(r / ing, 4), CFG['resultado_max_pct'], es_pct=True, celda=e.ref(celda_r, p))
        x.ok(4, p)
    cerrar(); x.ok(4)


def margenes(e, meses):
    return {p: (e.val('ingresos', p) - e.cmv_ef(p)) / e.val('ingresos', p) for p in meses if e.val('ingresos', p) > 0}


def c05(x):
    e = x.e; mb = margenes(e, x.meses)
    if not mb: return x.ok(5)
    med = statistics.median(mb.values()); lo, hi = CFG['margen_banda']; d = CFG['margen_desvio_pts']
    sin_exist = not any(e.v(e.r_ei, p) for p in e.meses)
    prev = None
    for p in x.meses:
        if p not in mb: continue
        m, motivos = mb[p], []
        if not lo <= m <= hi: motivos.append(f"fuera de la banda esperada {pct(lo)}–{pct(hi)}")
        if prev is not None and abs(m - mb[prev]) > d: motivos.append(f"{'+' if m > mb[prev] else ''}{(m - mb[prev]) * 100:.1f} pts vs {mm(prev)} ({pct(mb[prev])})".replace('.', ','))
        if abs(m - med) > d: motivos.append(f"{'+' if m > med else ''}{(m - med) * 100:.1f} pts vs la mediana del ejercicio ({pct(med)})".replace('.', ','))
        if motivos:
            x.add(5, p, 'Margen bruto', 'Rojo' if m < 0 or m > 0.6 else 'Amarillo',
                  f"Margen bruto de {mm(p)}: {pct(m)}; " + '; '.join(motivos) + '.' +
                  (' Sin existencias cargadas el CMV es igual a las compras del mes.' if sin_exist else ''),
                  round(m, 4), round(med, 4), es_pct=True, celda=e.ref(e.r_cmv, p), comparar=round(m, 3))
        prev = p; x.ok(5, p)
    x.ok(5)


NO_C06 = {'iibb', 'iva prorrateo', 'sueldos y cargas sociales', 'amortizacion'}
def c06(x):
    e = x.e
    for r, l in e.gastos:
        if l in NO_C06: continue
        rat = {p: e.v(r, p) / e.val('ingresos', p) for p in x.meses if e.v(r, p) and e.val('ingresos', p) > 0}
        if len(rat) < 3: continue
        meses = sorted(rat)
        for prev, p in zip(meses, meses[1:]):
            # salto: se aparta de la mediana del resto del ejercicio Y del mes anterior (un cambio de nivel se marca una sola vez)
            q = rat[p]; med = statistics.median([v for k, v in rat.items() if k != p]); ing = e.val('ingresos', p); esp = med * ing; v = e.v(r, p)
            salta = lambda ref: ref > 0 and abs(q - ref) / ref > CFG['gasto_salto_pct']
            mat = max(CFG['materialidad_min'], CFG['materialidad_pct_ventas'] * ing)
            if salta(med) and salta(rat[prev]) and abs(v - esp) > mat and abs(v - rat[prev] * ing) > mat:
                x.add(6, p, nfc(e.wf.cell(r, 1).value), 'Amarillo',
                      f"{nfc(e.wf.cell(r, 1).value)} es {pct(q)} de las ventas de {mm(p)}; lo habitual es {pct(med)} (mediana del resto del ejercicio): "
                      f"{'sube' if q > med else 'baja'} {abs(q - med) / med * 100:.0f}% (en {mm(prev)} fue {pct(rat[prev])}).", v, round(esp, 2), celda=e.ref(r, p))
    for p in x.meses: x.ok(6, p)
    x.ok(6)


NO_PAPEL = re.compile(r'no va|no usar|agente|automatico')     # papeles que el equipo marcó para no usar y los generados

def total_cruze(ws):
    # fila TOTAL DIFERENCIAS, columna "Dif. Total": lo que el cruce no explica (sin las NC del controlador fiscal)
    col = None
    for r in ws.iter_rows(values_only=True):
        if not r: continue
        if col is None: col = next((i for i, v in enumerate(r) if nfc(v) == 'Dif. Total'), None)
        if nfc(r[0]).upper().startswith('TOTAL DIFERENCIAS'):
            v = r[col if col is not None else 13] if len(r) > (col if col is not None else 13) else None
            return float(v) if num(v) else None
    return None


def cruce_zetti(x, p):
    # Devuelve (diferencia Zetti - Portal IVA sin explicar, archivo, dónde está) o None si no se corrió el cruce.
    # 1) hoja "cruze" del papel del mes (la deja la skill cruze-zetti-portaliva-ventas en AAAAMM-papeldetrabajo-X.xlsx),
    # 2) cruce de 3 fuentes (AAAAMM - CRUCE VENTAS MC PORTAL ZETTI.xlsx), 3) archivo de cruce suelto (versiones anteriores)
    for f in x.en_mes(p, r'papel.*\.xlsx$'):
        if NO_PAPEL.search(norm(os.path.basename(f))) or os.path.basename(f).startswith('~$'): continue
        wb = abrir(f, read_only=True, data_only=True)
        h = next((n for n in wb.sheetnames if n.strip().lower() == 'cruze'), None)
        if h: return total_cruze(wb[h]), f, f'hoja "{h}" de {os.path.basename(f)}'
    for f in x.en_mes(p, rf'^{p} - CRUCE VENTAS MC PORTAL ZETTI\.xlsx$'):
        wb = abrir(f, read_only=True, data_only=True)
        if 'Resumen' not in wb.sheetnames: continue
        for r in wb['Resumen'].iter_rows(values_only=True):
            if r and nfc(r[0]) == 'Zetti - Portal IVA':
                v = [c for c in r if num(c)]
                return (v[-1] if v else None), f, os.path.basename(f)
    for f in x.en_mes(p, rf'^{p} - CRUCE (VENTAS )?ZETTI VS PORTAL IVA\.xlsx$'):
        return total_cruze(abrir(f, read_only=True, data_only=True).worksheets[0]), f, os.path.basename(f)
    return None


def c07(x):
    e = x.e
    for p in x.meses:
        f = next(iter(x.buscar(rf'^{p} - PORTAL IVA - VENTAS\.xlsx$', p)), None) or next(iter(x.buscar(rf'^{p} - PORTAL IVA - VENTAS.*\.xlsx$', p)), None)
        tot = None
        if not f:
            x.add(7, p, 'Libro de ventas', 'Rojo', f"No está el libro {p} - PORTAL IVA - VENTAS.xlsx en la carpeta del mes: no se pueden controlar las ventas.",
                  fuente=x.carpeta and f"{x.carpeta}/01-Impuestos Mensuales")
        else:
            try: gr, ex = ventas_libro(f)
            except Exception as err: x.log.append(f"7 {mm(p)}: no se pudo leer {os.path.basename(f)} ({err})"); continue
            tot = gr + ex
            for lab, lib in (('ventas gravadas', gr), ('ventas exentas y no gravadas', ex)):
                r = e.ix.get(lab); v = e.v(r, p)
                if abs(v - lib) > 1:
                    x.add(7, p, nfc(e.wf.cell(r, 1).value) if r else lab, 'Rojo' if tot and abs(v - lib) / tot > CFG['ventas_tol_pct'] else 'Amarillo',
                          f"{nfc(e.wf.cell(r, 1).value) if r else lab} de {mm(p)} no coincide con el Portal IVA ({os.path.basename(f)}).",
                          v, round(lib, 2), celda=e.ref(r, p), fuente=f)
        try: cru = cruce_zetti(x, p)
        except Exception as err: x.log.append(f"7 {mm(p)}: no se pudo leer el cruce Zetti ({err})"); continue
        if cru:
            dif, fc, donde = cru
            if dif is None:
                x.log.append(f"7 {mm(p)}: el cruce ({donde}) no tiene la fila TOTAL DIFERENCIAS con la columna Dif. Total")
            elif abs(dif) > 1000:
                x.add(7, p, 'Zetti vs Portal IVA', 'Rojo' if tot and abs(dif) / tot > CFG['ventas_tol_pct'] else 'Amarillo',
                      f"Las ventas de Zetti difieren {pesos(dif)} del Portal IVA en {mm(p)} y el cruce no lo explica (ver el detalle en la {donde}).",
                      round(tot + dif, 2) if tot else None, round(tot, 2) if tot else None, dif, fuente=fc, comparar=round(dif, 2))
        elif x.en_mes(p, rf'^{p} - ZETTI - VENTAS\.xlsx$'):
            x.add(7, p, 'Zetti vs Portal IVA', 'Amarillo', f"Está el subdiario de Zetti de {mm(p)} pero no se corrió el cruce con el Portal IVA "
                  "(no hay hoja \"cruze\" en el papel del mes).", fuente=x.en_mes(p, rf'^{p} - ZETTI - VENTAS\.xlsx$')[0])
        x.ok(7, p)
    x.ok(7)


def c08(x):
    e = x.e
    if e.r_ei and not any(e.v(e.r_ei, p) for p in e.meses):
        x.add(8, '-', 'Existencia inicial', 'Rojo', "No está cargada la existencia inicial: el CMV queda igual a las compras y el margen bruto no es confiable.",
              celda=e.ref(e.r_ei, e.meses[0] if e.meses else None))
    if e.r_ef and e.meses and e.meses[-1] <= x.hasta and not any(e.v(e.r_ef, p) for p in e.meses):
        x.add(8, '-', 'Existencia final', 'Rojo', f"El ejercicio cerró en {mm(e.meses[-1])} y no está cargada la existencia final.",
              celda=e.ref(e.r_ef, e.meses[-1]))
    x.ok(8)


def conciliacion(x, p):
    # AAAAMM-COMPRAS-SIN-NCR-*.xlsx de conciliacion-ncr: NCR eliminadas (en positivo), NC comerciales sin respaldo,
    # cantidad de NCR no encontradas y {(CUIT, PV, número): importe} de cada NCR. None = no hay; False = no se pudo leer.
    if p in x._conc: return x._conc[p]
    f = next(iter(x.buscar(rf'^{p}-COMPRAS-SIN-NCR-.*\.xlsx$', p)), None); out = None
    if f:
        try:
            wb = abrir(f, read_only=True, data_only=True)
            def hoja(h):
                filas = list(wb[h].iter_rows(values_only=True)) if h in wb.sheetnames else []
                return (filas[0], [r for r in filas[1:] if any(v is not None for v in r)]) if filas else ((), [])
            def col(h, nombre):
                cab, filas = hoja(h); i = next((k for k, v in enumerate(cab) if nfc(v) == nombre), None)
                return [r[i] for r in filas if i is not None and num(r[i])]
            cab, el = hoja('Eliminados'); i = {nfc(v): k for k, v in enumerate(cab)}
            ent = lambda v: int(float(v)) if re.match(r'^\s*\d+(\.0+)?\s*$', str(v if v is not None else '')) else None
            claves = {(ndoc(r[i['_cuit']]), ent(r[i['_pv']]), ent(r[i['_nro']])): float(r[i['_imp']]) for r in el
                      if {'_cuit', '_pv', '_nro', '_imp'} <= set(i) and num(r[i['_imp']])}
            out = dict(f=f, ncr=-sum(col('Eliminados', '_imp')), nccom=-sum(col('NC sin respaldo', 'Importe Total')),
                       no_enc=len(hoja('NCR no encontradas')[1]), claves=claves)
        except Exception as err:
            x.log.append(f"{mm(p)}: no se pudo leer {os.path.basename(f)} ({err})"); out = False
    x._conc[p] = out
    return out


def c09(x):
    e = x.e; datos, sin = {}, []
    for p in x.meses:
        cn = conciliacion(x, p)
        if cn is None: sin.append(p); x.ok(9, p); continue
        if cn is False: continue
        ncr, nccom, no_enc, f = cn['ncr'], cn['nccom'], cn['no_enc'], cn['f']
        datos[p] = (ncr, nccom, f)
        if no_enc:
            x.add(9, p, 'NCR no encontradas', 'Amarillo', f"{no_enc} NCR del SharePoint no se encontraron en el Portal IVA de compras de {mm(p)}.",
                  no_enc, fuente=f)
        x.ok(9, p)
    if sin:
        x.add(9, '-', 'Sin conciliación de NCR', 'Amarillo', f"No hay conciliación de NCR (AAAAMM-COMPRAS-SIN-NCR) en {len(sin)} meses: "
              f"{', '.join(mm(p) for p in sin)}. Sin ella no se puede comparar NCR contra NC comerciales.", len(sin), comparar=len(sin))
    rat = {p: n / (n + c) for p, (n, c, _) in datos.items() if n + c > 0}
    if len(rat) >= 3:
        med = statistics.median(rat.values())
        for p, q in rat.items():
            if med > 0 and abs(q - med) / med > CFG['ncr_desvio_pct']:
                n, c, f = datos[p]
                x.add(9, p, 'NCR vs NC comerciales', 'Amarillo',
                      f"En {mm(p)} las NCR ({pesos(n)}) son {pct(q)} del total de notas de crédito de droguerías; lo habitual es {pct(med)}. NC comerciales: {pesos(c)}.",
                      round(q, 4), round(med, 4), es_pct=True, fuente=f)
    elif datos:
        x.log.append(f"9: solo {len(datos)} meses con conciliación de NCR, no alcanza para comparar contra el histórico")
    x.ok(9)


def compras_imputadas(e):
    ws = e.wbv['Compras imputadas']; filas = list(ws.iter_rows(values_only=True))
    h = filas[0]; cols = {}
    for i, v in enumerate(h):
        if isinstance(v, str) and re.match(r'^\d{2}/\d{4}$', v): cols[f"{v[3:]}{v[:2]}"] = i
    out = []
    for r in filas[1:]:
        if not re.match(r'^\d{11}$', str(r[0] or '').strip()): continue      # fila TOTAL y renglones vacíos
        out.append(dict(cuit=str(r[0] or '').strip(), prov=nfc(r[1]), cuenta=nfc(r[2]), rubro=nfc(r[3]), estado=nfc(r[4]),
                        m={p: float(r[i]) for p, i in cols.items() if num(r[i])}))
    return out, cols


def c10(x):
    e = x.e
    if 'Compras imputadas' not in e.wbv.sheetnames:
        x.add(10, '-', 'Hoja Compras imputadas', 'Rojo', 'El EERR no tiene la hoja "Compras imputadas": no se puede controlar la clasificación.'); return x.ok(10)
    CI, ccol = compras_imputadas(e)
    for p in x.meses:
        f = next(iter(x.buscar(rf'^{p} - PORTAL IVA - COMPRAS\.xlsx$', p)), None) or next(iter(x.buscar(rf'^{p} - PORTAL IVA - COMPRAS.*\.xlsx$', p)), None)
        imp = sum(c['m'].get(p, 0) for c in CI)
        if not f:
            x.add(10, p, 'Libro de compras', 'Rojo', f"No está el libro {p} - PORTAL IVA - COMPRAS.xlsx: no se puede controlar que las compras clasificadas cierren.",
                  round(imp, 2), fuente=x.carpeta and f"{x.carpeta}/01-Impuestos Mensuales")
        else:
            try: lib = compras_libro(f)
            except Exception as err: x.log.append(f"10 {mm(p)}: no se pudo leer {os.path.basename(f)} ({err})"); continue
            if abs(imp - lib) > CFG['compras_tol']:
                x.add(10, p, 'Compras clasificadas vs Portal IVA', 'Rojo',
                      f"Las compras clasificadas por cuenta de {mm(p)} suman {pesos(imp)} y el Portal IVA de compras {pesos(lib)} (sin IVA; B/C por total menos percepciones). "
                      "La diferencia son compras sin clasificar, duplicadas o un libro distinto al usado para el EERR.",
                      round(imp, 2), round(lib, 2), fuente=f, celda=f"'Compras imputadas'!{CL(ccol[p] + 1)}1" if p in ccol else None)
        x.ok(10, p)
    # cada rubro tiene que llegar a una fila del EERR, y la fila tiene que valer lo mismo que sus compras
    por_rubro = collections.defaultdict(lambda: collections.defaultdict(float))
    for c in CI:
        for p, v in c['m'].items(): por_rubro[norm(c['rubro'])][p] += v
    for rub, ms in por_rubro.items():
        tot = sum(v for p, v in ms.items() if p in x.meses)
        if rub.startswith('excluido') or abs(tot) < 1: continue
        r = e.ix.get(rub)
        if r is None:
            x.add(10, '-', f"Rubro sin fila: {rub.upper()}", 'Rojo', f"Hay compras imputadas a \"{rub.upper()}\" ({pesos(tot)} en el ejercicio) que no tienen fila en el EERR: no llegan al resultado.",
                  round(tot, 2)); continue
        for p in x.meses:
            v, c = e.v(r, p), ms.get(p, 0)
            falla = (v < c - 1) if rub == 'publicidad' else abs(v - c) > 1
            if falla:
                x.add(10, p, f"Fila {nfc(e.wf.cell(r, 1).value)}", 'Rojo',
                      f"La fila {nfc(e.wf.cell(r, 1).value)} de {mm(p)} no coincide con lo imputado a ese rubro en \"Compras imputadas\" (¿se pisó la fórmula?).",
                      v, round(c, 2), celda=e.ref(r, p))
    M = maestro()
    for c in CI:
        tot = sum(v for p, v in c['m'].items() if p in x.meses)
        if abs(tot) < 1: continue
        est = norm(c['estado'])
        if not (est.startswith('clasificado') or est.startswith('propuesta')):
            x.add(10, '-', f"Proveedor sin clasificar: {c['cuit']}", 'Amarillo',
                  f"{c['prov']} ({c['cuit']}) está \"{c['estado'] or 'sin estado'}\" en el Maestro y se imputó a {c['rubro'] or 'ningún rubro'} ({pesos(tot)}).",
                  round(tot, 2), fuente=MAESTRO)
        actual = M.get(c['cuit'])
        if actual and actual[0] and c['cuenta'] and norm(actual[0]) != norm(c['cuenta']):
            x.add(10, '-', f"Clasificación cambiada: {c['cuit']}", 'Amarillo',
                  f"{c['prov']} ({c['cuit']}): en el EERR está imputado a {c['cuenta']}, pero el Maestro ahora dice {actual[0]} ({pesos(tot)}). Regenerar el EERR.",
                  round(tot, 2), fuente=MAESTRO)
    gv = e.ix.get('gastos varios')
    if gv and x.meses:
        # en el ejercicio, no mes a mes: GASTOS VARIOS sobre el total de gastos que vienen de compras (sin CMV ni excluidos)
        base = sum(abs(v) for rub, ms in por_rubro.items() if rub != 'cmv' and not rub.startswith('excluido') for p, v in ms.items() if p in x.meses)
        v = sum(e.v(gv, p) for p in x.meses)
        if base > 0 and v / base > CFG['gastos_varios_max_pct'] and v > CFG['materialidad_min']:
            x.add(10, '-', 'GASTOS VARIOS', 'Amarillo', f"GASTOS VARIOS suma {pesos(v)} en el ejercicio: {pct(v / base)} de los gastos que vienen de compras "
                  f"(más de {pct(CFG['gastos_varios_max_pct'])}). Revisar qué proveedores conviene reclasificar en el Maestro.",
                  round(v, 2), None, celda=e.ref(gv, x.meses[-1]), comparar=round(v / base, 3))
    x.ok(10)


def c11(x):
    e = x.e; r = e.ix.get('sueldos y cargas sociales')
    hay = any(e.v(r, p) for p in e.meses) or bool(x.buscar(r'^\d{6}-F931-'))
    if not hay: return x.ok(11)
    for p in x.meses:
        fs = x.buscar(rf'^{p}-F931-.*\.pdf$', p)
        if not fs:
            x.add(11, p, 'F.931', 'Rojo', f"No está el F.931 de {mm(p)} en 02-Sueldos y F.931: no se pueden controlar los sueldos.",
                  e.v(r, p), fuente=f"{x.carpeta}/02-Sueldos y F.931" if x.carpeta else None); x.ok(11, p); continue
        leidos = []
        for f in fs:
            try:
                d = f931(f)
                if d: leidos.append((d[1]['rect'], d[0], f))
            except Exception as err: x.log.append(f"11 {mm(p)}: no se pudo leer {os.path.basename(f)} ({err})")
        if not leidos: continue
        _, esp, f = max(leidos)
        v = e.v(r, p)
        if abs(v - esp) > 1:
            x.add(11, p, 'Sueldos y Cargas Sociales', 'Rojo',
                  f"Sueldos y cargas sociales de {mm(p)} no coincide con el F.931 (remuneraciones + contribuciones 351, 352, 312 y 028).",
                  v, round(esp, 2), celda=e.ref(r, p), fuente=f)
        x.ok(11, p)
    rat = {p: e.v(r, p) / e.val('ingresos', p) for p in x.meses if e.v(r, p) and e.val('ingresos', p) > 0 and p[4:] not in ('06', '12')}
    if len(rat) >= 3:
        med = statistics.median(rat.values())
        for p, q in rat.items():
            if med > 0 and abs(q - med) / med > CFG['gasto_salto_pct']:
                x.add(11, p, 'Costo laboral / ventas', 'Amarillo',
                      f"El costo laboral es {pct(q)} de las ventas en {mm(p)}; lo habitual es {pct(med)} (sin contar junio y diciembre, que llevan aguinaldo).",
                      round(q, 4), round(med, 4), es_pct=True, celda=e.ref(r, p))
    x.ok(11)


def c12(x):
    e = x.e; r = e.ix.get('iibb'); sin = []
    for p in x.meses:
        fs = x.buscar(rf'^{p}-IIBB-DJ-.*\.pdf$', p)
        if not fs and e.cuit and os.path.exists(f"{ARBA_LOCAL}/{e.cuit}/{p}.pdf"): fs = [f"{ARBA_LOCAL}/{e.cuit}/{p}.pdf"]
        v = e.v(r, p)
        if not fs: sin.append(p)
        else:
            try: dj = dj_arba(fs[0])
            except Exception as err: dj = None; x.log.append(f"12 {mm(p)}: no se pudo leer {os.path.basename(fs[0])} ({err})")
            if dj is not None and abs(v - dj) > 1:
                x.add(12, p, 'IIBB', 'Rojo', f"IIBB de {mm(p)} no coincide con el impuesto determinado de la DJ de anticipo de ARBA.",
                      v, dj, celda=e.ref(r, p), fuente=fs[0])
        ing = e.val('ingresos', p)
        if ing > 0 and v:
            lo, hi = CFG['iibb_alic']
            if not lo <= v / ing <= hi:
                x.add(12, p, 'IIBB / ventas', 'Amarillo', f"IIBB es {pct(v / ing)} de las ventas de {mm(p)}; lo razonable es entre {pct(lo)} y {pct(hi)}.",
                      round(v / ing, 4), None, es_pct=True, celda=e.ref(r, p))
        elif ing > 0 and not v:
            x.add(12, p, 'IIBB', 'Rojo', f"{mm(p)} tiene ventas pero IIBB está en cero.", 0, celda=e.ref(r, p))
        x.ok(12, p)
    if sin:
        x.add(12, '-', 'Sin DJ de ARBA', 'Amarillo', f"No hay DJ de anticipo de ARBA archivada en {len(sin)} meses ({', '.join(mm(p) for p in sin)}): "
              "si es AGIP o Convenio Multilateral controlar contra esa DJ; si no, falta bajarla.", len(sin), comparar=len(sin))
    x.ok(12)


def c13(x):
    e = x.e; r = e.ix.get('iva prorrateo')
    if 'IVA prorrateo' not in e.wbv.sheetnames or r is None:
        x.add(13, '-', 'Hoja IVA prorrateo', 'Amarillo', 'No está la hoja de control "IVA prorrateo": no se puede controlar la fila IVA Prorrateo.'); return x.ok(13)
    ws = e.wbv['IVA prorrateo']; fila = {}
    for i in range(1, ws.max_row + 1):
        a = str(ws.cell(i, 1).value or '')
        if re.match(r'^\d{2}/\d{4}$', a): fila[f"{a[3:]}{a[:2]}"] = i
    for p in x.meses:
        if p not in fila: continue
        ctl = ws.cell(fila[p], 12).value; ctl = float(ctl) if num(ctl) else None
        v = e.v(r, p)
        if ctl is None: continue
        if ctl > 1 and not v:
            x.add(13, p, 'IVA Prorrateo', 'Rojo', f"IVA Prorrateo de {mm(p)} está en cero y el cálculo de control da {pesos(ctl)} (IVA de servicios × % de ventas exentas).",
                  v, round(ctl, 2), celda=e.ref(r, p), fuente=f"'IVA prorrateo'!L{fila[p]}")
        elif abs(v - ctl) > max(CFG['prorrateo_tol_min'], CFG['prorrateo_tol_pct'] * abs(ctl)):
            x.add(13, p, 'IVA Prorrateo', 'Amarillo', f"IVA Prorrateo de {mm(p)} (tomado del papel de trabajo) difiere del cálculo de control de la hoja \"IVA prorrateo\".",
                  v, round(ctl, 2), celda=e.ref(r, p), fuente=f"'IVA prorrateo'!L{fila[p]}")
        x.ok(13, p)
    x.ok(13)


def c14(x):
    e = x.e; r = e.ix.get('seguridad e higiene')
    if r:
        con = [p for p in x.meses if e.v(r, p)]
        if not con and x.meses:
            x.add(14, '-', 'Seguridad e Higiene', 'Amarillo', "No hay Seguridad e Higiene en todo el ejercicio. Si el municipio no la cobra (p. ej. CABA), aprobar.",
                  celda=e.ref(r, x.meses[0]))
        for p in x.meses:
            gr = e.val('ventas gravadas', p); v = e.v(r, p)
            if con and not v and gr > 0:
                x.add(14, p, 'Seguridad e Higiene', 'Amarillo', f"Falta Seguridad e Higiene en {mm(p)} (está en {len(con)} de {len(x.meses)} meses).", 0, celda=e.ref(r, p))
            elif v and gr > 0 and not CFG['sh_alic'][0] <= v / gr <= CFG['sh_alic'][1]:
                x.add(14, p, 'S&H / ventas gravadas', 'Amarillo', f"Seguridad e Higiene es {pct(v / gr)} de las ventas gravadas de {mm(p)}; lo razonable es entre "
                      f"{pct(CFG['sh_alic'][0])} y {pct(CFG['sh_alic'][1])} (salvo importe fijo del municipio).", round(v / gr, 4), es_pct=True, celda=e.ref(r, p))
            x.ok(14, p)
    m = e.ix.get('matricula colegio farmaceuticos')
    if m and x.meses and not any(e.v(m, p) for p in x.meses):
        x.add(14, '-', 'Matrícula Colegio Farmacéuticos', 'Amarillo', 'No hay matrícula del Colegio de Farmacéuticos en el ejercicio.', celda=e.ref(m, x.meses[0]))
    pub, mq = e.ix.get('publicidad'), e.ix.get('marquesinas')
    if (pub or mq) and x.meses and not any(e.v(pub, p) or e.v(mq, p) for p in x.meses):
        x.add(14, '-', 'Publicidad y marquesinas', 'Amarillo', 'No hay publicidad ni marquesinas en el ejercicio (derechos de publicidad y propaganda del municipio).',
              celda=e.ref(pub or mq, x.meses[0]))
    x.ok(14)


def c15(x):
    e = x.e; n = len(x.meses)
    for lab in CFG['fijos']:
        r = e.ix.get(lab)
        if not r or not n: continue
        con = [p for p in x.meses if e.v(r, p)]
        if not con and lab in CFG['fijos_obligatorios']:
            x.add(15, '-', nfc(e.wf.cell(r, 1).value), 'Amarillo', f"No hay {nfc(e.wf.cell(r, 1).value)} en todo el ejercicio: falta la factura o está clasificada en otra cuenta.",
                  celda=e.ref(r, x.meses[0]))
        elif len(con) >= CFG['fijos_presencia'] * n and len(con) < n:
            for p in x.meses:
                if p not in con:
                    x.add(15, p, nfc(e.wf.cell(r, 1).value), 'Amarillo', f"Falta {nfc(e.wf.cell(r, 1).value)} en {mm(p)} (está en {len(con)} de {n} meses).", 0, celda=e.ref(r, p))
    if 'Compras imputadas' in e.wbv.sheetnames:
        CI, ccol = compras_imputadas(e)
        for c in CI:
            if norm(c['rubro']) not in CFG['fijos']: continue
            ms = {p: v for p, v in c['m'].items() if p in x.meses and v > 0}
            if len(ms) < 3: continue
            for p, v in ms.items():
                med = statistics.median([w for k, w in ms.items() if k != p])
                if v > CFG['duplicado_factor'] * med and v - med > CFG['materialidad_min']:
                    x.add(15, p, f"Posible duplicado: {c['cuit']}", 'Amarillo',
                          f"{c['prov']} ({c['rubro']}) facturó {pesos(v)} en {mm(p)}, {str(round(v / med, 1)).replace('.', ',')} veces lo habitual ({pesos(med)}): posible factura duplicada o doble.",
                          round(v, 2), round(med, 2), celda=f"'Compras imputadas'!{CL(ccol[p] + 1)}1" if p in ccol else None)
    for lab in CFG['financieros']:
        r = e.ix.get(lab)
        if r and n and not any(e.v(r, p) for p in x.meses):
            x.add(15, '-', nfc(e.wf.cell(r, 1).value), 'Amarillo', f"{nfc(e.wf.cell(r, 1).value)} no tiene datos en el ejercicio: confirmar contra extractos bancarios, "
                  "cuadro de marcha de préstamos, planes de facilidades o resúmenes de FCI.", celda=e.ref(r, x.meses[0]))
    for p in x.meses: x.ok(15, p)
    x.ok(15)


# ------------------------------------------------------------------ papel de trabajo (controles 16 a 21)
# El papel de liquidaciones del equipo (01-Impuestos Mensuales/202600-liquidaciones-IVA, IIBB y Tasa Municipal-X.xlsx)
# tiene una hoja por mes (MM-AAAA) con el mismo modelo en todas las sociedades: columna B = etiqueta, C a G = Neto
# gravado, No gravado, Exento, IVA, Total; H = clasificación de cada compra (MERCADERIA / SERVICIOS AL x%); J y K =
# cuadro de prorrateo. Se lee por etiqueta, no por número de fila. No se usa el "- AUTOMATICO" (lo genera el estudio).
def leer_mes(rows):
    d = dict(ventas=[], compras=[], sub={}, pct={}, k={}, det={}, ncr=None, total_nc=None, fila={})
    sec, parte = None, 'cf'
    for i, r in enumerate(rows, 1):
        r = list(r) + [None] * max(0, 12 - len(r)); b = norm(r[1])
        for c in range(1, len(r)):
            if norm(r[c]) in ('ncr eliminadas', 'total nc') and num(r[c - 1]):
                d['ncr' if norm(r[c]) == 'ncr eliminadas' else 'total_nc'] = float(r[c - 1]); d['fila'][norm(r[c])] = f"{CL(c)}{i}"
        if norm(r[9]) and num(r[10]): d['k'][norm(r[9])] = float(r[10])
        if b == 'ventas' and sec is None: sec = 'ventas'; continue
        if b.startswith('prorrateo para credito fiscal'): sec = 'prorrateo'; continue
        if b == 'compras' and sec == 'prorrateo': sec = 'compras'; continue
        if b.startswith('determinacion del saldo de iva'): sec = 'det'; continue
        if b.startswith('ingresos brutos') or b.startswith('retenciones y percepciones'): sec = 'fin'
        if not b or b == 'tipo de comprobante' or b == 'concepto': continue
        vals = [float(v) if num(v) else 0.0 for v in r[2:7]]
        if sec in ('ventas', 'compras'):
            if b.startswith(('subtotal', 'totales', 'ventas netas')):
                d['sub'][(sec, b)] = vals; d['fila'][(sec, b)] = i
                if b == 'subtotal credito fiscal': parte = 'nc'
            else: d[sec].append(dict(lab=nfc(r[1]), b=b, vals=vals, clase=norm(r[7]), parte=parte, fila=i))
        elif sec == 'prorrateo' and b in ('venta gravada', 'venta exenta'):
            d['pct'][b] = float(r[2]) if num(r[2]) else None; d['fila'][('pct', b)] = i
        elif sec == 'det':
            d['det'][b] = float(r[2]) if num(r[2]) else None; d['fila'][('det', b)] = i
    return d


def papel_anual(x):
    # {AAAAMM: datos de la hoja del mes, con 'f' = archivo y 'hoja'} para los meses auditados con datos cargados
    if x._papel is not None: return x._papel
    x._papel = {}
    fs = sorted({q for n, ps in x.idx.items() if 'liquidaciones' in norm(n) and n.lower().endswith(('.xlsx', '.xls')) and
                 not n.startswith('~$') and not NO_PAPEL.search(norm(n)) for q in ps if os.sep + '01-Impuestos Mensuales' in q})
    for f in fs:
        if f.lower().endswith('.xls'): x.log.append(f"papel {os.path.basename(f)}: es .xls, no se puede leer (guardarlo como .xlsx)"); continue
        try: wb = abrir(f, read_only=True, data_only=True)
        except Exception as err: x.log.append(f"papel: no se pudo abrir {os.path.basename(f)} ({err})"); continue
        for n in wb.sheetnames:
            m = re.match(r'^(\d{2})-(\d{4})$', n.strip())
            if not m or f"{m[2]}{m[1]}" not in x.meses: continue
            p = f"{m[2]}{m[1]}"
            if p in x._papel and os.path.basename(x._papel[p]['f']).startswith(p[:4]): continue   # manda el papel de ese año
            d = leer_mes(wb[n].iter_rows(max_col=12, values_only=True))
            if not any(any(r['vals']) for r in d['ventas']): continue                              # mes sin cargar
            if d['det'].get('subtotal debito fiscal') is None:
                x.log.append(f"papel {mm(p)}: la hoja {n} no tiene los valores de las fórmulas (abrirlo y guardarlo con Excel)"); continue
            d.update(f=f, hoja=n); x._papel[p] = d
    return x._papel


def celda_papel(d, clave, col='C'):
    i = d['fila'].get(clave)
    return f"{col}{i}" if isinstance(i, int) else i


DDJJ_RX = dict(df=r'Total del d[ée]bito fiscal del per[ií]odo \$ ?(-?[\d\.]+,\d{2})',
               cf=r'Total del cr[ée]dito fiscal del per[ií]odo \$ ?(-?[\d\.]+,\d{2})')
def ddjj_iva(x, p):
    # débito y crédito fiscal de la DDJJ de IVA presentada: AAAAMM-IVA-DDJJ-*.pdf de la carpeta del mes (el más nuevo,
    # por las rectificativas) o, si no, el JSON de ddjj-iva.js en BVA-salidas/ddjj-iva/<CUIT>/AAAAMM.json
    for f in sorted(x.en_mes(p, rf'^{p}-IVA-DDJJ-.*\.pdf$'), key=os.path.getmtime, reverse=True):
        try: t = pdf_texto(f)
        except Exception as err: x.log.append(f"16 {mm(p)}: no se pudo leer {os.path.basename(f)} ({err})"); continue
        g = {k: re.search(rx, t, re.I) for k, rx in DDJJ_RX.items()}
        if all(g.values()): return {k: float(v[1].replace('.', '').replace(',', '.')) for k, v in g.items()}, f
    j = os.path.join(SALIDAS, 'ddjj-iva', x.e.cuit or '-', f'{p}.json')
    if os.path.exists(j):
        d = json.load(open(j, encoding='utf-8'))
        if num(d.get('df')) and num(d.get('cf')): return dict(df=d['df'], cf=d['cf']), j
    return None, None


def sin_papel(x, control):
    # un solo aviso (en el control 16) si la sociedad no tiene papel de liquidaciones legible; el resto no corre
    if x.papel_avisado: return
    x.papel_avisado = True
    x.add(control, '-', 'Papel de trabajo', 'Amarillo', 'No hay papel de liquidaciones (202600-liquidaciones-IVA, IIBB y Tasa Municipal-X.xlsx, '
          'una hoja MM-AAAA por mes) con meses del ejercicio cargados: no se pueden correr los controles del papel (16 a 21).',
          fuente=x.carpeta and f"{x.carpeta}/01-Impuestos Mensuales")


def c16(x):
    P = papel_anual(x)
    if not P: sin_papel(x, 16); return x.ok(16)
    rat = {}
    for p in x.meses:
        d = P.get(p)
        if not d: continue
        df, cf = d['det'].get('subtotal debito fiscal'), d['det'].get('total credito fiscal')
        dj, fdj = ddjj_iva(x, p)
        if dj:
            for k, nombre, v, et in (('df', 'Débito fiscal', df, 'subtotal debito fiscal'), ('cf', 'Crédito fiscal', cf, 'total credito fiscal')):
                if num(v) and abs(v - dj[k]) > 1:
                    x.add(16, p, f"{nombre} vs DDJJ", 'Rojo', f"El {nombre.lower()} del papel de {mm(p)} (celda {celda_papel(d, ('det', et))} de la hoja "
                          f"{d['hoja']}) no coincide con la DDJJ de IVA presentada ({os.path.basename(fdj)}).", round(v, 2), dj[k], fuente=d['f'])
        if num(df) and num(cf) and df > 0: rat[p] = cf / df
        x.ok(16, p)
    if len(rat) >= 3:
        med = statistics.median(rat.values())
        for p, q in rat.items():
            motivos = []
            if med > 0 and abs(q - med) / med > CFG['cf_df_desvio_pct']: motivos.append(f"lo habitual es {pct(med)} (mediana del ejercicio)")
            if q > 1 and med <= 1: motivos.append('el crédito supera al débito (saldo técnico a favor)')
            if motivos:
                d = P[p]
                x.add(16, p, 'Crédito / débito fiscal', 'Amarillo', f"En {mm(p)} el crédito fiscal es {pct(q)} del débito fiscal; " + '; '.join(motivos) +
                      '. Revisar compras de mercadería sin cargar, duplicadas o mal clasificadas (mercadería / servicios).',
                      round(q, 4), round(med, 4), es_pct=True, fuente=d['f'], comparar=round(q, 3))
    x.ok(16)


def es_nc(r):
    m = re.match(r'^(\d+)', r['b'])
    return 'credito' in r['b'] or bool(m and int(m[1]) in NC_COD)


def c17(x):
    P = papel_anual(x)
    if not P: sin_papel(x, 17); return x.ok(17)
    for p in x.meses:
        d = P.get(p)
        if not d: continue
        for sec in ('ventas', 'compras'):
            for r in d[sec]:
                if not es_nc(r) or not any(v > 0.5 for v in r['vals'][:4]): continue
                aviso = '' if 'credito' in r['b'] else f" (el código {r['b'].split()[0]} es de nota de crédito aunque la fila diga \"{r['lab']}\")"
                x.add(17, p, f"NC en positivo: {r['lab']} ({sec})", 'Amarillo' if aviso else 'Rojo',
                      f"En el papel de {mm(p)} la fila \"{r['lab']}\" de {sec} (fila {r['fila']} de la hoja {d['hoja']}) está en positivo{aviso}: las notas "
                      "de crédito van en negativo; así suman en lugar de restar en el subtotal, el prorrateo y la liquidación.",
                      round(sum(r['vals'][:4]), 2), fuente=d['f'])
        x.ok(17, p)
    x.ok(17)


def c18(x):
    # NRF = NC de recupero de las droguerías (NCR, SharePoint "NC por NRF"). Lo conciliado en COMPRAS-SIN-NCR tiene que
    # estar en las NC de compras del papel y coincidir con lo que el papel informa como "NCR eliminadas".
    P = papel_anual(x)
    if not P: sin_papel(x, 18); return x.ok(18)
    sin_bloque = []
    for p in x.meses:
        d, cn = P.get(p), conciliacion(x, p)
        if not d or not cn: continue
        ncr = cn['ncr']; nc = abs(d['sub'].get(('compras', 'subtotal notas de credito s/compras'), [0] * 5)[4])
        if ncr >= 1:
            if nc < ncr - 1:
                x.add(18, p, 'NCR sin computar', 'Amarillo', f"Las NC de compras del papel de {mm(p)} suman {pesos(nc)} y las NCR conciliadas {pesos(ncr)} "
                      f"({os.path.basename(cn['f'])}): el papel no incluye las NCR en las NC de compras (en los demás papeles se incluyen y se separan "
                      "en \"NCR eliminadas\").", round(nc, 2), round(ncr, 2), fuente=d['f'])
            if d['ncr'] is None: sin_bloque.append(p)
            elif abs(abs(d['ncr']) - ncr) > 1:
                x.add(18, p, 'NCR eliminadas', 'Rojo', f"El papel de {mm(p)} informa NCR eliminadas por {pesos(abs(d['ncr']))} (celda {d['fila']['ncr eliminadas']} "
                      f"de la hoja {d['hoja']}) y la conciliación de NCR da {pesos(ncr)}.", round(abs(d['ncr']), 2), round(ncr, 2), fuente=d['f'])
        x.ok(18, p)
    if sin_bloque:
        x.add(18, '-', 'NCR no identificadas en el papel', 'Amarillo', f"En {len(sin_bloque)} meses con conciliación de NCR el papel no separa las NCR de las "
              f"NC comerciales (falta \"NCR eliminadas\" / \"TOTAL NC\" debajo de las compras): {', '.join(mm(p) for p in sin_bloque)}.",
              len(sin_bloque), comparar=len(sin_bloque))
    x.ok(18)


def c19(x):
    P = papel_anual(x)
    if not P: sin_papel(x, 19); return x.ok(19)
    r_eerr = x.e.ix.get('iva prorrateo')
    for p in x.meses:
        d = P.get(p)
        if not d: continue
        vn = d['sub'].get(('ventas', 'ventas netas (s/ notas de credito)')); pe = d['pct'].get('venta exenta')
        cel = celda_papel(d, ('pct', 'venta exenta'))
        if vn and num(pe) and vn[0] + vn[2]:
            esp = vn[2] / (vn[0] + vn[2])
            if abs(pe - esp) > 0.0005:
                x.add(19, p, '% venta exenta', 'Rojo', f"El % de venta exenta del papel de {mm(p)} ({cel} de la hoja {d['hoja']}) no sale de las Ventas Netas "
                      f"del mismo papel ({pct(esp)}): ¿se pisó la fórmula?", round(pe, 4), round(esp, 4), es_pct=True, fuente=d['f'])
        f = next(iter(x.buscar(rf'^{p} - PORTAL IVA - VENTAS\.xlsx$', p)), None)
        if f and num(pe):
            try:
                gr, ex = ventas_libro(f)
                if gr + ex and abs(pe - ex / (gr + ex)) > CFG['prorrateo_pct_tol']:
                    x.add(19, p, '% venta exenta vs Portal IVA', 'Amarillo', f"El papel de {mm(p)} prorratea con {pct(pe)} de venta exenta y el libro de "
                          f"ventas del Portal IVA da {pct(ex / (gr + ex))} ({os.path.basename(f)}).", round(pe, 4), round(ex / (gr + ex), 4), es_pct=True, fuente=d['f'])
            except Exception as err: x.log.append(f"19 {mm(p)}: no se pudo leer {os.path.basename(f)} ({err})")
        serv = sum(r['vals'][3] for r in d['compras'] if r['parte'] == 'cf' and r['clase'].startswith('servicios'))
        comp, nocomp = d['k'].get('total iva computable (servicios)'), d['k'].get('total iva no computable (servicios)')
        if num(nocomp) and num(pe) and abs(nocomp - pe * serv) > 1:
            x.add(19, p, 'IVA no computable', 'Rojo', f"El IVA no computable por prorrateo del papel de {mm(p)} ({pesos(nocomp)}) no es el % de venta exenta "
                  f"({pct(pe)}) por el IVA de compras de servicios ({pesos(serv)}).", round(nocomp, 2), round(pe * serv, 2), fuente=d['f'])
        if num(comp) and num(nocomp) and serv - comp - nocomp > 1:
            sin = sorted({r['lab'] + ' / ' + r['clase'].upper() for r in d['compras'] if r['parte'] == 'cf' and r['clase'].startswith('servicios')
                          and not re.search(r'\d+(,\d+)?\s*%', r['clase']) and r['vals'][3]})
            x.add(19, p, 'IVA de servicios fuera del prorrateo', 'Rojo', f"En {mm(p)} hay {pesos(serv - comp - nocomp)} de IVA de servicios que no entra como "
                  f"computable ni como no computable: la clasificación no dice la alícuota (\"SERVICIOS AL x%\")" + (f": {'; '.join(sin)}." if sin else '.'),
                  round(comp + nocomp, 2), round(serv, 2), fuente=d['f'])
        v = x.e.v(r_eerr, p) if r_eerr else 0
        if r_eerr and num(nocomp) and abs(v - nocomp) > max(CFG['prorrateo_tol_min'], CFG['prorrateo_tol_pct'] * abs(nocomp)):
            x.add(19, p, 'IVA Prorrateo del EERR vs papel', 'Rojo' if nocomp > 1 and not v else 'Amarillo',
                  f"La fila IVA Prorrateo del EERR de {mm(p)} no es el IVA no computable del papel de trabajo (hoja {d['hoja']})" +
                  (": está en cero." if not v else ": suele ser la clasificación mercadería / servicios del Maestro distinta de la del papel."),
                  v, round(nocomp, 2), celda=x.e.ref(r_eerr, p), fuente=d['f'])
        x.ok(19, p)
    x.ok(19)


def c20(x):
    P = papel_anual(x)
    if not P: sin_papel(x, 20); return x.ok(20)
    rat = {'exentas': {}, 'gravadas': {}}; ventas = {}
    for p in x.meses:
        d = P.get(p)
        if not d: continue
        # compras de mercadería sin las NC: las NCR restan exento en montos que cambian mucho de un mes a otro
        vn = d['sub'].get(('ventas', 'ventas netas (s/ notas de credito)'))
        merc = [r['vals'] for r in d['compras'] if r['clase'] == 'mercaderia' and r['parte'] == 'cf']
        if not vn: continue
        vg, ve = vn[0], vn[1] + vn[2]; cg, ce = sum(v[0] for v in merc), sum(v[1] + v[2] for v in merc)
        if vg + ve <= 0 or cg + ce <= 0: x.ok(20, p); continue
        sv, sc = ve / (vg + ve), ce / (cg + ce)
        if abs(sv - sc) > CFG['mix_exento_pts']:
            x.add(20, p, 'Mezcla exento / gravado', 'Amarillo', f"En {mm(p)} las ventas exentas son {pct(sv)} de las ventas y las compras exentas de mercadería "
                  f"{pct(sc)} de las compras de mercadería (sin NC): revisar compras gravadas o exentas mal cargadas en el papel.", round(sc, 4), round(sv, 4),
                  es_pct=True, fuente=d['f'], comparar=round(sc - sv, 2))
        if ve > 0: rat['exentas'][p] = ce / ve
        if vg > 0: rat['gravadas'][p] = cg / vg
        ventas[p] = (ve, vg); x.ok(20, p)
    for tipo, rr in rat.items():
        if len(rr) < 3: continue
        med = statistics.median(rr.values())
        for p, q in rr.items():
            v = ventas[p][0 if tipo == 'exentas' else 1]
            if med > 0 and abs(q - med) / med > CFG['compra_venta_desvio_pct'] and abs(q - med) * v > CFG['materialidad_min']:
                x.add(20, p, f"Compras / ventas {tipo}", 'Amarillo', f"En {mm(p)} las compras {tipo[:-1]}s de mercadería (sin NC) son {pct(q)} de las ventas {tipo}; "
                      f"lo habitual es {pct(med)} (mediana del ejercicio).", round(q, 4), round(med, 4), es_pct=True, fuente=P[p]['f'], comparar=round(q, 2))
    x.ok(20)


def c21(x):
    P = papel_anual(x)
    if not P: sin_papel(x, 21); return x.ok(21)
    sin_ident = []
    for p in x.meses:
        d = P.get(p)
        f = next(iter(x.buscar(rf'^{p} - PORTAL IVA - COMPRAS\.xlsx$', p)), None) or next(iter(x.buscar(rf'^{p} - PORTAL IVA - COMPRAS.*\.xlsx$', p)), None)
        if not d or not f: continue
        try: lib, claves = nc_libro(f)
        except Exception as err: x.log.append(f"21 {mm(p)}: no se pudo leer {os.path.basename(f)} ({err})"); continue
        # el libro puede tener las NCR o no (después del paso 4 de conciliacion-ncr se eliminan): el papel las incluye
        cn = conciliacion(x, p) or {}; lib = abs(lib)
        pap = abs(d['sub'].get(('compras', 'subtotal notas de credito s/compras'), [0] * 5)[4])
        tol = lambda esp: max(CFG['nc_tol'], 0.001 * esp)
        if cn.get('claves'):
            ya = abs(sum(v for k, v in cn['claves'].items() if k not in claves)); ncr = cn['ncr']   # NCR ya eliminadas del libro
        elif d['ncr'] is not None:
            ncr = abs(d['ncr']); ya = ncr if abs(pap - ncr - lib) < abs(pap - lib) else 0      # el libro sin las NCR
        elif pap > lib + tol(lib):
            sin_ident.append((p, pap - lib)); x.ok(21, p); continue
        else: ya = ncr = 0
        esp = lib + ya
        if ncr and abs(pap - (esp - ncr)) <= tol(esp): x.ok(21, p); continue        # papel sin las NCR: lo informa el 18
        if abs(pap - esp) > tol(esp):
            partes = f" NCR: {pesos(ncr)} · NC comerciales: {pesos(pap - ncr)}." if ncr and pap >= ncr else ''
            x.add(21, p, 'NC de compras vs libro', 'Rojo' if esp and abs(pap - esp) / esp > CFG['ventas_tol_pct'] else 'Amarillo',
                  f"Las NC de compras del papel de {mm(p)} suman {pesos(pap)} y las del libro de compras del Portal IVA {pesos(lib)}" +
                  (f" más {pesos(ya)} de NCR ya eliminadas del libro" if ya else '') + f" ({os.path.basename(f)})." + partes,
                  round(pap, 2), round(esp, 2), fuente=d['f'])
        x.ok(21, p)
    if sin_ident:
        x.add(21, '-', 'NC de compras sin NCR identificadas', 'Amarillo',
              f"En {len(sin_ident)} meses las NC de compras del papel superan a las del libro del Portal IVA y no hay cómo saber si la diferencia son "
              "las NCR eliminadas del libro (no hay conciliación de NCR ni \"NCR eliminadas\" en el papel): " +
              ', '.join(f"{mm(p)} {pesos(v)}" for p, v in sin_ident) + '.', len(sin_ident), comparar=len(sin_ident))
    x.ok(21)


CHECKS = [c01, c02, c03, c04, c05, c06, c07, c08, c09, c10, c11, c12, c13, c14, c15, c16, c17, c18, c19, c20, c21]


def analizar(args):
    f, hasta = args
    x = Ctx(f, hasta)
    nombre = norm(re.sub(r'^EERR .*? - ', '', os.path.basename(f))[:-5])
    exc = {int(c) for k, cs in CFG['excepciones'].items() if norm(k) == nombre or re.sub(r'\D', '', k) == (x.e.cuit or '-') for c in cs}
    for i, chk in enumerate(CHECKS, 1):
        if i in exc: x.log.append(f"Control {i} no corre para esta sociedad (excepción en la configuración)"); continue
        try: chk(x)
        except Exception as err:
            import traceback
            x.log.append(f"Control {i} no pudo correr: {err!r}")
            x.scope = {s for s in x.scope if s[0] != i}
            x.log.append(traceback.format_exc(limit=2))
    e = x.e
    dup = {}
    for lab in DUP:
        r = e.ix.get(lab)
        for p in x.meses:
            v = round(e.v(r, p), 2)
            if abs(v) >= 10000: dup[(lab, p, v)] = e.ref(r, p)
    return dict(f=f, cuit=e.cuit, denominacion=e.denominacion, meses=x.meses, h=x.h, scope=sorted(x.scope), log=x.log, dup=dup,
                etiquetas={lab: nfc(e.wf.cell(e.ix[lab], 1).value) for lab in DUP if e.ix.get(lab)})


# ------------------------------------------------------------------ hoja Auditorias
# Estado va en la columna B para que se vea sin scroll (pedido de Javi 09/10/2026). La hoja se lee por el nombre del
# encabezado, así que las hojas escritas con el orden anterior (Estado en M) se migran solas en la próxima auditoría.
COLS = ['N°', 'Estado', 'Nivel', 'Fecha detección', 'Control', 'Mes', 'Concepto', 'Descripción', 'Valor EERR', 'Valor de control',
        'Diferencia', 'Ir a la celda', 'Documento fuente', 'Fecha estado', 'Observaciones', 'Última auditoría', 'Clave', 'Comparar']
CI = {n: i for i, n in enumerate(COLS, 1)}            # nombre -> número de columna
def CC(n): return CL(CI[n])                            # nombre -> letra de columna
ANCHOS = {'N°': 5, 'Estado': 11, 'Nivel': 9, 'Fecha detección': 11, 'Control': 24, 'Mes': 9, 'Concepto': 26, 'Descripción': 70,
          'Valor EERR': 15, 'Valor de control': 15, 'Diferencia': 15, 'Ir a la celda': 10, 'Documento fuente': 26,
          'Fecha estado': 11, 'Observaciones': 30, 'Última auditoría': 11, 'Clave': 10, 'Comparar': 10}
HIST = ['Fecha', 'Meses hasta', 'Controles corridos', 'Nuevos', 'Siguen pendientes', 'Resueltos solos', 'Reabiertos', 'Aprobados (no se informan)', 'Avisos']
HC = len(COLS) + 2   # columna donde empieza el historial
ROJO, AMAR = PatternFill('solid', fgColor='FFC7CE'), PatternFill('solid', fgColor='FFEB9C')
GRIS = PatternFill('solid', fgColor='D9E1F2')
FMT = '#,##0.00;[Red]-#,##0.00'


def clave(h): return f"C{h['control']:02d}|{h['mes']}|{h['concepto']}"
# Filas cargadas a mano por el equipo (pedido de Javi 09/10/2026, T-011): sin Clave o con Clave "M|n". Se conservan tal
# cual y su Estado lo maneja solo el equipo: el auditor no las resuelve ni las reabre.
MANUAL = ('Control', 'Mes', 'Concepto', 'Descripción', 'Valor EERR', 'Valor de control', 'Observaciones')
def es_manual(d): return str(d.get('Clave') or '').startswith('M|')


def leer_hoja(wb):
    if HOJA not in wb.sheetnames: return [], []
    ws = wb[HOJA]
    hdr = next((r for r in range(1, 15) if nfc(ws.cell(r, 1).value) == 'N°'), None)
    if not hdr: return [], []
    nombres = [nfc(ws.cell(hdr, c).value) for c in range(1, len(COLS) + 1)]
    filas, hist = [], []
    for r in range(hdr + 1, ws.max_row + 1):
        d = {n: ws.cell(r, c + 1).value for c, n in enumerate(nombres) if n}
        if not d.get('Clave') and any(nfc(d.get(n)) for n in MANUAL): d['Clave'] = 'M|'
        if d.get('Clave'):
            for k, n in (('_link', 'Ir a la celda'), ('_fuente', 'Documento fuente')):
                d[k] = ws.cell(r, nombres.index(n) + 1).hyperlink if n in nombres else None
            filas.append(d)
        h = [ws.cell(r, HC + i).value for i in range(len(HIST))]
        if h[0]: hist.append(h)
    return filas, hist


def mismo(a, b):
    if num(a) and num(b): return abs(a - b) <= max(1, 0.005 * abs(b))
    return nfc(a) == nfc(b)


def fusionar(previas, nuevos, scope, hoy):
    # devuelve filas finales + contadores
    sc = {tuple(s) for s in scope}
    por_clave = {}
    for i, d in enumerate(previas):
        if not es_manual(d): por_clave[d['Clave']] = i                # la última fila de cada clave manda
    filas = [dict(d) for d in previas]
    cont = collections.Counter(); vistas = set()
    for d in filas:
        if not es_manual(d): continue
        cont['manuales'] += 1
        if not nfc(d.get('Estado')): d['Estado'] = 'Pendiente'        # sin estado no entraría en los contadores
        if not d.get('Fecha detección'): d['Fecha detección'] = hoy
    for h in nuevos:
        k = clave(h); vistas.add(k)
        base = dict(h)
        if k in por_clave:
            d = filas[por_clave[k]]; est = norm(d.get('Estado'))
            if est == 'aprobado':
                if mismo(d.get('Comparar'), h['comparar']): cont['aprobados'] += 1; continue
                fa = d.get('Fecha estado'); fa = f"{fa:%d/%m/%Y}" if hasattr(fa, 'strftime') else (fa or 's/f')
                base['desc'] = f"El valor cambió desde la aprobación del {fa} (antes {d.get('Comparar')}). " + h['desc']
                filas.append(dict(_nuevo=base, Estado='Pendiente', Clave=k)); por_clave[k] = len(filas) - 1; cont['nuevos'] += 1; continue
            d['_nuevo'] = base; d['Última auditoría'] = hoy
            if est == 'resuelto':
                d['Estado'] = 'Pendiente'; d['Fecha estado'] = hoy
                d['Observaciones'] = ((d.get('Observaciones') or '') + f" Reabierto el {hoy:%d/%m/%Y}: se volvió a detectar.").strip()
                cont['reabiertos'] += 1
            else:
                d['Estado'] = 'Pendiente'; cont['pendientes'] += 1
        else:
            filas.append(dict(_nuevo=base, Estado='Pendiente', Clave=k, **{'Fecha detección': hoy, 'Última auditoría': hoy}))
            por_clave[k] = len(filas) - 1; cont['nuevos'] += 1
    for k, i in por_clave.items():
        d = filas[i]
        if k in vistas or norm(d.get('Estado')) != 'pendiente': continue
        c, mes = int(k[1:3]), k.split('|')[1]
        if (c, mes) in sc or (c, '*') in sc and mes == '-':
            d['Estado'] = 'Resuelto'; d['Fecha estado'] = hoy
            d['Observaciones'] = ((d.get('Observaciones') or '') + f" Resuelto solo: no se detectó en la auditoría del {hoy:%d/%m/%Y}.").strip()
            cont['resueltos'] += 1
    return filas, cont


def escribir(wb, res, hasta, hoy):
    previas, hist = leer_hoja(wb)
    filas, cont = fusionar(previas, res['h'], res['scope'], hoy)
    if HOJA in wb.sheetnames: del wb[HOJA]
    ws = wb.create_sheet(HOJA)
    H = 5; H0 = H + 1                                   # fila de encabezados y primera fila de hallazgos
    ws['A1'] = f"Auditoría del EERR — {res['denominacion']} (CUIT {res['cuit']})"; ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = ("Estado: Pendiente = a revisar · Resuelto = se corrigió (el auditor lo marca solo cuando deja de detectarlo) · "
                "Aprobado = revisado y aceptado: no se vuelve a informar mientras el valor no cambie. Cambiar el estado en la columna B.")
    ws['A2'].font = Font(italic=True, size=9)
    E = CC('Estado')
    ws['A3'] = 'Pendientes'; ws['B3'] = f'=COUNTIF({E}{H0}:{E}100000,"Pendiente")'
    ws['C3'] = 'Resueltos'; ws['D3'] = f'=COUNTIF({E}{H0}:{E}100000,"Resuelto")'
    ws['E3'] = 'Aprobados'; ws['F3'] = f'=COUNTIF({E}{H0}:{E}100000,"Aprobado")'
    ws['G3'] = f"Última auditoría: {hoy:%d/%m/%Y %H:%M} — meses hasta {mm(hasta)}"
    for c in ('A3', 'C3', 'E3'): ws[c].font = Font(bold=True)
    for i, n in enumerate(COLS, 1):
        c = ws.cell(H, i, n); c.font = Font(bold=True); c.fill = GRIS; c.alignment = Alignment(wrap_text=True, vertical='center')
    r = H
    for n, d in enumerate(filas, 1):
        r += 1; h = d.get('_nuevo')
        g = lambda k, kh: (h[kh] if h else d.get(k))
        vals = {'N°': n, 'Estado': d.get('Estado'), 'Nivel': g('Nivel', 'nivel'), 'Fecha detección': d.get('Fecha detección'),
                'Control': (f"{h['control']}. {CONTROLES[h['control']]}" if h else d.get('Control')), 'Mes': g('Mes', 'mes'),
                'Concepto': g('Concepto', 'concepto'), 'Descripción': g('Descripción', 'desc'), 'Valor EERR': g('Valor EERR', 'valor'),
                'Valor de control': g('Valor de control', 'control_val'), 'Diferencia': g('Diferencia', 'dif'),
                'Fecha estado': d.get('Fecha estado'), 'Observaciones': d.get('Observaciones'),
                'Última auditoría': d.get('Última auditoría'), 'Clave': f"M|{n}" if es_manual(d) else d['Clave'],
                'Comparar': g('Comparar', 'comparar')}
        for k, v in vals.items(): ws.cell(r, CI[k], v)
        ve = vals['Valor EERR']
        es_pct = h['es_pct'] if h else (ve is not None and num(ve) and abs(ve) < 5 and '%' in str(d.get('Descripción') or ''))
        for k in ('Valor EERR', 'Valor de control', 'Diferencia'): ws.cell(r, CI[k]).number_format = '0.0%' if es_pct else FMT
        for k in ('Fecha detección', 'Fecha estado', 'Última auditoría'): ws.cell(r, CI[k]).number_format = 'dd/mm/yyyy'
        ws.cell(r, CI['Descripción']).alignment = Alignment(wrap_text=True, vertical='top')
        nivel = vals['Nivel']
        ws.cell(r, CI['Nivel']).fill = ROJO if nivel == 'Rojo' else AMAR if nivel == 'Amarillo' else PatternFill()
        # links: a la celda del EERR (o de otra hoja del archivo) y al documento fuente en Drive
        celda = h['celda'] if h else None; fuente = h['fuente'] if h else None
        CK, CF = CI['Ir a la celda'], CI['Documento fuente']
        if h:
            if celda:
                ws.cell(r, CK, 'Ver celda').hyperlink = Hyperlink(ref=f"{CL(CK)}{r}", location=celda)
            if fuente:
                if fuente.startswith("'"):
                    ws.cell(r, CF, 'Ver hoja de control').hyperlink = Hyperlink(ref=f"{CL(CF)}{r}", location=fuente)
                else:
                    url = drive_url(fuente)
                    if url: ws.cell(r, CF, os.path.basename(fuente.rstrip('/\\')) or fuente).hyperlink = url
                    else: ws.cell(r, CF, os.path.relpath(fuente, U) if fuente.startswith(U) else fuente)
        else:
            for col, hl in ((CK, d.get('_link')), (CF, d.get('_fuente'))):
                if hl is not None:
                    ws.cell(r, col, d.get(COLS[col - 1]) or ('Ver celda' if col == CK else 'Documento'))
                    ws.cell(r, col).hyperlink = Hyperlink(ref=f"{CL(col)}{r}", location=hl.location, target=hl.target) if (hl.location or hl.target) else None
                else: ws.cell(r, col, d.get(COLS[col - 1]))
        for col in (CK, CF):
            if ws.cell(r, col).hyperlink: ws.cell(r, col).font = Font(color='0563C1', underline='single')
    ultima = max(r, H + 1)
    dv = DataValidation(type='list', formula1='"Pendiente,Resuelto,Aprobado"', allow_blank=False)
    rango = f"{E}{H0}:{E}{max(ultima + 500, 1000)}"
    ws.add_data_validation(dv); dv.add(rango)
    for txt, color in (('Pendiente', 'F8CBAD'), ('Resuelto', 'C6EFCE'), ('Aprobado', 'BDD7EE')):
        ws.conditional_formatting.add(rango, FormulaRule(formula=[f'${E}{H0}="{txt}"'], fill=PatternFill('solid', fgColor=color)))
    for n, w in ANCHOS.items(): ws.column_dimensions[CC(n)].width = w
    for n in ('Clave', 'Comparar'): ws.column_dimensions[CC(n)].hidden = True
    ws.freeze_panes = f"C{H0}"; ws.auto_filter.ref = f"A{H}:{CL(len(COLS))}{ultima}"
    # historial de ejecuciones
    ws.cell(H - 1, HC, 'Historial de auditorías').font = Font(bold=True)
    for i, n in enumerate(HIST):
        c = ws.cell(H, HC + i, n); c.font = Font(bold=True); c.fill = GRIS; c.alignment = Alignment(wrap_text=True)
    avisos = ' | '.join(l for l in res['log'] if not l.startswith('Traceback') and 'File "' not in l)[:1000]
    corridos = sorted({s[0] for s in res['scope'] if s[1] == '*'})
    n = len(CONTROLES)
    hist.append([hoy, mm(hasta), f"{len(corridos)} de {n}" + ('' if len(corridos) == n else f" (faltan {', '.join(str(i) for i in range(1, n + 1) if i not in corridos)})"),
                 cont['nuevos'], cont['pendientes'], cont['resueltos'], cont['reabiertos'], cont['aprobados'], avisos or None])
    for j, h in enumerate(hist):
        for i, v in enumerate(h):
            c = ws.cell(H + 1 + j, HC + i, v)
            if i == 0: c.number_format = 'dd/mm/yyyy hh:mm'
    for i, w in enumerate([16, 10, 14, 8, 10, 10, 10, 12, 60]): ws.column_dimensions[CL(HC + i)].width = w
    return cont


# ------------------------------------------------------------------ main
PS = r"""
$ErrorActionPreference = 'Stop'
$xl = New-Object -ComObject Excel.Application
$xl.Visible = $false; $xl.DisplayAlerts = $false
try {
  $wb = $xl.Workbooks.Open($env:AUDITOR_XLSX, 0)
  $xl.CalculateFull(); $wb.Save(); $wb.Close($false)
} finally { $xl.Quit(); [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($xl) }
"""
OSA = """with timeout of 600 seconds
tell application "Microsoft Excel"
  set display alerts to false
  set wb to open workbook workbook file name (POSIX file "{f}" as text) update links do not update links
  calculate full
  save wb
  close wb saving no
end tell
end timeout"""


def recalcular(f):
    # abre el archivo con Excel, recalcula todo y guarda (openpyxl no calcula fórmulas)
    f = os.path.abspath(f)
    for i in range(3):
        try:
            if WIN:
                r = subprocess.run(['powershell', '-NoProfile', '-NonInteractive', '-Command', PS], capture_output=True, text=True,
                                   timeout=900, env={**os.environ, 'AUDITOR_XLSX': f})
            elif MAC:
                r = subprocess.run(['osascript', '-e', OSA.format(f=f.replace('"', '\\"'))], capture_output=True, text=True, timeout=900)
            else:
                print('  ! sin Excel en este sistema: el archivo queda sin valores calculados'); return False
            if r.returncode == 0: return True
            err = (r.stderr or r.stdout).strip()
        except Exception as e: err = str(e)
        time.sleep(10)
    print('  ! Excel no pudo recalcular', os.path.basename(f), err[:200]); return False


def necesita_recalculo(f):
    wf, wv = abrir(f)['EERR'], abrir(f, data_only=True)['EERR']
    for row in wf.iter_rows(min_row=1, max_row=wf.max_row, max_col=14):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('=') and wv[c.coordinate].value is None and c.column > 1:
                return True
    return False


def main():
    try: sys.stdout.reconfigure(encoding='utf-8')      # consola de Windows
    except Exception: pass
    a = {k: v for k, _, v in (s.lstrip('-').partition('=') for s in sys.argv[1:])}
    hasta = a.get('hasta') or hasta_default()
    simular = 'simular' in a
    archivos = sorted(f for f in glob.glob(f"{EERR_DIR}/EERR *.xlsx") if not os.path.basename(f).startswith('~$'))
    if 'nombre' in a:
        q = norm(a['nombre'])
        archivos = [f for f in archivos if norm(re.sub(r'^EERR .*? - ', '', os.path.basename(f))[:-5]) == q] or \
                   [f for f in archivos if q in norm(os.path.basename(f))]
    elif 'todas' not in a:
        sys.exit('Uso: auditor_eerr.py --nombre="Sociedad" | --todas  [--hasta=AAAAMM] [--simular] [--procesos=4]')
    if not archivos: sys.exit('No se encontró el EERR pedido.')
    hoy = datetime.datetime.now().replace(microsecond=0)
    trabajo = f"{SALIDA}/trabajo"; os.makedirs(trabajo, exist_ok=True)
    resp = f"{SALIDA}/_respaldos/{hoy:%Y%m%d-%H%M}"
    print(f"Auditoría de {len(archivos)} EERR, meses hasta {mm(hasta)}{' (simulación)' if simular else ''}")
    # 1) copia de trabajo (y recálculo si faltan valores de fórmulas)
    copias, saltear = {}, []
    for f in archivos:
        w = os.path.join(trabajo, os.path.basename(f))
        shutil.copy2(f, w); copias[f] = (w, os.path.getmtime(f))
        if necesita_recalculo(w):
            print('  recalculando con Excel', os.path.basename(f))
            if not recalcular(w) or necesita_recalculo(w):
                # sin los valores de las fórmulas la auditoría daría todo en cero
                print('  ! se saltea: las fórmulas no tienen valores calculados y Excel no pudo recalcular', os.path.basename(f))
                saltear.append(f)
    archivos = [f for f in archivos if f not in saltear]
    if not archivos: sys.exit('No quedó ningún EERR para auditar.')
    # 2) análisis en paralelo
    res = {}
    with ProcessPoolExecutor(int(a.get('procesos', 4))) as ex:
        for f, r in zip(archivos, ex.map(analizar, [(copias[f][0], hasta) for f in archivos])):
            r['original'] = f; res[f] = r
    # 3) valores idénticos entre sociedades (copias de planillas)
    if len(archivos) > 1:
        idx = collections.defaultdict(list)
        for f, r in res.items():
            for (lab, p, v), ref in r['dup'].items(): idx[(lab, p, v)].append((f, ref))
        for (lab, p, v), lst in idx.items():
            if len({f for f, _ in lst}) < 2: continue
            for f, ref in lst:
                otras = ', '.join(sorted(re.sub(r'^EERR .*? - ', '', os.path.basename(g))[:-5] for g, _ in lst if g != f))
                res[f]['h'].append(dict(control=2, mes=mm(p), concepto=f"{res[f]['etiquetas'].get(lab, lab)} igual a otra sociedad", nivel='Rojo',
                                        desc=f"{res[f]['etiquetas'].get(lab, lab)} de {mm(p)} ({pesos(v)}) es idéntico al de: {otras}. Posible planilla copiada.",
                                        valor=v, control_val=None, dif=None, celda=ref, fuente=None, es_pct=False, comparar=v))
    # 4) escribir la hoja, recalcular y devolver a la unidad
    resumen = []
    for f in archivos:
        r = res[f]; w, mt = copias[f]
        wb = abrir(w)
        cont = escribir(wb, r, hasta, hoy)
        wb.save(w); recalcular(w)
        nombre = re.sub(r'^EERR .*? - ', '', os.path.basename(f))[:-5]
        estado = 'simulación'
        if simular:
            os.makedirs(f"{SALIDA}/simulacion", exist_ok=True); shutil.copy2(w, f"{SALIDA}/simulacion/{os.path.basename(f)}")
        elif os.path.getmtime(f) != mt:
            estado = 'NO GUARDADO: el archivo cambió en la unidad mientras se auditaba'
            os.makedirs(f"{SALIDA}/no-guardados", exist_ok=True); shutil.copy2(w, f"{SALIDA}/no-guardados/{os.path.basename(f)}")
        else:
            os.makedirs(resp, exist_ok=True); shutil.copy2(f, f"{resp}/{os.path.basename(f)}")
            shutil.copy2(w, f); estado = 'guardado'
        niveles = collections.Counter(h['nivel'] for h in r['h'])
        resumen.append((nombre, cont, niveles, estado, r['log']))
        print(f"  {nombre:28} nuevos {cont['nuevos']:3}  siguen {cont['pendientes']:3}  resueltos {cont['resueltos']:3}  "
              f"reabiertos {cont['reabiertos']:2}  aprobados omitidos {cont['aprobados']:3}  manuales {cont['manuales']:2}  (rojo {niveles['Rojo']}, amarillo {niveles['Amarillo']})  {estado}")
        for l in r['log']:
            if not l.startswith('Traceback') and 'File "' not in l: print('     aviso:', l)
    json.dump([dict(sociedad=n, **c, rojo=v['Rojo'], amarillo=v['Amarillo'], estado=e) for n, c, v, e, _ in resumen],
              open(f"{SALIDA}/ultima-ejecucion.json", 'w', encoding='utf-8'), ensure_ascii=False, indent=1, default=str)


if __name__ == '__main__':
    main()
