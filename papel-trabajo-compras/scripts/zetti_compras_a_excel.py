"""Subdiario de IVA Compras de Zetti (reporte 5.6.4, CSV) -> Excel con el formato
de la solapa "Compras Zetti" del papel de trabajo.

    python3 zetti_compras_a_excel.py <entrada.csv> <salida.xlsx> [--periodo=AAAAMM]

El CSV de T&S Web viene en UTF-8 (a veces latin-1), separado por ";" y con
columnas vacias intercaladas:
    filas 1-5  encabezado (T&S Web, fecha de emision, "Subdiario de IVA Compras"
               con Fecha des/has y Nodo, Razon social y C.U.I.T.)
    fila 6     titulos (Fecha;;TC;M;Nro. Comp.; Proveedor;;;;CUIT;;;;R;Exen;...;P/R IG;Total)
    datos      una fila por comprobante, ordenadas por fecha
    ultima     ";;...Total: ;;<Exen>;;;;<Grav>;;;<IVA>;;<P/R IIBB>;;;<P/R IVA>;;;<P/R IG>;<Total>"

Salida: hoja "Compras Zetti" igual a la que arma el equipo en los papeles
(202608-papelesdetrabajo-*.xlsx: el reporte abierto en Excel, sin las columnas
vacias):
    fila 1  A "T&S Web" | D razon social | F CUIT
    fila 2  D "Fecha des:" | E fecha
    fila 3  D "Fecha has:" | E fecha
    fila 4  A "  Subdiario de IVA Compras"
    fila 5  A "Razón Social:"
    fila 6  Fecha | TC | M | Nro. Comp. | " Proveedor" | CUIT | R | Exen | Grav
            | IVA | P/R IIBB | P/R IVA | Total
    datos   desde la fila 7, en el orden del reporte (por fecha)
    total   fila siguiente a los datos, importes de H a M sin rotulo
Como al abrir el CSV en Excel: lo que es solo digitos (CUIT sin guiones,
"Nro. Comp." sin guion) queda numero, el resto texto tal cual (el proveedor con
su espacio adelante). P/R IG no va (el papel no la tiene); si viniera con
importe se avisa en el JSON.

Imprime por stdout un JSON con razon social, CUIT, fechas, cantidad por TC,
totales y el control contra la fila "Total:" del reporte. Sale con 1 (y el JSON
con "error") si el Total no coincide al centavo (o alguna otra columna difiere
en mas de $0,05), si hay filas que no reconoce o fechas fuera del periodo
pedido: en ese caso no hay que archivar.
"""
import csv
import json
import re
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# Titulo del reporte -> clave interna.
COLUMNAS = [
    ('Fecha', 'fecha'), ('TC', 'tc'), ('M', 'm'), ('Nro. Comp.', 'nro'),
    ('Proveedor', 'proveedor'), ('CUIT', 'cuit'), ('R', 'r'),
    ('Exen', 'exen'), ('Grav', 'grav'), ('IVA', 'iva'),
    ('P/R IIBB', 'pib'), ('P/R IVA', 'piva'), ('P/R IG', 'pig'), ('Total', 'total'),
]
IMPORTES = ['exen', 'grav', 'iva', 'pib', 'piva', 'pig', 'total']
TOLERANCIA = 0.05
TEXTO = ['tc', 'm', 'nro', 'proveedor', 'cuit', 'r']


def fallar(msg, **extra):
    print(json.dumps({'error': msg, **extra}, ensure_ascii=False))
    sys.exit(1)


def importe(s):
    s = (s or '').strip()
    if not s:
        return 0.0
    try:
        return float(s.replace('.', '').replace(',', '.'))
    except ValueError:
        fallar(f'Importe que no entiendo: "{s}"')


def fecha(s):
    s = (s or '').strip()
    for fmt in ('%d/%m/%y', '%d/%m/%Y'):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None


def tras(fila, etiqueta, crudo=False):
    """Primer valor no vacio a la derecha de la celda que dice `etiqueta`."""
    for i, c in enumerate(fila):
        if c.strip().lower().startswith(etiqueta.lower()):
            for v in fila[i + 1:]:
                if v.strip():
                    return v if crudo else v.strip()
    return ''


def como_excel(v):
    """Lo que hace Excel al abrir el CSV: solo digitos -> numero."""
    return int(v) if re.fullmatch(r'\d{1,15}', v.strip()) else v


def main():
    pos = [a for a in sys.argv[1:] if not a.startswith('--')]
    opts = dict(a[2:].split('=', 1) for a in sys.argv[1:] if a.startswith('--') and '=' in a)
    if len(pos) != 2:
        sys.exit(__doc__)
    entrada, salida = pos
    periodo = opts.get('periodo')

    crudo = open(entrada, 'rb').read()
    try:
        texto = crudo.decode('utf-8-sig')
    except UnicodeDecodeError:
        texto = crudo.decode('latin-1')
    filas = list(csv.reader(texto.splitlines(), delimiter=';'))

    i_tit = next((i for i, f in enumerate(filas) if f and f[0].strip() == 'Fecha' and 'TC' in [c.strip() for c in f]), None)
    if i_tit is None:
        fallar('No encontre la fila de titulos (Fecha;TC;...) en el CSV.')
    titulos = [c.strip() for c in filas[i_tit]]
    pos_tit = {}
    for nombre, clave in COLUMNAS:
        if nombre not in titulos:
            fallar(f'Falta la columna "{nombre}" en el CSV.', titulos=[t for t in titulos if t])
        pos_tit[clave] = titulos.index(nombre)
    # Los datos no siempre caen bajo su titulo: "Grav" esta en la columna 17 y
    # los importes (y el del Total) en la 18. Cada columna es la primera con
    # datos entre su titulo y el titulo siguiente.
    con_datos = set()
    for f in filas[i_tit + 1:]:
        if f and fecha(f[0]):
            con_datos.update(i for i, c in enumerate(f) if c.strip())
    orden = sorted(pos_tit.items(), key=lambda kv: kv[1])
    idx = {}
    for n, (clave, p) in enumerate(orden):
        hasta_p = orden[n + 1][1] if n + 1 < len(orden) else p + 1
        idx[clave] = next((i for i in range(p, hasta_p) if i in con_datos), p)
    ancho = max(len(f) for f in filas) + 1

    razon = cuit_soc = desde = hasta = nodo = ''
    for f in filas[:i_tit]:
        razon = razon or tras(f, 'Raz', crudo=True)   # el papel la deja con el espacio final
        cuit_soc = cuit_soc or tras(f, 'C.U.I.T')
        desde = desde or tras(f, 'Fecha des')
        hasta = hasta or tras(f, 'Fecha has')
        nodo = nodo or tras(f, 'Nodo')
    cuit_fmt = cuit_soc
    cuit_soc = re.sub(r'\D', '', cuit_soc)

    datos, total_reporte, raras = [], None, []
    for n, f in enumerate(filas[i_tit + 1:], start=i_tit + 2):
        if not any(c.strip() for c in f):
            continue
        f = f + [''] * (ancho - len(f))
        d = fecha(f[idx['fecha']])
        if d is None:
            if any(c.strip().lower().startswith('total') for c in f):
                total_reporte = {k: importe(f[idx[k]]) for k in IMPORTES}
            elif any(c.strip().lower().startswith('página') or c.strip().lower().startswith('pagina') for c in f):
                pass   # pie de pagina "; :   -  ;...Página ;1"
            else:
                raras.append({'fila': n, 'texto': ';'.join(c for c in f if c.strip())[:120]})
            continue
        r = {'fecha': d}
        for k in TEXTO:
            r[k] = f[idx[k]]
        for k in IMPORTES:
            r[k] = importe(f[idx[k]])
        datos.append(r)

    if raras:
        fallar('Hay filas que no son comprobantes ni el total: revisar el CSV antes de archivar.', filas=raras)
    if not datos:
        fallar('El reporte no trae comprobantes.', cuit=cuit_soc, desde=desde, hasta=hasta)
    if total_reporte is None:
        fallar('El CSV no trae la fila "Total:": puede haber bajado cortado.')

    fuera = []
    if periodo:
        fuera = sorted({r['fecha'].strftime('%d/%m/%Y') for r in datos if r['fecha'].strftime('%Y%m') != periodo})

    suma = {k: round(sum(r[k] for r in datos), 2) for k in IMPORTES}
    difs = {k: round(suma[k] - total_reporte[k], 2) for k in IMPORTES if abs(suma[k] - total_reporte[k]) >= 0.005}
    graves = {k: v for k, v in difs.items() if k == 'total' or abs(v) > TOLERANCIA}

    wb = Workbook()
    ws = wb.active
    ws.title = 'Compras Zetti'
    numero = '#,##0.00'
    fdate = 'mm-dd-yy'   # el formato que deja Excel al abrir el CSV (se ve dd/mm/aa en es-AR)

    ws['A1'] = 'T&S Web'
    ws['D1'] = razon
    ws['F1'] = cuit_fmt
    ws['D2'] = 'Fecha des:'
    ws['D3'] = 'Fecha has:'
    for celda, valor in (('E2', desde), ('E3', hasta)):
        d = fecha(valor)
        ws[celda] = d or valor
        ws[celda].number_format = fdate
    ws['A4'] = '  Subdiario de IVA Compras'
    ws['A5'] = 'Razón Social:'

    columnas = [('Fecha', 'fecha'), ('TC', 'tc'), ('M', 'm'), ('Nro. Comp.', 'nro'), (' Proveedor', 'proveedor'),
                ('CUIT', 'cuit'), ('R', 'r'), ('Exen', 'exen'), ('Grav', 'grav'), ('IVA', 'iva'),
                ('P/R IIBB', 'pib'), ('P/R IVA', 'piva'), ('Total', 'total')]
    for j, (t, _) in enumerate(columnas, start=1):
        ws.cell(6, j, t)

    def poner(c, k, v):
        if k in IMPORTES:
            # Excel le pone separador de miles solo a lo que vino con punto de
            # miles ("9.595,36"); "0,00" o "239,88" quedan en General.
            c.value = v
            if abs(v) >= 1000:
                c.number_format = numero
        elif k == 'fecha':
            c.value = v
            c.number_format = fdate
        elif k in ('nro', 'cuit'):
            c.value = como_excel(v) if v.strip() else None
        else:
            c.value = v if v.strip() else None

    for i, r in enumerate(datos, start=7):
        for j, (_, k) in enumerate(columnas, start=1):
            poner(ws.cell(i, j), k, r[k])
    fila_total = 7 + len(datos)
    for j, (_, k) in enumerate(columnas, start=1):
        if k in IMPORTES:
            poner(ws.cell(fila_total, j), k, total_reporte[k])

    anchos = {'B': 5.33, 'C': 2.78, 'D': 18.78, 'E': 19.44, 'F': 13.89, 'G': 4.33}
    for letra, w in anchos.items():
        ws.column_dimensions[letra].width = w
    wb.save(salida)

    por_tc = {}
    for r in datos:
        clave = f"{r['tc'].strip()} {r['m'].strip()}".strip()
        por_tc[clave] = por_tc.get(clave, 0) + 1
    res = {
        'razon_social': razon.strip(), 'cuit': cuit_soc, 'nodo': nodo, 'desde': desde, 'hasta': hasta,
        'comprobantes': len(datos), 'por_tc': por_tc, 'totales': suma,
        'control_total_reporte': 'ok' if not difs else ('redondeo' if not graves else 'no coincide'),
        **({'diferencias_total_reporte': difs} if difs else {}),
        'fechas_fuera_de_periodo': fuera, 'fila_total': fila_total, 'xlsx': salida,
    }
    # Mismo comprobante dos veces en Zetti (paso con NC de Suizo y Monroe): se
    # deja como viene, igual que en el papel, pero se avisa.
    vistos = {}
    for r in datos:
        k = (r['tc'].strip(), r['m'].strip(), r['nro'].strip(), r['cuit'].strip(), r['total'])
        vistos[k] = vistos.get(k, 0) + 1
    dups = [{'tc': k[0], 'm': k[1], 'nro': k[2], 'cuit': k[3], 'total': k[4], 'veces': n} for k, n in vistos.items() if n > 1]
    if dups:
        res['aviso_duplicados'] = dups
    if suma['pig']:
        res['aviso_pig'] = f'Hay P/R IG (percepciones de Ganancias) por {suma["pig"]:.2f}: el papel no tiene esa columna (esta incluida en Total).'
    if graves:
        res['error'] = 'La suma de los comprobantes no coincide con el Total del reporte.'
    elif suma['iva'] and not suma['grav']:
        res['error'] = 'Hay IVA pero Grav suma 0: las columnas del CSV no se leyeron bien.'
    elif fuera:
        res['error'] = f'Hay comprobantes fuera del periodo {periodo}.'
    print(json.dumps(res, ensure_ascii=False))
    sys.exit(1 if 'error' in res else 0)


if __name__ == '__main__':
    main()
