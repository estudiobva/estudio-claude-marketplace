# armar_papel_compras.py — agrega las compras del mes al papel de trabajo de una sociedad.
#
#   python scripts/armar_papel_compras.py --papel=<copia del papel> --salida=<out.xlsx> --borrador=<csv|zip>
#       --zetti=<AAAAMM - ZETTI - COMPRAS.xlsx> --sin-ncr=<AAAAMM-COMPRAS-SIN-NCR-*.xlsx> --ncr-dir=<carpeta ncr>
#       --unidad=<raiz de la unidad> --maestro=<Maestro de Proveedores.xlsx> --cuit=<CUIT> --periodo=AAAAMM
#       [--zetti-unidad=<el de la unidad, para comparar>] [--portal-anterior=<compras del mes anterior>]
#       [--no-es-ncr=503-86152;3116-1234] [--si-existe=abortar|reemplazar] [--fecha-borrador="21/10/2026 15:33"]
#
# Hojas que agrega antes de "liquidacion" (formato de Soy 59 / Sanar 3 / Soy 51 Solano):
#   Compras Zetti            copia del 5.6.4
#   Compras Portal IVA       borrador ACTUAL del Libro Compras (sin NCR), clasificado con el Maestro (AG:AL) y, al
#                            final, el cuadro de Compras + "Prorrateo IVA Compras de Servicios"
#   Cruce compras            Zetti vs Portal por CUIT + tipo + letra + PV + numero
#   Resumen cruce compras    estados, control contra las hojas, proveedores con diferencias, hallazgos y metodologia
#   Clasificacion compras    por rubro del EERR y por proveedor
#
# Corta (ok=false) si alguna NCR conciliada o de los listados sigue en el borrador, si las hojas ya existen
# (salvo --si-existe=reemplazar), si el papel no tiene "liquidacion" con "Venta Gravada" o si aparece un tipo
# de comprobante que el cuadro no contempla. Las hojas que ya tenia el papel no se tocan.
#
# Salida: JSON por stdout; progreso por stderr.
import sys, os, re, io, csv, glob, json, copy, zipfile, argparse
from collections import OrderedDict, defaultdict, Counter
from datetime import datetime
import warnings
warnings.filterwarnings('ignore')
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import clasificar_compras as cl   # copia de clasificacion-compras/scripts/clasificar-compras.py

HOJAS = ('Compras Zetti', 'Compras Portal IVA', 'Cruce compras', 'Resumen cruce compras', 'Clasificacion compras')
TIPOS = {1: '1 - Factura A', 2: '2 - Nota de Débito A', 3: '3 - Nota de Crédito A', 4: '4 - Recibo A', 6: '6 - Factura B',
         7: '7 - Nota de Débito B', 8: '8 - Nota de Crédito B', 11: '11 - Factura C', 12: '12 - Nota de Débito C',
         13: '13 - Nota de Crédito C', 63: '63 - Liquidación A', 81: '81 - Tique Factura A', 111: '111 - Tique Factura C',
         201: '201 - Factura de Crédito Electrónica MiPyMEs (FCE) A', 202: '202 - Nota de Débito Electrónica MiPyMEs (FCE) A',
         203: '203 - Nota de Crédito Electrónica MiPyMEs (FCE) A'}
TC = {1: ('FC', 'A'), 2: ('ND', 'A'), 3: ('NC', 'A'), 4: ('RC', 'A'), 6: ('FC', 'B'), 7: ('ND', 'B'), 8: ('NC', 'B'),
      11: ('FC', 'C'), 12: ('ND', 'C'), 13: ('NC', 'C'), 63: ('LQ', 'A'), 81: ('FC', 'A'), 111: ('FC', 'C'),
      201: ('FC', 'A'), 202: ('ND', 'A'), 203: ('NC', 'A')}
SIGLA = {1: 'FC-A', 2: 'ND-A', 3: 'NC-A', 4: 'RC-A', 6: 'FC-B', 7: 'ND-B', 8: 'NC-B', 11: 'FC-C', 12: 'ND-C', 13: 'NC-C',
         63: 'LQ-A', 81: 'TF-A', 111: 'TF-C', 201: 'FCE-A', 202: 'NDE-A', 203: 'NCE-A'}
# tipos que el cuadro agrupa en cada fila
T_FA = [1, 63, 81, 4, 201]          # "1 - Factura A"
T_ND = [2, 202]                     # "2 - Nota de Débito A"
T_NCA = [3, 203]                    # "3 - Nota de Crédito A"
T_FC = [11, 111]                    # "11 - Factura C"
T_CUADRO = set(T_FA + T_ND + T_NCA + T_FC + [6, 13])
SERV_CUITS = {'30709565075': 'AYSA', '30655116512': 'EDESUR', '30657863676': 'METROGAS'}
DROG = {'30516968431': 'SUIZO', '30517059095': 'MONROE'}
NC_GRANDE = 1000000                 # NC de drogueria solo en el Portal por mas de esto -> revisar si es NCR

MON = '_-"$"\\ * #,##0.00_-;\\-"$"\\ * #,##0.00_-;_-"$"\\ * "-"??_-;_-@_-'
NUM = '#,##0.00'
thin = Side(style='thin', color='FFBFBFBF')
BORDE = Border(left=thin, right=thin, top=thin, bottom=thin)
def F(name='Calibri', **k): return Font(name=name, size=k.pop('size', 11), **k)
HDR_FILL = PatternFill('solid', fgColor='FF2E75B6')
TIT_FILL = PatternFill('solid', fgColor='FF1F4E79')
SUB_FILL = PatternFill('solid', fgColor='FFD9E1F2')
TOT_FILL = PatternFill('solid', fgColor='FF9BC2E6')
AMARILLO = PatternFill('solid', fgColor='FFFFF2CC')
ROJO = PatternFill('solid', fgColor='FFFCE4D6')
NARANJA = PatternFill('solid', fgColor='FFF8CBAD')
digitos = lambda s: re.sub(r'\D', '', str(s or ''))
fcuit = lambda c: f'{c[:2]}-{c[2:10]}-{c[10:]}' if len(c) == 11 else c
pesos = lambda x: f'$ {x:,.2f}'


def log(m):
    sys.stderr.write(f'[armar-compras] {m}\n')


class Corte(Exception):
    def __init__(self, code, msg, **extra):
        super().__init__(msg); self.code = code; self.extra = extra


def argumentos():
    p = argparse.ArgumentParser()
    for a in ('papel', 'salida', 'borrador', 'zetti', 'sin-ncr', 'ncr-dir', 'unidad', 'maestro', 'cuit', 'periodo'):
        p.add_argument('--' + a, required=True)
    for a in ('zetti-unidad', 'portal-anterior', 'no-es-ncr', 'fecha-borrador', 'fuente'):
        p.add_argument('--' + a, default='')
    p.add_argument('--si-existe', default='abortar', choices=['abortar', 'reemplazar'])
    return p.parse_args()


# ── lectura de fuentes ─────────────────────────────────────────────────────────────────────────────
def leer_borrador(ruta):
    """CSV del Libro Compras del borrador (pelado o dentro del ZIP de ARCA) -> filas de 32 columnas."""
    if ruta.lower().endswith('.zip'):
        with zipfile.ZipFile(ruta) as z:
            csvs = [n for n in z.namelist() if n.lower().endswith('.csv')]
            if len(csvs) != 1:
                raise Corte('borrador', f'El ZIP del borrador tiene {len(csvs)} CSV: {csvs}')
            raw = z.read(csvs[0])
    else:
        raw = open(ruta, 'rb').read()
    txt = raw.decode('utf-8-sig') if raw.startswith(b'\xef\xbb\xbf') else raw.decode('latin-1')
    filas = list(csv.reader(txt.splitlines(), delimiter=';'))
    num = lambda x: float(x.replace('.', '').replace(',', '.')) if x not in ('', None) else None
    out = []
    for r in filas[1:]:
        if not r or not r[0]:
            continue
        out.append([datetime.strptime(r[0], '%Y-%m-%d'), int(r[1]), int(r[2]), int(r[3]), int(r[4]), digitos(r[5]), r[6]]
                   + [num(r[7]), r[8], num(r[9])] + [num(x) for x in r[10:32]])
    return out


def _n(s):
    return cl.norm(str(s or '')).lower().strip()


def leer_zetti(ruta):
    """Hoja del 5.6.4: ubica el encabezado (Fecha/TC) y mapea columnas por nombre (dos formatos conocidos:
    'Nro. Comp.' 0000-00000000, o PV y numero en columnas separadas con P/R IG)."""
    wb = openpyxl.load_workbook(ruta)
    ws = wb['Compras Zetti'] if 'Compras Zetti' in wb.sheetnames else wb.worksheets[0]
    hdr = None
    for r in range(1, min(ws.max_row, 30) + 1):
        vals = [_n(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if 'fecha' in vals and 'tc' in vals:
            hdr, cab = r, vals; break
    if not hdr:
        raise Corte('zetti_formato', f'No encuentro el encabezado (Fecha / TC) en {os.path.basename(ruta)}')
    def col(*pats, req=True):
        for j, h in enumerate(cab, 1):
            if any(re.fullmatch(p, h) for p in pats):
                return j
        if req:
            raise Corte('zetti_formato', f'Falta la columna {pats[0]} en Zetti. Encabezado: {cab}')
        return None
    c = {'fecha': col('fecha'), 'tc': col('tc'), 'm': col('m', 'letra'),
         'comp': col(r'nro\.? ?comp\.?', r'comprobante', req=False),
         'prov': col(r'proveedor', r'razon social', r'raz.n social'), 'cuit': col('cuit'),
         'exen': col(r'exen\.?', r'exento'), 'grav': col(r'grav\.?', r'gravado'), 'iva': col('iva'),
         'piibb': col(r'p/r iibb'), 'piva': col(r'p/r iva'), 'pig': col(r'p/r ig', r'p/r gan.*', req=False),
         'total': col('total')}
    if not c['comp']:
        c['pv'] = col(r'p\.? ?v\.?', r'pto\.? ?(de )?vta\.?', r'punto de venta')
        c['nro'] = col(r'nro\.?', r'n.mero', r'numero')
    filas = []
    fin = None
    for r in range(hdr + 1, ws.max_row + 2):
        g = lambda k: ws.cell(r, c[k]).value if c.get(k) else None
        if g('fecha') is None:
            fin = r - 1; break
        if c['comp']:
            raw = str(g('comp') or '')
        else:
            pv, nro = g('pv'), g('nro')
            raw = f'{int(pv):04d}-{int(nro):08d}' if str(pv or '').isdigit() and str(nro or '').isdigit() else f'{pv or ""}{nro or ""}'
        filas.append({'fila': r, 'fecha': g('fecha'), 'tc': str(g('tc') or '').strip(), 'm': str(g('m') or '').strip(),
                      'raw': raw, 'prov': str(g('prov') or '').strip(), 'cuit': digitos(g('cuit')),
                      'exen': g('exen') or 0, 'grav': g('grav') or 0, 'iva': g('iva') or 0, 'piibb': g('piibb') or 0,
                      'piva': (g('piva') or 0) + (g('pig') or 0), 'total': g('total') or 0})
    tot_existente = fin + 1 <= ws.max_row and isinstance(ws.cell(fin + 1, c['total']).value, (int, float, str)) \
        and ws.cell(fin + 1, c['total']).value not in (None, '')
    encab = ' '.join(str(ws.cell(r, cc).value or '') for r in range(1, hdr) for cc in range(1, ws.max_column + 1))
    return ws, filas, {'hdr': hdr, 'ini': hdr + 1, 'fin': fin, 'cols': c, 'tot_existente': tot_existente, 'encabezado': encab}


def comparar_zetti(a, b):
    if not b or not os.path.exists(b):
        return None
    _, fa, _ = leer_zetti(a); _, fb, _ = leer_zetti(b)
    k = lambda z: (z['fecha'], z['tc'], z['m'], z['raw'], z['cuit'], round(z['total'], 2))
    sa, sb = Counter(map(k, fa)), Counter(map(k, fb))
    return {'igual': sa == sb, 'filas_hoy': len(fa), 'filas_unidad': len(fb),
            'solo_hoy': [list(map(str, x)) for x in (sa - sb)][:20], 'solo_unidad': [list(map(str, x)) for x in (sb - sa)][:20]}


def listados_ncr(unidad, periodo, ncr_dir):
    anios = {periodo[:4], str(int(periodo[:4]) - 1)}
    fs = []
    for a in anios:
        fs += glob.glob(os.path.join(unidad, '*', '01-Impuestos Mensuales', f'{a}*', 'IVA', 'ncr', '*.xlsx'))
        fs += glob.glob(os.path.join(unidad, '*', '01-Impuestos Mensuales', a, f'{a}*', 'IVA', 'ncr', '*.xlsx'))
        fs += glob.glob(os.path.join(unidad, '*', '01-Impuestos Mensuales', f'{a}*', 'ncr', '*.xlsx'))
    fs += glob.glob(os.path.join(ncr_dir, '*.xlsx'))
    return sorted({os.path.normcase(os.path.abspath(f)): f for f in fs
                   if 'COMPRAS' not in os.path.basename(f).upper() and not os.path.basename(f).startswith('~$')}.values())


def buscar_en_listados(archivos, buscadas):
    """(pv, nro) de NC de droguerias buscadas en todos los listados de NCR de la unidad."""
    hall = []
    for f in archivos:
        try:
            ws = openpyxl.load_workbook(f, data_only=True, read_only=True).active
        except Exception as e:
            log(f'no pude leer {f}: {e}'); continue
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r:
                continue
            cand = set()
            for v in r[:6]:
                mm = re.match(r'^0*(\d+)-0*(\d+)$', str(v or ''))
                if mm:
                    cand.add((int(mm.group(1)), int(mm.group(2))))
            if len(r) > 3 and isinstance(r[3], (int, float)) and r[3]:
                cand.add((int(r[2]) if isinstance(r[2], (int, float)) and r[2] else 503, int(r[3])))
            for k in cand & buscadas:
                hall.append({'pv': k[0], 'nro': k[1], 'archivo': f})
    return hall


def rango_monroe(ncr_dir):
    """Fechas del listado de NCR de Monroe de la farmacia (export DMA con 'Numero Formateado')."""
    for f in glob.glob(os.path.join(ncr_dir, '*.xlsx')):
        try:
            ws = openpyxl.load_workbook(f, data_only=True, read_only=True).active
            cab = [_n(v) for v in next(ws.iter_rows(max_row=1, values_only=True))]
        except Exception:
            continue
        if 'numero formateado' in cab and 'fecha' in cab:
            i = cab.index('fecha')
            fs = [r[i] for r in ws.iter_rows(min_row=2, values_only=True) if r and isinstance(r[i], datetime)]
            if fs:
                return os.path.basename(f), min(fs), max(fs), len(fs)
    return None


def servicios_mes_anterior(ruta):
    if not ruta or not os.path.exists(ruta):
        return None
    ws = openpyxl.load_workbook(ruta, data_only=True, read_only=True).active
    return {SERV_CUITS[digitos(r[5])] for r in ws.iter_rows(min_row=2, values_only=True)
            if r and len(r) > 5 and digitos(r[5]) in SERV_CUITS}


# ── armado ─────────────────────────────────────────────────────────────────────────────────────────
def armar(a):
    PERIODO = a.periodo; MM_AAAA = f'{PERIODO[4:]}/{PERIODO[:4]}'
    CUIT = digitos(a.cuit)
    NO_ES_NCR = {tuple(int(x) for x in s.split('-')) for s in re.split(r'[;,\s]+', a.no_es_ncr) if re.fullmatch(r'\d+-\d+', s)}
    FECHA_B = a.fecha_borrador or datetime.now().strftime('%d/%m/%Y %H:%M')
    # de donde salio el libro: el borrador (default) o la DDJJ presentada (lo arma el orquestador)
    FUENTE = a.fuente or f'borrador del Libro Compras exportado de ARCA el {FECHA_B}'

    # 1. Libro Compras del borrador actual (sin NCR)
    portal = leer_borrador(a.borrador)
    PORTAL_ARCA, PORTAL_ARCA_TOT = len(portal), round(sum(r[7] for r in portal), 2)
    fuera = sorted({r[1] for r in portal} - (T_CUADRO & set(TIPOS)))
    if fuera:
        raise Corte('tipo_desconocido', f'Tipos de comprobante que el cuadro de compras no contempla: {fuera}. '
                    'Agregarlos en armar_papel_compras.py (TIPOS / T_*).')
    wsn = openpyxl.load_workbook(a.sin_ncr, data_only=True)
    cab = [c.value for c in wsn['Compras S.NCR'][1]][:32]
    ncr_elim, ncr_imp, elim_drog = set(), [], Counter()
    for r in wsn['Eliminados'].iter_rows(min_row=2, values_only=True):
        if r[0]:
            ncr_elim.add((digitos(r[2]), int(r[3]), int(r[4]))); ncr_imp.append((digitos(r[2]), round(r[5], 2), r[6]))
            elim_drog[str(r[0]).upper()] += 1
    ctrl = {str(r[0]): r for r in wsn['Control'].iter_rows(min_row=2, values_only=True) if r and r[0]} if 'Control' in wsn.sheetnames else {}
    PORTAL_ORIG = ctrl.get('Compras Portal IVA', (None, None))[1]
    siguen = [k for k in ((r[5], r[2], r[3]) for r in portal) if k in ncr_elim]
    if siguen:
        raise Corte('ncr_en_borrador', f'{len(siguen)} NCR conciliadas siguen en el borrador del Libro Compras '
                    f'(ej. {siguen[:5]}). Correr la eliminacion de conciliacion-ncr antes de armar el papel.',
                    ncr=[list(map(str, s)) for s in siguen])
    nc_drog = [r for r in portal if r[1] in T_NCA + [8, 13] and r[5] in DROG]
    listados = listados_ncr(a.unidad, PERIODO, a.ncr_dir)
    hall = buscar_en_listados(listados, {(r[2], r[3]) for r in nc_drog})
    if hall:
        raise Corte('ncr_en_borrador', f'{len(hall)} NC de Suizo/Monroe del borrador figuran en listados de NCR: '
                    + '; '.join(f"{h['pv']}-{h['nro']} ({os.path.basename(h['archivo'])})" for h in hall[:8])
                    + '. Conciliarlas y eliminarlas antes de armar el papel.', ncr=hall)
    log(f'borrador: {PORTAL_ARCA} comprobantes, {pesos(PORTAL_ARCA_TOT)}; NCR eliminadas {len(ncr_elim)}, ninguna en el borrador; '
        f'{len(nc_drog)} NC de droguerias, ninguna en los {len(listados)} listados de NCR')

    # Zetti
    wz, zetti, zi = leer_zetti(a.zetti)
    Z_INI, Z_FIN, ZC = zi['ini'], zi['fin'], zi['cols']
    SOC = re.sub(r'\s+', ' ', next((str(wz.cell(r, cc).value).strip() for r in range(1, zi['hdr'])
                                    for cc in range(2, wz.max_column + 1)
                                    if isinstance(wz.cell(r, cc).value, str) and len(wz.cell(r, cc).value.strip()) > 8
                                    and not re.search(r'fecha|subdiario|raz', wz.cell(r, cc).value, re.I)), '')).strip(' ,')
    m = re.search(r'(\d{2})-?(\d{8})-?(\d)', zi['encabezado'])
    if m and ''.join(m.groups()) != CUIT:
        raise Corte('zetti_otra_sociedad', f'El Zetti es del CUIT {"".join(m.groups())}, no de {CUIT}')
    SOC = SOC or CUIT
    zcmp = comparar_zetti(a.zetti, a.zetti_unidad) if a.zetti_unidad and os.path.abspath(a.zetti_unidad) != os.path.abspath(a.zetti) else None

    # Servicios publicos: los de Zetti tienen que estar en el borrador; si faltan, se agregan a mano (naranja)
    SERVICIOS, MANUALES, serv_dif = [], [], []
    for z in zetti:
        if z['cuit'] not in SERV_CUITS:
            continue
        mm = re.match(r'^0*(\d+)-0*(\d+)$', z['raw'].strip())
        if not mm:
            continue
        pv, nro = int(mm.group(1)), int(mm.group(2))
        p_ = [r for r in portal if r[5] == z['cuit'] and r[2] == pv and r[3] == nro]
        if p_:
            if abs(p_[0][7] - z['total']) > 1 or abs((p_[0][31] or 0) - z['iva']) > 1:
                serv_dif.append(f"{SERV_CUITS[z['cuit']]} {pv:04d}-{nro:08d}: Zetti total {pesos(z['total'])} IVA {pesos(z['iva'])}, "
                                f"Portal total {pesos(p_[0][7])} IVA {pesos(p_[0][31] or 0)}")
            SERVICIOS.append((z['cuit'], pv, nro, z['fecha'], p_[0][0], False))
            continue
        ali = min((27, 21, 10.5), key=lambda x: abs((z['iva'] / z['grav'] * 100 if z['grav'] else 0) - x))
        neto = {10.5: (24, 25), 21: (26, 27), 27: (28, 29)}[ali]
        nueva = [z['fecha'], 1, pv, nro, 80, z['cuit'], SERV_CUITS[z['cuit']] + ' (agregado desde Zetti)', z['total'], 'PES', 1,
                 0, z['exen'], z['iva'], 0, z['piibb'], 0, z['piva'], 0, 0] + [None] * 11 + [z['grav'], z['iva']]
        nueva[neto[0]], nueva[neto[1]] = z['grav'], z['iva']
        portal.append(nueva); MANUALES.append((z['cuit'], pv, nro))
        SERVICIOS.append((z['cuit'], pv, nro, z['fecha'], z['fecha'], True))
    serv_ant = servicios_mes_anterior(a.portal_anterior)
    serv_hoy = {SERV_CUITS[r[5]] for r in portal if r[5] in SERV_CUITS} | {SERV_CUITS[z['cuit']] for z in zetti if z['cuit'] in SERV_CUITS}
    serv_faltan = sorted(serv_ant - serv_hoy) if serv_ant is not None else []

    maestro = cl.leer_maestro(a.maestro)
    for r in portal:
        for i in (10, 11):
            if r[i] is None:
                r[i] = 0
    cuenta_de = {}
    for r in portal:
        if r[5] not in cuenta_de:
            cuenta_de[r[5]] = cl.imputar(maestro.get(r[5]))
    estados = Counter(v[2] for v in cuenta_de.values())
    excluidos = {c: v for c, v in cuenta_de.items() if v[1].startswith('EXCLUIDO')}

    # papel
    wb = openpyxl.load_workbook(a.papel)
    ya = [n for n in HOJAS if n in wb.sheetnames]
    if ya and a.si_existe != 'reemplazar':
        raise Corte('hojas_existen', f'El papel ya tiene {ya}. Usar --si-existe=reemplazar para rehacerlas.')
    for n in ya:
        del wb[n]
    if 'liquidacion' not in wb.sheetnames:
        raise Corte('papel_sin_liquidacion', 'El papel no tiene la hoja "liquidacion" (correr antes el cruce de ventas).')
    VG = next((c.row for c in wb['liquidacion']['A'] if c.value == 'Venta Gravada'), None)
    if not VG:
        raise Corte('papel_sin_liquidacion', 'La hoja "liquidacion" no tiene la fila "Venta Gravada" (prorrateo).')
    pos = wb.sheetnames.index('liquidacion')

    # 2. Compras Zetti: copia tal cual del 5.6.4 (+ fila TOTAL si el reporte no la trae)
    sz = wb.create_sheet('Compras Zetti', pos)
    for row in wz.iter_rows():
        for c in row:
            d = sz.cell(c.row, c.column, c.value)
            if c.has_style:
                d.font = copy.copy(c.font); d.number_format = c.number_format
                d.alignment = copy.copy(c.alignment); d.border = copy.copy(c.border); d.fill = copy.copy(c.fill)
    for col, dim in wz.column_dimensions.items():
        sz.column_dimensions[col].width = dim.width
    num_cols = [ZC[k] for k in ('exen', 'grav', 'iva', 'piibb', 'piva', 'pig', 'total') if ZC.get(k)]
    if not zi['tot_existente']:
        zt = Z_FIN + 2
        sz.cell(zt, 1, 'TOTAL').font = F(b=True)
        sz.cell(zt, 2, f'=COUNTA(A{Z_INI}:A{Z_FIN})').font = F(b=True)
        for j in num_cols:
            c = sz.cell(zt, j, f'=SUM({L(j)}{Z_INI}:{L(j)}{Z_FIN})'); c.font = F(b=True); c.number_format = NUM
    sz.freeze_panes = f'A{Z_INI}'

    # 3. Compras Portal IVA
    sp = wb.create_sheet('Compras Portal IVA', pos + 1)
    P_INI = 3; P_FIN = P_INI + len(portal) - 1
    txt_serv = ' y '.join(SERV_CUITS[s[0]] for s in SERVICIOS) or 'sin servicios publicos'
    sp['A1'] = (f'Comprobantes de Compras - CUIT {CUIT} - Portal IVA ARCA - {MM_AAAA} - DEPURADO ({FUENTE}: '
                f'{PORTAL_ARCA} comprobantes, sin las {len(ncr_elim)} NCR '
                f'({", ".join(f"{n} de {d.title()}" for d, n in elim_drog.items()) or "ninguna"}) eliminadas por la conciliacion; '
                + (f'+ {" y ".join(SERV_CUITS[m_[0]] for m_ in MANUALES)} agregados a mano desde Zetti (en naranja) porque NO estan en el borrador)'
                   if MANUALES else f'servicios: {txt_serv})'))
    sp['A1'].font = F(b=True, size=12)
    cab_ext = cab + ['Cuenta contable', 'Rubro EERR', 'Estado Maestro', 'Importe imputado', 'Clasificacion', 'Fila cuadro (No Grav./Exento)']
    for j, h in enumerate(cab_ext, 1):
        c = sp.cell(2, j, h); c.font = F(b=True, color='FFFFFFFF'); c.fill = HDR_FILL
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = BORDE
    sp.row_dimensions[2].height = 45
    for i, r in enumerate(sorted(portal, key=lambda r: (r[0], r[1], r[2], r[3])), P_INI):
        vals = list(r); vals[1] = TIPOS[int(r[1])]; vals[5] = int(r[5])
        for j, v in enumerate(vals, 1):
            c = sp.cell(i, j, v)
            if j == 1: c.number_format = 'dd/mm/yyyy'
            elif j >= 8 and j not in (9, 10): c.number_format = MON
        cta, rub, est = cuenta_de[r[5]]
        sp.cell(i, 33, cta); sp.cell(i, 34, rub); sp.cell(i, 35, est)
        if est != 'Clasificado':
            for j in (33, 34, 35): sp.cell(i, j).fill = AMARILLO
        if (r[5], int(r[2]), int(r[3])) in MANUALES:
            for j in range(1, 8): sp.cell(i, j).fill = NARANJA
            sp.cell(i, 39, 'AGREGADO A MANO desde Zetti: no esta en el Libro Compras del Portal IVA. Cargarlo en el borrador.')
        sp.cell(i, 36, f'=IF(OR(ABS(SUM(K{i},L{i},T{i},U{i},W{i},Y{i},AA{i},AC{i}))>0.004,ABS(AF{i})>0.004),'
                       f'ROUND(SUM(K{i},L{i},T{i},U{i},W{i},Y{i},AA{i},AC{i}),2),ROUND(H{i}-SUM(N{i}:S{i}),2))').number_format = MON
        sp.cell(i, 37, f'=IF(LEFT(AH{i},8)="EXCLUIDO","EXCLUIDO",IF(AH{i}="CMV","MERCADERIA","SERVICIOS"))')
        sp.cell(i, 38, f'=IF(AK{i}<>"SERVICIOS",AK{i},IF(ABS(AD{i})>0.004,"SERVICIOS AL 27%",IF(ABS(AB{i})>0.004,"SERVICIOS AL 21%",'
                       f'IF(ABS(Z{i})>0.004,"SERVICIOS AL 10,5%","SERVICIOS AL 21%"))))')
    T = P_FIN + 2
    sp.cell(T, 1, 'TOTAL').font = F(b=True)
    sp.cell(T, 2, f'=COUNTA(A{P_INI}:A{P_FIN})').font = F(b=True)
    for j in [8] + list(range(11, 33)) + [36]:
        c = sp.cell(T, j, f'=SUM({L(j)}{P_INI}:{L(j)}{P_FIN})'); c.font = F(b=True); c.number_format = MON
    for j in range(1, 39):
        sp.cell(T, j).fill = SUB_FILL
    anchos = {'A': 11, 'B': 22, 'C': 8, 'D': 11, 'E': 7, 'F': 13, 'G': 38, 'H': 16, 'I': 7, 'J': 7}
    for j in range(1, 38):
        sp.column_dimensions[L(j)].width = anchos.get(L(j), 15)
    sp.column_dimensions['AG'].width = 30; sp.column_dimensions['AH'].width = 26; sp.column_dimensions['AI'].width = 22
    sp.column_dimensions['AK'].width = 14; sp.column_dimensions['AL'].width = 20; sp.column_dimensions['AM'].width = 60
    sp.freeze_panes = 'C3'
    sp.auto_filter.ref = f'A2:AL{P_FIN}'

    prt = []
    for i in range(P_INI, P_FIN + 1):
        g = lambda col: sp[f'{col}{i}'].value or 0
        tipo = int(sp[f'B{i}'].value.split(' - ')[0])
        perc = sum(g(c) for c in ('N', 'O', 'P', 'Q', 'R', 'S'))
        netos = g('K') + g('L') + g('AE')
        neto = netos if (abs(netos) > 0.004 or abs(g('AF')) > 0.004) else round(g('H') - perc, 2)
        prt.append({'fila': i, 'fecha': sp[f'A{i}'].value, 'tipo': tipo, 'pv': int(sp[f'C{i}'].value),
                    'nro': int(sp[f'D{i}'].value), 'cuit': str(sp[f'F{i}'].value), 'prov': sp[f'G{i}'].value,
                    'total': g('H'), 'iva': g('AF'), 'neto': round(neto, 2), 'perc': round(perc, 2),
                    'piibb': g('O'), 'piva': g('Q'), 'pnac': g('N'), 'potros': g('P') + g('R') + g('S')})

    # 4. Cruce
    def pv_nro(s):
        p = str(s or '').split('-')
        if len(p) == 2 and p[0].strip().isdigit() and p[1].strip().isdigit():
            return int(p[0]), int(p[1])
        return None, None
    for z in zetti:
        z['pv'], z['nro'] = pv_nro(z['raw'])
        z['neto'] = round(z['exen'] + z['grav'], 2); z['perc'] = round(z['piibb'] + z['piva'], 2)

    usados_p, estado_z = set(), {}
    idx1 = defaultdict(list); idx2 = defaultdict(list)
    for p in prt:
        tc, le = TC[p['tipo']]
        idx1[(p['cuit'], tc, le, p['pv'], p['nro'])].append(p['fila'])
        idx2[(p['cuit'], p['pv'], p['nro'])].append(p['fila'])
    pfila = {p['fila']: p for p in prt}
    claves_z = set()
    for paso in (1, 2):
        for z in zetti:
            if z['fila'] in estado_z or z['pv'] is None:
                continue
            key = (z['cuit'], z['tc'], z['m'], z['pv'], z['nro']) if paso == 1 else (z['cuit'], z['pv'], z['nro'])
            cand = [f for f in (idx1 if paso == 1 else idx2).get(key, []) if f not in usados_p]
            if cand:
                usados_p.add(cand[0]); estado_z[z['fila']] = ('par', cand[0], paso); claves_z.add(key)
            elif paso == 1 and key in claves_z:
                estado_z[z['fila']] = ('Solo Zetti - duplicado', None, 1)
    for z in zetti:   # 3: NC sin numero util -> CUIT + tipo + importe, fecha mas cercana
        if z['fila'] in estado_z:
            continue
        cand = [p for p in prt if p['fila'] not in usados_p and p['cuit'] == z['cuit']
                and TC[p['tipo']][0] == z['tc'] and abs(p['total'] - z['total']) <= 1]
        if cand:
            p = min(cand, key=lambda p: abs((p['fecha'] - z['fecha']).days))
            usados_p.add(p['fila']); estado_z[z['fila']] = ('par', p['fila'], 3)
    for z in zetti:   # 4: Zetti sin CUIT -> PV + numero + importe
        if z['fila'] in estado_z or z['cuit'] or z['pv'] is None:
            continue
        cand = [p for p in prt if p['fila'] not in usados_p and p['pv'] == z['pv'] and p['nro'] == z['nro']
                and abs(p['total'] - z['total']) <= 1]
        if cand:
            usados_p.add(cand[0]['fila']); estado_z[z['fila']] = ('par', cand[0]['fila'], 4)
    for z in zetti:   # 5: mismo PV + numero + importe con otro CUIT en Zetti
        if z['fila'] in estado_z or z['pv'] is None:
            continue
        cand = [p for p in prt if p['fila'] not in usados_p and p['pv'] == z['pv'] and p['nro'] == z['nro']
                and TC[p['tipo']][0] == z['tc'] and abs(p['total'] - z['total']) <= 1]
        if len(cand) == 1:
            usados_p.add(cand[0]['fila']); estado_z[z['fila']] = ('par', cand[0]['fila'], 5)
    dup_obs = {}
    par_de = {st[1]: f for f, st in estado_z.items() if st[0] == 'par'}
    zfila = {z['fila']: z for z in zetti}
    for z in zetti:   # duplicados en Zetti con el numero incompleto
        if z['fila'] in estado_z:
            continue
        n = digitos(str(z['raw']).split('-')[-1])
        if not n:
            continue
        cand = [p for p in prt if p['fila'] in usados_p and p['nro'] == int(n) and TC[p['tipo']][0] == z['tc']
                and abs(p['total'] - z['total']) <= max(1, abs(p['total']) * 0.01)]
        if len(cand) == 1:
            p = cand[0]; z2 = zfila[par_de[p['fila']]]
            txt = (f"Cargado dos veces en Zetti: ya esta como {z2['tc']}-{z2['m']} {z2['raw']} de {z2['prov']} (fila {z2['fila']}), "
                   f"emparejado con el Portal {SIGLA[p['tipo']]} {p['pv']:05d}-{p['nro']:08d} de {p['prov']}. Esta carga "
                   + (f"tiene letra {z['m']} (no fiscal)" if z['m'] != z2['m'] else 'tiene el numero incompleto'))
            if z['cuit'] != p['cuit']:
                txt += f" y otro proveedor ({z['prov']} {z['cuit']})"
            if abs(z['total'] - p['total']) > 0.01:
                txt += f" e importe distinto (Zetti {z['total']:,.2f} vs Portal {p['total']:,.2f})"
            estado_z[z['fila']] = ('Solo Zetti - duplicado', None, 1)
            dup_obs[z['fila']] = txt + '. Dar de baja en Zetti.'

    def dif(x, y): return abs(round(x - y, 2)) > 1
    filas_cruce = []
    for z in zetti:
        st = estado_z.get(z['fila'])
        comp_z = f"{z['tc']}-{z['m']} {z['raw']}"
        if st and st[0] == 'par':
            p = pfila[st[1]]
            obs = []
            if st[2] == 2:
                obs.append(f"Tipo/letra distinto: Zetti {z['tc']}-{z['m']}, Portal {SIGLA[p['tipo']]}")
            if st[2] == 3:
                obs.append('Emparejado por CUIT + tipo + importe (Zetti sin numero de comprobante util)')
            if st[2] == 4:
                obs.append(f"Zetti sin CUIT ni proveedor: emparejado por PV + numero + importe. Completar el proveedor en Zetti ({p['prov']})")
            if st[2] == 5:
                obs.append(f"CUIT distinto: en Zetti esta como {z['prov']} ({z['cuit']}) y en el Portal es {p['prov']} ({p['cuit']}); "
                           "mismo PV, numero e importe. Corregir el proveedor en Zetti")
            if p['tipo'] == 111 and st[2] == 1:
                obs.append('Portal: 111 - Tique Factura C; Zetti la carga como FC-C')
            d = []
            if dif(z['neto'], p['neto']): d.append(f"Neto {z['neto'] - p['neto']:,.2f}")
            if dif(z['iva'], p['iva']): d.append(f"IVA {z['iva'] - p['iva']:,.2f}")
            if dif(z['perc'], p['perc']): d.append(f"Percep. {z['perc'] - p['perc']:,.2f}")
            if dif(z['total'], p['total']): d.append(f"Total {z['total'] - p['total']:,.2f}")
            cuadra_z = abs(z['neto'] + z['iva'] + z['perc'] - z['total']) <= 1
            if d:
                txt = 'Diferencias (Zetti - Portal): ' + '; '.join(d)
                if not dif(z['total'], p['total']):
                    txt += '. El total coincide: es distinta apertura entre conceptos'
                if not cuadra_z:
                    txt += f". Zetti no cierra: Exen+Grav+IVA+Percep - Total = {z['neto'] + z['iva'] + z['perc'] - z['total']:,.2f}"
                obs.append(txt)
            estado = 'Con diferencia' if d or st[2] == 5 else 'Coincide'
            filas_cruce.append([estado, p['fecha'], int(p['cuit']), p['prov'], f"{SIGLA[p['tipo']]} {p['pv']:05d}-{p['nro']:08d}",
                                comp_z, z['fila'], p['fila'], z['neto'], p['neto'], z['iva'], p['iva'], z['perc'], p['perc'],
                                z['total'], p['total'], z['piibb'], z['piva'], p['piibb'], p['piva'], p['pnac'], p['potros'],
                                ' | '.join(obs)])
        else:
            if st:
                estado, obs = st[0], dup_obs.get(z['fila'], 'El mismo comprobante ya esta cargado en Zetti y emparejado con el Portal: dar de baja en Zetti')
            elif (z['pv'] is not None and (z['cuit'], z['pv'], z['nro']) in ncr_elim) or (
                    z['pv'] is None and z['tc'] == 'NC' and any(c == z['cuit'] and abs(abs(i) - abs(z['total'])) <= 1
                                                                 for c, i, _ in ncr_imp)):
                estado, obs = 'Solo Zetti - NCR eliminada del Portal', ('NCR de recupero eliminada del Libro Compras '
                                                                        '(conciliacion NCR); en Zetti sigue cargada')
            elif z['m'] in ('X', 'R'):
                estado, obs = f"Solo Zetti - no fiscal ({z['m']})", 'Comprobante no fiscal: no va al Libro IVA'
            else:
                estado, obs = 'Solo en Zetti', 'No esta en el Libro Compras del Portal IVA'
            filas_cruce.append([estado, z['fecha'], int(z['cuit']) if z['cuit'] else None, z['prov'],
                                f"{z['tc']}-{z['m']} {z['raw']}", comp_z, z['fila'], None, z['neto'], 0, z['iva'], 0,
                                z['perc'], 0, z['total'], 0, z['piibb'], z['piva'], 0, 0, 0, 0, obs])
    for p in prt:
        if p['fila'] in usados_p:
            continue
        obs = ('Liquidacion bancaria: normalmente no se carga en Zetti, se toma del Portal' if p['tipo'] == 63
               else 'No esta cargado en Zetti')
        filas_cruce.append(['Solo en Portal', p['fecha'], int(p['cuit']), p['prov'],
                            f"{SIGLA[p['tipo']]} {p['pv']:05d}-{p['nro']:08d}", None, None, p['fila'], 0, p['neto'], 0,
                            p['iva'], 0, p['perc'], 0, p['total'], 0, 0, p['piibb'], p['piva'], p['pnac'], p['potros'], obs])
    ORDEN = ['Con diferencia', 'Solo en Portal', 'Solo en Zetti', 'Solo Zetti - NCR eliminada del Portal',
             'Solo Zetti - duplicado', 'Solo Zetti - no fiscal (X)', 'Solo Zetti - no fiscal (R)', 'Coincide']
    filas_cruce.sort(key=lambda f: (ORDEN.index(f[0]), f[1], f[4]))
    assert sum(1 for f in filas_cruce if f[7]) == len(prt)
    assert sum(1 for f in filas_cruce if f[6]) == len(zetti)

    sc = wb.create_sheet('Cruce compras', pos + 2)
    sc['A1'] = f'Cruce Compras Zetti vs Portal IVA - {SOC} ({fcuit(CUIT)}) - Período {MM_AAAA}'
    sc['A1'].font = F(b=True, size=14, color='FFFFFFFF'); sc['A1'].fill = TIT_FILL
    sc['A2'] = ('Neto = Exen + Grav (Zetti) vs No Gravado + Exento + Neto Gravado (Portal). Diferencia = Zetti - Portal. '
                f'Tolerancia $1. Portal = {FUENTE} (sin las {len(ncr_elim)} NCR).')
    sc['A2'].font = F(color='FF404040')
    hc = ['Estado', 'Fecha', 'CUIT', 'Proveedor', 'Comprobante', 'Comp. Zetti', 'Fila Zetti', 'Fila Portal',
          'Neto Zetti', 'Neto Portal', 'Dif. Neto', 'IVA Zetti', 'IVA Portal', 'Dif. IVA', 'Percep. Zetti',
          'Percep. Portal', 'Dif. Percep.', 'Total Zetti', 'Total Portal', 'Dif. Total', 'P/R IIBB Zetti',
          'P/R IVA Zetti', 'P.IIBB Portal', 'P.IVA Portal', 'P.Gan/Otros Imp.Nac. Portal', 'Otros trib. Portal',
          'Observaciones']
    for j, h in enumerate(hc, 1):
        c = sc.cell(3, j, h); c.font = F(b=True, color='FFFFFFFF'); c.fill = HDR_FILL; c.border = BORDE
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    sc.row_dimensions[3].height = 32
    C_INI = 4
    for i, f in enumerate(filas_cruce, C_INI):
        vals = f[:10] + [f'=N(I{i})-N(J{i})'] + f[10:12] + [f'=N(L{i})-N(M{i})'] + f[12:14] + [f'=N(O{i})-N(P{i})'] \
            + f[14:16] + [f'=N(R{i})-N(S{i})'] + f[16:]
        for j, v in enumerate(vals, 1):
            c = sc.cell(i, j, v); c.border = BORDE
            if j == 2: c.number_format = 'dd/mm/yyyy'
            elif 9 <= j <= 26: c.number_format = NUM
        if f[0] != 'Coincide':
            sc.cell(i, 1).fill = ROJO if f[0] == 'Con diferencia' else AMARILLO
    C_FIN = C_INI + len(filas_cruce) - 1
    CT = C_FIN + 1
    sc.cell(CT, 1, 'TOTAL')
    for j in range(9, 27):
        sc.cell(CT, j, f'=SUBTOTAL(9,{L(j)}{C_INI}:{L(j)}{C_FIN})').number_format = NUM
    for j in range(1, 28):
        c = sc.cell(CT, j); c.font = F(b=True); c.fill = SUB_FILL; c.border = BORDE
    for j, w in enumerate([24, 11, 13, 36, 22, 20, 8, 8] + [14] * 18 + [80], 1):
        sc.column_dimensions[L(j)].width = w
    sc.freeze_panes = 'E4'
    sc.auto_filter.ref = f'A3:AA{C_FIN}'

    # 5. Resumen cruce compras
    sr = wb.create_sheet('Resumen cruce compras', pos + 3)
    sr['A1'] = f'Resumen del cruce Compras Zetti vs Portal IVA - {SOC} {MM_AAAA}'
    sr['A1'].font = F(b=True, size=14, color='FFFFFFFF'); sr['A1'].fill = TIT_FILL
    def hdr(ws, r, vals, c0=1):
        for j, h in enumerate(vals, c0):
            c = ws.cell(r, j, h); c.font = F(b=True, color='FFFFFFFF'); c.fill = HDR_FILL; c.border = BORDE
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    def fila(ws, r, vals, fmt=NUM, bold=False, fill=None, c0=1):
        for j, v in enumerate(vals, c0):
            c = ws.cell(r, j, v); c.border = BORDE; c.font = F(b=bold)
            if j > c0: c.number_format = fmt
            if fill: c.fill = fill
    rng = lambda col: f"'Cruce compras'!${col}${C_INI}:${col}${C_FIN}"
    hdr(sr, 3, ['Estado', 'Cant.', 'Total Zetti', 'Total Portal', 'Dif. Total'])
    r = 4
    for e in [e for e in ORDEN if any(f[0] == e for f in filas_cruce)]:
        fila(sr, r, [e, f"=COUNTIF({rng('A')},A{r})", f"=SUMIF({rng('A')},A{r},{rng('R')})",
                     f"=SUMIF({rng('A')},A{r},{rng('S')})", f'=C{r}-D{r}'])
        sr.cell(r, 2).number_format = '0'; r += 1
    E_FIN = r - 1
    fila(sr, r, ['TOTAL', f'=SUM(B4:B{E_FIN})', f'=SUM(C4:C{E_FIN})', f'=SUM(D4:D{E_FIN})', f'=C{r}-D{r}'],
         bold=True, fill=SUB_FILL); sr.cell(r, 2).number_format = '0'
    E_TOT = r
    r += 2
    sr.cell(r, 1, 'Control contra las hojas de compras').font = F(b=True); r += 1
    hdr(sr, r, ['Fuente', 'Cant.', 'Exen / No Grav.+Exento', 'Neto Gravado', 'IVA', 'P.IIBB', 'P.IVA',
                'P.Gan/Otros Imp.Nac.', 'Otros tributos', 'Total']); r += 1
    zr = lambda k: f"'Compras Zetti'!${L(ZC[k])}${Z_INI}:${L(ZC[k])}${Z_FIN}"
    pr = lambda col: f"'Compras Portal IVA'!${col}${P_INI}:${col}${P_FIN}"
    R_Z = r
    fila(sr, r, ['Zetti (5.6.4)', f"=COUNTA({zr('fecha')})", f"=SUM({zr('exen')})", f"=SUM({zr('grav')})", f"=SUM({zr('iva')})",
                 f"=SUM({zr('piibb')})", f"=SUM({zr('piva')})", f"=SUM({zr('pig')})" if ZC.get('pig') else 0, 0,
                 f"=SUM({zr('total')})"]); r += 1
    R_P = r
    fila(sr, r, ['Portal IVA (libro depurado)', f"=COUNTA({pr('A')})", f"=SUM({pr('K')})+SUM({pr('L')})",
                 f"=SUM({pr('AE')})", f"=SUM({pr('AF')})", f"=SUM({pr('O')})", f"=SUM({pr('Q')})", f"=SUM({pr('N')})",
                 f"=SUM({pr('P')})+SUM({pr('R')})+SUM({pr('S')})", f"=SUM({pr('H')})"]); r += 1
    fila(sr, r, ['Diferencia Zetti - Portal'] + [f'={L(j)}{R_Z}-{L(j)}{R_P}' for j in range(2, 11)], bold=True, fill=SUB_FILL)
    for rr in (R_Z, R_P, r): sr.cell(rr, 2).number_format = '0'
    r += 2
    fila(sr, r, ['Control: TOTAL del cruce - hojas (tiene que dar 0)', None, f'=ROUND(C{E_TOT}-J{R_Z},2)',
                 f'=ROUND(D{E_TOT}-J{R_P},2)'], bold=True)
    R_CTRL = r
    r += 2
    sr.cell(r, 1, 'Proveedores con diferencias o faltantes').font = F(b=True); r += 1
    hdr(sr, r, ['Proveedor (CUIT)', 'Cant.', 'Dif. Total (Zetti - Portal)', 'Estados']); r += 1
    agr = OrderedDict()
    for f in filas_cruce:
        if f[0] == 'Coincide': continue
        g_ = agr.setdefault(f[2], {'prov': f[3], 'n': 0, 'dif': 0.0, 'est': Counter()})
        g_['n'] += 1; g_['dif'] += f[14] - f[15]; g_['est'][f[0]] += 1
    for cuit, g_ in sorted(agr.items(), key=lambda x: -abs(x[1]['dif'])):
        fila(sr, r, [f"{g_['prov']} ({cuit})", g_['n'], round(g_['dif'], 2),
                     ', '.join(f'{e}: {n}' for e, n in g_['est'].most_common())])
        sr.cell(r, 2).number_format = '0'; sr.cell(r, 4).number_format = 'General'; r += 1
    r += 1

    # hallazgos
    cnt = Counter(f[0] for f in filas_cruce)
    tot = lambda e, i: sum(f[i] for f in filas_cruce if f[0] == e)
    h = []
    def ej(e, n=4):
        xs = [f for f in filas_cruce if f[0] == e][:n]
        return '; '.join(f"{f[3]} {f[4]} ${f[14] or f[15]:,.2f}" for f in xs) + ('...' if cnt[e] > n else '')
    serv_f = [f for f in filas_cruce if str(f[2]) in SERV_CUITS and f[0] != 'Solo en Portal']
    man = [f for f in serv_f if (str(f[2]), *map(int, f[4].split()[-1].split('-'))) in MANUALES]
    if man:
        h.append('FALTAN CARGAR EN EL BORRADOR DEL PORTAL IVA los servicios del mes que estan en Zetti: ' + '; '.join(
            f"{SERV_CUITS[str(f[2])]} {f[4]} del {f[1]:%d/%m/%Y} total {pesos(f[14])} (IVA {pesos(f[10])})" for f in man) +
            ". En este papel se agregaron a mano en 'Compras Portal IVA' (filas en naranja) para que entren al cuadro y al "
            "prorrateo de servicios; hay que cargarlos en el Libro Compras de ARCA.")
    cargados = [f for f in serv_f if f not in man]
    if cargados:
        h.append('Servicios del mes en el Libro Compras del Portal IVA y controlados contra Zetti: ' + '; '.join(
            f"{SERV_CUITS[str(f[2])]} {f[4]} total {pesos(f[14])} (IVA {pesos(f[10])})" for f in cargados) +
            f". IVA {pesos(sum(f[10] for f in cargados))}, que entra al prorrateo de servicios.")
    for t in serv_dif:
        h.append(f'Servicio con importes distintos entre Zetti y el Portal: {t}. Revisar la factura.')
    for c_, pv, nro, fz, fp, manual in SERVICIOS:
        if not manual and fz.date() != fp.date():
            h.append(f"{SERV_CUITS[c_]} {pv:04d}-{nro:08d}: en el Portal tiene fecha {fp:%d/%m/%Y} y en Zetti {fz:%d/%m/%Y}. "
                     "Importes iguales; revisar la fecha de emision en la factura (no cambia el periodo de computo).")
    if serv_faltan:
        h.append(f"Servicios que estaban en el mes anterior y este mes no estan ni en Zetti ni en el Portal: {', '.join(serv_faltan)}. "
                 "Pedir la factura del mes.")
    sin_serv = sorted(set(SERV_CUITS.values()) - serv_hoy - set(serv_faltan))
    if sin_serv:
        h.append(f"Sin {', '.join(sin_serv)} en el mes" + (' (tampoco en el mes anterior).' if serv_ant is not None else '.'))
    revisar, confirmadas = [], []
    for f in filas_cruce:
        if f[0] == 'Solo en Portal' and str(f[2]) in DROG and f[4].startswith('NC') and f[15] < -NC_GRANDE:
            k = tuple(int(x) for x in f[4].split()[-1].split('-'))
            (confirmadas if k in NO_ES_NCR else revisar).append(f)
    pvnro = lambda f: '-'.join(str(int(x)) for x in f[4].split()[-1].split('-'))
    for f in revisar:
        h.append(f"REVISAR NC de {DROG[str(f[2])]} {f[4]} del {f[1]:%d/%m/%Y} por {pesos(f[15])}: sigue en el Libro Compras, no esta en "
                 "Zetti y no figura en ningun listado de NCR de la unidad (ningun mes, ninguna farmacia). Por el monto puede ser una "
                 "NCR de recupero: pedir el comprobante y, si es NCR, eliminarla del borrador (cambia el cuadro de compras: NC de "
                 f"mercaderia). Si no lo es, volver a correr con --no-es-ncr={pvnro(f)}.")
    for f in confirmadas:
        h.append(f"NC de {DROG[str(f[2])]} {f[4]} del {f[1]:%d/%m/%Y} por {pesos(f[15])}: confirmado por el estudio que NO es NCR "
                 "(tampoco figura en ningun listado de NCR de la unidad). Queda en el Libro Compras como NC de mercaderia; "
                 "no esta en Zetti: cargarla en Zetti.")
    if nc_drog:
        enz = [p for p in prt if p['tipo'] in T_NCA + [8, 13] and p['cuit'] in DROG and p['fila'] in usados_p]
        rm = rango_monroe(a.ncr_dir)
        txt_m = (f"Monroe {rm[0]} del {rm[1]:%d/%m} al {rm[2]:%d/%m}: {rm[3]} NCR" if rm else 'sin listado de Monroe en ncr/')
        h.append(f"NC de Suizo y Monroe que quedan en el Libro Compras ({len(nc_drog)}, {pesos(sum(r[7] for r in nc_drog))}): ninguna "
                 f"figura en los {len(listados)} listados de NCR de la unidad ({txt_m}; eliminadas por la conciliacion: "
                 f"{', '.join(f'{n} de {d.title()}' for d, n in elim_drog.items()) or 'ninguna'}, ninguna de vuelta en el borrador), "
                 f"asi que se tratan como NC comerciales (devoluciones / bonificaciones) y se dejan. {len(enz)} estan tambien en "
                 f"Zetti; {len(nc_drog) - len(enz)} solo en el Portal: cargarlas en Zetti.")
    if cnt['Solo Zetti - NCR eliminada del Portal']:
        h.append(f"NCR eliminadas del Libro Compras que Zetti SI tiene cargadas ({cnt['Solo Zetti - NCR eliminada del Portal']}, "
                 f"{pesos(tot('Solo Zetti - NCR eliminada del Portal', 14))}): {ej('Solo Zetti - NCR eliminada del Portal')}. "
                 "En el Portal se sacaron por la conciliacion NCR; en Zetti siguen (es esperable si Zetti las usa para la cuenta corriente).")
    if cnt['Solo Zetti - duplicado']:
        h.append(f"Comprobantes duplicados en Zetti ({cnt['Solo Zetti - duplicado']}, {pesos(tot('Solo Zetti - duplicado', 14))}): "
                 f"{ej('Solo Zetti - duplicado')}. Hay que darlos de baja en Zetti.")
    if cnt['Solo en Zetti']:
        h.append(f"Solo en Zetti, no estan en el Portal ({cnt['Solo en Zetti']}, {pesos(tot('Solo en Zetti', 14))}): {ej('Solo en Zetti', 6)}.")
    bancos = [f for f in filas_cruce if f[0] == 'Solo en Portal' and f[4].startswith('LQ')]
    otros_p = [f for f in filas_cruce if f[0] == 'Solo en Portal' and not f[4].startswith('LQ')]
    if bancos:
        h.append(f"Liquidaciones bancarias solo en el Portal ({len(bancos)}, {pesos(sum(f[15] for f in bancos))}): "
                 "normalmente no se cargan en Zetti, se toman del Portal.")
    if otros_p:
        h.append(f"Otros comprobantes solo en el Portal ({len(otros_p)}, {pesos(sum(f[15] for f in otros_p))}): "
                 + '; '.join(f"{f[3]} {f[4]} ${f[15]:,.2f}" for f in otros_p[:8]) + ('...' if len(otros_p) > 8 else '') + '.')
    cd = [f for f in filas_cruce if f[0] == 'Con diferencia']
    if cd:
        soloap = [f for f in cd if abs(f[14] - f[15]) <= 1]
        h.append(f"Con diferencia ({len(cd)}): en {len(soloap)} el total coincide y solo cambia la apertura entre neto, IVA y "
                 f"percepciones; en {len(cd) - len(soloap)} difiere el total (dif. {pesos(sum(f[14] - f[15] for f in cd))}). "
                 "Detalle en la columna Observaciones de 'Cruce compras'.")
    mon = [f for f in cd if str(f[2]) == '30517059095']
    if mon:
        _dp = sum(f[12] - f[13] for f in mon)
        h.append(f"Monroe ({len(mon)} comprobante/s con diferencia): Zetti abre distinto el neto, el IVA y las percepciones "
                 f"(percepciones {pesos(abs(_dp))} {'mas' if _dp > 0 else 'menos'} que el Portal). Es la apertura del reporte de Zetti; "
                 "para el IVA y las percepciones vale el Portal.")
    cero = [f for f in filas_cruce if f[0] == 'Solo en Zetti' and abs(f[14]) < 0.01]
    if cero:
        h.append('En Zetti hay comprobantes en $ 0 que no son fiscales: ' + '; '.join(f"{f[3].strip()} {f[4]}" for f in cero) + '.')
    if zcmp and not zcmp['igual']:
        h.insert(0, f"El Zetti bajado hoy ({zcmp['filas_hoy']} comprobantes) NO es igual al de la unidad ({zcmp['filas_unidad']}): "
                 "el papel usa el de hoy. Revisar que cambio en Zetti.")
    sr.cell(r, 1, 'HALLAZGOS').font = F(b=True); r += 1
    for n, t in enumerate(h, 1):
        sr.cell(r, 1, f'{n}. {t}').alignment = Alignment(wrap_text=True, vertical='top')
        sr.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
        sr.row_dimensions[r].height = max(30, 15 * (len(t) // 150 + 1)); r += 1
    r += 1
    sr.cell(r, 1, 'METODOLOGIA').font = F(b=True); r += 1
    txt_z = ('bajado de nuevo y sin cambios contra el de la unidad' if zcmp and zcmp['igual'] else
             'bajado de nuevo; DISTINTO al de la unidad' if zcmp else 'bajado de Zetti')
    for t in [f'Fuentes: hoja "Compras Zetti" (T&S Web 5.6.4 de {SOC}, {len(zetti)} comprobantes del {MM_AAAA}, {txt_z}) y hoja '
              f'"Compras Portal IVA" (Libro Compras {MM_AAAA} del Portal IVA'
              + (f': {PORTAL_ORIG} comprobantes' if PORTAL_ORIG else '')
              + f' menos las {len(ncr_elim)} NCR conciliadas y eliminadas del borrador. Se usa el {FUENTE} '
              f'(solo lectura, sin importar): {PORTAL_ARCA} comprobantes, {pesos(PORTAL_ARCA_TOT)}, ninguna de las NCR'
              + (f'; + {len(MANUALES)} servicios agregados a mano desde Zetti' if MANUALES else '') + ').',
              'Clave: CUIT emisor + Tipo (FC/NC/ND/LQ) + Letra + Punto de Venta + Numero. La razon social no se usa. '
              'Si no empareja: 2) CUIT + PV + Numero sin mirar el tipo; 3) CUIT + tipo + importe total, fecha mas cercana; '
              '4) Zetti sin CUIT: PV + Numero + importe; 5) mismo tipo + PV + Numero + importe con otro CUIT en Zetti. Las cargas de Zetti con numero incompleto cuyo comprobante del Portal ya '
              'esta emparejado con otra carga se marcan como duplicado.',
              'Neto = Exen + Grav (Zetti) contra No Gravado + Exento + Neto Gravado (Portal); en Facturas C y tiques del Portal '
              'el neto es el total menos percepciones. Percepciones: Zetti P/R IIBB + P/R IVA (+ P/R IG) contra todas las del Portal '
              '(IIBB, IVA, Otros Imp. Nac., municipales, internos y otros tributos). Tolerancia $1.']:
        sr.cell(r, 1, t).alignment = Alignment(wrap_text=True, vertical='top')
        sr.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10); sr.row_dimensions[r].height = 32; r += 1
    for j, w in enumerate([52, 8, 20, 20, 16, 15, 15, 18, 15, 18], 1):
        sr.column_dimensions[L(j)].width = w

    # 6. Clasificacion compras
    sk = wb.create_sheet('Clasificacion compras', pos + 4)
    sk['A1'] = f'Clasificacion de compras por proveedor - {SOC} - {MM_AAAA}'
    sk['A1'].font = F(b=True, size=14, color='FFFFFFFF'); sk['A1'].fill = TIT_FILL
    sk['A2'] = ('Maestro de Proveedores (0 - ESTUDIO, hoja MAESTRO) aplicado al Libro Compras depurado. Importe imputado = '
                'neto (No Gravado + Exento + Netos Gravados); sin IVA discriminado, total menos percepciones. '
                'Rubro CMV = MERCADERIA, resto = SERVICIOS. Amarillo: proveedor sin clasificacion firme en el Maestro.')
    sk['A2'].font = F(color='FF404040', size=10); sk['A2'].alignment = Alignment(wrap_text=True)
    sk.merge_cells('A2:K2'); sk.row_dimensions[2].height = 30
    r = 4
    sk.cell(r, 1, 'Por rubro del EERR').font = F(b=True); r += 1
    hdr(sk, r, ['Rubro EERR', 'Merc./Serv.', 'Comprobantes', 'Importe imputado', 'IVA', 'Total']); r += 1
    orden_rub = sorted({v[1] for v in cuenta_de.values()}, key=lambda x: (x.startswith('EXCLUIDO'), x != 'CMV', x))
    R0 = r
    for rub in orden_rub:
        fila(sk, r, [rub, f'=IF(A{r}="CMV","MERCADERIA","SERVICIOS")', f"=COUNTIF({pr('AH')},A{r})",
                     f"=SUMIF({pr('AH')},A{r},{pr('AJ')})", f"=SUMIF({pr('AH')},A{r},{pr('AF')})",
                     f"=SUMIF({pr('AH')},A{r},{pr('H')})"])
        sk.cell(r, 2).number_format = 'General'; sk.cell(r, 3).number_format = '0'; r += 1
    fila(sk, r, ['TOTAL', None, f'=SUM(C{R0}:C{r-1})', f'=SUM(D{R0}:D{r-1})', f'=SUM(E{R0}:E{r-1})', f'=SUM(F{R0}:F{r-1})'],
         bold=True, fill=SUB_FILL); sk.cell(r, 3).number_format = '0'
    RK_TOT = r; r += 1
    for clase in ('MERCADERIA', 'SERVICIOS'):
        fila(sk, r, [f'  de los cuales {clase}', None] + [f'=SUMIF(B{R0}:B{RK_TOT-1},"{clase}",{col}{R0}:{col}{RK_TOT-1})' for col in 'CDEF'])
        sk.cell(r, 3).number_format = '0'; r += 1
    fila(sk, r, ['Control vs hoja Compras Portal IVA (tiene que dar 0)', None, f"=C{RK_TOT}-COUNTA({pr('A')})",
                 f"=ROUND(D{RK_TOT}-SUM({pr('AJ')}),2)", f"=ROUND(E{RK_TOT}-SUM({pr('AF')}),2)",
                 f"=ROUND(F{RK_TOT}-SUM({pr('H')}),2)"], bold=True); sk.cell(r, 3).number_format = '0'
    RK_CTRL = r
    r += 3
    sk.cell(r, 1, 'Por proveedor').font = F(b=True); r += 1
    hdr(sk, r, ['Proveedor', 'CUIT', 'Cuenta contable', 'Rubro EERR', 'Estado Maestro', 'Merc./Serv.', 'Comprobantes',
                'Importe imputado', 'IVA', 'Total', 'Que hacer']); r += 1
    agp = defaultdict(lambda: {'n': 0, 'imp': 0.0, 'den': None})
    kk = {cl.norm(c).lower(): c for c in cab}
    for p in portal:
        g_ = agp[p[5]]; g_['n'] += 1; g_['den'] = g_['den'] or p[6]
        g_['imp'] += cl.importe_imputado(dict(zip(cab, p)), kk)
    RP0 = r
    for c_, g_ in sorted(agp.items(), key=lambda x: (orden_rub.index(cuenta_de[x[0]][1]), -abs(x[1]['imp']))):
        cta, rub, est = cuenta_de[c_]
        den = (maestro.get(c_) or {}).get('razon') or g_['den']
        fila(sk, r, [den, int(c_), cta, rub, est, 'MERCADERIA' if rub == 'CMV' else 'SERVICIOS',
                     f"=COUNTIF({pr('F')},B{r})", f"=SUMIF({pr('F')},B{r},{pr('AJ')})",
                     f"=SUMIF({pr('F')},B{r},{pr('AF')})", f"=SUMIF({pr('F')},B{r},{pr('H')})",
                     None if est == 'Clasificado' else cl.accion(est)])
        for j in (2, 3, 4, 5, 6, 11): sk.cell(r, j).number_format = 'General'
        sk.cell(r, 2).number_format = '0'; sk.cell(r, 7).number_format = '0'
        if est != 'Clasificado':
            for j in range(1, 12): sk.cell(r, j).fill = AMARILLO
        r += 1
    fila(sk, r, ['TOTAL', None, None, None, None, None, f'=SUM(G{RP0}:G{r-1})', f'=SUM(H{RP0}:H{r-1})', f'=SUM(I{RP0}:I{r-1})',
                 f'=SUM(J{RP0}:J{r-1})', None], bold=True, fill=SUB_FILL); sk.cell(r, 7).number_format = '0'
    for j, w in enumerate([46, 13, 30, 26, 22, 13, 13, 18, 16, 18, 50], 1):
        sk.column_dimensions[L(j)].width = w
    sk.freeze_panes = 'A4'

    # 7. Cuadro de Compras + Prorrateo IVA Compras de Servicios, al final de 'Compras Portal IVA'
    s0 = T + 3
    H0 = s0 + 1
    def cel(r, col, v, b=False, fmt=None, fill=None, borde=True, color=None, it=False, sz_=11):
        c = sp[f'{col}{r}']; c.value = v; c.font = F(b=b, color=color, i=it, size=sz_)
        if fmt: c.number_format = fmt
        if fill: c.fill = fill
        if borde: c.border = BORDE
        return c
    cel(s0, 'A', 'Compras', b=True, color='FFFFFFFF', fill=HDR_FILL, borde=False)
    for col in 'BCDEFG': sp[f'{col}{s0}'].fill = HDR_FILL
    for j, hh in enumerate(['Tipo de Comprobante', 'Neto Gravado', 'No Gravado', 'Exento', 'IVA', 'Total', 'Clasificación'], 1):
        cel(H0, L(j), hh, b=True)
    sp[f'G{H0}'].alignment = Alignment(horizontal='center')
    cel(H0, 'I', 'Prorrateo IVA Compras de Servicios', b=True); cel(H0, 'J', None, b=True)
    sp.merge_cells(f'I{H0}:J{H0}'); sp[f'I{H0}'].alignment = Alignment(horizontal='center')
    P = lambda col: f'${col}${P_INI}:${col}${P_FIN}'
    def S(col, tipo, clas=None):
        return f'SUMIFS({P(col)},{P("B")},"{TIPOS[tipo]}"' + (f',{P("AK")},"{clas}")' if clas else ')')
    def SG(col, tipos, clas=None): return '+'.join(S(col, t, clas) for t in tipos)
    def SAL(col, fila_):
        return '+'.join(f'SUMIFS({P(col)},{P("B")},"{TIPOS[t]}",{P("AL")},"{fila_}")' for t in T_FA)
    NETOS_C, IVAS_C = ('T', 'U', 'W', 'Y', 'AA', 'AC'), ('V', 'X', 'Z', 'AB', 'AD')
    def todos(cols, tipos, clas=None): return '+'.join(SG(c, tipos, clas) for c in cols)
    r1 = H0 + 1
    filas = [
        ('1 - Factura A  al 10,5%', '=' + SG('Y', T_FA, 'MERCADERIA'), None, None, '=' + SG('Z', T_FA, 'MERCADERIA'), 'MERCADERIA'),
        ('1 - Factura A  al  21%', '=' + SG('AA', T_FA, 'MERCADERIA'), '=' + SG('K', T_FA, 'MERCADERIA'), '=' + SG('L', T_FA, 'MERCADERIA'),
         '=' + SG('AB', T_FA, 'MERCADERIA'), 'MERCADERIA'),
        ('1 - Factura A', '=' + SG('Y', T_FA, 'SERVICIOS'), '=' + SAL('K', 'SERVICIOS AL 10,5%'), '=' + SAL('L', 'SERVICIOS AL 10,5%'),
         '=' + SG('Z', T_FA, 'SERVICIOS'), 'SERVICIOS AL 10,5%'),
        ('1 - Factura A', '=' + SG('AA', T_FA, 'SERVICIOS'), '=' + SAL('K', 'SERVICIOS AL 21%'), '=' + SAL('L', 'SERVICIOS AL 21%'),
         '=' + SG('AB', T_FA, 'SERVICIOS'), 'SERVICIOS AL 21%'),
        ('1 - Factura A', '=' + SG('AC', T_FA, 'SERVICIOS'), '=' + SAL('K', 'SERVICIOS AL 27%'), '=' + SAL('L', 'SERVICIOS AL 27%'),
         '=' + SG('AD', T_FA, 'SERVICIOS'), 'SERVICIOS AL 27%'),
        ('6 - Factura B', None, None, None, None, None),
        ('2 - Nota de Débito A', '=' + todos(NETOS_C, T_ND), '=' + SG('K', T_ND), '=' + SG('L', T_ND), '=' + todos(IVAS_C, T_ND), None),
        ('11 - Factura C', None, None, None, None, None),
    ]
    for k, (a_, b, c, d, e, g) in enumerate(filas):
        rr = r1 + k
        cel(rr, 'A', a_)
        for col, v in zip('BCDE', (b, c, d, e)):
            cel(rr, col, v, fmt=NUM)
        cel(rr, 'F', f'=SUM(B{rr}:E{rr})', fmt=NUM)
        cel(rr, 'G', g)
    rFB, rFC, rN = r1 + 5, r1 + 7, r1 + 7
    sp[f'F{rFB}'] = f'=SUMIFS({P("AJ")},{P("B")},"{TIPOS[6]}")'
    sp[f'F{rFC}'] = '=' + '+'.join(f'SUMIFS({P("AJ")},{P("B")},"{TIPOS[t]}")' for t in T_FC)
    rSUB = r1 + 8
    for col in 'ABCDEFG':
        c = sp[f'{col}{rSUB}']; c.font = F(b=True); c.fill = SUB_FILL; c.border = BORDE
    sp[f'A{rSUB}'] = 'Subtotal Crédito Fiscal'
    for col in 'BCDEF':
        sp[f'{col}{rSUB}'] = f'=SUM({col}{r1}:{col}{rN})'; sp[f'{col}{rSUB}'].number_format = NUM
    rNCC, rNCAS, rNCA = rSUB + 1, rSUB + 2, rSUB + 3
    cel(rNCC, 'A', '13 - Nota de Crédito C')
    for col in 'BCDE': cel(rNCC, col, None, fmt=NUM)
    cel(rNCC, 'F', f'=SUMIFS({P("AJ")},{P("B")},"{TIPOS[13]}")', fmt=NUM)
    cel(rNCC, 'G', 'SERVICIOS', color='FFFF0000')
    for rr, clas in ((rNCAS, 'SERVICIOS'), (rNCA, 'MERCADERIA')):
        cel(rr, 'A', '3 - Nota de Crédito A')
        cel(rr, 'B', '=' + todos(NETOS_C, T_NCA, clas), fmt=NUM); cel(rr, 'C', '=' + SG('K', T_NCA, clas), fmt=NUM)
        cel(rr, 'D', '=' + SG('L', T_NCA, clas), fmt=NUM); cel(rr, 'E', '=' + todos(IVAS_C, T_NCA, clas), fmt=NUM)
        cel(rr, 'F', f'=SUM(B{rr}:E{rr})', fmt=NUM); cel(rr, 'G', clas, color='FFFF0000')
    rSNC, rTOT = rNCA + 1, rNCA + 2
    for rr, txt, fml in ((rSNC, 'Subtotal Notas de Crédito s/Compras', '=SUM({c}%d:{c}%d)' % (rNCC, rNCA)),
                         (rTOT, 'TOTALES COMPRAS (Crédito Fiscal + NC)', '={c}%d+{c}%d' % (rSUB, rSNC))):
        for col in 'ABCDEF':
            c = sp[f'{col}{rr}']; c.font = F(b=True); c.fill = SUB_FILL if rr == rSNC else TOT_FILL; c.border = BORDE
            if col != 'A': c.value = fml.format(c=col); c.number_format = NUM
        sp[f'A{rr}'] = txt
    G_R = f'$G${r1}:$G${rN}'; E_R = f'$E${r1}:$E${rN}'
    p = r1
    der = [(p, 'IVA Computable al 10,5%', f'=SUMIF({G_R},"SERVICIOS AL 10,5%",{E_R})*liquidacion!$B${VG}', False),
           (p + 1, 'IVA Computable al 21%', f'=SUMIF({G_R},"SERVICIOS AL 21%",{E_R})*liquidacion!$B${VG}', False),
           (p + 2, 'IVA Computable al 27%', f'=SUMIF({G_R},"SERVICIOS AL 27%",{E_R})*liquidacion!$B${VG}', False),
           (p + 3, 'Total IVA Computable (Servicios)', f'=SUM(J{p}:J{p + 2})', True),
           (p + 4, 'Total IVA No Computable (Servicios)',
            f'=SUMIF({G_R},"SERVICIOS AL 10,5%",{E_R})+SUMIF({G_R},"SERVICIOS AL 21%",{E_R})+SUMIF({G_R},"SERVICIOS AL 27%",{E_R})-J{p + 3}', True),
           (p + 6, 'Crédito Fiscal Mercadería', f'=SUMIF({G_R},"MERCADERIA",{E_R})', False),
           (p + 7, 'SUBTOTAL CF (Merc.+Servicios)', f'=J{p + 3}+J{p + 6}', True)]
    for rr, t, fml, b in der:
        cel(rr, 'I', t, b=b); cel(rr, 'J', fml, b=b, fmt=NUM)
        if b:
            sp[f'I{rr}'].fill = SUB_FILL; sp[f'J{rr}'].fill = SUB_FILL
    cel(p + 9, 'I', f'Prorrateo: % de Venta Gravada de la hoja liquidacion (B{VG}).', it=True, sz_=9, color='FF595959', borde=False)
    q = p + 11
    cel(q, 'I', 'Servicios públicos incluidos (IVA)', b=True, fill=SUB_FILL); cel(q, 'J', None, b=True, fill=SUB_FILL)
    serv_list = list(OrderedDict.fromkeys(s[0] for s in SERVICIOS))
    for k, cu in enumerate(serv_list, 1):
        _ali = [x for x in prt if x['cuit'] == cu]
        _a = ('27%' if any(abs(sp[f'AD{x["fila"]}'].value or 0) > 0.004 for x in _ali) else
              '21%' if any(abs(sp[f'AB{x["fila"]}'].value or 0) > 0.004 for x in _ali) else '10,5%')
        manual = any(m_[0] == cu for m_ in MANUALES)
        cel(q + k, 'I', f'{SERV_CUITS[cu]} ({cu}) - en SERVICIOS AL {_a}' + (' (agregado a mano, falta en ARCA)' if manual else ''))
        if manual: sp[f'I{q + k}'].fill = NARANJA
        cel(q + k, 'J', '=' + '+'.join(f'SUMIFS({P(c)},{P("F")},{cu})' for c in ('Z', 'AB', 'AD')), fmt=NUM)
    nS = max(len(serv_list), 1)
    cel(q + nS + 1, 'I', 'Total servicios públicos', b=True, fill=SUB_FILL)
    cel(q + nS + 1, 'J', f'=SUM(J{q + 1}:J{q + nS})', b=True, fmt=NUM, fill=SUB_FILL)
    notas = ['Los cuadros de Ventas y Compras reproducen el detalle por tipo de comprobante utilizado para la carga en el Portal IVA de ARCA.',
             'Clasificación según el Maestro de Proveedores (rubro CMV = MERCADERIA, resto = SERVICIOS), columnas AG a AK. '
             '1 - Factura A incluye 63 - Liquidación A (bancos), 81 - Tique Factura A, 4 - Recibo A y 201 - FCE A; 11 - Factura C incluye '
             '111 - Tique Factura C; Factura C y tiques sin IVA discriminado van por su total sin percepciones. Notas de crédito en negativo.']
    for n, t in enumerate(notas, rTOT + 1):
        cel(n, 'A', t, it=True, sz_=9, color='FF595959', borde=False)
    rCTL = rTOT + 4
    cel(rCTL, 'A', 'Control: total del libro sin percepciones ni impuestos', b=True, borde=False)
    cel(rCTL, 'F', f"=SUM({P('H')})-SUM($N${P_INI}:$S${P_FIN})", fmt=NUM, b=True)
    cel(rCTL + 1, 'A', 'Diferencia (debe ser 0; centavos = redondeo de ARCA)', b=True, borde=False)
    cel(rCTL + 1, 'F', f'=ROUND(F{rTOT}-F{rCTL},2)', fmt=NUM, b=True)
    for col, w in (('A', 38), ('I', 36), ('J', 18)):
        sp.column_dimensions[col].width = max(w, sp.column_dimensions[col].width or 0)

    wb.save(a.salida)
    return {
        'ok': True, 'sociedad': SOC, 'cuit': CUIT, 'periodo': PERIODO, 'salida': a.salida, 'hojas': wb.sheetnames,
        'reemplazadas': ya,
        'portal': {'comprobantes': PORTAL_ARCA, 'total': PORTAL_ARCA_TOT, 'filas': [P_INI, P_FIN], 'original': PORTAL_ORIG},
        'zetti': {'comprobantes': len(zetti), 'total': round(sum(z['total'] for z in zetti), 2), 'filas': [Z_INI, Z_FIN],
                  'comparacion_unidad': zcmp},
        'ncr': {'eliminadas': len(ncr_elim), 'por_drogueria': dict(elim_drog), 'en_borrador': 0,
                'nc_droguerias_en_borrador': len(nc_drog), 'listados_revisados': len(listados),
                'revisar_posible_ncr': [f[4] for f in revisar], 'confirmadas_no_ncr': [f[4] for f in confirmadas]},
        'servicios': {'en_borrador': [SERV_CUITS[s[0]] for s in SERVICIOS if not s[5]],
                      'agregados_a_mano': [SERV_CUITS[m_[0]] for m_ in MANUALES], 'diferencias': serv_dif,
                      'faltan_vs_mes_anterior': serv_faltan},
        'cruce': dict(Counter(f[0] for f in filas_cruce)),
        'maestro': {'estados': dict(estados), 'excluidos': excluidos,
                    'pendientes': {c: v for c, v in cuenta_de.items() if v[2] != 'Clasificado'}},
        'filas': {'venta_gravada': VG, 'cuadro_total': rTOT, 'cuadro_dif': rCTL + 1, 'resumen_ctrl': R_CTRL, 'clasif_ctrl': RK_CTRL},
        'hallazgos': h,
    }


if __name__ == '__main__':
    a = argumentos()
    try:
        out = armar(a)
    except Corte as e:
        out = {'ok': False, 'code': e.code, 'error': str(e), **e.extra}
    except Exception as e:
        import traceback; traceback.print_exc()
        out = {'ok': False, 'code': 'error', 'error': f'{type(e).__name__}: {e}'}
    print(json.dumps(out, ensure_ascii=False, indent=1, default=str))
    sys.exit(0 if out.get('ok') else 1)
