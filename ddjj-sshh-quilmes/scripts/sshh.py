#!/usr/bin/env python3
"""sshh.py — helper de ddjj-sshh-quilmes.js.

  python sshh.py papel <papel.xlsx> [<otro.xlsx> ...]  -> Neto Gravado de Ventas Netas (hoja liquidacion)
  python sshh.py arba <dj.pdf>                         -> actividades de la DJ de anticipo IIBB de ARBA
  python sshh.py ddjj <DDJJ.pdf>                       -> datos de la Declaracion Jurada de ARQui
  python sshh.py recibo <Recibo.pdf>                   -> datos del Recibo de Pago de ARQui
  python sshh.py excel <datos.json> <salida.xlsx>      -> resumen

Todo sale como una linea JSON por stdout. Los importes vuelven como numeros
(punto decimal); si algo no se puede leer, viene {"error": "..."}.
"""

import json
import os
import re
import sys


def numero(s):
    """'1.234.567,89' / '1234567,89' / '87872.06' -> float"""
    s = s.strip()
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    return float(s)


def texto_pdf(pdf, paginas=1):
    from pypdf import PdfReader

    r = PdfReader(pdf)
    return '\n'.join((p.extract_text() or '') for p in r.pages[:paginas])


# ── Papel de trabajo ─────────────────────────────────────────────────────────

def papel(archivos):
    """Neto Gravado de la fila "Ventas Netas (s/ Notas de Credito)" de la hoja
    liquidacion. Prueba los archivos en orden y devuelve el primero con valor."""
    from openpyxl import load_workbook

    motivos = []
    for f in archivos:
        try:
            wb = load_workbook(f, data_only=True, read_only=True)
        except Exception as e:  # archivo abierto, corrupto, etc.
            motivos.append(f'{os.path.basename(f)}: no se pudo abrir ({e})')
            continue
        hoja = next((n for n in wb.sheetnames if n.strip().lower() == 'liquidacion'), None)
        if not hoja:
            motivos.append(f'{os.path.basename(f)}: no tiene la hoja liquidacion')
            continue
        col = None
        for fila in wb[hoja].iter_rows(values_only=True):
            celdas = [str(c).strip() if c is not None else '' for c in fila]
            if col is None and 'Neto Gravado' in celdas:
                col = celdas.index('Neto Gravado')
            if celdas and celdas[0].startswith('Ventas Netas'):
                if col is None:
                    motivos.append(f'{os.path.basename(f)}: la hoja liquidacion no tiene la columna Neto Gravado')
                    break
                v = fila[col]
                if isinstance(v, (int, float)):
                    return {'archivo': f, 'neto_gravado': float(v)}
                motivos.append(f'{os.path.basename(f)}: la celda de Ventas Netas / Neto Gravado esta vacia (formulas sin calcular: abrir y guardar el papel)')
                break
        else:
            motivos.append(f'{os.path.basename(f)}: la hoja liquidacion no tiene la fila Ventas Netas')
    return {'error': '; '.join(motivos) or 'no hay papel de trabajo'}


# ── DJ de anticipo IIBB (ARBA) ───────────────────────────────────────────────

def arba(pdf):
    t = texto_pdf(pdf)
    if not t.strip():
        return {'archivo': pdf, 'error': 'el PDF no tiene texto (escaneado o impreso como imagen): no se puede leer'}
    if 'Ingresos Brutos' not in t and 'INGRESOS BRUTOS' not in t:
        return {'archivo': pdf, 'error': 'el PDF no es una DJ de anticipo de IIBB de ARBA'}
    # "477320 $ 11.127.543,00 3,50 $ 389.464,00"
    act = {}
    for cod, ing in re.findall(r'\b(\d{6}) \$ ([\d.]+,\d\d) [\d,]+ \$', t):
        act[cod] = act.get(cod, 0) + numero(ing)
    per = re.search(r'(?<![\d/])(\d{1,2})/(20\d\d)(?![\d/])', t)
    cuit = re.search(r'\b(\d{2})-(\d{8})-(\d)\b', t)
    out = {'archivo': pdf, 'actividades': act}
    if per:
        out['periodo'] = f'{per.group(2)}{int(per.group(1)):02d}'
    if cuit:
        out['cuit'] = ''.join(cuit.groups())
    if not act:
        out['error'] = 'no encontre las actividades (codigo $ ingreso) en la DJ'
    return out


# ── Declaracion Jurada de ARQui ──────────────────────────────────────────────

def ddjj(pdf):
    t = texto_pdf(pdf)
    lineas = [l.strip() for l in t.splitlines() if l.strip()]
    if not any('Seguridad e Higiene' in l or 'TISEH' in l for l in lineas):
        return {'archivo': pdf, 'error': 'el PDF no es una DDJJ de Seguridad e Higiene de ARQui'}
    out = {'archivo': pdf, 'caracter': lineas[0] if lineas else None, 'actividades': []}
    for l in lineas:
        if re.fullmatch(r'\d{11}', l) and 'cuit' not in out:
            out['cuit'] = l
        elif re.fullmatch(r'\d{9}-\d', l) and 'cuenta' not in out:
            out['cuenta'] = l
    per = re.search(r'Per[ií]odo\n(\d{1,2})/(\d{4})', '\n'.join(lineas))
    if per:
        out['periodo'] = f'{per.group(2)}{int(per.group(1)):02d}'
    # Tabla INGRESOS DEVENGADOS: Tipo, codigo, descripcion (1 o mas lineas),
    # alicuota (0,013), ingreso (11127543,00) y tasa determinada (144658.06).
    i = next((k for k, l in enumerate(lineas) if l.startswith('INGRESOS DEVENGADOS')), None)
    if i is not None:
        k = i + 1
        while k < len(lineas) and not lineas[k].startswith('TOTAL'):
            if re.fullmatch(r'\d{6}', lineas[k]):
                cod = lineas[k]
                desc = []
                j = k + 1
                while j < len(lineas) and not re.fullmatch(r'0,\d+', lineas[j]):
                    desc.append(lineas[j])
                    j += 1
                if j + 2 < len(lineas):
                    out['actividades'].append({
                        'codigo': cod,
                        'descripcion': ' '.join(desc),
                        'alicuota': numero(lineas[j]),
                        'ingreso': numero(lineas[j + 1]),
                        'tasa': numero(lineas[j + 2]),
                    })
                k = j + 3
                continue
            k += 1
    tot = re.search(r'TOTAL \$:\n([\d.,]+)', '\n'.join(lineas))
    if tot:
        out['total'] = numero(tot.group(1))
    rec = re.search(r'RECIBO N°:\nFECHA VTO:\nEMISIÓN:\n(\d+)\n(\d\d/\d\d/\d{4})\n(\d\d/\d\d/\d{4})', '\n'.join(lineas))
    if rec:
        out['recibo'], out['vencimiento'], out['emision'] = rec.groups()
    if not out['actividades']:
        out['error'] = 'no encontre la tabla de ingresos devengados en la DDJJ'
    return out


# ── Recibo de Pago de ARQui (la boleta que va a TSH) ─────────────────────────

def recibo(pdf):
    t = texto_pdf(pdf)
    lineas = [l.strip() for l in t.splitlines() if l.strip()]
    if not any('SEGURIDAD E HIGIENE' in l for l in lineas):
        return {'archivo': pdf, 'error': 'el PDF no es un recibo de Seguridad e Higiene de ARQui'}
    out = {'archivo': pdf, 'lineas': []}
    cta = next((l for l in lineas if re.fullmatch(r'\d{9}-\d', l)), None)
    if cta:
        out['cuenta'] = cta
    vistos = set()
    for k, l in enumerate(lineas):
        m = re.fullmatch(r'(\d{4})\.(\d\d)\.M\.\d{3}\.\d\d', l)
        if m and l not in vistos and k + 1 < len(lineas):
            vistos.add(l)
            out['lineas'].append({'periodo': l, 'deuda': numero(lineas[k + 1])})
    out['total'] = round(sum(x['deuda'] for x in out['lineas']), 2)
    rec = re.search(r'RECIBO N °:\nFECHA VTO:\nEMISIÓN:\n(\d+)\n(\d\d/\d\d/\d{4})\n(\d\d/\d\d/\d{4})', '\n'.join(lineas))
    if rec:
        out['recibo'], out['vencimiento'], out['emision'] = rec.groups()
    if not out['lineas']:
        out['error'] = 'no encontre las lineas de periodo del recibo'
    return out


# ── Resumen ──────────────────────────────────────────────────────────────────

COLUMNAS = [
    ('sociedad', 'Sociedad', 26), ('cuit', 'CUIT', 13), ('cuenta', 'Cuenta', 13), ('periodo', 'Período', 9),
    ('estado', 'Estado', 16), ('fecha', 'Presentada', 11), ('caracter', 'Carácter', 13),
    ('declarado', 'Declarado 477320', 15), ('otras', 'Otras actividades', 18), ('tasa', 'Tasa determinada', 14),
    ('papel', 'Papel (Neto Gravado)', 17), ('papel_entero', 'Papel (entero)', 14), ('arba', 'ARBA 477320', 14),
    ('dif_papel', 'Dif. vs papel', 12), ('dif_arba', 'Dif. vs ARBA', 12), ('boleta', 'Boleta', 14),
    ('vencimiento', 'Vencimiento', 11), ('recibo', 'Nº recibo', 14), ('observacion', 'Observación', 70),
    ('archivo', 'Archivo', 60),
]


def excel(datos_json, salida):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    datos = json.load(open(datos_json, encoding='utf-8'))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Control'
    ws.append([c[1] for c in COLUMNAS])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDEBF7')
    amarillo = PatternFill('solid', fgColor='FFF2CC')
    for f in datos.get('filas', []):
        ws.append([f.get(k) for k, _, _ in COLUMNAS])
        if f.get('estado') not in ('OK',):
            for c in ws[ws.max_row]:
                c.fill = amarillo
    money = '#,##0.00'
    for k, (clave, _, ancho) in enumerate(COLUMNAS, start=1):
        letra = ws.cell(row=1, column=k).column_letter
        ws.column_dimensions[letra].width = ancho
        if clave in ('declarado', 'tasa', 'papel', 'papel_entero', 'arba', 'dif_papel', 'dif_arba'):
            for (c,) in ws.iter_rows(min_row=2, min_col=k, max_col=k):
                c.number_format = money
        if clave in ('observacion',):
            for (c,) in ws.iter_rows(min_row=2, min_col=k, max_col=k):
                c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions

    if datos.get('control'):
        wc = wb.create_sheet('Control planilla')
        cols = [('cuenta', 'Cuenta ARQui', 13), ('cuit_arqui', 'CUIT (ARQui)', 13), ('sociedad', 'Sociedad (carpeta)', 26),
                ('nombre_planilla', 'Nombre en la planilla', 22), ('cuit_planilla', 'CUIT planilla', 13), ('resultado', 'Resultado', 70)]
        wc.append([c[1] for c in cols])
        for c in wc[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill('solid', fgColor='DDEBF7')
        for f in datos['control']:
            wc.append([f.get(k) for k, _, _ in cols])
            if f.get('resultado') != 'OK':
                for c in wc[wc.max_row]:
                    c.fill = amarillo
        for k, (_, _, ancho) in enumerate(cols, start=1):
            wc.column_dimensions[wc.cell(row=1, column=k).column_letter].width = ancho
        wc.freeze_panes = 'A2'
    wb.save(salida)
    return {'archivo': salida}


def main():
    if len(sys.argv) < 3:
        print(json.dumps({'error': 'uso: sshh.py papel|arba|ddjj|recibo|excel <archivo> ...'}))
        sys.exit(1)
    cmd, args = sys.argv[1], sys.argv[2:]
    try:
        if cmd == 'papel':
            r = papel(args)
        elif cmd == 'arba':
            r = arba(args[0])
        elif cmd == 'ddjj':
            r = ddjj(args[0])
        elif cmd == 'recibo':
            r = recibo(args[0])
        elif cmd == 'excel':
            r = excel(args[0], args[1])
        else:
            r = {'error': f'comando desconocido: {cmd}'}
    except Exception as e:
        r = {'error': f'{type(e).__name__}: {e}'}
    print(json.dumps(r, ensure_ascii=False))


if __name__ == '__main__':
    main()
