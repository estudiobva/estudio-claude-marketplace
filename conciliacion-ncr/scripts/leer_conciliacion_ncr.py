# leer_conciliacion_ncr.py <AAAAMM-COMPRAS-SIN-NCR-SOCIEDAD.xlsx>
#
# Lee la salida de la skill conciliacion-ncr y devuelve por stdout (JSON) lo que
# necesita eliminar-ncr.js: la lista de NCR conciliadas (hoja "Eliminados") y los
# totales esperados de "Compras S.NCR" para verificar al final.
#
# Solo se toma la hoja Eliminados. "NC sin respaldo" y "NCR no encontradas" NO se
# eliminan nunca: ese es el criterio conservador de la conciliacion.
#
# Solo usa openpyxl (sin pandas) para que corra en cualquier maquina del estudio.
import sys, json, datetime
from openpyxl import load_workbook

CUITS_DROGUERIAS = {'30517059095': 'MONROE', '30516968431': 'SUIZO'}


def hoja(wb, nombre):
    if nombre not in wb.sheetnames:
        raise SystemExit(f'El archivo no tiene la hoja "{nombre}": no es una salida de conciliacion-ncr.')
    filas = list(wb[nombre].iter_rows(values_only=True))
    if not filas:
        return []
    cab = [str(c).strip() if c is not None else '' for c in filas[0]]
    return [dict(zip(cab, f)) for f in filas[1:] if any(v not in (None, '') for v in f)]


def numero(v):
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or '').strip()
    if ',' in s:                      # texto del CSV de ARCA: 1.234,56
        s = s.replace('.', '').replace(',', '.')
    return float(s)


def fecha_iso(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d')
    return str(v or '')[:10]


wb = load_workbook(sys.argv[1], read_only=True, data_only=True)
elim = hoja(wb, 'Eliminados')
compras = hoja(wb, 'Compras S.NCR')

items, errores = [], []
if elim:
    faltan = [c for c in ('_drog', '_farmacia', '_cuit', '_pv', '_nro', '_imp', '_fecha') if c not in elim[0]]
    if faltan:
        raise SystemExit(f'A la hoja Eliminados le faltan columnas: {faltan}')

for n, r in enumerate(elim, start=2):
    cuit = ''.join(ch for ch in str(r['_cuit']) if ch.isdigit())
    try:
        pv, nro = int(numero(r['_pv'])), int(numero(r['_nro']))
        imp = round(numero(r.get('portal_importe') if r.get('portal_importe') not in (None, '') else r['_imp']), 2)
    except ValueError:
        errores.append(f'fila {n}: PV/numero/importe no numericos'); continue
    if cuit not in CUITS_DROGUERIAS:
        errores.append(f'fila {n}: CUIT {cuit} no es Suizo ni Monroe'); continue
    fecha = fecha_iso(r.get('portal_fecha') or r['_fecha'])
    items.append({'fila': n, 'drogueria': r['_drog'], 'farmacia': r['_farmacia'],
                  'cuit': cuit, 'tipo': 3, 'pv': pv, 'nro': nro, 'importe': imp, 'fecha': fecha})

vistas, dup = set(), set()
for i in items:
    k = (i['cuit'], i['pv'], i['nro'])
    (dup if k in vistas else vistas).add(k)
if dup:
    errores.append(f'comprobantes repetidos en Eliminados: {sorted(dup)[:5]}')

esperado = None
if compras and 'Importe Total' in compras[0]:
    esperado = {'comprobantes': len(compras),
                'total': round(sum(numero(c['Importe Total'] or 0) for c in compras), 2)}

print(json.dumps({'items': items, 'errores': errores, 'compras_sin_ncr': esperado}, ensure_ascii=False))
